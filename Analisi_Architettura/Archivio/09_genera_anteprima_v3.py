"""Anteprima v3 dello stile: affinamenti dopo la valutazione della v2.

Uso:
    python Analisi_Architettura/09_genera_anteprima_v3.py

Scrive in Analisi_Architettura/ 09_Anteprima_v3_registro.xlsx (report e registro in un solo
file) e 09_Anteprima_v3_schede.xlsx (schede dei rischi a parte). Dati di esempio del pacchetto,
calcolati in Python e scritti come valori: NON e' collegato al registro. Riprende la
valutazione della v2 (09_Anteprima_v2_valutata.xlsx). Riusa gli strumenti della v2 e della v1.
"""

import importlib.util
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

QUI = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location("anteprima_v2", QUI / "09_genera_anteprima_v2.py")
v2 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(v2)
v1, g = v2.v1, v2.g

USCITA_REGISTRO = QUI / "09_Anteprima_v3_registro.xlsx"
USCITA_SCHEDE = QUI / "09_Anteprima_v3_schede.xlsx"
BLU_NOTTE, BLU_MEDIO, BLU_CHIARO = v2.BLU_NOTTE, v2.BLU_MEDIO, v2.BLU_CHIARO
ROSSO, AMBRA, GRIGIO, BLU_DIM = v2.ROSSO, v2.AMBRA, v2.GRIGIO, v2.BLU_DIM
TESTO, TESTO_TENUE, FONDO_TENUE = v2.TESTO, v2.TESTO_TENUE, v2.FONDO_TENUE
ROSSO_TENUE, AMBRA_TENUE, VERDE_TENUE = v2.ROSSO_TENUE, v2.AMBRA_TENUE, v2.VERDE_TENUE
SCALA, SCALA_CONF, PROCESSO, TAB = v2.SCALA, v2.SCALA_CONF, v2.PROCESSO, v2.TAB
fill, cella, fascia, guida, num = v2.fill, v2.cella, v2.fascia, v2.guida, v2.num
bordo_sottile = v2.bordo_sottile
PESI = {"sev": 0.30, "prob": 0.20, "gap": 0.20, "vel": 0.15, "pro": 0.15}


def score(sev, vel, pro, gap, prob):
    return round(PESI["sev"] * sev + PESI["vel"] * vel + PESI["pro"] * pro + PESI["gap"] * gap + PESI["prob"] * prob, 2)


def dati():
    rischi, kri = v2.dati()
    for r in rischi:
        storia = []
        for ciclo in ("2025Q4", "2026Q1", "2026Q2"):
            s = g.STORIA[ciclo].get(r["id"])
            if s:
                storia.append((ciclo, s[0], g.PROBABILITA[r["id"]], s[3], s[1], s[2], s[4], score(s[0], s[1], s[2], s[3], g.PROBABILITA[r["id"]])))
        r["storia"] = storia
        r["delta6"] = round(r["score"] - storia[0][-1], 2) if storia[0][0] == "2025Q4" else None
    return rischi, kri


def legenda(ws, riga, ultima_col, testo, altezza=None):
    """Legenda delle abbreviazioni in fondo a un foglio del report (nota D della v2)."""
    ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=ultima_col)
    cella(ws, riga, 1, "Legenda — " + testo, font=Font(size=8, italic=True, color=TESTO_TENUE),
          alignment=Alignment(wrap_text=True, vertical="top"))
    larghezza = sum((ws.column_dimensions[get_column_letter(j)].width or 8.43) for j in range(1, ultima_col + 1))
    ws.row_dimensions[riga].height = altezza or 11 * (len(testo) // int(larghezza * 1.45) + 1) + 4


LEG_RISCHIO = "NFR-nn = codice del rischio; ◆ = tema all'attenzione del supervisore (BCE, EBA, Banca d'Italia), indicato dal validatore"
LEG_MOSCA = ("X = anni per decidere e avviare la mitigazione (un trimestre); Y = anni per mitigare (azioni o adeguamento strutturale); "
             "Z = anni prima che il rischio possa manifestarsi (dalla Prossimita); AGIRE ORA = X + Y > Z")
LEG_PXI = "P = Probabilita, S = Severity, P x I = Probabilita x Severity (impatto), da 1 a 25"


# ============================================================== radar: legenda dentro il grafico (nota 9 della v2)
_grafico_radar_v2 = v2.grafico_radar


def grafico_radar(ws, info, piccolo=False):
    """Come la v2, con la linea di Mosca spiegata in legenda sotto il campione giallo."""
    col, r0 = info["col"], info["r0"]
    # Campione della linea rossa: due punti sotto il campione giallo (colonne di appoggio 22, 23).
    ws.cell(row=r0, column=col(21), value=-6.4)          # campione giallo un po' piu' in alto
    ws.cell(row=r0 + 2, column=col(22), value=-7.35)
    ws.cell(row=r0 + 2, column=col(23), value=-7.45)
    ws.cell(row=r0 + 3, column=col(22), value=-6.45)
    ws.cell(row=r0 + 3, column=col(23), value=-7.45)
    for c_ in (col(22), col(23)):
        ws.column_dimensions[get_column_letter(c_)].hidden = True
    ch = _grafico_radar_v2(ws, info, piccolo)
    # L'etichetta sul cerchio rosso diventa corta: la spiegazione sta in legenda.
    ch.series[0].tx = SeriesLabel(v="Mosca")
    s = Series(Reference(ws, min_col=col(23), min_row=r0 + 2, max_row=r0 + 3), xvalues=Reference(ws, min_col=col(22), min_row=r0 + 2, max_row=r0 + 3))
    s.tx = SeriesLabel(v="AGIRE ORA: X + Y > Z" if piccolo else "dentro il cerchio rosso: AGIRE ORA (X + Y > Z)")
    if piccolo:
        ch.series[-1].tx = SeriesLabel(v="attenzione supervisore")
    s.marker.symbol = "none"
    s.graphicalProperties.line.solidFill = ROSSO
    s.graphicalProperties.line.width = 28575
    base_l = dict(showSerName=True, showVal=False, showCatName=False, showLegendKey=False, showPercent=False, dLblPos="r")
    s.dLbls = DataLabelList(dLbl=[DataLabel(idx=1, txPr=v1.testo_rich(650 if piccolo else 800, True, ROSSO), **base_l)],
                            showVal=False, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False)
    ch.series.append(s)
    if piccolo:
        ch.height = ch.width = 8.2
    return ch


v2.grafico_radar = grafico_radar      # usato anche dal foglio Radar della v2


# ============================================================== Sintesi: un solo foglio (note 3, 4, F della v2)
def perche_breve(r, kri):
    return v2.perche_segnalato(r, kri)


def foglio_sintesi(ws, rischi, kri):
    ws.sheet_properties.tabColor = TAB["report"]
    fascia(ws, 15, "Quadro d'insieme")
    for j in range(1, 16):
        ws.column_dimensions[get_column_letter(j)].width = 10.5
    ws.column_dimensions["I"].width = 3
    agire = [r for r in rischi if r["urgenza"] == "AGIRE ORA"]
    variati = [r for r in rischi if r["esito"] == "Variato"]
    nuovi = [r for r in rischi if r["esito"] == "Nuovo"]
    rossi = [k for k in kri if k["stato"] == "Rosso"]
    riquadri = [
        ("Rischi valutati", str(len(rischi)), f"su {len(rischi)} in anagrafica · completo", FONDO_TENUE, BLU_NOTTE),
        ("Da affrontare subito", str(len(agire)), "X + Y > Z: " + ", ".join(r["id"] for r in agire), ROSSO_TENUE, ROSSO),
        ("Variati · nuovi", f"{len(variati)} · {len(nuovi)}", "rispetto al 2026Q1", FONDO_TENUE, BLU_NOTTE),
        ("KRI in rosso", str(len(rossi)), f"su {len(kri)} in vigore", ROSSO_TENUE if rossi else FONDO_TENUE, ROSSO),
        ("Validazione", "✓", "validato il 15/07/2026", VERDE_TENUE, "375623"),
    ]
    for (da, a), (didascalia, numero, sotto, fondo, colore) in zip([(1, 3), (4, 6), (7, 9), (10, 12), (13, 15)], riquadri):
        l1, l2 = get_column_letter(da), get_column_letter(a)
        for r, testo, font in [(4, didascalia, Font(size=9, bold=True, color=TESTO_TENUE)),
                               (5, numero, Font(size=26, bold=True, color=colore)), (6, sotto, Font(size=9, color=TESTO_TENUE))]:
            ws.merge_cells(f"{l1}{r}:{l2}{r}")
            cella(ws, r, da, testo, font=font, alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))
            for x in range(da, a + 1):
                ws.cell(row=r, column=x).fill = fill(fondo)
            ws.cell(row=r, column=a).border = Border(right=Side(style="thick", color="FFFFFF"))
    ws.row_dimensions[5].height = 36
    ws.row_dimensions[6].height = 22

    v1.sezione(ws, "A8", "H8", "Priorita' — primi cinque per score composito")
    v1.sezione(ws, "J8", "O8", "Gravita': Probabilita x Severity")
    for j, t in [(1, "#"), (2, "Rischio"), (6, "Score"), (7, "Variazione"), (8, "Urgenza")]:
        cella(ws, 9, j, t, font=Font(size=9, bold=True, color=TESTO_TENUE))
    ws.merge_cells("B9:E9")
    for k, r in enumerate(rischi[:5], start=1):
        riga = 9 + k
        cella(ws, riga, 1, k, font=Font(bold=True, color=BLU_NOTTE), alignment=Alignment(vertical="center"))
        ws.merge_cells(f"B{riga}:E{riga}")
        cella(ws, riga, 2, f"{r['id']}  {r['nome']}" + ("  ◆" if r["vigilanza"] else ""), font=Font(size=10),
              alignment=Alignment(wrap_text=True, vertical="center"))
        cella(ws, riga, 6, r["score"], number_format="0.00", font=Font(bold=True), alignment=Alignment(vertical="center"))
        cella(ws, riga, 7, v1.freccia(r["delta"]), alignment=Alignment(vertical="center"),
              font=Font(size=10, color=ROSSO if (r["delta"] or 0) > 0 else BLU_DIM if (r["delta"] or 0) < 0 else TESTO_TENUE))
        v1.chip_urgenza(ws.cell(row=riga, column=8), r["urgenza"])
        ws.row_dimensions[riga].height = 25
        for j in range(1, 9):
            ws.cell(row=riga, column=j).border = Border(bottom=bordo_sottile)
    ws.conditional_formatting.add("F10:F14", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=5, color="8EA9DB"))
    cella(ws, 15, 1, "Score 1-5, 5 = peggiore; barra su scala fissa 0-5. Variazione sul ciclo precedente.",
          font=Font(size=8, italic=True, color=TESTO_TENUE))
    cella(ws, 9, 10, "P \\ S", font=Font(size=8, bold=True, color=TESTO_TENUE))
    for s in range(1, 6):
        cella(ws, 9, 10 + s, s, font=Font(size=8, bold=True, color=TESTO_TENUE), alignment=Alignment(horizontal="center"))
    bianco = Side(style="thin", color="FFFFFF")
    for i, p in enumerate(range(5, 0, -1)):
        riga = 10 + i
        cella(ws, riga, 10, p, font=Font(size=8, bold=True, color=TESTO_TENUE), alignment=Alignment(horizontal="center", vertical="center"))
        for s in range(1, 6):
            ids = [r["id"] for r in rischi if r["prob"] == p and r["sev"] == s]
            testo = "\n".join(" ".join(ids[i_:i_ + 2]) for i_ in range(0, len(ids), 2))
            cella(ws, riga, 10 + s, testo or None, fill=fill(v1.fascia_pxi(p * s)), font=Font(size=7, bold=True),
                  alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
                  border=Border(left=bianco, right=bianco, top=bianco, bottom=bianco))
    cella(ws, 15, 10, "Verde P x I fino a 6, ambra 8-12, rosso da 15.", font=Font(size=8, italic=True, color=TESTO_TENUE))

    # Radar di Mosca: dati di appoggio nel foglio stesso (il grafico non dipende da un altro foglio).
    v1.sezione(ws, "A17", "H17", "Urgenza: radar di Mosca")
    v1.sezione(ws, "J17", "O17", "I rischi del ciclo")
    info = v2.dati_radar(ws, rischi, base=40)
    ws.add_chart(grafico_radar(ws, info, piccolo=True), "B18")
    for i, r in enumerate(sorted(rischi, key=lambda x: x["id"])):
        riga = 18 + i
        ws.merge_cells(f"K{riga}:O{riga}")
        cella(ws, riga, 10, r["id"], font=Font(size=8, bold=True, color=BLU_NOTTE))
        cella(ws, riga, 11, r["nome"] + ("  ◆" if r["vigilanza"] else ""), font=Font(size=8), alignment=Alignment(vertical="center"))
        ws.row_dimensions[riga].height = 19.5
    cella(ws, 31, 1, "Radar: piu' vicino al centro = agire prima; area del cerchio = P x I; bordo giallo = attenzione del supervisore. "
                     "Numeri nel radar = NFR-nn.", font=Font(size=8, italic=True, color=TESTO_TENUE))
    ws.row_breaks.append(Break(id=31))

    # Pagina 2: il commento (sotto i grafici) e poi i rischi segnalati (nota F della v2).
    riga = 33
    v1.sezione(ws, f"A{riga}", f"O{riga}", "Il commento dell'analisi (agente A1, scritto prima della validazione)")
    riga += 1
    _, a5, _ = g.SINTESI_ESEMPIO["2026Q2"]
    commento = [("Messaggio chiave", v2.COMMENTO_ESTESO[0][1], None),
                ("Decisioni richieste", v2.COMMENTO_ESTESO[3][1], None),
                ("Controllo indipendente (A5)", a5 + " Osservazioni non bloccanti: 1.", None),
                ("Dopo il commento", "Il validatore ha applicato 3 rettifiche o respingimenti: numeri e grafici del report sono quelli "
                                     "validati e prevalgono sul commento.", AMBRA_TENUE)]
    for titolo, testo, fondo in commento:
        ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=3)
        ws.merge_cells(start_row=riga, start_column=4, end_row=riga, end_column=15)
        cella(ws, riga, 1, titolo, font=Font(size=10, bold=True, color=BLU_NOTTE), alignment=Alignment(vertical="top"))
        cella(ws, riga, 4, testo, font=Font(size=10), alignment=Alignment(wrap_text=True, vertical="top"))
        if fondo:
            for x in range(4, 16):
                ws.cell(row=riga, column=x).fill = fill(fondo)
        ws.row_dimensions[riga].height = 14 * (len(testo) // 135 + 1) + 4
        riga += 1
    riga += 1
    v1.sezione(ws, f"A{riga}", f"O{riga}", "Rischi segnalati e perche'")
    riga += 1
    cella(ws, riga, 1, "Rischi fra i primi tre, da affrontare subito o variati nel ciclo. Tutto calcolato dai dati validati del registro.",
          font=Font(size=8, italic=True, color=TESTO_TENUE))
    riga += 1
    for (da, a), t in [((1, 3), "Rischio"), ((4, 4), "Urgenza"), ((5, 8), "Perche' lo segnaliamo"), ((9, 15), "Cosa e' cambiato, e su quale evidenza")]:
        if da != a:
            ws.merge_cells(start_row=riga, start_column=da, end_row=riga, end_column=a)
        for x in range(da, a + 1):
            cella(ws, riga, x, t if x == da else None, font=Font(bold=True, color="FFFFFF"), fill=fill(BLU_MEDIO))
    riga += 1
    for r in [r for r in rischi if r["rank"] <= 3 or r["urgenza"] == "AGIRE ORA" or r["esito"] == "Variato"]:
        for (da, a), testo, font in [((1, 3), f"{r['id']}\n{r['nome']}" + ("\n◆ attenzione del supervisore" if r["vigilanza"] else ""),
                                      Font(size=9, bold=True, color=BLU_NOTTE)),
                                     ((5, 8), perche_breve(r, kri), Font(size=9)),
                                     ((9, 15), v2.cosa_cambiato(r), Font(size=9))]:
            ws.merge_cells(start_row=riga, start_column=da, end_row=riga, end_column=a)
            cella(ws, riga, da, testo, font=font, alignment=Alignment(wrap_text=True, vertical="top"))
        # L'urgenza e' un'etichetta esplicita (la v2 colorava di rosa la cella del rischio senza spiegarlo).
        v1.chip_urgenza(ws.cell(row=riga, column=4), r["urgenza"])
        ws.cell(row=riga, column=4).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        for x in range(1, 16):
            ws.cell(row=riga, column=x).border = Border(bottom=bordo_sottile)
        cambiato = v2.cosa_cambiato(r)
        righe = max(-(-len(perche_breve(r, kri)) // 52), sum(-(-len(x) // 100) for x in cambiato.split("\n")),
                    4 if r["vigilanza"] else 3)
        ws.row_dimensions[riga].height = 12 * righe + 6
        riga += 1
    riga += 1
    ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=15)
    cella(ws, riga, 1, "Limiti: " + g.leggi_limiti(), font=Font(size=8, italic=True, color=TESTO_TENUE),
          alignment=Alignment(wrap_text=True, vertical="top"))
    ws.row_dimensions[riga].height = 34
    legenda(ws, riga + 1, 15, f"{LEG_RISCHIO}. {LEG_PXI}. {LEG_MOSCA}. KRI = indicatore di rischio (Key Risk Indicator). "
                              "A1 = agente che valuta i rischi ogni trimestre; A5 = agente di controllo indipendente.")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.page_setup.scale = 80
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.left.text = "Anteprima di stile — dati di esempio"
    ws.oddFooter.right.text = "Pagina &P di &N"


# ============================================================== Delta con scala spiegata e motivazioni (nota 12 della v2)
def motivo_delta(r):
    if not r["var"]:
        return "Invariato: nessuna evidenza nuova dopo il 31/03/2026; restano valide le fonti del ciclo prima."
    return "\n".join(f"{v2.NOMI_DIM[d]} {a} → {b}: {ev} ({v2.data_it(dt)})" for d, a, b, ev, dt in r["var"])


def foglio_delta(ws, rischi):
    ws.sheet_properties.tabColor = TAB["report"]
    fascia(ws, 16, "che cosa e' cambiato dal ciclo precedente")
    ordinati = sorted([r for r in rischi if r["delta"] is not None], key=lambda r: -r["delta"])
    base = 30          # appoggio del grafico (nascosto)
    for j, t in enumerate(["Rischio", "In aumento", "In diminuzione", "Invariato"]):
        cella(ws, 4, base + j, t)
    for i, r in enumerate(ordinati):
        riga = 5 + i
        cella(ws, riga, base, f"{r['id']}  {r['nome'][:28]}{'…' if len(r['nome']) > 28 else ''}")
        colonna = base + 1 if r["delta"] > 0 else base + 2 if r["delta"] < 0 else base + 3
        cella(ws, riga, colonna, r["delta"], number_format="+0.00;-0.00;0.00")
    for c_ in range(base, base + 4):
        ws.column_dimensions[get_column_letter(c_)].hidden = True
    ultima = 4 + len(ordinati)
    su, giu = sum(r["delta"] > 0 for r in ordinati), sum(r["delta"] < 0 for r in ordinati)
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "clustered"
    ch.overlap = 100
    ch.gapWidth = 50
    ch.visible_cells_only = False
    ch.title = v1.titolo_grafico(f"{su} in aumento, {giu} in diminuzione, {len(ordinati) - su - giu} invariati")
    for c_, colore in ((base + 1, ROSSO), (base + 2, BLU_DIM), (base + 3, GRIGIO)):
        ch.add_data(Reference(ws, min_col=c_, min_row=4, max_row=ultima), titles_from_data=True)
        s = ch.series[-1]
        s.graphicalProperties.solidFill = colore
        s.graphicalProperties.line.noFill = True
        s.invertIfNegative = False
        if c_ != base + 3:
            v1.etichette_valore(s, formato="+0.00;-0.00;0.00", colore=colore)
    ch.set_categories(Reference(ws, min_col=base, min_row=5, max_row=ultima))
    ch.x_axis.scaling.orientation = "maxMin"
    ch.x_axis.tickLblPos = "low"
    ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = -1, 1, 0.25
    ch.y_axis.numFmt = "+0.00;-0.00;0"
    ch.varyColors = False
    v1.pulisci(ch)
    ch.height, ch.width = 11.5, 14.5
    ws.add_chart(ch, "A4")
    for j in range(1, 17):
        ws.column_dimensions[get_column_letter(j)].width = 9
    ws.column_dimensions["J"].width = 9
    ws.column_dimensions["K"].width = 7
    # A fianco: il perche' della variazione o della non variazione, nello stesso ordine delle barre.
    v1.sezione(ws, "J4", "P4", "Perche' varia (o no), nello stesso ordine")
    for j, t in [(10, "Rischio"), (11, "Δ"), (12, "Motivo, con l'evidenza")]:
        cella(ws, 5, j, t, font=Font(size=9, bold=True, color="FFFFFF"), fill=fill(BLU_MEDIO))
    for x in range(12, 17):
        ws.cell(row=5, column=x).fill = fill(BLU_MEDIO)
    ws.merge_cells("L5:P5")
    for i, r in enumerate(ordinati):
        riga = 6 + i
        ws.merge_cells(start_row=riga, start_column=12, end_row=riga, end_column=16)
        cella(ws, riga, 10, r["id"], font=Font(size=8.5, bold=True, color=BLU_NOTTE), alignment=Alignment(vertical="top"))
        colore = ROSSO if r["delta"] > 0 else BLU_DIM if r["delta"] < 0 else TESTO_TENUE
        cella(ws, riga, 11, r["delta"], number_format="+0.00;-0.00;0.00", font=Font(size=8.5, bold=True, color=colore),
              alignment=Alignment(vertical="top"))
        motivo = motivo_delta(r)
        cella(ws, riga, 12, motivo, font=Font(size=8, color=TESTO if r["var"] else TESTO_TENUE),
              alignment=Alignment(wrap_text=True, vertical="top"))
        for x in range(10, 17):
            ws.cell(row=riga, column=x).border = Border(bottom=bordo_sottile)
        ws.row_dimensions[riga].height = 11 * (sum(len(p) // 58 + 1 for p in motivo.split("\n"))) + 3
    fine = 6 + len(ordinati)
    ws.merge_cells(start_row=fine + 1, start_column=1, end_row=fine + 1, end_column=16)
    cella(ws, fine + 1, 1,
          "Come si legge la scala: il Δ e' la variazione dello score composito sul ciclo precedente. Ogni dimensione puo' cambiare al "
          "massimo di un punto fra due cicli, quindi un solo punto vale: Severity ±0,30, Probabilita ±0,20, Gap_Controlli ±0,20, "
          "Velocity ±0,15, Prossimita ±0,15. +1,00 si ha solo se tutte e cinque le dimensioni peggiorano di un punto; -1,00 se "
          "migliorano tutte. Un Δ oltre ±0,40 (soglia di Parametri) va guardato. Rosso = in aumento, blu = in diminuzione.",
          font=Font(size=9, color=TESTO), fill=fill(FONDO_TENUE), alignment=Alignment(wrap_text=True, vertical="top"))
    ws.row_dimensions[fine + 1].height = 44
    legenda(ws, fine + 2, 16, f"Δ = variazione sul ciclo precedente. {LEG_RISCHIO}.")
    v1.stampa(ws)
    ws.print_area = f"A1:P{fine + 2}"


# ============================================================== Registro (nota B della v2)
def foglio_registro(ws, rischi):
    ws.sheet_properties.tabColor = TAB["lavoro"]
    larghezze = [9, 38, 7, 6, 6, 6, 6, 6, 6, 8, 8, 8, 13, 8, 5, 4, 4, 4, 11, 28]
    for j, w in enumerate(larghezze, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    fascia(ws, 20, "Registro: il ciclo corrente")
    guida(ws, 4, 20,
          "COSA E': la fotografia del ciclo corrente, valori validati, ordinata per score. L'andamento di un rischio negli ultimi "
          "trimestri e' nella sua scheda (file delle schede).\n"
          "COME SI CALCOLA: Score = 0,30 Severity + 0,20 Probabilita + 0,20 Gap_Controlli + 0,15 Velocity + 0,15 Prossimita (pesi in "
          "Parametri). Δ ciclo prima = score meno score del trimestre precedente; Δ 6 mesi = score meno score di due trimestri prima. "
          "Margine di Mosca = Z - (X + Y); urgenza: margine < 0 AGIRE ORA, < 1 entro l'anno. KRI: quanti in rosso, ambra, verde. "
          "Notizie: segnali peggiorativi meno migliorativi delle ultime 4 settimane (non una sola settimana); oltre ±3 e' una "
          "tendenza (▲ o ▼).\n"
          "COLORI: intestazioni come chi scrive (azzurro trimestrale, verde settimanale, giallo validatore, grigio calcolo); "
          "dimensioni dal chiaro (1) allo scuro (5). Ogni colonna e' spiegata nel foglio Dizionario.", "calcolo")
    gruppi = [((1, 3), "Rischio", "calcolo"), ((4, 9), "Dimensioni 1-5 (A1, validate)", "trimestrale"),
              ((10, 14), "Calcolati dal foglio", "calcolo"), ((15, 15), "Sup.", "validatore"),
              ((16, 18), "KRI del ciclo", "trimestrale"), ((19, 19), "Notizie · 4 settimane", "settimanale"),
              ((20, 20), "Anagrafica", "validatore")]
    for (da, a), t, p in gruppi:
        if da != a:
            ws.merge_cells(start_row=6, start_column=da, end_row=6, end_column=a)
        for x in range(da, a + 1):
            cella(ws, 6, x, t if x == da else None, fill=fill(PROCESSO[p][0]), font=Font(size=9, bold=True, color=PROCESSO[p][1]),
                  alignment=Alignment(horizontal="center", wrap_text=True))
    ws.row_dimensions[6].height = 26
    intest = ["ID", "Nome", "Event type", "Sev", "Prob", "Gap", "Vel", "Pross", "Conf", "Score", "Δ ciclo prima", "Δ 6 mesi",
              "Urgenza", "Margine (anni)", "◆", "R", "A", "V", "Segnali netti", "Owner"]
    for j, t in enumerate(intest, start=1):
        cella(ws, 7, j, t, font=Font(size=9, bold=True, color=TESTO), border=Border(bottom=Side(style="medium", color=BLU_NOTTE)),
              alignment=Alignment(horizontal="center" if j > 2 else "left", wrap_text=True, vertical="bottom"))
    ws.row_dimensions[7].height = 28
    freccia6 = lambda d: "nuovo" if d is None else v1.freccia(d)
    for i, r in enumerate(rischi):
        riga = 8 + i
        valori = [r["id"], r["nome"], r["et"], r["sev"], r["prob"], r["gap"], r["vel"], r["pro"], r["conf"], r["score"],
                  v1.freccia(r["delta"]), freccia6(r["delta6"]), None, r["margine"], "◆" if r["vigilanza"] else None,
                  r["kri_r"] or None, r["kri_a"] or None, r["kri_v"] or None,
                  ("▲ +" if r["netti"] >= 3 else "▼ " if r["netti"] <= -3 else "● ") + str(r["netti"]), r["owner"]]
        for j, v in enumerate(valori, start=1):
            c = cella(ws, riga, j, v, font=Font(size=10), alignment=Alignment(horizontal="center" if j > 2 else "left", vertical="center"),
                      border=Border(bottom=bordo_sottile))
            if i % 2:
                c.fill = fill("FAFAFA")
        for j, v in zip(range(4, 9), (r["sev"], r["prob"], r["gap"], r["vel"], r["pro"])):
            ws.cell(row=riga, column=j).fill = fill(SCALA[v])
            ws.cell(row=riga, column=j).font = Font(size=10, bold=True, color="FFFFFF" if v == 5 else TESTO)
        ws.cell(row=riga, column=9).fill = fill(SCALA_CONF[r["conf"]])
        ws.cell(row=riga, column=10).number_format = "0.00"
        ws.cell(row=riga, column=10).font = Font(size=10, bold=True)
        for j, d in ((11, r["delta"]), (12, r["delta6"])):
            ws.cell(row=riga, column=j).font = Font(size=9, color=ROSSO if (d or 0) > 0 else BLU_DIM if (d or 0) < 0 else TESTO_TENUE)
        v1.chip_urgenza(ws.cell(row=riga, column=13), r["urgenza"])
        ws.cell(row=riga, column=14).number_format = "+0.00;-0.00;0.00"
        ws.cell(row=riga, column=15).font = Font(bold=True, color="BF8F00")
        for j, colore in ((16, ROSSO), (17, AMBRA), (18, "70AD47")):
            if ws.cell(row=riga, column=j).value:
                ws.cell(row=riga, column=j).fill = fill(colore)
                ws.cell(row=riga, column=j).font = Font(size=10, bold=True, color="FFFFFF" if j != 17 else TESTO)
        ws.cell(row=riga, column=19).font = Font(size=9, color=ROSSO if r["netti"] >= 3 else BLU_DIM if r["netti"] <= -3 else TESTO_TENUE)
        ws.row_dimensions[riga].height = 20
    fine = 7 + len(rischi)
    ws.conditional_formatting.add(f"J8:J{fine}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=5, color="8EA9DB"))
    legenda(ws, fine + 2, 20,
            "Sev = Severity (impatto); Prob = Probabilita; Gap = Gap_Controlli (quanto manca a un presidio adeguato); Vel = Velocity "
            "(quanto in fretta arriva il danno); Pross = Prossimita (quando puo' manifestarsi); Conf = Confidence (solidita' delle "
            "fonti, fuori dallo score); Δ = variazione dello score; Sup. / ◆ = tema all'attenzione del supervisore; R, A, V = KRI in "
            "rosso, ambra, verde; Segnali netti = notizie peggiorative meno migliorative nelle ultime 4 settimane (▲ ▼ oltre la "
            f"soglia di 3). Event type L1-01..L1-07 = categorie di perdita di Basilea (foglio Dizionario). {LEG_MOSCA}.")
    ws.freeze_panes = "C8"
    v1.stampa(ws)


# ============================================================== Dizionario completo (nota D della v2)
CHI_TESTO = {"A1": "A1, ciclo trimestrale; validato", "A3": "A3, ciclo settimanale", "A5": "A5, ciclo trimestrale",
             "script": "Script, quando scrive il ciclo", "formula": "Formula del foglio",
             "validazione": "Comando valida, dalla decisione del validatore", "validatore": "Validatore, a mano"}
COME = {
    "Score_Composito": "0,30 Sev + 0,20 Prob + 0,20 Gap + 0,15 Vel + 0,15 Pross, pesi in Parametri; 2 decimali; vuoto se manca una dimensione.",
    "Chiave": "Ciclo e ID uniti da \"|\": serve alle viste per trovare la riga.",
    "Mesi_Orizzonte": "Il numero di mesi dell'orizzonte (3, 6, 12); il piu' lungo fra le azioni non respinte e' la Y di Mosca.",
    "Stato@tblKRI": "Rosso se valore >= allerta; Ambra se >= attenzione; Verde sotto; Non rilevato senza valore; SOGLIE INVERTITE se allerta < attenzione.",
    "Settimana": "Lunedi' della settimana della data della notizia.",
    "Owner@tblRegister": "Copiato dall'anagrafica (Parametri, colonna L) al momento della scrittura.",
    "Esito": "Confermato se nessun valore cambia; Variato se ne cambia almeno uno; Nuovo per un rischio nuovo. Ricalcolato dopo la validazione.",
    "Variazioni": "Per ogni valore cambiato: da, a, evidenza e sua data (dopo il ciclo precedente); al massimo un punto per dimensione.",
    "Severity": "Descrittore della rubrica piu' aderente all'evidenza; al massimo un punto di variazione fra cicli.",
    "Probabilita": "Come Severity; un numero senza fonte e' dichiarato come valutazione.",
    "Gap_Controlli": "Come Severity; scende solo con un controllo testato.",
    "Confidence": "4 richiede almeno 2 fonti, 5 almeno 3.",
    "KRI_ID": "Riuso del codice se il KRI resta; altrimenti primo progressivo libero del rischio (NFR-nn-Kn), mai riusato.",
    "ID_Azione": "<ciclo>-<rischio>-A<n>, n = posizione dell'azione nel rischio.",
    "Valore_Corrente": "Rilevato dalla fonte interna; vuoto se non rilevato, mai ricopiato dal ciclo prima.",
    "Rischio_ID": "Rischio indicato da A3; NON PRESIDIATO se la notizia non tocca rischi censiti.",
    "Osservazioni_Controlli": "Conteggio dei rilievi 'Osservazione' di A5 e dei controlli dello script.",
    "Stato@tblAzioni": "Proposta — da valutare; Approvata o Respinta con la validazione; In corso e Completata a mano; Annullata se il rischio esce.",
    "Stato@tblValidazione": "Da decidere quando la scrive lo script; Applicata dopo il comando valida. Non si tocca a mano.",
    "Data_Scrittura": "Data del giorno in cui lo script scrive la riga.",
    "Approvato_Da": "Il nome passato al comando valida (--validatore).",
    "Data_Validazione": "Giorno in cui si lancia il comando valida.",
    "Validazione": "Da validare alla scrittura; il comando valida scrive Accettato, Rettificato o Respinto (un rischio non si respinge).",
    "Data_Decisione": "Giorno del comando valida che ha deciso l'azione.",
    "Decisa_Da": "Il nome passato al comando valida (--validatore).",
    "Data_Rilevazione": "Data di riferimento del ciclo se c'e' un valore rilevato; vuota altrimenti.",
    "In_Vigore": "Sì per i KRI proposti nel ciclo; No nella riga del ciclo in cui un KRI non e' riproposto o e' respinto.",
    "Ciclo_Definizione": "Ciclo in cui il KRI e' nato; non cambia quando se ne aggiorna la definizione.",
    "Periodo": "Ciclo (per il trimestrale) o lunedi' della settimana (per il settimanale) della scrittura.",
    "Oggetto": "Tipo della riga proposta dall'agente.",
    "ID": "Codice dell'oggetto proposto (rischio, KRI, azione, notizia).",
    "Nome@tblValidazione": "Copiato dall'oggetto, per riconoscerlo senza cambiare foglio.",
    "Campo": "Uno dei campi da validare: 6 per rischio, 3 per KRI, 2 per azione, 2 per notizia.",
    "Proposta_Agente": "Valore scritto dall'agente, copiato qui alla scrittura; resta per sempre.",
    "Decisione": "Scelta dall'elenco; lasciata vuota, con --accetta-il-resto diventa Accettato.",
    "Valore_Validato": "Rischio: intero 1-5 a non piu' di un punto dal ciclo prima; KRI: numero; azione: owner o 3/6/12 mesi.",
    "Motivazione": "Obbligatoria con Rettificato o Respinto: almeno 10 caratteri, con l'evidenza.",
    "Validato_Da": "Il nome passato al comando valida (--validatore).",
    "Data_Validazione@tblValidazione": "Giorno in cui si lancia il comando valida.",
    "Ciclo": "Ciclo da valutare, dalla baseline.",
    "Data_Riferimento": "Fine del trimestre valutato, dalla baseline.",
}
ESEMPI_EXTRA = {"Score_Composito": "3,80", "Chiave": "2026Q2|NFR-01", "Mesi_Orizzonte": "6", "Stato@tblKRI": "Ambra",
                "Settimana": "08/06/2026", "Variazioni": "Severity 3 → 4: report frodi (10/06/2026)", "Fonti": "4 fonti, 2 interne"}
VALORI = {"Ciclo": "AAAAQn", "Data_Riferimento": "data", "Severity": "1-5", "Velocity": "1-5", "Prossimita": "1-5",
          "Gap_Controlli": "1-5", "Probabilita": "1-5", "Confidence": "1-5", "Score_Composito": "1,00-5,00",
          "Esito": "Confermato, Variato, Nuovo", "Event_Type": "L1-01..L1-07", "Orizzonte": "3, 6, 12 mesi",
          "Segnale": "Peggiorativo, Migliorativo, Neutro, Non rilevante", "Decisione": "Accettato, Rettificato, Respinto",
          "Validazione": "Da validare, Accettato, Rettificato, Respinto", "In_Vigore": "Sì, No", "Attenzione_Supervisore": "Sì, No"}
CALCOLATI = [
    ("Δ ciclo prima", "Registro, Delta", "Formula del foglio", "Variazione dello score sul trimestre precedente.", "Score del ciclo meno score del ciclo prima, stessi pesi.", "-1..+1", "+0,30"),
    ("Δ 6 mesi", "Registro", "Formula del foglio", "Variazione dello score su due trimestri.", "Score del ciclo meno score di due cicli prima; \"nuovo\" se allora il rischio non c'era.", "-2..+2", "+0,45"),
    ("P x I", "Matrice, Radar", "Formula del foglio", "Gravita': probabilita' per impatto.", "Probabilita x Severity.", "1-25", "16"),
    ("Z", "Radar, Registro", "Formula del foglio", "Anni prima che il rischio possa manifestarsi.", "Dalla Prossimita: 5 → 0; 4 → 0,5; 3 → 1,5; 2 → 2,5; 1 → 4 (Parametri).", "anni", "0,5"),
    ("X", "Radar, Registro", "Parametro", "Anni per decidere e avviare.", "Costante in Parametri: un trimestre.", "anni", "0,25"),
    ("Y", "Radar, Registro", "Formula del foglio", "Anni per mitigare.", "Orizzonte piu' lungo delle azioni non respinte (1 anno se non ce ne sono); se maggiori, gli anni di adeguamento.", "anni", "0,5"),
    ("Margine / Urgenza", "Radar, Registro, Classifica, Sintesi", "Formula del foglio", "Quanto presto agire.", "Margine = Z - (X + Y). Sotto 0: AGIRE ORA; da 0 a 1: entro l'anno; oltre: margine in anni.", "anni", "-0,25 · AGIRE ORA"),
    ("KRI R, A, V", "Registro", "Formula del foglio", "Semaforo dei KRI del rischio.", "Conteggio dei KRI in vigore per Stato.", "interi", "0, 2, 1"),
    ("Segnali netti", "Registro", "Formula del foglio, dopo il settimanale", "Direzione delle notizie sul rischio nelle ultime 4 settimane.", "Peggiorative meno migliorative; ▲ o ▼ oltre la soglia di Parametri (3). Non cambia lo score.", "intero", "+3 ▲"),
    ("Posizione in classifica", "Classifica, Sintesi", "Formula del foglio", "Ordine per score composito.", "Dal piu' alto; a parita', ordine dell'anagrafica.", "1..n", "1"),
    ("Perche' lo segnaliamo", "Sintesi", "Formula del foglio", "Motivo per cui un rischio e' nella Sintesi.", "Primi tre per score, AGIRE ORA, variati; con urgenza, supervisore, KRI rossi, notizie.", "testo", "1° per score..."),
]
PARAMETRI = [
    ("Pesi dello score", "Parametri B4:B8", "Funzione di rischio (DECISIONI D-19)", "Peso di ogni dimensione nello score.", "Somma controllata = 1.", "0-1", "0,30 / 0,20 / 0,20 / 0,15 / 0,15"),
    ("Soglia di direzione", "Parametri B11", "Funzione di rischio", "Segnali netti oltre cui c'e' una tendenza.", "Confronto con i segnali netti di 4 settimane.", "intero", "3"),
    ("Soglia del Delta", "Parametri B12", "Funzione di rischio", "Variazione dello score da guardare.", "Evidenzia |Δ| >= soglia.", "decimale", "0,40"),
    ("Anni per Prossimita (Z)", "Parametri B51:B55", "Funzione di rischio", "Tempo prima della manifestazione per livello.", "Letti da Z.", "anni", "4 / 2,5 / 1,5 / 0,5 / 0"),
    ("Owner", "Parametri, anagrafica L", "Validatore, a mano", "Funzione responsabile del rischio.", "Resta da un ciclo all'altro; Da assegnare per un rischio nuovo.", "testo", "Funzione Antifrode"),
    ("Anni_Adeguamento", "Parametri, anagrafica M", "Validatore, a mano", "Durata di un adeguamento strutturale.", "Diventa la Y minima di Mosca.", "anni o vuoto", "NFR-11: 3"),
    ("Attenzione_Supervisore", "Parametri, anagrafica N", "Validatore, a mano", "Tema nelle priorita' di BCE, EBA o Banca d'Italia.", "Bordo giallo nel radar, ◆ nelle tabelle; non cambia score ne' urgenza.", "Sì, No", "Sì"),
]


def esempi():
    register, azioni, kri, news, validazione = g.dati_esempio()
    trova = lambda righe, **k: next(r for r in righe if all(r.get(a) == b for a, b in k.items()))
    return {"tblRegister": trova(register, Ciclo="2026Q2", ID_Rischio="NFR-01"), "tblAzioni": trova(azioni, ID_Rischio="NFR-01"),
            "tblKRI": trova(kri, Ciclo="2026Q2", KRI_ID="NFR-01-K1"), "tblNews": news[1],
            "tblSintesi": {"Ciclo": "2026Q2", "Data_Riferimento": "30/06/2026", "Sintesi_A1": "Trimestre con quattro rischi variati...",
                           "Sintesi_A5": "Controlli superati...", "Osservazioni_Controlli": 1, "Data_Scrittura": "30/06/2026"},
            "tblValidazione": trova(validazione, ID="NFR-10", Campo="Probabilita")}


def foglio_dizionario(ws):
    ws.sheet_properties.tabColor = TAB["lavoro"]
    for j, w in enumerate([22, 19, 24, 38, 56, 14, 22], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    fascia(ws, 7, "Dizionario dei dati")
    guida(ws, 4, 7, "Tutti i dati del registro, tabella per tabella, piu' i valori calcolati nelle viste e i parametri: dove "
                    "stanno, chi li scrive e quando, cosa significano, come si calcolano, che valori possono avere, un esempio "
                    "(NFR-01, ciclo 2026Q2). Nel file vero si genera dalla mappa delle colonne e dalla rubrica: una sola fonte.", "calcolo")
    intest = ["Dato", "Dove", "Chi lo scrive, quando", "Che cosa significa", "Come si calcola / regola", "Valori", "Esempio"]
    for j, t in enumerate(intest, start=1):
        cella(ws, 6, j, t, font=Font(bold=True, color="FFFFFF"), fill=fill(BLU_NOTTE))
    riga = 7
    es = esempi()

    def sezione(titolo):
        nonlocal riga
        ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=7)
        cella(ws, riga, 1, titolo, font=Font(bold=True, color=BLU_NOTTE), fill=fill(BLU_CHIARO))
        riga += 1

    def voce(valori, chi_proc):
        nonlocal riga
        for j, v in enumerate(valori, start=1):
            c = cella(ws, riga, j, v, font=Font(size=9, bold=j == 1), alignment=Alignment(wrap_text=True, vertical="top", horizontal="left"),
                      border=Border(bottom=bordo_sottile))
            if j == 3:
                c.fill = fill(PROCESSO[chi_proc][0])
        ws.row_dimensions[riga].height = 12 * max(2, len(str(valori[4] or "")) // 52 + 1, len(str(valori[3] or "")) // 36 + 1) + 4
        riga += 1

    proc_di = {"A1": "trimestrale", "A5": "trimestrale", "A3": "settimanale", "script": "calcolo", "formula": "calcolo",
               "validazione": "validatore", "validatore": "validatore"}
    tabelle = [("tblRegister", "RiskRegister — valutazioni dei rischi"), ("tblAzioni", "Azioni — azioni di mitigazione"),
               ("tblKRI", "KRI — storico degli indicatori"), ("tblNews", "NewsLog — notizie"),
               ("tblSintesi", "Sintesi_Cicli — commenti di A1 e A5"), ("tblValidazione", "Validazione — decisioni del validatore")]
    for nome_t, titolo in tabelle:
        sezione(titolo)
        spec = g.spec_tabella(nome_t)
        for c in spec["colonne"]:
            chi = g.chi_colonna(nome_t, c)
            n = c["nome"]
            come = COME.get(f"{n}@{nome_t}") or COME.get(n) or ("Dal payload dell'agente, scritto dallo script." if chi in ("A1", "A3", "A5") else "—")
            esempio = ESEMPI_EXTRA.get(f"{n}@{nome_t}") or ESEMPI_EXTRA.get(n) or es.get(nome_t, {}).get(n)
            if isinstance(esempio, (int, float)) and not isinstance(esempio, bool):
                esempio = num(esempio) if isinstance(esempio, float) else str(esempio)
            if isinstance(esempio, str) and len(esempio) > 38:
                esempio = esempio[:36] + "…"
            if hasattr(esempio, "strftime"):
                esempio = esempio.strftime("%d/%m/%Y")
            if isinstance(esempio, float):
                esempio = num(esempio)
            voce([n, spec["foglio"], CHI_TESTO[chi], c["significato"], come, VALORI.get(n, "—"), esempio], proc_di[chi])
    sezione("Valori calcolati nelle viste")
    for v in CALCOLATI:
        voce(list(v), "calcolo")
    sezione("Parametri e anagrafica (input)")
    for v in PARAMETRI:
        voce(list(v), "validatore")
    ws.freeze_panes = "B7"
    v1.stampa(ws)
    ws.page_setup.fitToHeight = 0


# ============================================================== Validazione: Stato a elenco (nota E della v2)
def foglio_validazione(ws):
    v2.foglio_validazione(ws)
    dv = DataValidation(type="list", formula1='"Da decidere,Applicata"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("L8:L200")
    ws.merge_cells("A14:L14")
    cella(ws, 14, 1, "Stato: lo scrive lo script (Da decidere) e lo aggiorna il comando valida (Applicata). E' a elenco per filtrare, "
                     "non va cambiato a mano. Decisione e' a elenco: Accettato, Rettificato, Respinto.",
          font=Font(size=9, italic=True, color=TESTO_TENUE), alignment=Alignment(wrap_text=True))
    ws.row_dimensions[14].height = 24


# ============================================================== Piano delle azioni (nota G della v2)
def foglio_piano_azioni(ws):
    ws.sheet_properties.tabColor = TAB["lavoro"]
    for j, w in enumerate([9, 9, 17, 54, 26, 13, 9, 17, 16, 12, 30], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    fascia(ws, 11, "Piano delle azioni")
    guida(ws, 4, 11,
          "COSA E': tutte le azioni di mitigazione, proposte e decise, con il loro stato. E' il registro delle azioni da validare "
          "e confermare.\nCOME FUNZIONA: A1 propone da 1 a 5 azioni per rischio a ogni trimestre (Stato: Proposta — da valutare). "
          "Il validatore le decide nel foglio Validazione (owner e orizzonte: Accettato, Rettificato o Respinto): e' la delibera, "
          "e il comando valida porta lo Stato ad Approvata o Respinta. Poi la funzione di rischio aggiorna a mano lo Stato "
          "(In corso, Completata) quando l'owner lo comunica. Scadenza = data della decisione + orizzonte. Le azioni aperte "
          "passano ad A1 nel trimestre dopo.", "validatore")
    gruppi = [((1, 7), "La proposta (A1)", "trimestrale"), ((8, 10), "La decisione (validazione)", "validatore"), ((11, 11), "Avanzamento (a mano)", "validatore")]
    for (da, a), t, p in gruppi:
        if da != a:
            ws.merge_cells(start_row=6, start_column=da, end_row=6, end_column=a)
        for x in range(da, a + 1):
            cella(ws, 6, x, t if x == da else None, fill=fill(PROCESSO[p][0]), font=Font(size=9, bold=True, color=PROCESSO[p][1]))
    intest = ["Ciclo", "Rischio", "ID azione", "Azione", "Owner", "Riduce", "Orizzonte", "Stato", "Decisa il", "Scadenza", "Nota"]
    for j, t in enumerate(intest, start=1):
        cella(ws, 7, j, t, font=Font(size=9, bold=True, color="FFFFFF"), fill=fill(BLU_NOTTE))
    colori_stato = {"Proposta — da valutare": (AMBRA_TENUE, "7F6000"), "Approvata": (BLU_CHIARO, BLU_NOTTE), "Respinta": ("F2F2F2", TESTO_TENUE),
                    "In corso": ("E2EFDA", "375623"), "Completata": ("70AD47", "FFFFFF")}
    righe = []
    for rid, descr, owner, dim, oriz, _, stato in g.AZIONI_ESEMPIO:
        mesi = int(oriz.split()[0])
        scad = "—" if stato == "Respinta" else f"15/{7 + mesi - 12 if 7 + mesi > 12 else 7 + mesi:02d}/{2027 if 7 + mesi > 12 else 2026}"
        stato_vista = {"NFR-06": "In corso", "NFR-08": "Completata"}.get(rid, stato)
        nota = {"NFR-06": "DLP esteso a 3 servizi su 5 (settembre)", "NFR-08": "Revisione automatica attiva da agosto",
                "NFR-05": "Respinta: duplicata del presidio PMO"}.get(rid)
        righe.append(("2026Q2", rid, f"2026Q2-{rid}-A1", descr, owner, dim, oriz, stato_vista, "15/07/2026", scad, nota))
    righe += [("2026Q3", "NFR-01", "2026Q3-NFR-01-A2", "Limite giornaliero sui bonifici istantanei verso beneficiari nuovi per i clienti retail.",
               "Funzione Antifrode", "probabilita", "3 mesi", "Proposta — da valutare", None, None, "Da decidere nella validazione del 2026Q3"),
              ("2026Q3", "NFR-11", "2026Q3-NFR-11-A1", "Avvio del programma pluriennale di sostituzione delle componenti fuori supporto.",
               "Funzione IT Operations", "gap_controlli", "12 mesi", "Proposta — da valutare", None, None, "Da decidere nella validazione del 2026Q3")]
    for i, rr in enumerate(sorted(righe, key=lambda x: (x[7] != "Proposta — da valutare", x[0], x[1])), start=8):
        for j, v in enumerate(rr, start=1):
            cella(ws, i, j, v, font=Font(size=9), alignment=Alignment(wrap_text=True, vertical="top"), border=Border(bottom=bordo_sottile))
        f_, t_ = colori_stato[rr[7]]
        c = ws.cell(row=i, column=8)
        c.fill = fill(f_)
        c.font = Font(size=9, bold=True, color=t_)
        ws.row_dimensions[i].height = 26
    dv = DataValidation(type="list", formula1='"Proposta — da valutare,Approvata,Respinta,In corso,Completata,Annullata"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"H8:H{7 + len(righe)}")
    legenda(ws, 9 + len(righe), 11, "Riduce = dimensione del rischio che l'azione riduce. Orizzonte = tempo per completarla. "
                                    "Stato: giallo = da decidere nella validazione; azzurro = approvata; verde = in corso o completata.")
    ws.freeze_panes = "A8"
    v1.stampa(ws)


# ============================================================== schede: andamento negli ultimi cicli (nota B della v2)
_scheda_v2 = v2.scheda


def scheda(ws, r, kri):
    _scheda_v2(ws, r, kri)
    riga = ws.max_row + 2
    v1.sezione(ws, f"A{riga}", f"H{riga}", "Andamento negli ultimi trimestri")
    riga += 1
    intest = ["Ciclo", "Sev", "Prob", "Gap", "Vel", "Pross", "Conf", "Score"]
    for j, t in enumerate(intest, start=1):
        cella(ws, riga, j, t, font=Font(size=9, bold=True, color=TESTO_TENUE), alignment=Alignment(horizontal="center"))
    inizio = riga + 1
    for ciclo, sev, prob, gap, vel, pro, conf, sc in r["storia"]:
        riga += 1
        cella(ws, riga, 1, ciclo, font=Font(size=10, bold=True, color=BLU_NOTTE), alignment=Alignment(horizontal="center"))
        for j, v in enumerate((sev, prob, gap, vel, pro), start=2):
            cella(ws, riga, j, v, fill=fill(SCALA[v]), font=Font(size=10, bold=True, color="FFFFFF" if v == 5 else TESTO),
                  alignment=Alignment(horizontal="center"))
        cella(ws, riga, 7, conf, fill=fill(SCALA_CONF[conf]), alignment=Alignment(horizontal="center"))
        cella(ws, riga, 8, sc, number_format="0.00", font=Font(size=10, bold=True), alignment=Alignment(horizontal="center"))
    graf = LineChart()
    graf.add_data(Reference(ws, min_col=8, min_row=inizio - 1, max_row=riga), titles_from_data=True)
    graf.set_categories(Reference(ws, min_col=1, min_row=inizio, max_row=riga))
    s = graf.series[0]
    s.smooth = False
    s.marker.symbol = "circle"
    s.marker.size = 7
    s.graphicalProperties.line.solidFill = BLU_NOTTE
    s.marker.graphicalProperties = GraphicalProperties(solidFill=BLU_NOTTE)
    v1.etichette_valore(s, posizione="t")
    graf.y_axis.scaling.min, graf.y_axis.scaling.max, graf.y_axis.majorUnit = 1, 5, 1
    graf.legend = None
    graf.title = v1.titolo_grafico("Score composito per trimestre (scala fissa 1-5)", sz=1000)
    v1.pulisci(graf)
    graf.height, graf.width = 5.5, 11
    ws.add_chart(graf, f"A{riga + 2}")
    if len(r["storia"]) < 3:
        cella(ws, riga + 1, 1, f"Rischio entrato nel registro nel {r['storia'][0][0]}.", font=Font(size=8, italic=True, color=TESTO_TENUE))
    legenda(ws, riga + 14, 8, "Sev = Severity, Prob = Probabilita, Gap = Gap_Controlli, Vel = Velocity, Pross = Prossimita, "
                               f"Conf = Confidence (fuori dallo score). {LEG_MOSCA}. KRI = indicatore di rischio.")


v2.scheda = scheda
v2.USCITA_SCHEDE = USCITA_SCHEDE


# ============================================================== guida alla valutazione e struttura
def foglio_valuta(ws):
    ws.sheet_properties.tabColor = TAB["guida"]
    fascia(ws, 7, "anteprima v3: come valutare")
    for l, w in zip("ABCDEFG", [7, 16, 42, 50, 40, 15, 30]):
        ws.column_dimensions[l].width = w
    guida(ws, 4, 7, "Terza anteprima, dati di esempio, non collegata al registro. Recepisce le tue note sulla v2 (conservata in "
                    "09_Anteprima_v2_valutata.xlsx) e le tue decisioni: struttura approvata; il commento di A1 resta breve "
                    "(messaggio e decisioni), il perche' di ogni rischio lo calcola il foglio. Se va bene, si passa al generatore vero.",
          "validatore")
    note = [
        ("3-4, F", "Sintesi", "Titolo \"Quadro d'insieme\"; radar non visibile; un solo foglio (anche 2 pagine); perche' NFR-01 e 11 rosa?; "
                              "commento sotto i grafici, rischi segnalati ancora sotto",
         "Titolo cambiato. Il radar ora ha i dati nel foglio stesso (prima li prendeva dal foglio Radar) ed e' piu' grande. Tutto in "
         "un foglio, due pagine in stampa: grafici; poi commento di A1 (messaggio e decisioni), A5, avviso; poi i rischi segnalati. "
         "Il rosa segnava AGIRE ORA senza dirlo: ora c'e' una colonna Urgenza con l'etichetta."),
        ("9", "Radar", "Legenda del cerchio rosso sotto il campione giallo",
         "In basso a sinistra: campione giallo (attenzione del supervisore) e, sotto, la linea rossa \"dentro il cerchio rosso: AGIRE "
         "ORA (X + Y > Z)\". Sul cerchio resta solo \"Mosca\". Anche nel radar piccolo della Sintesi."),
        ("12", "Delta", "Spiegare la scala, quando si ha +1 e -1; a fianco i motivi della variazione o non variazione",
         "Nota sotto il grafico: quanto vale un punto di ogni dimensione, +1 solo se peggiorano tutte e cinque. A fianco, nello stesso "
         "ordine delle barre, il motivo con l'evidenza e la data, o \"invariato: nessuna evidenza nuova\"."),
        ("B", "Registro", "Togliere \"un rischio per riga\"; andamento nei trimestri nella scheda; \"Settim.\" non chiaro; variazione a 6 mesi",
         "Titolo \"Registro: il ciclo corrente\". Gruppo \"Notizie · 4 settimane\", colonna \"Segnali netti\" (4 settimane, non una), "
         "spiegata in guida e legenda. Nuova colonna Δ 6 mesi. Ogni scheda ha \"Andamento negli ultimi trimestri\" (tabella e grafico)."),
        ("D", "Dizionario, report", "Verificare che ci siano tutti i dati; legenda delle abbreviazioni in fondo ai report",
         "Dizionario completo: ogni colonna di ogni tabella (dalla mappa), i valori calcolati nelle viste, parametri e anagrafica. "
         "Legenda per esteso in fondo a Quadro d'insieme, Classifica, Radar, Matrice, Delta, Registro, Piano azioni e schede."),
        ("E", "Validazione", "Il campo Stato non si compila in automatico: almeno una lista di valori",
         "Stato a elenco (Da decidere, Applicata) con una nota: lo scrive lo script e lo aggiorna il comando valida."),
        ("G", "Piano azioni", "Esiste un registro delle azioni proposte da validare/confermare?",
         "Si': la tabella Azioni (una riga per azione) piu' le righe di Validazione su owner e orizzonte. Nuovo foglio di lavoro "
         "\"Piano azioni\": proposte da decidere in cima, poi approvate, in corso, completate, con scadenza."),
    ]
    for j, t in enumerate(["Nota", "Foglio", "Cosa hai scritto", "Cosa ho cambiato", "", "Va bene ora?", "Nota"], start=1):
        cella(ws, 6, j, t or None, font=Font(bold=True, color="FFFFFF"), fill=fill(BLU_NOTTE))
    for i, (n_, foglio, scritto, cambiato) in enumerate(note, start=7):
        ws.merge_cells(start_row=i, start_column=4, end_row=i, end_column=5)
        # colonne: A nota, B foglio, C cosa hai scritto, D-E cosa ho cambiato
        cella(ws, i, 1, n_, font=Font(bold=True, color=BLU_NOTTE), alignment=Alignment(vertical="top"))
        cella(ws, i, 2, foglio, alignment=Alignment(wrap_text=True, vertical="top"))
        cella(ws, i, 3, scritto, alignment=Alignment(wrap_text=True, vertical="top"))
        cella(ws, i, 4, cambiato, alignment=Alignment(wrap_text=True, vertical="top"))
        for x in range(1, 8):
            ws.cell(row=i, column=x).border = Border(bottom=bordo_sottile)
        for x in (6, 7):
            ws.cell(row=i, column=x).fill = fill(AMBRA_TENUE)
        ws.row_dimensions[i].height = 14 * (max(len(cambiato) // 100, len(scritto) // 46) + 1) + 6
    dv = DataValidation(type="list", formula1='"Sì,No,Sì con modifiche"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F7:F{6 + len(note)}")
    v1.stampa(ws)


def foglio_struttura(ws):
    ws.sheet_properties.tabColor = TAB["guida"]
    fascia(ws, 5, "struttura del file unico (approvata)")
    for l, w in zip("ABCDE", [22, 26, 22, 58, 12]):
        ws.column_dimensions[l].width = w
    guida(ws, 4, 5, "Tre gruppi di fogli con colori di linguetta diversi — REPORT (blu notte, si stampano per il comitato), LAVORO "
                    "(blu, per chi analizza e valida), DATI (grigio, li scrive solo lo script) — piu' la Guida. Le schede dei rischi "
                    "escono in un file a parte. Nessun dato si perde: cambiano solo le viste.", "calcolo")
    for j, t_ in enumerate(["Oggi", "Proposta", "Gruppo", "Perche'", "In anteprima"], start=1):
        cella(ws, 6, j, t_, font=Font(bold=True, color="FFFFFF"), fill=fill(BLU_NOTTE))
    righe = [
        ("Leggimi", "Guida", "guida", "Come si usa, chi aggiorna cosa, quando si lancia A5, dove si scrive a mano.", "—"),
        ("Sintesi", "Quadro d'insieme", "report", "Un foglio, due pagine: grafici; poi commento di A1, A5 e avviso; poi i rischi "
                                                   "segnalati e perche'.", "Sì"),
        ("Classifica", "Classifica", "report", "Invariata nella sostanza; stile nuovo.", "Sì"),
        ("Radar_Mosca", "Radar", "report", "Grafico piu' leggibile, legenda nel grafico; la tabella di calcolo va nel Registro.", "Sì"),
        ("—", "Matrice", "report", "Nuova: gravita' P x I, con l'elenco dei rischi.", "Sì"),
        ("Delta", "Delta", "report", "Scala spiegata; a fianco il motivo di ogni variazione o non variazione.", "Sì"),
        ("Metodologia", "Metodologia", "report", "Resta per il comitato; rimanda al Dizionario per i dettagli.", "—"),
        ("Schede, Approfondimento", "file Schede_<ciclo>.xlsx", "file a parte", "Testi lunghi per rischio e andamento negli ultimi "
                                                                               "trimestri: a tendere li scrive un agente dedicato.", "file a parte"),
        ("Dashboard", "Registro", "lavoro", "Una riga per rischio con tutto il ciclo, Δ sul trimestre prima e su 6 mesi.", "Sì"),
        ("Trend", "Storico", "lavoro", "Score degli ultimi otto cicli, una riga per rischio.", "—"),
        ("KRI_Correnti + Andamento_KRI", "KRI", "lavoro", "I KRI del ciclo e, scelto un KRI, la sua storia: un foglio invece di due.", "—"),
        ("Analisi_News", "Notizie", "lavoro", "Pressione delle notizie e aree non presidiate.", "—"),
        ("—", "Piano azioni", "lavoro", "Nuovo: tutte le azioni, da decidere, approvate, in corso, completate, con scadenza.", "Sì"),
        ("Validazione", "Validazione", "lavoro", "Istruzioni e criteri in testa, colonne da compilare in giallo, Stato a elenco.", "Sì"),
        ("Parametri", "Parametri", "lavoro", "Invariato; stile nuovo.", "—"),
        ("—", "Dizionario", "lavoro", "Nuovo: ogni dato con chi lo scrive, quando, significato, calcolo, valori, esempio.", "Sì"),
        ("RiskRegister, Azioni, KRI, NewsLog, Sintesi_Cicli", "stessi nomi, in fondo", "dati", "Tabelle ridisegnate, con gruppi di "
         "colonne e la riga \"chi · cosa\". Lo script non cambia.", "RiskRegister"),
    ]
    colori = {"report": BLU_NOTTE, "lavoro": "2E75B6", "dati": "A6A6A6", "guida": "BF8F00", "file a parte": "7030A0"}
    for i, riga in enumerate(righe, start=7):
        for j, v in enumerate(riga, start=1):
            c = cella(ws, i, j, v, alignment=Alignment(wrap_text=True, vertical="top"), border=Border(bottom=bordo_sottile))
            if j == 3:
                c.font = Font(bold=True, color=colori[v])
            if j == 2:
                c.font = Font(bold=True)
        ws.row_dimensions[i].height = 30
    fine = 7 + len(righe) + 1
    ws.merge_cells(start_row=fine, start_column=1, end_row=fine, end_column=5)
    ws.row_dimensions[fine].height = 45
    cella(ws, fine, 1, "20 fogli come oggi, ma senza doppioni: Dashboard, Approfondimento, KRI_Correnti, Andamento_KRI e Schede "
                       "spariscono; entrano Guida, Matrice, Dizionario, Registro e Piano azioni; le schede vanno in un file a parte. Lo "
                       "script di campo scrive solo nelle tabelle dati, che non cambiano: il registro in uso passa alla struttura nuova "
                       "con aggiorna-registro.", font=Font(bold=True, color=BLU_NOTTE), alignment=Alignment(wrap_text=True, vertical="top"))
    v1.stampa(ws)


# ============================================================== main
def main():
    rischi, kri = dati()
    wb = Workbook()
    nomi = ["Come valutare", "Struttura proposta", "Quadro d’insieme", "Classifica", "Radar", "Matrice", "Delta",
            "Registro", "Piano azioni", "Validazione", "Dizionario", "Dati_RiskRegister"]
    fogli = {nomi[0]: wb.active}
    wb.active.title = nomi[0]
    for n in nomi[1:]:
        fogli[n] = wb.create_sheet(n)
    foglio_valuta(fogli["Come valutare"])
    foglio_struttura(fogli["Struttura proposta"])
    foglio_sintesi(fogli["Quadro d’insieme"], rischi, kri)
    v2.foglio_radar(fogli["Radar"], rischi)
    legenda(fogli["Radar"], 38, 16, f"{LEG_PXI}. {LEG_MOSCA}. L1-01..L1-07 = categorie di perdita di Basilea (event type). {LEG_RISCHIO}.")
    fogli["Radar"].print_area = "A1:P38"
    v2.foglio_matrice(fogli["Matrice"], rischi)
    legenda(fogli["Matrice"], 31, 7, f"{LEG_PXI}. {LEG_RISCHIO}.")
    foglio_delta(fogli["Delta"], rischi)
    foglio_registro(fogli["Registro"], rischi)
    foglio_piano_azioni(fogli["Piano azioni"])
    foglio_validazione(fogli["Validazione"])
    foglio_dizionario(fogli["Dizionario"])
    v2.foglio_dati(fogli["Dati_RiskRegister"], rischi)
    wb.remove(fogli["Classifica"])
    v1.foglio_classifica(wb, rischi)
    wb["Classifica"].sheet_properties.tabColor = TAB["report"]
    legenda(wb["Classifica"], 31, 20, f"{LEG_MOSCA}. Frecce: ▲ score in aumento, ▼ in diminuzione, ● invariato sul ciclo prima. {LEG_RISCHIO}.")
    wb["Classifica"].print_area = "A1:T31"
    wb._sheets.sort(key=lambda s: nomi.index(s.title))
    wb.active = 2
    wb.save(USCITA_REGISTRO)
    v1.trasparenza(USCITA_REGISTRO, v1.RIEMPIMENTO_CERCHIO)
    v2.file_schede(rischi, kri)
    print(f"Anteprima v3 scritta:\n  {USCITA_REGISTRO}\n  {USCITA_SCHEDE}")


if __name__ == "__main__":
    main()
