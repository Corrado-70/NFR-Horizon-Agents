"""Presentazione del lavoro: agenti e passi Python, schede dei rischi, report trimestrale foglio per foglio, parametri.

Uso:  python genera_presentazione.py            (immagini rifatte dal registro demo + presentazione)
      python genera_presentazione.py --solo-pptx (riusa le immagini gia' preparate)

Le immagini sono pagine vere dell'Excel: si copia Demo_W1_W2/NFR_Registro.xlsx in una cartella temporanea, si lanciano
i comandi `report --con-lavoro` e `schede` dello script di campo su quella copia (il repository non si tocca), si
esportano in PDF gli altri fogli con Excel desktop e si rendono in PNG con PyMuPDF.
Richiede: Excel desktop, pywin32, python-pptx, pymupdf, pillow.
"""
import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

QUI = Path(__file__).resolve().parent
RADICE = QUI.parent
SCRIPT = RADICE / "NFR_Emerging_Risk_Agents" / "07_Senza_Studio"
KERNEL = RADICE / "NFR_Emerging_Risk_Agents" / "01_Kernel"
DEMO = RADICE / "Demo_W1_W2"
LAVORO = Path(tempfile.gettempdir()) / "nfr_presentazione"
IMG = LAVORO / "img"
USCITA = QUI / "Horizon_Scan_NFR_presentazione.pptx"


# ====================================================================== immagini dal registro demo
def prepara_immagini():
    import pymupdf
    import win32com.client as w
    if LAVORO.exists():
        shutil.rmtree(LAVORO)
    (LAVORO / "arch").mkdir(parents=True)
    pdf = LAVORO / "pdf"
    pdf.mkdir()
    IMG.mkdir()
    registro = LAVORO / "Registro.xlsx"
    shutil.copy(DEMO / "NFR_Registro.xlsx", registro)
    base = [sys.executable, "esegui_ciclo_locale.py", "--workbook", str(registro), "--kernel", str(KERNEL),
            "--cartella-uscita", str(LAVORO / "arch")]
    for comando in (["report", "--con-lavoro"], ["schede"]):
        subprocess.run(base + comando, cwd=SCRIPT, check=True)
    ciclo = next((LAVORO / "arch" / "Cicli").iterdir())
    shutil.copy(next(ciclo.glob("Report_comitato_*.pdf")), pdf / "Report.pdf")
    shutil.copy(DEMO / "Grafici" / "Validazione_2026Q3.pdf", pdf / "Validazione_Q3.pdf")
    schede = next(ciclo.glob("Schede_*.xlsx"))
    # Fogli fuori dal report: area di stampa ridotta, orizzontale, una pagina di larghezza.
    aree = {"Storico": "A1:J20", "Indicatori": "A1:AA40", "Notizie": "A1:S32", "Piano_azioni": "A1:J30",
            "Osservazione": "A1:Z12", "Guida": "A1:E61", "Dizionario": "A1:G40", "RiskRegister": "A1:AF22",
            "Sintesi_Cicli": "A1:J11", "Sintesi_Settimane": "A1:F10", "Storico_Validazioni": "A1:L30", "Azioni": "A1:L25",
            "KRI": "A1:T25", "NewsLog": "A1:S21"}
    xl = w.DispatchEx("Excel.Application")
    xl.Visible = xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(str(registro), ReadOnly=True)
        xl.CalculateFull()
        xl.PrintCommunication = False
        for f, a in aree.items():
            ps = wb.Worksheets(f).PageSetup
            ps.PrintArea, ps.Orientation, ps.Zoom, ps.FitToPagesWide, ps.FitToPagesTall = a, 2, False, 1, False
            for parte in ("LeftHeader", "CenterHeader", "RightHeader", "LeftFooter", "CenterFooter", "RightFooter"):
                setattr(ps, parte, "")
        xl.PrintCommunication = True
        for f in aree:
            wb.Worksheets(f).ExportAsFixedFormat(0, str(pdf / f"f_{f}.pdf"))
        wb.Close(False)
        wb = xl.Workbooks.Open(str(schede), ReadOnly=True)
        for f in ("Indice", "NFR-01", "NFR-11"):
            wb.Worksheets(f).ExportAsFixedFormat(0, str(pdf / f"Schede_{f}.pdf"))
        wb.Close(False)
    finally:
        xl.Quit()

    def rendi(file, pagina, nome, dpi=170):
        """La pagina senza intestazione e pie' di pagina, ritagliata sul contenuto."""
        pg = pymupdf.open(file)[pagina]
        alto = pg.rect.height
        r = pymupdf.Rect()
        for b in pg.get_text("blocks"):
            if b[1] > 42 and b[3] < alto - 40:
                r |= pymupdf.Rect(b[:4])
        for d in pg.get_drawings():
            if d["rect"].y0 > 42 and d["rect"].y1 < alto - 40:
                r |= d["rect"]
        for im in pg.get_image_info():
            r |= pymupdf.Rect(im["bbox"])
        pg.get_pixmap(dpi=dpi, clip=(r + (-4, -4, 4, 4)) & pg.rect).save(IMG / f"{nome}.png")

    for i, n in enumerate(["Quadro_1", "Quadro_2", "Classifica", "Radar_Mosca", "Matrice", "Delta_1", "Delta_2", "Fenomeni",
                           "Metodologia_1", "Metodologia_2", "Metodologia_3", "Metodologia_4", "Metodologia_5",
                           "Registro", "Parametri_1", "Parametri_2"]):
        rendi(pdf / "Report.pdf", i, n)
    for f in aree:
        rendi(pdf / f"f_{f}.pdf", 0, f"f_{f}")
    rendi(pdf / "Schede_Indice.pdf", 0, "Schede_Indice")
    for p in (0, 1):
        rendi(pdf / "Schede_NFR-01.pdf", p, f"Scheda_NFR01_{p + 1}")
        rendi(pdf / "Schede_NFR-11.pdf", p, f"Scheda_NFR11_{p + 1}")
    for p in (0, 1, 2):
        rendi(pdf / "Validazione_Q3.pdf", p, f"Validazione_Q3_{p + 1}")


# ====================================================================== stile e mattoni delle slide
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLU = RGBColor(0x2E, 0x75, 0xB6)
AZZURRO = RGBColor(0xDD, 0xEB, 0xF7)
GRIGIO = RGBColor(0x59, 0x59, 0x59)
GRIGIO_CH = RGBColor(0xF2, 0xF2, 0xF2)
LINEA = RGBColor(0xBF, 0xBF, 0xBF)
TESTO = RGBColor(0x26, 0x26, 0x26)
BIANCO = RGBColor(0xFF, 0xFF, 0xFF)
ROSSO = RGBColor(0xC0, 0x00, 0x00)
AMBRA = RGBColor(0xFF, 0xC0, 0x00)
VERDE = RGBColor(0x54, 0x82, 0x35)
VERDE_CH = RGBColor(0xE2, 0xEF, 0xDA)
AMBRA_CH = RGBColor(0xFF, 0xF2, 0xCC)
ROSSO_CH = RGBColor(0xF8, 0xD7, 0xDA)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
VUOTO = prs.slide_layouts[6]
stato = {"sezione": "", "n": 0}


def _run(p, testo, size, bold=False, colore=TESTO, italic=False):
    r = p.add_run()
    r.text = testo
    f = r.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    f.color.rgb = colore
    return r


def _ricco(p, testo, size, colore=TESTO, bold=False, italic=False):
    """**grassetto** dentro il testo."""
    for i, pezzo in enumerate(testo.split("**")):
        if pezzo:
            _run(p, pezzo, size, bold or i % 2 == 1, colore, italic)


def _punto(p, livello, car="•"):
    pPr = p._p.get_or_add_pPr()
    mar = 0.18 + 0.22 * livello
    pPr.set("marL", str(int(Inches(mar))))
    pPr.set("indent", str(int(-Inches(0.16))))
    for t in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(t)):
            pPr.remove(e)
    bu = pPr.makeelement(qn("a:buChar"), {"char": car})
    pPr.append(bu)


def testo(slide, x, y, w, h, righe, size=14, colore=TESTO, ancora=MSO_ANCHOR.TOP, spazio=4):
    """righe: '## titoletto', '- punto', '  - sottopunto', '> nota in corsivo', altro = paragrafo."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    prima = True
    for riga in righe:
        p = tf.paragraphs[0] if prima else tf.add_paragraph()
        prima = False
        p.space_after = Pt(spazio)
        if riga.startswith("## "):
            p.space_before = Pt(4)
            _run(p, riga[3:], size + 1, True, NAVY)
        elif riga.startswith("  - "):
            _punto(p, 1, "–")
            _ricco(p, riga[4:], size - 1, colore)
        elif riga.startswith("- "):
            _punto(p, 0)
            _ricco(p, riga[2:], size, colore)
        elif riga.startswith("> "):
            _ricco(p, riga[2:], size - 1.5, GRIGIO, italic=True)
        else:
            _ricco(p, riga, size, colore)
    return tb


def rettangolo(slide, x, y, w, h, fondo, bordo=None, forma=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fondo
    if bordo is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = bordo
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    return s


def nuova(titolo, sezione=None):
    s = prs.slides.add_slide(VUOTO)
    stato["n"] += 1
    if sezione:
        stato["sezione"] = sezione
    testo(s, 0.5, 0.22, 9, 0.3, [stato["sezione"].upper()], 10.5, BLU)
    testo(s, 0.5, 0.48, 12.3, 0.62, [titolo], 25, NAVY).text_frame.paragraphs[0].runs[0].font.bold = True
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(1.12), Inches(12.83), Inches(1.12))
    ln.line.color.rgb = NAVY
    ln.line.width = Pt(1.5)
    testo(s, 0.5, 7.08, 8, 0.3, ["Horizon Scan NFR · agenti, report e parametri · Uso interno"], 9, GRIGIO)
    testo(s, 11.83, 7.08, 1, 0.3, [str(stato["n"])], 9, GRIGIO).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    return s


def divisore(numero, titolo, sottotitolo):
    s = prs.slides.add_slide(VUOTO)
    stato["n"] += 1
    stato["sezione"] = titolo
    rettangolo(s, 0, 0, 13.333, 7.5, NAVY)
    rettangolo(s, 0.8, 2.55, 0.12, 2.1, AMBRA)
    testo(s, 1.15, 2.35, 10, 0.8, [numero], 22, AMBRA)
    testo(s, 1.15, 2.95, 11, 1.0, [titolo], 40, BIANCO).text_frame.paragraphs[0].runs[0].font.bold = True
    testo(s, 1.15, 3.95, 11, 1.2, sottotitolo, 17, RGBColor(0xDD, 0xE6, 0xF2))
    return s


def scheda(slide, x, y, w, h, titolo, righe, colore=NAVY, size=12.5, fondo=None):
    """Riquadro con una fascia colorata del titolo."""
    rettangolo(slide, x, y, w, h, fondo or BIANCO, LINEA)
    rettangolo(slide, x, y, w, 0.36, colore)
    testo(slide, x + 0.08, y + 0.02, w - 0.16, 0.34, [titolo], 12.5, BIANCO, MSO_ANCHOR.MIDDLE).text_frame.paragraphs[0].runs[0].font.bold = True
    testo(slide, x + 0.1, y + 0.44, w - 0.2, h - 0.5, righe, size)


def tabella(slide, x, y, w, intest, righe, larg, size=11, alt_riga=None, h=None, colori_prima=None):
    """Tabella con intestazione blu; larg = frazioni della larghezza. Il testo puo' usare **grassetto**."""
    nr, nc = len(righe) + 1, len(intest)
    alt = alt_riga or 0.3
    tot = sum(larg)
    cw = [w * l / tot for l in larg]

    def stima(celle):
        n = 1
        for j, v in enumerate(celle):
            per_riga = int(max(1, (cw[j] - 0.14) * 72 / (size * 0.5)))
            n = max(n, sum(max(1, -(-len(q.replace("**", "")) // per_riga)) for q in str(v).split("\n")))
        return n * size * 1.22 / 72 + 0.08

    altezze = [max(0.36, stima(intest))] + [max(alt, stima(r)) for r in righe]
    t = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(sum(altezze))).table
    for j, l in enumerate(larg):
        t.columns[j].width = Emu(int(Inches(w) * l / tot))
    for i in range(nr):
        t.rows[i].height = Inches(altezze[i])
        for j in range(nc):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.025)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            v = intest[j] if i == 0 else righe[i - 1][j]
            if i == 0:
                _run(p, v, size, True, BIANCO)
                c.fill.solid()
                c.fill.fore_color.rgb = NAVY
            else:
                for k, parte in enumerate(str(v).split("\n")):
                    q = p if k == 0 else tf.add_paragraph()
                    _ricco(q, parte, size, TESTO, bold=(j == 0))
                c.fill.solid()
                c.fill.fore_color.rgb = BIANCO if i % 2 else GRIGIO_CH
                if j == 0 and colori_prima and colori_prima.get(i - 1):
                    c.fill.fore_color.rgb = colori_prima[i - 1]
    return y + sum(altezze)


def immagine(slide, nome, x, y, w, h, taglio=None, bordo=True, allinea="centro"):
    """L'immagine dentro il riquadro, proporzioni intatte; taglio = frazioni (sx, alto, dx, basso)."""
    p = IMG / f"{nome}.png"
    if taglio:
        im = Image.open(p)
        W, H = im.size
        im = im.crop((int(W * taglio[0]), int(H * taglio[1]), int(W * taglio[2]), int(H * taglio[3])))
        p = LAVORO / f"_{nome}_{'_'.join(str(t) for t in taglio)}.png"
        im.save(p)
    W, H = Image.open(p).size
    k = min(w / W, h / H)
    ww, hh = W * k, H * k
    xx = x + (w - ww) / 2 if allinea == "centro" else x
    pic = slide.shapes.add_picture(str(p), Inches(xx), Inches(y), Inches(ww), Inches(hh))
    if bordo:
        pic.line.color.rgb = LINEA
        pic.line.width = Pt(0.75)
    return xx, ww, hh


def freccia(slide, x1, y1, x2, y2, colore=GRIGIO):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = colore
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))


def blocco(slide, x, y, w, h, righe, tipo, size=11):
    """Passo di un flusso: agente (blu), Python (verde), persona (ambra)."""
    fondo, col = {"agente": (BLU, BIANCO), "python": (VERDE, BIANCO), "persona": (AMBRA, TESTO),
                  "file": (AZZURRO, NAVY)}[tipo]
    rettangolo(slide, x, y, w, h, fondo, forma=MSO_SHAPE.ROUNDED_RECTANGLE)
    tb = testo(slide, x + 0.04, y, w - 0.08, h, righe, size, col, MSO_ANCHOR.MIDDLE, spazio=0)
    for p in tb.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER


def nota_fondo(slide, righe, y=6.55, size=11):
    testo(slide, 0.5, y, 12.3, 0.5, righe, size, GRIGIO)


def foglio(titolo, img, spiega, taglio=None, larg_img=7.2, size=12, sezione=None, h_img=5.55):
    """Slide tipo per un foglio: immagine a sinistra, spiegazione a destra."""
    s = nuova(titolo, sezione)
    W, H = Image.open(IMG / f"{img}.png").size
    if W / H >= 2.2:
        _, _, hh = immagine(s, img, 0.5, 1.35, 12.33, 3.3, taglio)
        y = 1.35 + hh + 0.2
        meta = (len(spiega) + 1) // 2
        while meta < len(spiega) and spiega[meta - 1].startswith("## "):
            meta += 1
        testo(s, 0.5, y, 6.0, 6.95 - y, spiega[:meta], size)
        testo(s, 6.83, y, 6.0, 6.95 - y, spiega[meta:], size)
        return s
    immagine(s, img, 0.5, 1.35, larg_img, h_img, taglio, allinea="sx")
    testo(s, 0.5 + larg_img + 0.25, 1.3, 12.83 - (0.5 + larg_img + 0.25), 5.7, spiega, size)
    return s


# ====================================================================== contenuto
def copertina():
    s = prs.slides.add_slide(VUOTO)
    stato["n"] += 1
    rettangolo(s, 0, 0, 13.333, 7.5, NAVY)
    rettangolo(s, 0.8, 2.0, 0.12, 2.9, AMBRA)
    testo(s, 1.15, 1.85, 11, 1, ["Horizon Scan NFR"], 46, BIANCO).text_frame.paragraphs[0].runs[0].font.bold = True
    testo(s, 1.15, 2.85, 11, 1.2, ["Il sistema di agenti per i rischi non finanziari emergenti:",
                                   "che cosa fanno gli agenti, che cosa dicono i report, con quali parametri"], 21,
          RGBColor(0xDD, 0xE6, 0xF2))
    testo(s, 1.15, 4.2, 11, 0.8, ["Esempi tratti dal registro della demo (ciclo straordinario 2026Q4-S1, dati simulati) · "
                                  "stato del pacchetto al 16/10/2026"], 14, AMBRA)
    testo(s, 1.15, 6.6, 11, 0.4, ["Uso interno"], 12, RGBColor(0xBF, 0xCF, 0xE6))


def indice():
    s = nuova("Di che cosa parliamo", "Introduzione")
    voci = [("1", "Il sistema in breve", "componenti, cicli, principi, chi fa che cosa"),
            ("2", "Gli agenti e i passi Python", "un agente per slide: che cosa fa, quando si usa, quale comando Python lo precede e lo segue"),
            ("3", "Le schede dei rischi", "a che cosa servono, come sono impostate, il significato di ogni informazione"),
            ("4", "Il report trimestrale", "dal Quadro d'insieme a ogni foglio del registro Excel: uso e significato di ogni dato"),
            ("5", "I parametri", "gerarchia di rischio, scale, pesi, soglie, urgenza, livello di allerta, valori ammessi")]
    for i, (n, t, d) in enumerate(voci):
        y = 1.55 + i * 1.02
        rettangolo(s, 0.7, y, 0.75, 0.75, NAVY, forma=MSO_SHAPE.ROUNDED_RECTANGLE)
        tb = testo(s, 0.7, y, 0.75, 0.75, [n], 26, BIANCO, MSO_ANCHOR.MIDDLE)
        tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        testo(s, 1.7, y - 0.02, 10.5, 0.4, [f"**{t}**"], 19, NAVY)
        testo(s, 1.7, y + 0.37, 10.8, 0.4, [d], 14, GRIGIO)


def sistema():
    s = nuova("Il sistema in una pagina: tre ritmi di lavoro", "1 · Il sistema in breve")
    testo(s, 0.5, 1.3, 2.2, 0.4, ["## Ogni lunedì"], 13)
    passi = [(["raccogli_notizie.py", "notizie dai feed"], "python"), (["A3 Interim Update", "valuta le notizie"], "agente"),
             (["settimanale", "scrive il NewsLog"], "python"), (["valida (facoltativa)", "il validatore"], "persona")]
    for i, (r, t) in enumerate(passi):
        x = 0.5 + i * 2.55
        blocco(s, x, 1.75, 2.2, 0.72, r, t)
        if i:
            freccia(s, x - 0.35, 2.11, x - 0.02, 2.11)
    testo(s, 0.5, 2.7, 3, 0.4, ["## Ogni trimestre"], 13)
    riga1 = [(["baseline", "fotografia + allegato di A1"], "python"), (["A1 Horizon Scanner", "valuta i rischi"], "agente"),
             (["precontrollo", "facoltativo"], "python"), (["A5 Challenger", "facoltativo"], "agente"),
             (["trimestrale", "controlla e scrive"], "python")]
    riga2 = [(["Validazione + valida", "accetta, rettifica, respinge"], "persona"), (["schede-input", "dati validati"], "python"),
             (["A6 Schede", "testi per chi legge"], "agente"), (["schede", "Schede_<ciclo>.xlsx"], "python"),
             (["report", "PDF al comitato"], "python")]
    for k, riga in enumerate((riga1, riga2)):
        y = 3.15 + k * 1.05
        for i, (r, t) in enumerate(riga):
            x = 0.5 + i * 2.5
            blocco(s, x, y, 2.15, 0.72, r, t)
            if i:
                freccia(s, x - 0.33, y + 0.36, x - 0.02, y + 0.36)
    freccia(s, 11.07, 3.87, 11.07, 4.0)
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(11.07), Inches(4.0), Inches(1.57), Inches(4.0))
    c.line.color.rgb = GRIGIO
    c.line.width = Pt(1.5)
    freccia(s, 1.57, 4.0, 1.57, 4.18)
    testo(s, 0.5, 5.05, 3, 0.4, ["## Quando serve"], 13)
    altri = [(["A4 News Assessor", "parere su una notizia"], "agente"), (["A7 Assistente", "il perché in chat"], "agente"),
             (["A8 Evento critico", "ciclo straordinario"], "agente"), (["riapri · dismetti", "correzioni e uscite"], "python"),
             (["celle gialle", "presidio, owner, supervisore"], "persona")]
    for i, (r, t) in enumerate(altri):
        blocco(s, 0.5 + i * 2.5, 5.5, 2.15, 0.72, r, t)
    for i, (t, n) in enumerate([("agente", "Agente Copilot (giudizio, testi)"), ("python", "Comando Python (controlli, scrittura, calcolo)"),
                                ("persona", "Persona (decisione)")]):
        blocco(s, 0.5 + i * 4.1, 6.5, 0.35, 0.3, [""], t)
        testo(s, 0.92 + i * 4.1, 6.47, 3.6, 0.35, [n], 11, GRIGIO)


def principi():
    s = nuova("Sette principi che reggono tutto il disegno")
    pr = [("Giudizio all'agente, calcolo al codice", "Gli agenti danno valori 1-5, fonti e testi. Score, esito, urgenza e livello di allerta li calcola il foglio: una media fatta a mente sbaglia in modo plausibile."),
          ("La storia non si riscrive", "RiskRegister, Azioni, KRI e NewsLog crescono solo in coda: una riga per soggetto per ciclo. Ogni ciclo lascia un Excel congelato in archivio."),
          ("Solo un ciclo cambia uno score", "Il trimestrale, o lo straordinario deciso dalla funzione di rischio. Le notizie settimanali danno segnali, mai variazioni."),
          ("Ogni regola ha un controllo", "Dove un vincolo si può scrivere come regola, lo verifica lo script; il giudizio resta ad A5 (facoltativo) e al validatore."),
          ("Una persona valida sempre", "Nessuna valutazione vale finché il validatore non la accetta, la rettifica o la respinge. Il trimestre dopo non parte prima."),
          ("La proposta resta", "Ogni rettifica conserva proposta dell'agente, valore deciso, motivazione, chi e quando: è la traccia di audit e la base per migliorare gli agenti."),
          ("Una regola, un posto", "Scale, tassonomia, fonti e colonne vivono in un solo file; prompt, Excel e controlli ne derivano e il collaudo verifica che siano allineati.")]
    for i, (t, d) in enumerate(pr):
        col, rig = i % 4, i // 4
        x, y = 0.5 + col * 3.1, 1.4 + rig * 2.75
        scheda(s, x, y, 2.95, 2.55, t, [d], NAVY if i % 2 == 0 else BLU, 12)


def chi_fa_cosa():
    s = nuova("Chi fa che cosa: agenti, Python, Excel, persone")
    tabella(s, 0.5, 1.4, 12.33, ["", "Che cosa fa", "Che cosa NON fa"], [
        ["Agenti Copilot\n(A1, A3, A4, A5, A6, A7, A8)",
         "Leggono documenti, notizie e siti; danno i valori 1-5 con le fonti; scrivono razionali, azioni, KRI, sintesi; "
         "rispondono in chat", "Non scrivono in Excel, non salvano file, non calcolano score né urgenza, non decidono"],
        ["Script Python sul PC\n(esegui_ciclo_locale.py, raccogli_notizie.py)",
         "Prepara gli allegati per gli agenti; controlla schema e regole (FERMATO se un difetto renderebbe incoerente il "
         "registro); fa il backup, scrive, ricalcola Excel, archivia; applica la regola del radar; produce schede e PDF",
         "Non giudica la plausibilità dei valori (controllo C4): è giudizio"],
        ["Registro Excel\n(NFR_Registro.xlsx)",
         "Conserva la storia; calcola con le formule score composito, Δ, urgenza di Mosca, livello di allerta, stato dei "
         "KRI, segnali netti; mostra report e viste", "Nessuna costante nascosta: pesi e soglie sono celle gialle di Parametri"],
        ["Persone\n(analista, validatore, funzione di rischio)",
         "L'analista lancia agenti e comandi; il validatore decide ogni proposta nella griglia; la funzione di rischio fissa "
         "parametri, decide gli straordinari e le dismissioni", "—"]],
        [2.3, 6.2, 3.8], 12.5, 1.15)


def agenti_panoramica():
    s = nuova("Gli agenti in una tabella", "2 · Gli agenti e i passi Python")
    tabella(s, 0.5, 1.35, 12.33, ["Agente", "Mansione", "Quando", "Che cosa gli si dà", "Che cosa restituisce"], [
        ["A1 Horizon Scanner", "valuta i rischi, censisce i nuovi, propone azioni e due KRI per rischio del radar", "ogni trimestre",
         "allegato Input_A1_<ciclo>.txt; legge da solo cartelle e siti", "JSON → A1_<ciclo>.json"],
        ["A3 Interim Update", "collega le notizie della settimana ai rischi, intensità sulle nove categorie, flag", "ogni lunedì",
         "niente: legge News_Inbox, siti e baseline", "JSON → A3_<lunedì>.json"],
        ["A4 News Assessor", "parere su UNA notizia rispetto al registro, senza numeri nuovi", "su richiesta",
         "la notizia (testo, link, file)", "parere in sei sezioni, in chat"],
        ["A5 Challenger", "nove controlli C1-C9 sull'output di A1", "trimestre, facoltativo",
         "Input_A5_<ciclo>.txt (output di A1 e baseline)", "JSON → A5_<ciclo>.json"],
        ["A6 Schede dei rischi", "testi delle schede sui dati validati, per chi non conosce il tema", "dopo la validazione",
         "solo Input_A6_<ciclo>.json", "JSON → A6_<ciclo>.json"],
        ["A7 Assistente del registro", "spiega il perché di valori, urgenza, livello, temi non censiti", "su richiesta",
         "niente: legge Kernel\\Spiegazioni", "risposte in chat; sintesi per il Quadro"],
        ["A8 Evento critico", "analisi dell'evento, rivaluta i rischi indicati, censisce i nuovi", "ciclo straordinario",
         "allegato Input_A8_<AAAAQn-Sk>.txt e documenti dell'evento", "JSON → A8_<AAAAQn-Sk>.json"]],
        [1.9, 4.0, 1.5, 2.9, 2.0], 11.5, 0.5)
    nota_fondo(s, ["Regole per tutti: il messaggio aggiunge indicazioni per quella volta, non cambia le regole dell'agente; per "
                   "rifare una valutazione si apre una chat nuova; gli agenti JSON rispondono solo con il JSON. Ogni prompt sta "
                   "nel limite di 8.000 caratteri di Agent Builder."], 6.3, 11)


def slide_agente(nome, sotto, cosa_fa, quando, python, limiti, colore=BLU):
    s = nuova(nome)
    testo(s, 0.5, 1.2, 12.3, 0.4, [sotto], 14, GRIGIO)
    scheda(s, 0.5, 1.7, 6.3, 5.25, "Che cosa fa quando viene usato", cosa_fa, colore, 12)
    scheda(s, 7.0, 1.7, 5.83, 1.75, "Quando e come si usa", quando, NAVY, 11.5)
    scheda(s, 7.0, 3.55, 5.83, 2.25, "Passi Python prima e dopo", python, VERDE, 11.5)
    scheda(s, 7.0, 5.9, 5.83, 1.05, "Che cosa non fa", limiti, GRIGIO, 11.5)


def agenti():
    slide_agente("A1 Horizon Scanner", "Il cuore del trimestre: la valutazione dei rischi, con le fonti e il perché di ogni numero.", [
        "- Valuta **ogni rischio del radar**: Impatto, Probabilità, Prossimità (1-5) e Velocity solo se le fonti ne documentano i tempi; stima gli **anni di adeguamento** strutturale quando servono.",
        "- Rivaluta i **rischi in osservazione** quando trova evidenza nuova; **censisce i rischi nuovi**, anche lontani (con orizzonte e segnale di avvicinamento).",
        "- Per ogni valore cita le **fonti** (data, tipo Interna/Esterna, rilevanza 1-3) e scrive il **razionale in parti**: RISCHIO, IMPATTO, PROBABILITA, PROSSIMITA, VELOCITY, ADEGUAMENTO.",
        "- Per i rischi del radar propone **1-5 azioni** (owner = funzione, orizzonte 3/6/12 mesi) e **due KRI** (esposizione ed evoluzione), con il perché e l'EFFETTO atteso; **rilegge i valori** dei KRI del ciclo prima.",
        "- Registra i **fenomeni considerati e non censiti**, con il motivo; scrive il **commento per il comitato** (MESSAGGIO / DECISIONI).",
        "- Prima di rispondere verifica il proprio output con i controlli C1-C9 e la guida di scrittura contenuti nell'allegato."],
        ["Chat nuova, allegato **Input_A1_<ciclo>.txt**, messaggio «Esegui il ciclo».",
         "Variazioni: rinforzato (leggere tutte le cartelle), documenti precisi, tema da verificare, fenomeno da osservare, simulazione di un ciclo passato, correzione."],
        ["- **Prima: baseline** → Baseline_corrente.json e allegato (baseline + controlli + guida di scrittura).",
         "- **Dopo: precontrollo** (facoltativo, senza Excel) → **trimestrale**: controlli, backup, scrittura, ricalcolo, regola del radar, archivio, knowledge di A7.",
         "- Se FERMATO: si incolla ad A1 il file **Correzione_A1_<ciclo>.txt** e si rilancia.",
         "- Poi la validazione e il comando **valida**."],
        ["Non calcola score, esito (Nuovo/Variato/Confermato), urgenza né livello; non muove un valore di più di un punto; non cambia senza un'evidenza datata."])

    slide_agente("Precontrollo e scrittura (A2): i passi Python del trimestre",
                 "Non sono agenti: sono il codice che si mette fra la risposta di A1 e il registro.", [
        "## precontrollo (facoltativo, pochi secondi, non scrive nulla)",
        "- Schema JSON dell'output, controlli meccanizzati C1, C2, C3, C5, C7, C9 (con le osservazioni di C5, C6, C8) e parti del razionale.",
        "- Divide i rilievi in **da correggere** (fermerebbero il ciclo) e **da verificare** (andranno in validazione).",
        "- Scrive il messaggio di correzione per A1 e l'impronta dell'output (serve alla prova di Copilot Studio).",
        "## trimestrale (la scrittura, detta A2)",
        "- Rifà gli stessi controlli: **FERMATO** solo per i difetti di integrità (schema, C1 rischio mancante o categoria diversa, C2 variazioni che non tornano, C7 e identificativi dei KRI).",
        "- C3, date di C5, coppia e soglie dei KRI (C9) e i Bloccanti di A5 diventano **segnalazioni da motivare ⚑** nella griglia di validazione.",
        "- Calcola esito e «da → a» delle variazioni, assegna gli ID di KRI e azioni, applica la **regola del registro unico** (radar o osservazione).",
        "- Backup, scrittura, ricalcolo di Excel, verifica zero errori di formula, Excel congelato del ciclo, cartella Spiegazioni per A7."],
        ["Dopo A1 (e A5 se usato):",
         "python esegui_ciclo_locale.py precontrollo --payload A1_<ciclo>.json",
         "python esegui_ciclo_locale.py trimestrale --payload A1_<ciclo>.json [--challenger A5_<ciclo>.json]"],
        ["- Codici di uscita: 0 eseguito; 1 **FERMATO** (registro non toccato, si corregge e si rilancia); 2 **ERRORE TECNICO** (il messaggio dice se prima o dopo il salvataggio).",
         "- Il registro deve essere chiuso: Excel lavora invisibile.",
         "- Il trimestre non parte se quello prima non è validato."],
        ["Non giudica la plausibilità dei valori (C4): la giudicano A5, se usato, e il validatore."], VERDE)

    slide_agente("A5 Challenger (facoltativo)", "Un secondo agente che controlla A1 con nove controlli, su un input ristretto.", [
        "- **C1 Completezza**: ogni rischio del radar valutato, stessa categoria, nessun ID ripetuto.",
        "- **C2 Variazioni**: una voce per ogni valore cambiato, e solo per quelli.",
        "- **C3 Evidenza**: data dopo il ciclo prima e non dopo la data di riferimento; evidenza specifica; al massimo un punto.",
        "- **C4 Plausibilità**: i punteggi corrispondono ai descrittori dati razionale e fonti; un rischio = un meccanismo (solo giudizio: nessun codice lo sostituisce).",
        "- **C5 Fonti**: rilevanza coerente; P5 o Prossimità 4-5 con una fonte di rilevanza alta; nessuna fonte futura.",
        "- **C6 Deriva**: aumenti dallo stesso fatto ripetuto o su più di metà dei rischi senza fonti nuove.",
        "- **C7 Rilettura KRI** · **C8 Azioni** (owner = funzione, pertinente, non ripetuta) · **C9 KRI** (coppia esposizione/evoluzione, soglie, fonte dati).",
        "- Esito Approvato o Da rivedere, con i rilievi Bloccante / Osservazione."],
        ["Dopo A1, chat nuova, allegato SOLO **Input_A5_<ciclo>.txt** (output di A1 e baseline, dal precontrollo): «Esegui i nove controlli».",
         "Va rifatto ogni volta che A1 viene rifatto; mai per le settimane."],
        ["- Prima: (facoltativo) **precontrollo**, che segnala già la parte meccanizzabile.",
         "- Dopo: **trimestrale --challenger A5_<ciclo>.json**: i rilievi vanno nella colonna Segnalazioni della griglia; un Bloccante rende obbligatoria la Motivazione.",
         "- La sintesi di A5 compare nel Quadro e in Sintesi_Cicli."],
        ["Non ferma il ciclo (D-71): i suoi rilievi li decide il validatore. Non corregge A1."])

    slide_agente("A3 Interim Update", "Ogni lunedì: le notizie della settimana lette contro il registro, senza toccare gli score.", [
        "- Prende le notizie degli ultimi 7 giorni da tre canali: **News_Inbox** (feed raccolti e file Manuali), **siti** pubblici dell'agente, cartelle **Analisi_Interne / Esterne**.",
        "- Più articoli sullo stesso fatto = **una notizia**; codice fonte del catalogo (FNT-nn) o MAN.",
        "- Per ogni notizia: **intensità 0-3 su tutte e nove le categorie** e direzione (Peggiorativo, Migliorativo, Neutro, Non rilevante).",
        "- **Rischio collegato**: del radar o in osservazione; nessuno → possibile **area non presidiata**.",
        "- **Flag proposti** quando un descrittore potrebbe non essere più quello vigente (o un rischio in osservazione realizza il suo segnale).",
        "- **Proposta di ciclo straordinario** solo con la regola del segnale critico.",
        "- **Sintesi** della settimana: notizie rilevanti, aree non presidiate, direzioni per categoria."],
        ["Chat nuova, allegato **Input_A3_A4.txt** (la baseline): «Esegui l'aggiornamento settimanale». Senza allegato: «Baseline mancante».",
         "Variazioni: correzione dopo un fermo, insistere su un rischio."],
        ["- **Prima: raccogli_notizie.py** (lunedì alle 7, Utilità di pianificazione) → News_Inbox\\<lunedì>\\Raccolta_<data>.txt; a mano i file in Manuali.",
         "- **Dopo: settimanale --payload A3_<lunedì>.json** → NewsLog senza duplicati, Sintesi_Settimane, segnali netti, aree non presidiate, proposta di straordinario a video.",
         "- Facoltativa: validazione delle notizie (**valida --periodo <lunedì>**)."],
        ["Non cambia score, non propone numeri, non decide lo straordinario: lo propone."])

    slide_agente("A4 News Assessor", "Su richiesta di un analista: un parere su una sola notizia, ancorato alla valutazione deliberata.", [
        "Risponde sempre con sei sezioni, in quest'ordine:",
        "- **1. Il fatto** — due o tre frasi, con fonte e data.",
        "- **2. Impatto sulle nove categorie** — tabella con intensità 0-3, direzione e motivazione per ognuna.",
        "- **3. Rischi del registro toccati** — ID, nome, score vigenti, quale descrittore la notizia mette in discussione; altrimenti «Area non presidiata».",
        "- **4. Coerenza con la valutazione deliberata** — conferma, indebolisce o non tocca (un segnale non è una tendenza).",
        "- **5. Cosa suggerisco** — verifiche per la funzione di rischio; in testa «PROPOSTA DI CICLO STRAORDINARIO» se vale la regola del segnale critico.",
        "- **6. Limiti del parere** — che cosa non sa, quali fonti mancano.",
        "Distingue il fatto dalla propria interpretazione; non nomina persone."],
        ["Allegato **Input_A3_A4.txt** (la baseline) e la notizia: «Dammi il parere su questa notizia: [testo o link]».",
         "Variazioni: rispetto a un rischio preciso [NFR-nn]; con attenzione a una categoria [9.n]."],
        ["- **Nessun passo Python**: il parere resta in chat, non entra nel registro.",
         "- Se la notizia è rilevante deve entrare nel prossimo aggiornamento settimanale (A3 + settimanale).",
         "- Se propone uno straordinario: decide la funzione di rischio, poi **baseline --straordinario** e A8."],
        ["Mai una cifra nuova: niente score, stime di perdita o probabilità numeriche. Non scrive nel registro."])

    slide_agente("A6 Schede dei rischi", "Dopo la validazione: rende le schede leggibili a chi non conosce il tema, senza aggiungere numeri.", [
        "- Riceve i **dati validati** del ciclo (valori, fonti, razionale in parti di A1, azioni, KRI, urgenza, storia).",
        "- Scrive per ogni rischio la **descrizione estesa**: il fenomeno in termini di rischio (causa → evento → danno).",
        "- Spiega **ogni dimensione**: perché quel valore, con i fatti e i documenti.",
        "- Dice **a che cosa serve ogni fonte** («a sostegno di»).",
        "- Non valuta: spiega i valori che il validatore ha già deciso.",
        "## Senza A6",
        "- Le schede escono lo stesso con i testi di A1 dal registro (il razionale è già diviso in parti, D-68): A6 è un miglioramento della leggibilità, non un passo obbligato."],
        ["Agente A6, allegato SOLO **Input_A6_<ciclo>.json**: «Scrivi le schede».",
         "Variazione: correzione dopo un fermo del comando schede."],
        ["- **Prima: schede-input** → Input_A6_<ciclo>.json (si ferma se il ciclo non è validato).",
         "- **Dopo: schede --payload A6_<ciclo>.json** → controlla stessi rischi, quattro dimensioni, stesse fonti, **nessun numero fuori dai dati**; poi produce Schede_<ciclo>.xlsx.",
         "- Senza A6: **schede** senza payload."],
        ["Non cambia valori, fonti o urgenza; un numero che non sta nei dati del rischio ferma il comando."])

    slide_agente("A7 Assistente del registro", "In chat, quando serve: il perché delle valutazioni, distinguendo registrato, calcolato e letto.", [
        "- Risponde a domande come: perché NFR-01 è primo? perché ha Prossimità 5, da quale documento? perché è AGIRE ORA? perché Attenzione e non Attivazione? che cosa ha cambiato il validatore? perché il tema X non è un rischio? perché NFR-14 è in osservazione?",
        "- Tiene separati: ciò che è **registrato** (motivazioni di A1, decisioni del validatore), ciò che è **calcolato** (score, urgenza, livello) e la **sua lettura** dei documenti.",
        "- Per i temi non censiti usa i «fenomeni considerati e non censiti» di A1 e i rischi in osservazione.",
        "- A richiesta scrive la **sintesi per il Quadro** (al massimo 600 caratteri) da incollare nella cella gialla Note dell'analista.",
        "- Variazioni: «spiegamelo come a un membro del comitato, in cinque righe»; confronto fra due rischi."],
        ["Nessun allegato, nessun messaggio fisso: si fanno domande.",
         "Legge la cartella **Kernel\\Spiegazioni**: Ciclo.txt e Rischio_01..20.txt, sempre con gli stessi nomi."],
        ["- I file della knowledge li riscrivono da soli **trimestrale**, **valida** e **dismetti**: nessun comando dedicato.",
         "- La nota incollata nel Quadro va nel PDF del **report**; al ciclo dopo **trimestrale** la archivia in Sintesi_Cicli."],
        ["Non propone valori nuovi e non modifica nulla (D-76): i valori si cambiano con riapri + valida, i testi a mano."])

    slide_agente("A8 Evento critico", "Nel ciclo straordinario deciso dalla funzione di rischio: un evento che non può aspettare il trimestre.", [
        "- Legge **per intero** i documenti dell'evento e scrive l'**analisi dell'evento**: fatto, catena causale, rischi toccati e non toccati, documenti, limiti.",
        "- **Rivaluta solo i rischi indicati** dalla funzione di rischio, con le stesse regole di A1 (al massimo un punto rispetto al periodo validato prima, evidenza datata).",
        "- **Censisce i rischi nuovi** che l'evento crea; senza rischi indicati lo straordinario è di **solo censimento**.",
        "- Rilegge i valori dei KRI dei rischi rivalutati.",
        "- Un rischio vicino che l'evento tocca ma non è indicato lo segnala al validatore: non lo riscrive.",
        "- Il periodo si chiama come il trimestre chiuso più -S1, -S2… (es. 2026Q4-S1) e conta come ciclo: il trimestre dopo parte dai suoi valori."],
        ["Chat nuova, allegato **Input_A8_<AAAAQn-Sk>.txt** e i documenti dell'evento: «Esegui il ciclo straordinario».",
         "Lo propongono A3 o A4 (segnale critico) o l'analista; decide la funzione di rischio."],
        ["- **Prima: baseline --straordinario --rischi NFR-.. --motivo \"…\" --decisore \"…\"** → decisione archiviata e allegato.",
         "- **Dopo**: precontrollo (facoltativo), **trimestrale --payload A8_<…>.json**, validazione e **valida**; poi schede e report come nel trimestre.",
         "- L'analisi dell'evento va in Sintesi_Cicli (il fatto) e nella knowledge di A7."],
        ["Non sceglie quali rischi rivalutare; l'analisi dell'evento spiega, non porta valori."])


def comandi():
    s = nuova("I comandi Python, uno per azione")
    fondo = tabella(s, 0.5, 1.35, 12.33, ["Comando", "Quando", "Che cosa fa"], [
        ["raccogli_notizie.py", "ogni lunedì (automatico)", "scarica i feed delle fonti approvate in News_Inbox\\<lunedì>\\"],
        ["baseline [--straordinario]", "inizio trimestre / straordinario", "fotografia del registro validato + allegati: di A1 (baseline, controlli, guida), di A3 e A4 (baseline)"],
        ["precontrollo", "dopo A1, facoltativo", "controlli senza Excel, messaggio di correzione per A1"],
        ["trimestrale", "dopo A1 (e A5)", "controlla, scrive il ciclo, regola del radar, archivio, knowledge di A7"],
        ["settimanale", "dopo A3", "scrive le notizie, segnali netti, sintesi della settimana"],
        ["valida", "dopo la griglia", "applica le decisioni del validatore, Storico_Validazioni; rigenera il report al comitato"],
        ["griglia · riapri", "quando serve", "ricostruisce il foglio Validazione · riapre un ciclo già validato"],
        ["schede-input · schede", "dopo la validazione", "dati per A6 · controlli su A6 e Schede_<ciclo>.xlsx"],
        ["report [--con-lavoro]", "dopo la validazione", "PDF del report al comitato dall'Excel com'è ora"],
        ["dismetti · sblocca", "decisione della funzione di rischio · lucchetto rimasto", "toglie un rischio, la storia resta · toglie il lucchetto e lo registra"],
        ["inizializza · aggiorna-registro", "primo uso · pacchetto nuovo", "registro vuoto · porta il registro sul template nuovo"],
        ["diagnosi", "dopo la validazione, quando serve", "solo lettura: valori dei rischi nuovi, verso delle rettifiche, rischi lontani senza anni di adeguamento"]],
        [2.7, 2.7, 6.9], 12, 0.36)
    nota_fondo(s, ["Ogni scrittura: lucchetto (più analisti), controlli, copia di lavoro, ricalcolo e verifica in Excel, "
                   "registro sostituito solo se nessuno l'ha cambiato nel frattempo; backup e archivio."], fondo + 0.1)


def app_nfr():
    """L'app di orchestrazione (D-83): la finestra (immagine Presentazione/app_finestra.png) e come si lavora."""
    IMG.mkdir(parents=True, exist_ok=True)
    shutil.copy2(QUI / "app_finestra.png", IMG / "app_finestra.png")
    s = nuova("L'app NFR: un doppio clic, il prossimo passo")
    _, _, hh = immagine(s, "app_finestra", 0.5, 1.35, 7.9, 3.3, taglio=(0, 0, 1, 0.75))
    scheda(s, 0.5, 1.35 + hh + 0.15, 7.9, 5.95 - hh - 0.15, "Che cosa fa (e che cosa no)", [
        "- **Stato e prossimo passo** letti dall'archivio condiviso: uguali per tutti i colleghi.",
        "- Un pulsante per processo: settimana, trimestre, straordinario, validazione, schede, report, riapertura…",
        "- Lancia i comandi Python nell'ordine giusto e mostra l'esito (OK, FERMATO, ERRORE TECNICO).",
        "- **Non decide niente**: controlli e scrittura restano allo script, le decisioni al validatore."], NAVY, 11.5)
    scheda(s, 8.65, 1.35, 4.18, 5.95, "Con un agente: tre gesti", [
        "L'app prepara gli allegati, apre l'agente e mette la richiesta negli appunti. Tu:",
        "- **chat nuova**, trascini i file, incolli la richiesta;",
        "- **copi la risposta** e la **incolli nell'app** (o premi Copia: l'app la prende da sola);",
        "- se serve una correzione, il messaggio è già negli appunti: **stessa chat**.",
        "L'app salva il JSON con il nome giusto in In_arrivo e passa al passo dopo.",
        "## Agenti",
        "- Agent Builder o Copilot Studio: stessi file, stesso registro, origine annotata.",
        "## Più analisti",
        "- Lucchetto e copia di lavoro: il registro si sostituisce solo se nessuno l'ha cambiato."], VERDE, 11.5)


def cicli():
    s = nuova("Il trimestre passo per passo: chi fa che cosa")
    tabella(s, 0.5, 1.35, 7.7, ["#", "Chi", "Passo"], [
        ["1", "analista", "Registro, riga 4: il trimestre prima è «Ciclo validato»"],
        ["2", "Python", "baseline → Input_A1_<ciclo>.txt, Input_A3_A4.txt"],
        ["3", "—", "niente da copiare nel Kernel: A3 e A4 hanno l'allegato (AU7)"],
        ["4", "A1", "«Esegui il ciclo» con l'allegato → A1_<ciclo>.json"],
        ["5", "Python · A5", "precontrollo e A5, entrambi facoltativi"],
        ["6", "Python", "trimestrale → ciclo scritto (o FERMATO + correzione)"],
        ["7", "validatore", "griglia del foglio Validazione, poi valida"],
        ["8", "Python · A6", "schede-input, A6, schede → Schede_<ciclo>.xlsx"],
        ["9", "Python", "report → Report_comitato_<ciclo>.pdf (lo fa già valida)"]],
        [0.4, 1.4, 5.9], 12, 0.43)
    scheda(s, 8.45, 1.35, 4.38, 2.6, "Ogni lunedì (≈ 20 minuti)", [
        "- raccogli_notizie.py (automatico alle 7)",
        "- file a mano in Manuali",
        "- A3 con Input_A3_A4.txt → A3_<lunedì>.json",
        "- settimanale → notizie nel registro",
        "- facoltativo: valida le notizie",
        "> Nessuno score cambia di settimana in settimana."], VERDE, 12)
    scheda(s, 8.45, 4.1, 4.38, 2.85, "Ciclo straordinario (D-75, D-77)", [
        "- deciso dalla funzione di rischio",
        "- baseline --straordinario con rischi, motivo, decisore",
        "- A8 con i documenti dell'evento",
        "- trimestrale, validazione, schede, report come nel trimestre",
        "> Rivaluta solo i rischi indicati; gli altri restano validati."], ROSSO, 12)


def schede_rischi():
    divisore("3", "Le schede dei rischi", ["Un file a parte, Schede_<ciclo>.xlsx: una scheda per ogni rischio del radar,",
                                           "con il perché di ogni numero."])
    s = nuova("A che cosa servono le schede")
    scheda(s, 0.5, 1.4, 4.0, 5.5, "Per chi", [
        "- **Comitato**: allegato al report, spiega i rischi uno per uno.",
        "- **Owner del rischio**: che cosa si chiede alla sua funzione (azioni, KRI, scadenze) e perché.",
        "- **Validatore e audit**: su quali documenti poggia ogni valore.",
        "- **Chi non conosce il tema**: il fenomeno raccontato in termini di rischio, senza gergo.",
        "- **Il ciclo dopo**: la memoria del perché, quando i valori cambiano."], NAVY, 12.5)
    scheda(s, 4.67, 1.4, 4.0, 5.5, "Che cosa garantiscono", [
        "- **Solo valori validati**: si producono dopo la validazione, dal registro.",
        "- **Numeri, fonti e urgenza vengono dal registro**: nessun agente li cambia.",
        "- Con A6, il comando **schede** controlla stessi rischi, quattro dimensioni, stesse fonti e **nessun numero fuori dai dati**.",
        "- Ogni punteggio ha accanto il suo perché e i documenti (D-68).",
        "- Ogni azione e ogni KRI hanno il perché e l'effetto atteso."], BLU, 12.5)
    scheda(s, 8.84, 1.4, 4.0, 5.5, "Come si producono", [
        "- **schede-input** → dati validati per A6.",
        "- **A6** scrive i testi (facoltativo).",
        "- **schede** controlla e produce **Schede_<ciclo>.xlsx** in Archivio_Cicli\\Cicli\\<ciclo>.",
        "- Un foglio **Indice** e un foglio per rischio, nell'ordine della classifica (score composito).",
        "- Pronte per la stampa: A4 verticale, intestazione «Uso interno», pagine numerate.",
        "- Senza A6: testi di A1 dal registro."], VERDE, 12.5)

    s = nuova("L'indice delle schede")
    immagine(s, "Schede_Indice", 0.5, 1.35, 8.0, 5.3, allinea="sx")
    testo(s, 8.75, 1.3, 4.1, 5.7, [
        "## Una riga per rischio, dal più alto per score",
        "- **#** posizione in classifica.",
        "- **ID** collegamento alla scheda (clic).",
        "- **Rischio** nome breve; **◆** tema all'attenzione del supervisore.",
        "- **Score** composito 1-5 validato.",
        "- **Urgenza** di Mosca: AGIRE ORA (rosso), Entro l'anno (ambra), Margine n anni (grigio).",
        "- **Livello** di allerta: Attivazione, Attenzione, Monitoraggio.",
        "- **Esito** del ciclo: Nuovo, Variato, Confermato.",
        "- **Owner**: la funzione responsabile.",
        "## In testa",
        "- Ciclo, data dei dati, stato di validazione, origine dei testi (A6 o A1).",
        "## In fondo",
        "- La legenda di tutte le abbreviazioni."], 12)

    s = nuova("Com'è impostata una scheda: dall'alto in basso")
    immagine(s, "Scheda_NFR01_1", 0.5, 1.3, 4.4, 5.7, allinea="sx")
    immagine(s, "Scheda_NFR01_2", 5.0, 1.3, 4.2, 3.1, allinea="sx")
    testo(s, 9.4, 1.3, 3.45, 5.7, [
        "## Pagina 1",
        "- **1** Intestazione: posizione, ID, nome, score, variazione, urgenza, livello.",
        "- **2** Classificazione: categoria, driver, owner, presidio, ◆, esito.",
        "- **3** Il rischio · **4** Razionale.",
        "- **5** Valutazione dimensione per dimensione, con la formula dello score.",
        "- **6** Urgenza e livello di allerta.",
        "- **7** Cosa è cambiato nel ciclo.",
        "- **8** Andamento negli ultimi trimestri.",
        "- **9** Fonti.",
        "## Pagina 2",
        "- **10** Azioni · **11** KRI in vigore, ognuno con perché ed effetto atteso.",
        "- **12** Legenda."], 12)
    testo(s, 5.0, 4.55, 4.2, 2.4, [
        "> Esempio: NFR-01 «Frodi su pagamenti istantanei abilitate da AI generativa», ciclo straordinario 2026Q4-S1, "
        "testi di A1 (senza A6). Lo stesso ordine vale per tutte le schede: chi legge sa sempre dove trovare un'informazione."], 12)

    s = nuova("Intestazione e classificazione: che cosa significa ogni dato")
    tabella(s, 0.5, 1.35, 12.33, ["Informazione", "Che cosa significa", "Perché è importante"], [
        ["Posizione · ID · Nome", "ordine per score composito; codice stabile NFR-nn, uguale in tutti i cicli", "si ritrova lo stesso rischio in classifica, radar, registro e storia"],
        ["Score composito", "media pesata delle dimensioni valutate, 1-5 (5 = peggiore)", "la priorità: quanto è grande il rischio nel complesso"],
        ["Variazione ▲ ▼ ● / nuovo", "differenza di score rispetto al ciclo prima", "dice subito se il rischio sale, scende o è fermo"],
        ["Urgenza", "esito della disuguaglianza di Mosca: AGIRE ORA, Entro l'anno, Margine n anni", "quanto presto agire, indipendente dalla grandezza"],
        ["Livello di allerta", "Attivazione, Attenzione, Monitoraggio, calcolato dal registro", "che cosa si fa adesso: comitato con un piano, verifica dell'owner, ciclo ordinario"],
        ["Categoria · Driver", "categoria della tassonomia interna (9.x) · perché oggi è emergente (DRV-…)", "dove ricade la perdita e che cosa la alimenta"],
        ["Owner · Presidio", "funzione responsabile · giudizio della Banca sui controlli (Adeguato, Parziale, Assente, Non verificato)", "chi risponde; il presidio non entra nello score ma pesa sul livello"],
        ["◆ Attenzione del supervisore", "tema nelle priorità di BCE, EBA o Banca d'Italia, indicato dal validatore", "porta il rischio nel radar anche sotto le soglie"],
        ["Esito · validato", "Nuovo, Variato (cambia una dimensione) o Confermato, calcolato dallo script", "cosa è successo nel ciclo; i valori sono quelli validati"]],
        [2.4, 5.2, 4.7], 12, 0.55)

    s = nuova("Il rischio, il razionale e la valutazione dimensione per dimensione")
    immagine(s, "Scheda_NFR01_1", 0.5, 1.35, 6.2, 5.55, taglio=(0, 0.08, 1, 0.48), allinea="sx")
    testo(s, 6.95, 1.3, 5.9, 5.7, [
        "## Il rischio",
        "- Il **meccanismo di perdita**: che cosa succede, a chi, con quale danno. Con A6 diventa la descrizione estesa.",
        "## Razionale: il fenomeno in termini di rischio",
        "- La parte **RISCHIO** del razionale di A1: causa → evento → danno per la Banca. Fa capire perché è un rischio della Banca e non solo una notizia.",
        "## Valutazione, dimensione per dimensione",
        "- Una riga per **Impatto, Probabilità, Prossimità, Velocity**: il valore validato (colore più scuro = peggiore; n.v. = non valutata) e, accanto, **perché quel valore**, con i fatti e i documenti.",
        "- Sotto, in corsivo, **la formula dello score** con i pesi in vigore (es. 4,65 = (0,35×4 + 0,25×5 + 0,20×5 + 0,20×5) / 1,00); senza Velocity i pesi si ridistribuiscono.",
        "> Importanza: chi legge può verificare ogni numero contro la rubrica e i documenti, senza fidarsi dell'agente."], 12.5)

    s = nuova("Urgenza, livello di allerta, cambiamenti e andamento")
    immagine(s, "Scheda_NFR01_1", 0.5, 1.35, 6.2, 5.55, taglio=(0, 0.36, 1, 0.78), allinea="sx")
    testo(s, 6.95, 1.3, 5.9, 5.7, [
        "## Urgenza (Mosca) e livello di allerta",
        "- **Livello** con il **motivo** (es. «AGIRE ORA; KRI in rosso»): fondo rosso, giallo o verde.",
        "- Il conto di Mosca per esteso: tempo necessario **X + Y** (decidere + mitigare, con gli eventuali anni di adeguamento) contro tempo disponibile **Z** (dalla Prossimità); **margine Z − (X + Y)**. Se negativo: AGIRE ORA.",
        "## Cosa è cambiato nel ciclo",
        "- Per ogni valore cambiato: **da → a**, l'**evidenza** e la sua **data**; altrimenti «confermato sulle fonti del ciclo precedente».",
        "## Andamento negli ultimi trimestri",
        "- Le quattro dimensioni e lo score per ciclo (straordinari compresi) e il grafico dello score su **scala fissa 1-5**.",
        "> Importanza: si vede se un'urgenza viene dalla vicinanza (Z piccolo) o dalla lentezza della mitigazione (Y grande), e se il rischio cresce nel tempo."], 12.5)

    s = nuova("Fonti, azioni e KRI")
    immagine(s, "Scheda_NFR01_2", 0.5, 1.35, 6.0, 5.55, allinea="sx")
    fondo = tabella(s, 6.75, 1.35, 6.08, ["Blocco", "Informazioni e significato"], [
        ["Fonti", "riferimento, data, tipo (Interna/Esterna), rilevanza per la previsione (Alta, Media, Bassa); con A6 anche «a sostegno di» quale valore"],
        ["Azioni", "che cosa fare, owner (una funzione), orizzonte 3/6/12 mesi, stato (Proposta, Approvata, In corso, Completata…) e ciclo; sotto, **perché** e **effetto atteso**"],
        ["KRI in vigore", "[tipo] nome (unità), valore del ciclo, soglie attenzione / allerta, valore del ciclo prima, stato Verde/Ambra/Rosso; sotto, **perché lo misuriamo** ed **effetto atteso**"]],
        [1.3, 4.8], 11.5, 0.5)
    testo(s, 6.75, fondo + 0.2, 6.08, 6.95 - fondo - 0.2, [
        "> Importanza: le fonti dicono quanto è solido il giudizio; le azioni sono gli impegni chiesti alle funzioni e "
        "alimentano la Y di Mosca; i KRI dicono come la Banca saprà, nei cicli dopo, se il rischio cresce (esposizione) e "
        "se il fenomeno accelera (evoluzione)."], 12)

    s = nuova("Due rischi AGIRE ORA per ragioni diverse")
    immagine(s, "Scheda_NFR01_1", 0.5, 1.35, 6.0, 2.6, taglio=(0, 0, 1, 0.42))
    immagine(s, "Scheda_NFR11_1", 6.83, 1.35, 6.0, 2.6, taglio=(0, 0, 1, 0.42))
    scheda(s, 0.5, 4.15, 6.0, 2.8, "NFR-01 · il rischio è già qui", [
        "- Prossimità 5 (in corso presso la Banca) → **Z = 0** anni.",
        "- Azioni a 3 mesi → X + Y = 0,25 + 0,25 = 0,50.",
        "- Margine −0,50: **AGIRE ORA**; con un KRI in rosso → **Attivazione**.",
        "- Score 4,65: primo in classifica."], ROSSO, 12.5)
    scheda(s, 6.83, 4.15, 6.0, 2.8, "NFR-11 · la mitigazione è lunga", [
        "- Prossimità 3 → **Z = 1,5** anni: c'è tempo prima che si manifesti.",
        "- Ma servono **3 anni di adeguamento** strutturale (sostituire sistemi fuori supporto) → Y = 3.",
        "- X + Y = 3,25 > 1,5: **AGIRE ORA** anche con score 2,80, undicesimo in classifica.",
        "- È il caso per cui Mosca ha scritto la disuguaglianza."], AMBRA, 12.5)


def report_intro():
    divisore("4", "Il report trimestrale", ["Il registro Excel NFR_Registro.xlsx: dal Quadro d'insieme",
                                            "al significato di ogni foglio e di ogni dato."])
    s = nuova("Il registro: un solo file, quattro gruppi di fogli")
    gruppi = [("Guida", ["Guida"], GRIGIO, "come si usa il workbook, chi aggiorna cosa, dove si scrive a mano, regole"),
              ("Report al comitato", ["Quadro", "Classifica", "Radar_Mosca", "Matrice", "Delta", "Fenomeni", "Metodologia"], NAVY,
               "le pagine del PDF per il comitato"),
              ("Lavoro", ["Registro", "Storico", "Indicatori", "Notizie", "Piano_azioni", "Osservazione", "Validazione",
                          "Parametri", "Dizionario"], BLU, "viste per l'analista e il validatore; l'unico foglio da compilare è Validazione"),
              ("Dati (li scrive solo lo script)", ["RiskRegister", "Azioni", "KRI", "NewsLog", "Sintesi_Cicli", "Sintesi_Settimane",
                                                   "Storico_Validazioni"], VERDE, "la storia, una riga per soggetto per ciclo o per notizia")]
    y = 1.4
    for nome, fogli, col, d in gruppi:
        h = 0.3 + 0.42 * (-(-len(fogli) // 5))
        rettangolo(s, 0.5, y, 3.2, h, col)
        testo(s, 0.6, y, 3.0, h, [f"**{nome}**", d], 11, BIANCO, MSO_ANCHOR.MIDDLE, spazio=1)
        for i, f in enumerate(fogli):
            blocco(s, 3.85 + (i % 5) * 1.82, y + 0.1 + (i // 5) * 0.42, 1.72, 0.34, [f], "file", 11)
        y += h + 0.12
    testo(s, 0.5, y + 0.05, 12.3, 6.95 - y, [
        "- **Fascia blu** in cima a ogni foglio: ciclo, data dei dati, «validato» o «da validare». **Riga 2**: chi aggiorna il foglio "
        "(azzurro trimestrale, verde settimanale, giallo validatore, grigio calcolo), che cosa contiene, come si legge.",
        "- **Celle gialle** = input a mano; tutto il resto è scritto dallo script o è una formula. In fondo ai fogli del report, la legenda."], 12)

    s = nuova("Il report al comitato: che cos'è e come si usa")
    scheda(s, 0.5, 1.4, 5.6, 5.5, "Che cos'è", [
        "- Il **PDF** che il comando **report** produce dal registro validato: Report_comitato_<ciclo>.pdf.",
        "- Si fa **dopo la validazione**: con righe da validare il comando si ferma.",
        "- Si **rigenera** quando il validatore cambia a mano presidio, anni di adeguamento, supervisore o stato di un'azione (la versione prima resta, con la data nel nome).",
        "- **--con-lavoro** aggiunge Registro e Parametri.",
        "- Si accompagna con il file delle **schede**.",
        "## Come si legge",
        "- Dal generale al particolare: priorità e stato (Quadro) → grandezza e urgenza (Classifica, Radar, Matrice) → cambiamenti (Delta) → rischi fuori dal radar (Fenomeni) → metodo (Metodologia)."], NAVY, 12.5)
    tabella(s, 6.35, 1.4, 6.48, ["Pagina", "La domanda a cui risponde"], [
        ["Quadro (2 pagine)", "Il ciclo è completo e validato? Quali sono le priorità e perché?"],
        ["Classifica", "Quali rischi sono più grandi, e chi è in ritardo?"],
        ["Radar_Mosca", "Quanto presto bisogna agire, rischio per rischio?"],
        ["Matrice", "Quanto è grave ogni rischio (Probabilità × Impatto)?"],
        ["Delta (2 pagine)", "Che cosa è cambiato dal ciclo prima, e su quale evidenza?"],
        ["Fenomeni", "Quali rischi teniamo d'occhio fuori dal radar?"],
        ["Metodologia (5 pagine)", "Con quali scale, pesi e regole sono fatti i numeri?"],
        ["Registro, Parametri", "(con --con-lavoro) il dettaglio e i parametri in vigore"]],
        [2.0, 4.5], 12, 0.55)


def report_fogli():
    s = nuova("Quadro d'insieme (1/3): i cinque riquadri", "4 · Il report trimestrale")
    immagine(s, "Quadro_1", 0.5, 1.35, 7.1, 5.55, allinea="sx")
    testo(s, 7.85, 1.3, 5.0, 5.7, [
        "La prima pagina per il comitato. **Prima di leggere i numeri: sono definitivi?**",
        "- **Rischi valutati** — 13 su 13 in anagrafica, «completo»: nessun rischio dimenticato.",
        "- **Livello di allerta** — quanti in **attivazione** e in **attenzione** (sotto: quali in attivazione, quanti in monitoraggio): il carico di lavoro del comitato.",
        "- **Variati · nuovi (radar)** — quanti hanno cambiato valore e quanti sono nuovi rispetto al ciclo prima; sotto, entrati e usciti dal radar e quanti in osservazione.",
        "- **KRI in rosso** — su quanti in vigore: i segnali misurati, non stimati.",
        "- **Validazione** — ✓ ciclo validato, oppure quante righe restano: se non è validato, i numeri non valgono.",
        "> Colori: rosso = qualcosa chiede attenzione, verde = a posto."], 12.5)

    s = nuova("Quadro d'insieme (2/3): priorità, gravità, urgenza")
    immagine(s, "Quadro_1", 0.5, 1.35, 7.1, 5.55, taglio=(0, 0.27, 1, 1), allinea="sx")
    testo(s, 7.85, 1.3, 5.0, 5.7, [
        "- **Priorità — primi cinque per score composito**: posizione, rischio (◆ supervisore), score con la barra su scala fissa 0-5, variazione (▲ ▼ ●) e urgenza colorata.",
        "- **Gravità: Probabilità × Impatto**: la matrice in piccolo, con i numeri dei rischi (01 = NFR-01) nelle celle verdi, ambra, rosse.",
        "- **Note dell'analista** (cella gialla): testo libero, per esempio la sintesi di A7; al ciclo dopo va in Sintesi_Cicli.",
        "- **Urgenza: radar di Mosca**: lo stesso grafico del foglio Radar_Mosca, in piccolo.",
        "- **I rischi del ciclo**: la legenda dei codici, con ◆.",
        "> Uso: in un colpo d'occhio, i rischi più grandi (priorità), i più gravi (matrice) e i più urgenti (radar) — tre misure diverse che non sempre coincidono."], 12.5)

    s = nuova("Quadro d'insieme (3/3): il commento di A1 e i rischi segnalati")
    immagine(s, "Quadro_2", 0.5, 1.35, 7.1, 5.55, allinea="sx")
    testo(s, 7.85, 1.3, 5.0, 5.7, [
        "## Il commento dell'analisi (A1)",
        "- **Messaggio chiave** e **Decisioni richieste**: il commento per il comitato, scritto **prima** della validazione.",
        "- **Controlli**: la sintesi di A5 o «A5 non usato», con le osservazioni.",
        "- **Dopo il commento**: compilata se il validatore ha rettificato qualcosa dopo — valgono i numeri validati.",
        "- **Segnalazioni da motivare**: quante ⚑ sono andate in validazione.",
        "## Rischi segnalati e perché",
        "- Chi: i primi tre per score, quelli in attivazione o attenzione, i variati o nuovi.",
        "- **Livello** · **Perché lo segnaliamo** (calcolato: posizione, AGIRE ORA con X + Y e Z, supervisore, presidio, KRI rossi, notizie) · **Cosa è cambiato, e su quale evidenza**.",
        "- In fondo i **limiti** della valutazione e la **legenda**."], 12)

    foglio("Classifica e lente di Mosca", "Classifica", [
        "## Tabella",
        "- **Pos., ID, Nome**; **Score composito** e **score del ciclo prima**; **Δ score**; **Esito** (Nuovo, Variato, Confermato); **Urgenza** colorata; **Margine Z − (X + Y)** in anni; **Owner**.",
        "## A sinistra: la classifica",
        "- Una barra per rischio, dal più alto per score; colore = urgenza (rosso agire ora, ambra entro l'anno, grigio c'è margine); freccia = variazione.",
        "## A destra: la lente di Mosca",
        "- Dal più in ritardo: la colonna è il **tempo necessario** (grigio X, arancio Y entro Z, rosso Y oltre Z), il rombo il **tempo disponibile Z**.",
        "> Uso: la classifica dice la grandezza, la lente spiega il colore — perché un rischio piccolo può essere urgente (NFR-11)."],
        larg_img=7.3)

    s = nuova("Radar di Mosca: quanto presto agire, e quanto è grande ogni rischio")
    immagine(s, "Radar_Mosca", 0.5, 1.3, 6.3, 5.65, taglio=(0, 0.47, 0.6, 1), allinea="sx")
    testo(s, 6.95, 1.3, 5.9, 5.7, [
        "- **Distanza dal centro** = margine di Mosca + 2 (fra 0,6 e 6): **più vicino al centro = agire prima**.",
        "- **Cerchio rosso «Mosca»** = margine 0: dentro, X + Y > Z, si è già in ritardo anche se l'evento non è imminente.",
        "- **Cerchi grigi**: 1, 2 e 3 anni di margine.",
        "- **Area del cerchio** = Probabilità × Impatto (la grandezza).",
        "- **Colore** = livello di allerta: rosso Attivazione, giallo Attenzione, verde Monitoraggio.",
        "- **Settore** = categoria della tassonomia interna (nove spicchi); rischi dello stesso settore a distanze angolari uguali.",
        "- **◆** = tema all'attenzione del supervisore.",
        "> Uso: unisce in un'immagine urgenza (posizione), grandezza (area) e decisione (colore). Nessun agente stima l'urgenza: la calcola il foglio."], 12.5)

    s = nuova("Radar di Mosca: la tabella sotto il grafico")
    immagine(s, "Radar_Mosca", 0.5, 1.35, 7.2, 3.3, taglio=(0, 0.0, 1, 0.42), allinea="sx")
    fondo = tabella(s, 7.95, 1.35, 4.88, ["Colonne", "Significato"], [
        ["Impatto, Probabilità, P×I", "la grandezza (area)"],
        ["Prossimità → Z", "anni prima che si manifesti"],
        ["X", "anni per decidere e avviare (0,25)"],
        ["Y", "anni per mitigare: azioni o adeguamento"],
        ["Margine, Urgenza", "Z − (X + Y): AGIRE ORA, entro l'anno, margine"],
        ["Settore, angolo, raggio, x, y", "geometria del grafico"],
        ["F1-F3, A1-A3", "segnali forti e aggravanti (1 = sì)"],
        ["Livello, Motivo", "livello di allerta e perché"]],
        [2.1, 2.8], 10.5, 0.3)
    tabella(s, 0.5, max(fondo, 4.7) + 0.15, 12.33, ["Livello", "Quando (calcolato dal foglio)", "Conseguenza"], [
        ["Attivazione", "almeno 2 segnali forti: F1 AGIRE ORA, F2 un KRI in rosso, F3 notizie in peggioramento oltre la soglia", "al comitato con un piano d'azione"],
        ["Attenzione", "1 segnale forte, oppure margine sotto 1 anno con un aggravante: A1 P×I ≥ 12, A2 presidio assente, A3 metà dei KRI in ambra o rosso", "l'owner conferma azioni e KRI entro il ciclo"],
        ["Monitoraggio", "gli altri casi", "ciclo trimestrale ordinario"]],
        [1.5, 7.6, 3.2], 11.5, 0.36, colori_prima={0: ROSSO_CH, 1: AMBRA_CH, 2: VERDE_CH})

    foglio("Matrice di gravità", "Matrice", [
        "- Ogni rischio nella cella della sua **Probabilità** (righe, da Remota a Quasi certa) e del suo **Impatto** (colonne, da Limitato a Critico).",
        "- **Colore** = P × I: verde fino a 6, ambra da 8 a 12, rosso da 15.",
        "- **◆** = tema all'attenzione del supervisore.",
        "- Sotto: **i rischi dal più grave**, con P · I e il prodotto P × I.",
        "> Dice la **gravità**, non l'urgenza: un rischio grave ma lontano sta in alto a destra nella matrice e lontano dal centro nel radar.",
        "> Uso: il linguaggio più familiare al comitato, quello delle matrici di rischio operativo."],
        larg_img=5.4)

    s = nuova("Delta: che cosa è cambiato dal ciclo precedente")
    immagine(s, "Delta_1", 0.5, 1.35, 7.4, 2.75, allinea="sx")
    immagine(s, "Delta_2", 0.5, 4.2, 7.4, 2.75, allinea="sx")
    testo(s, 8.15, 1.3, 4.7, 5.7, [
        "## Pagina 1",
        "- **Una barra per rischio**, dalla variazione più alta dello score: rosso in aumento, blu in diminuzione, grigio invariato.",
        "- **Perché varia (o no)**, nello stesso ordine: il **motivo con l'evidenza** e la sua data, oppure «Invariato: nessuna evidenza nuova».",
        "## Pagina 2",
        "- **Quanto vale un punto** di ogni dimensione sullo score (con e senza Velocity).",
        "- **Dettaglio per dimensione**: Δ Impatto, Probabilità, Prossimità, Velocity e Δ score; la **Nota** segnala una violazione della regola del punto.",
        "- Evidenziate le variazioni di score oltre la soglia del Delta (0,30).",
        "> Uso: il comitato vede che ogni cambiamento ha un'evidenza datata."], 12)

    foglio("Fenomeni: i rischi in osservazione", "Fenomeni", [
        "- In testa: **quanti nel radar** (e la capienza) e **quanti in osservazione** (al massimo 20), entrati e usciti nel ciclo.",
        "- Per ogni rischio fuori dal radar, dal più alto per score: **categoria**, le quattro dimensioni, **score**.",
        "- **Perché fuori dal radar**: quale soglia manca (Prossimità o Probabilità sotto 2) o la capienza.",
        "- **Orizzonte** in anni (la Prossimità non distingue oltre i 3 anni).",
        "- **Segnale che lo porterebbe nel radar**: il fatto osservabile da tenere d'occhio.",
        "- **Driver** e **notizie di A3** collegate nelle ultime 4 settimane.",
        "- In fondo chi è **entrato o uscito** dal radar nel ciclo.",
        "> Uso: il comitato vede anche ciò che non è ancora urgente, e che cosa lo farebbe diventare tale (es. il calcolo quantistico)."],
        larg_img=7.3)

    s = nuova("Metodologia: come si leggono i numeri")
    immagine(s, "Metodologia_1", 0.5, 1.35, 7.9, 5.55, allinea="sx")
    testo(s, 8.65, 1.3, 4.2, 5.7, [
        "Cinque pagine, sempre allineate ai parametri in vigore (le legge dal foglio Parametri):",
        "- **Chi produce i numeri**: agenti, script, validatore.",
        "- **Fonti di evidenza** e loro rilevanza.",
        "- **Analisi per dimensione** e **scale** con i descrittori.",
        "- **Score composito**: formula e pesi.",
        "- **Regole di variazione** fra cicli.",
        "- **Urgenza** (Mosca) e **livello di allerta**.",
        "- **KRI** e **registro unico**.",
        "- **Come si leggono i grafici** e la **validazione**.",
        "> Uso: risponde alla domanda «con quale metodo?» senza aprire altri documenti."], 12)

    foglio("Registro: il ciclo corrente, un rischio per riga", "Registro", [
        "- **Righe 4-5**: righe da validare e rettifiche applicate → «Ciclo validato»; rischi in anagrafica contro valutati → «Ciclo completo». Prima di tutto: se non dice così, i numeri non sono definitivi.",
        "- **Rischio**: ID (fondo = livello di allerta), nome, categoria.",
        "- **Dimensioni** Imp, Prob, Pross, Vel (n.v. = non valutata).",
        "- **Calcolati**: score, Δ ciclo prima, Δ 6 mesi, urgenza, margine.",
        "- **◆** supervisore; **KRI** R · A · V del ciclo.",
        "- **Segnali netti** delle notizie (4 settimane, ▲ ▼ oltre la soglia): un allarme, non cambia lo score.",
        "- **Presidio** della Banca; **Rilievi** di A5 e dei controlli (⚑); **Owner**.",
        "> Uso: la tabella di lavoro dell'analista."], sezione=None, larg_img=7.4, size=12)

    foglio("Storico: lo score negli ultimi otto cicli", "f_Storico", [
        "- Una riga per rischio, una colonna per ciclo (straordinari compresi), dal più vecchio.",
        "- **#N/A = rischio non valutato in quel ciclo**, non score zero.",
        "- I cicli passati sono **ricalcolati con i pesi in vigore**: cambiando i pesi, il trend resta confrontabile.",
        "- Il grafico ha **scala fissa** 0-5: cicli diversi si confrontano a colpo d'occhio.",
        "> Uso: vedere da quanto tempo un rischio cresce, e se la crescita è continua o a salti."], larg_img=7.4)

    foglio("Indicatori: i KRI del ciclo e la storia di un KRI", "f_Indicatori", [
        "- In testa: quanti KRI **Verde, Ambra, Rosso, Non rilevato, Soglie da definire**.",
        "- **Tabella** dei KRI in vigore: ID, rischio, nome, unità, soglie, **valore** e data di rilevazione, **stato**, valore del ciclo prima, **tendenza** (▲ peggiora), **tipo** (Esposizione / Evoluzione).",
        "- **A destra** la storia di un KRI, da scegliere nella cella gialla N4 (anche uno uscito di vigore).",
        "- Stato: Verde sotto la soglia di attenzione, Ambra fino all'allerta, Rosso oltre; Non rilevato senza valore (mai zero).",
        "> Uso: sono le misure, non le stime: un KRI in rosso è un segnale forte del livello di allerta."], larg_img=7.4)

    foglio("Notizie: la pressione sulle categorie di rischio", "f_Notizie", [
        "- Per ognuna delle **nove categorie**, sulle ultime 4 settimane: **intensità massima** (0-3), notizie con intensità ≥ 2, rischi in anagrafica, **presidio**.",
        "- **AREA NON PRESIDIATA** = notizie di intensità ≥ 2 su una categoria senza rischi nel registro: tema per il prossimo trimestre. **In osservazione** = nessun rischio nel radar, almeno uno in osservazione.",
        "- Grafico: notizie per intensità 1, 2, 3.",
        "- A destra: **sintesi di A3** della settimana, **flag proposti**, in rosso l'eventuale **proposta di ciclo straordinario**.",
        "> Uso: il foglio del ciclo settimanale. Le notizie non cambiano score."], larg_img=7.4)

    foglio("Piano delle azioni", "f_Piano_azioni", [
        "- Tutte le azioni: prima le **proposte da decidere**, poi approvate, in corso, completate.",
        "- **La proposta (A1)**: ciclo, rischio, ID azione (<ciclo>-NFR-nn-An), azione, owner, orizzonte.",
        "- **La decisione e l'avanzamento**: stato, data della decisione, **scadenza** = decisione + orizzonte, esito della validazione.",
        "- La validazione è la **delibera**: Accettato o Rettificato = Approvata, Respinto = Respinta. In corso e Completata le scrive a mano chi tiene il piano (foglio Azioni, colonna Stato).",
        "- Le azioni aperte tornano ad A1 nel trimestre dopo.",
        "> Uso: il registro degli impegni e delle scadenze; l'orizzonte più lungo è la Y di Mosca."], larg_img=7.4)

    foglio("Osservazione: dove si concentrano i rischi fuori dal radar", "f_Osservazione", [
        "- **A sinistra**: i rischi in osservazione con dimensioni, score, **perché fuori dal radar**, orizzonte, **segnale** di avvicinamento, **meccanismo**, validazione, **notizie di A3** delle ultime 4 settimane.",
        "- **A destra**: incrocio **driver × categoria** dei rischi in osservazione e, accanto, dei rischi del radar.",
        "- Una casella piena nei rischi in osservazione e vuota nel radar = **un tema che il radar non gestisce**.",
        "> Uso: il lavoro dell'analista sui rischi lontani; per il comitato la stessa informazione è nella pagina Fenomeni."], larg_img=7.4)

    s = nuova("Validazione (1/2): la griglia, una riga per oggetto")
    immagine(s, "Validazione_Q3_1", 0.5, 1.35, 7.4, 5.55, allinea="sx")
    testo(s, 8.15, 1.3, 4.7, 5.7, [
        "L'**unico foglio da compilare**. Ci sono solo i periodi ancora da validare.",
        "- Una riga per **rischio** (nel radar o in osservazione), **KRI**, **azione** e **notizia**.",
        "- Celle **gialle** = valore proposto dall'agente (accanto, in grigio, la proposta): se d'accordo **non si toccano** (= Accettato).",
        "- Se non d'accordo: si scrive il valore giusto (la cella diventa arancio) e il perché in **Motivazione** (almeno 10 caratteri).",
        "- **Respingi = Sì** toglie un KRI, un'azione o una notizia (un rischio non si respinge).",
        "- **Segnalazioni dei controlli**: ⚑ = Motivazione obbligatoria anche senza cambiare nulla.",
        "- **Score** e **margine di Mosca** si ricalcolano mentre si scrive.",
        "- **Decisione sullo stato**: Automatico, Radar o Osservazione."], 12)

    s = nuova("Validazione (2/2): KRI, azioni e criteri per decidere")
    immagine(s, "Validazione_Q3_2", 0.5, 1.35, 7.4, 5.55, allinea="sx")
    fondo = tabella(s, 8.15, 1.35, 4.68, ["Oggetto", "Si cambia / si respinge se…"], [
        ["Rischio", "il valore non corrisponde al descrittore; intero 1-5, al massimo un punto, con evidenza; anni di adeguamento 0-30"],
        ["KRI", "fonte dati non identificabile o non misura il rischio; soglie (entrambe, allerta ≥ attenzione) o valore errato"],
        ["Azione", "duplicata o non pertinente; owner = una funzione, mai una persona; orizzonte 3, 6 o 12 mesi"],
        ["Notizia", "rischio o segnale sbagliati; non pertinente o già contata"]],
        [1.1, 3.6], 11, 0.4)
    testo(s, 8.15, fondo + 0.15, 4.68, 6.95 - fondo - 0.15, [
        "Poi, a file chiuso: **valida --periodo <ciclo> --validatore \"…\"**. Controlla tutte le righe prima di applicarne una; "
        "le righe escono dal foglio e il dettaglio resta in Storico_Validazioni. Per cambiare un ciclo già validato: **riapri**."], 11.5)

    s = nuova("Guida e Dizionario: il manuale dentro il file")
    immagine(s, "f_Guida", 0.5, 1.35, 6.0, 4.2)
    immagine(s, "f_Dizionario", 6.83, 1.35, 6.0, 4.2)
    testo(s, 0.5, 5.65, 6.0, 1.4, [
        "**Guida**: a cosa serve il registro, **chi aggiorna cosa e quando** (foglio per foglio), quando si lanciano A5 e A6, "
        "**dove si scrive a mano**, le regole, il primo uso."], 12)
    testo(s, 6.83, 5.65, 6.0, 1.4, [
        "**Dizionario**: **ogni dato** del workbook — dove sta, chi lo scrive e quando, che cosa significa, come si calcola, "
        "valori ammessi, un esempio. Viene dalla mappa dei campi: sempre quello in uso."], 12)

    foglio("Dati · RiskRegister: tutte le valutazioni, per sempre", "f_RiskRegister", [
        "Una riga per rischio per ciclo, solo in aggiunta, con i **valori validati**:",
        "- **Identità**: ciclo, data di riferimento, ID, nome, categoria, driver, owner.",
        "- **Valutazione**: Impatto, Probabilità, Prossimità, Velocity, score composito (formula), anni di adeguamento.",
        "- **Perché**: descrizione (meccanismo), razionale in parti, variazioni da → a con evidenza, fonti con rilevanza.",
        "- **Controlli e decisioni**: esito, rilievi (⚑), validazione, approvato da, data.",
        "- **Registro unico**: stato (Radar, Radar in uscita, Osservazione, Dismesso), motivo, movimento (Entrato/Uscito), decisione sullo stato, orizzonte, segnale.",
        "> Importanza: è la fonte di ogni vista e della baseline del ciclo dopo."], larg_img=7.4, size=11.5)

    s = nuova("Dati · Azioni, KRI, NewsLog")
    immagine(s, "f_Azioni", 0.5, 1.35, 4.0, 2.5)
    immagine(s, "f_KRI", 4.67, 1.35, 4.0, 2.5)
    immagine(s, "f_NewsLog", 8.84, 1.35, 4.0, 2.5)
    scheda(s, 0.5, 4.0, 4.0, 2.95, "Azioni", [
        "Una riga per azione: ciclo, rischio, **ID**, descrizione, owner, orizzonte (mesi → Y di Mosca), razionale con EFFETTO, **stato**, data e autore della decisione, validazione."], NAVY, 11.5)
    scheda(s, 4.67, 4.0, 4.0, 2.95, "KRI", [
        "Una riga per KRI per ciclo: **KRI_ID** stabile, rischio, nome, **tipo**, unità, fonte dati, **soglie**, **valore** (vuoto se non rilevato, mai ricopiato), stato, **in vigore**, ciclo di nascita."], BLU, 11.5)
    scheda(s, 8.84, 4.0, 4.0, 2.95, "NewsLog", [
        "Una riga per notizia: data, settimana, **ID**, titolo, fonte e codice, **rischio collegato** (o NON PRESIDIATO), **segnale**, intensità **Imp_9_x** sulle nove categorie, sintesi, validazione."], VERDE, 11.5)

    s = nuova("Dati · Sintesi_Cicli, Sintesi_Settimane, Storico_Validazioni")
    immagine(s, "f_Sintesi_Cicli", 0.5, 1.35, 4.0, 2.5)
    immagine(s, "f_Sintesi_Settimane", 4.67, 1.35, 4.0, 2.5)
    immagine(s, "f_Storico_Validazioni", 8.84, 1.35, 4.0, 2.5)
    scheda(s, 0.5, 4.0, 4.0, 2.95, "Sintesi_Cicli", [
        "Una riga per ciclo: **commento di A1** (MESSAGGIO / DECISIONI), **sintesi di A5**, osservazioni e rilievi del ciclo, segnalazioni da motivare, **Note dell'analista** archiviate, dati dello **straordinario** (chi, quando, perché)."], NAVY, 11.5)
    scheda(s, 4.67, 4.0, 4.0, 2.95, "Sintesi_Settimane", [
        "Una riga per settimana: notizie valutate, **sintesi di A3**, **proposta di straordinario**, **flag proposti** (rischio, direzione, motivazione). Un flag non cambia score: va al trimestre."], BLU, 11.5)
    scheda(s, 8.84, 4.0, 4.0, 2.95, "Storico_Validazioni", [
        "Una riga per oggetto e campo: **proposta dell'agente**, **decisione** (Accettato, Rettificato, Respinto), **valore validato**, **motivazione**, chi e quando. È l'**audit** e la base per migliorare gli agenti."], VERDE, 11.5)


def parametri():
    divisore("5", "I parametri", ["Gerarchia di rischio, scale, pesi, soglie e valori ammessi:",
                                  "che cosa vale oggi e chi lo decide."])
    s = nuova("Dove vivono i parametri e chi li decide")
    immagine(s, "Parametri_1", 0.5, 1.35, 7.2, 5.55, allinea="sx")
    testo(s, 7.95, 1.3, 4.9, 5.7, [
        "## Foglio Parametri (celle gialle)",
        "- Pesi dello score, soglie, parametri dell'urgenza e del livello di allerta, regola del radar: li decide la **funzione di rischio** e li registra in DECISIONI.md. Valgono per tutti i cicli, passati compresi.",
        "- **Anagrafica** dei rischi: owner, anni di adeguamento, attenzione del supervisore, presidio dei controlli: le decide il **validatore**.",
        "- Ciclo e settimana correnti e periodi del trend li scrive lo script.",
        "## File del kernel",
        "- **Tassonomia** (categorie e driver), **rubrica** (scale e regole), **fonti** (catalogo): entrano nei prompt degli agenti.",
        "> I descrittori sono segnaposto verosimili: la Banca li riscrive nel proprio linguaggio; i pesi sono provvisori (D-20)."], 12)

    s = nuova("La gerarchia: dalla famiglia di rischio al singolo indicatore")
    livelli = [("Famiglia", "Rischi Non Finanziari – Rischi Operativi (tassonomia interna della Banca)", NAVY),
               ("Categoria (9 in perimetro)", "dove ricade la perdita: 9.1 … 9.13 · è il settore del radar", BLU),
               ("Rischio NFR-nn", "un meccanismo di perdita, una sola categoria, uno o più driver · nel radar o in osservazione", RGBColor(0x41, 0x8A, 0xC9)),
               ("KRI e azioni", "due KRI (NFR-nn-K1, K2: esposizione ed evoluzione) · 1-5 azioni (<ciclo>-NFR-nn-An)", VERDE),
               ("Fonti ed evidenze", "documenti con data, tipo e rilevanza 1-3 · notizie collegate (N-AAAAMMGG-nn)", GRIGIO)]
    for i, (t, d, c) in enumerate(livelli):
        x, y = 0.5 + i * 0.45, 1.4 + i * 0.95
        rettangolo(s, x, y, 3.0, 0.8, c, forma=MSO_SHAPE.ROUNDED_RECTANGLE)
        testo(s, x + 0.05, y, 2.9, 0.8, [f"**{t}**"], 13, BIANCO, MSO_ANCHOR.MIDDLE).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        testo(s, x + 3.15, y, 12.83 - x - 3.15 - 3.3, 0.8, [d], 12.5, TESTO, MSO_ANCHOR.MIDDLE)
    scheda(s, 9.55, 1.4, 3.28, 4.6, "Trasversali", [
        "- **Driver emergenti** (8): perché il rischio è emergente oggi; un rischio ne ha uno o più.",
        "- **Owner**: la funzione che risponde del rischio e delle azioni.",
        "- **Stato nel registro unico**: Radar, Radar in uscita, Osservazione, Dismesso."], AMBRA, 12)
    nota_fondo(s, ["Regola di classificazione: un rischio ha UNA categoria e uno o più driver. Se sembra avere due categorie sono due "
                   "rischi; anche due meccanismi di perdita diversi nella stessa categoria sono due rischi."], 6.3, 12)

    s = nuova("Tassonomia: le nove categorie in perimetro e gli otto driver")
    tabella(s, 0.5, 1.35, 6.3, ["Codice", "Categoria", "Perdite da"], [
        ["9.1", "Condotta", "comportamenti non corretti verso clienti o mercati"],
        ["9.4", "Legale", "violazioni di leggi, responsabilità contrattuale, controversie"],
        ["9.5", "Fiscale", "violazioni di norme tributarie"],
        ["9.6", "Tecnologico", "sistemi e servizi ICT inadeguati o non sostituibili"],
        ["9.7", "Sicurezza informatica", "dati e sistemi violati o indisponibili, attacchi"],
        ["9.8", "Sicurezza fisica", "eventi fisici su persone, sedi, beni"],
        ["9.9", "Terze parti", "ricorso a fornitori di beni o servizi"],
        ["9.11", "Employer", "rapporti di lavoro, salute e sicurezza"],
        ["9.13", "Processo", "errori di processo o di esecuzione delle operazioni"]],
        [0.7, 1.7, 3.9], 11.5, 0.47)
    tabella(s, 7.05, 1.35, 5.78, ["Codice", "Driver", "Esempi"], [
        ["DRV-CYB", "Cyber", "nuove tecniche d'attacco, vulnerabilità diffuse"],
        ["DRV-TPR", "Terze parti", "fornitori ICT critici, concentrazione cloud"],
        ["DRV-AI", "Intelligenza artificiale", "uso interno di modelli, attacchi abilitati da AI"],
        ["DRV-CLI", "Clima e ambiente", "eventi fisici estremi, transizione"],
        ["DRV-CON", "Condotta", "pratiche commerciali, greenwashing"],
        ["DRV-GEO", "Geopolitica", "sanzioni, conflitti, catene di fornitura"],
        ["DRV-TEC", "Debito tecnologico", "sistemi obsoleti, migrazioni"],
        ["DRV-HUM", "Capitale umano", "carenza di competenze, turnover in ruoli chiave"]],
        [1.0, 1.6, 3.2], 11.5, 0.47)
    nota_fondo(s, ["Fuori perimetro: rischi finanziari, strategici, reputazionali e di modello; frode, financial crime, non "
                   "conformità, qualità dei dati; eventi già accaduti; temi senza un meccanismo di perdita per la Banca. "
                   "Una categoria senza rischi resta visibile: le notizie che la toccano escono come AREA NON PRESIDIATA."], 6.15, 11.5)

    s = nuova("Scale: Impatto e Probabilità (interi 1-5, 5 = peggiore)")
    f1 = tabella(s, 0.5, 1.35, 6.05, ["", "Impatto — danno se il rischio si manifesta"], [
        ["1", "**Limitato**: nessun cliente toccato, perdita entro 100.000 euro"],
        ["2", "**Contenuto**: alcuni clienti coinvolti o un processo non critico"],
        ["3", "**Rilevante**: processo critico degradato o clientela ampia; segnalazione all'autorità possibile"],
        ["4", "**Grave**: interruzione di una funzione essenziale, sanzione probabile, danno reputazionale nazionale"],
        ["5", "**Critico**: effetto sulla continuità, sui prodotti o sulla solvibilità; intervento dell'autorità certo"]],
        [0.4, 5.6], 12, 0.5)
    f2 = tabella(s, 6.78, 1.35, 6.05, ["", "Probabilità — plausibilità per la Banca"], [
        ["1", "**Remota**: solo teorica, nessun caso noto"],
        ["2", "**Improbabile**: casi solo in altri settori o paesi, o Banca poco esposta"],
        ["3", "**Possibile**: casi in altri settori anche non finanziari; Banca esposta"],
        ["4", "**Probabile**: casi ripetuti in aziende strutturate o banche, o crescita documentata"],
        ["5", "**Quasi certa**: già osservato presso istituzioni finanziarie o la loro filiera, o rapida accelerazione"]],
        [0.4, 5.6], 12, 0.5)
    nota_fondo(s, ["Si sceglie il descrittore più aderente all'evidenza, senza interpolare. Probabilità 4-5 solo se la Banca è esposta "
                   "(ha il prodotto, il processo o il fornitore colpito); se le fonti non lo dicono, A1 scrive «esposizione non nota» "
                   "con il valore «se esposta» e decide il validatore (D-81). Probabilità dice QUANTO è plausibile; Prossimità dice QUANDO. "
                   "Probabilità × Impatto è la magnitudo (area nel radar, cella della matrice)."], max(f1, f2) + 0.2, 12)

    s = nuova("Scale: Prossimità (con la Z di Mosca) e Velocity")
    f1 = tabella(s, 0.5, 1.35, 6.05, ["", "Prossimità — quando può manifestarsi", "Z (anni)"], [
        ["1", "Oltre 3 anni", "3,5"],
        ["2", "Da 2 a 3 anni", "2,5"],
        ["3", "Da 1 a 2 anni", "1,5"],
        ["4", "Entro 12 mesi", "0,5"],
        ["5", "In corso o imminente presso la Banca, la sua filiera o aziende ampie e complesse", "0"]],
        [0.4, 4.6, 1.0], 12, 0.62)
    tabella(s, 6.78, 1.35, 6.05, ["", "Velocity — tempo fra innesco e pieno impatto"], [
        ["1", "Oltre 12 mesi"], ["2", "Da 6 a 12 mesi"], ["3", "Da 3 a 6 mesi"], ["4", "Da 1 a 3 mesi"], ["5", "Entro 1 mese"]],
        [0.4, 5.6], 12, 0.62)
    testo(s, 0.5, f1 + 0.2, 6.05, 6.95 - f1 - 0.2, [
        "- Prossimità 4-5 solo con casi documentati (banca comparabile, fornitore, quasi-incidente) o un evento certo; deroga dichiarata per fenomeni ad altissimo impatto.",
        "- **Z** = metà della fascia (per «oltre 3 anni», 3,5): è il tempo disponibile di Mosca."], 11.5)
    testo(s, 6.78, f1 + 0.2, 6.05, 6.95 - f1 - 0.2, [
        "- Velocity **facoltativa**: vuota (n.v.) se nessuna fonte ne documenta i tempi; da n.v. entra con qualunque valore.",
        "- **Anni di adeguamento** (0-30, facoltativi): durata di un adeguamento strutturale oltre le azioni; li stima A1, li decide il validatore; sono la **Y minima** di Mosca, anche per i rischi in osservazione con Impatto 4-5 (con X + Y > Z entrano nel radar). Non entrano nello score."], 11.5)

    s = nuova("Score composito: pesi e calcolo")
    tabella(s, 0.5, 1.4, 6.0, ["Dimensione", "Peso", "Senza Velocity"], [
        ["Impatto", "0,35", "0,4375"], ["Probabilità", "0,25", "0,3125"], ["Prossimità", "0,20", "0,25"],
        ["Velocity (solo se valutata)", "0,20", "—"], ["Somma (controllata)", "1,00", "1,00"]],
        [3.0, 1.3, 1.7], 13, 0.5)
    testo(s, 6.8, 1.35, 6.0, 5.6, [
        "## La formula",
        "- **Media pesata delle dimensioni valutate**, divisa per la somma dei pesi presenti: senza Velocity i pesi si ridistribuiscono sulle altre tre.",
        "- Due decimali; vuoto se manca Impatto, Probabilità o Prossimità.",
        "## Esempio (NFR-01, 2026Q4-S1)",
        "- (0,35×4 + 0,25×5 + 0,20×5 + 0,20×5) / 1,00 = **4,65**.",
        "## Che cosa NON entra nello score",
        "- Presidio dei controlli, anni di adeguamento, rilevanza delle fonti, attenzione del supervisore, segnali delle notizie.",
        "> Lo calcola il foglio, non un agente; cambiando i pesi si ricalcolano anche i cicli passati, così il trend resta confrontabile."], 12.5)
    testo(s, 0.5, 4.3, 6.0, 2.6, [
        "## Soglie collegate",
        "- **Soglia del Delta**: 0,30 — variazioni di score da guardare.",
        "- **Matrice**: P × I verde fino a 6, ambra 8-12, rosso da 15.",
        "- **Trend**: ultimi 8 cicli, scala fissa."], 12.5)

    s = nuova("Fonti: rilevanza, catalogo e regole di variazione")
    f1 = tabella(s, 0.5, 1.35, 6.05, ["", "Rilevanza di una fonte per la previsione"], [
        ["1 Bassa", "stampa generalista, opinioni"],
        ["2 Media", "fonte specializzata con dati o casi (report di settore, threat intelligence)"],
        ["3 Alta", "regolatore o autorità, consorzio di dati (es. ORX); caso documentato presso la Banca, una banca comparabile o un fornitore"]],
        [1.1, 4.9], 12, 0.7)
    testo(s, 0.5, f1 + 0.2, 6.05, 6.95 - f1 - 0.2, [
        "- Stesso ente o stesso fatto = **una** fonte. Tipo **Interna / Esterna** (dalla cartella).",
        "- Probabilità 5 o Prossimità 4-5: almeno una fonte di rilevanza 3 (salvo deroga).",
        "- **Catalogo** (fonti.md): 36 fonti FNT-nn, 31 approvate, 1 esclusa; canali **rss** (19, raccolta automatica), **sito** (6, letti dall'agente), **share** (7, documenti interni e a licenza); **MAN** per un file portato a mano."], 12)
    f2 = tabella(s, 6.78, 1.35, 6.05, ["#", "Regole di variazione fra cicli"], [
        ["1", "Senza evidenza nuova il valore resta quello del ciclo prima"],
        ["2", "Al massimo un punto per dimensione fra due cicli consecutivi"],
        ["3", "Ogni valore cambiato cita un'evidenza specifica, datata dopo il ciclo prima"],
        ["4", "Una riduzione ha lo stesso onere di un aumento"],
        ["5", "Confermare = riprendere le fonti; se non valgono più, il valore va rivisto"],
        ["6", "La copertura mediatica non è evidenza"],
        ["7", "Un numero nel razionale cita la sua fonte o è dichiarato come stima"]],
        [0.4, 5.6], 11.5, 0.36)
    testo(s, 6.78, f2 + 0.15, 6.05, 6.95 - f2 - 0.15, [
        "**Esito** (lo calcola lo script): **Nuovo** se l'ID non c'era, **Variato** se cambia una dimensione (anche da n.v.), "
        "altrimenti **Confermato**."], 12)

    s = nuova("KRI: due per rischio, con soglie della Banca")
    f1 = tabella(s, 0.5, 1.35, 6.05, ["Regola", "Valore"], [
        ["Quanti", "esattamente due per rischio del radar: uno di **Esposizione** (quanto la Banca è esposta, di solito dato interno) e uno di **Evoluzione** (quanto cresce il fenomeno)"],
        ["Verso", "valore crescente = peggio (altrimenti si usa il complemento)"],
        ["Soglie", "attenzione e allerta, entrambe o nessuna, allerta ≥ attenzione; solo se una fonte le sostiene, altrimenti le fissa il validatore"],
        ["Valore", "rilevato nel ciclo; vuoto se non rilevato, mai ricopiato"],
        ["Codice", "NFR-nn-Kn stabile, mai riusato; In_Vigore = No chiude la serie"]],
        [1.1, 4.95], 11.5, 0.4)
    f2 = tabella(s, 6.78, 1.35, 6.05, ["Stato", "Quando"], [
        ["Verde", "valore sotto la soglia di attenzione"],
        ["Ambra", "da attenzione fino all'allerta"],
        ["Rosso", "da allerta in su"],
        ["Non rilevato", "nessun valore nel ciclo (mai zero)"],
        ["Soglie da definire", "manca una soglia: non conta nel semaforo"],
        ["SOGLIE INVERTITE", "allerta < attenzione: va rettificato"]],
        [1.9, 4.15], 12, 0.55, colori_prima={0: VERDE_CH, 1: AMBRA_CH, 2: ROSSO_CH})
    nota_fondo(s, ["Un KRI in rosso è il segnale forte F2 del livello di allerta; metà o più dei KRI in ambra o rosso è l'aggravante A3. "
                   "Tendenza: ▲ peggiora sul ciclo prima."], max(f1, f2) + 0.2, 12)

    s = nuova("Notizie: intensità, direzione, segnali e segnale critico")
    f1 = tabella(s, 0.5, 1.35, 6.05, ["Intensità", "Su ciascuna delle nove categorie"], [
        ["0", "nessun nesso con la categoria"],
        ["1", "nesso indiretto o solo teorico per la Banca"],
        ["2", "nesso diretto con un meccanismo di perdita presente nella Banca"],
        ["3", "nesso diretto e immediato: evento in corso su controparti, sistemi o prodotti che la Banca usa"]],
        [1.1, 4.95], 12, 0.4)
    testo(s, 0.5, f1 + 0.2, 6.05, 6.95 - f1 - 0.2, [
        "- **Direzione**: solo quattro valori — Peggiorativo, Migliorativo, Neutro, Non rilevante (un sinonimo vale zero).",
        "- **Segnali netti** = peggiorative − migliorative sullo stesso rischio in **4 settimane**; **soglia di direzione 2** → ▲ ▼ e flag proposto, mai una variazione di score.",
        "- Una notizia ripresa da più testate è un solo segnale."], 12)
    scheda(s, 6.78, 1.35, 6.05, 3.0, "Segnale critico → proposta di ciclo straordinario", [
        "Solo se una notizia:",
        "- **(a)** ha intensità 3 su un rischio con Impatto 4 o 5;",
        "- **(b)** riporta un evento in corso presso la Banca, un suo fornitore o una controparte;",
        "- **(c)** mostra una capacità nuova di una minaccia che rende superato il descrittore vigente.",
        "Propongono A3 o A4; **decide la funzione di rischio**."], ROSSO, 12)
    scheda(s, 6.78, 4.5, 6.05, 2.4, "Presidio delle categorie (foglio Notizie)", [
        "- **AREA NON PRESIDIATA**: notizie di intensità ≥ 2 su una categoria senza rischi nel registro.",
        "- **In osservazione**: nessun rischio nel radar, almeno uno in osservazione.",
        "- **Presidiata**: almeno un rischio nel radar."], NAVY, 12)

    s = nuova("Urgenza: i parametri della disuguaglianza di Mosca")
    f1 = tabella(s, 0.5, 1.35, 6.4, ["Termine", "Significato", "Valore"], [
        ["X", "anni per decidere e avviare la mitigazione", "0,25 (un trimestre)"],
        ["Y", "anni per completare la mitigazione: il maggiore fra l'orizzonte più lungo delle azioni del ciclo non respinte (3/6/12 mesi) e gli anni di adeguamento", "dalle azioni; 1 se non ce ne sono"],
        ["Z", "anni prima che il rischio possa manifestarsi: metà della fascia della Prossimità", "5→0 · 4→0,5 · 3→1,5 · 2→2,5 · 1→3,5"],
        ["Margine", "Z − (X + Y)", "anni"]],
        [0.9, 3.7, 1.8], 12, 0.4)
    f2 = tabella(s, 7.15, 1.35, 5.68, ["Margine", "Urgenza"], [
        ["sotto 0", "AGIRE ORA (X + Y > Z): già in ritardo"],
        ["da 0 a 1 anno", "Agire entro l'anno"],
        ["oltre 1 anno", "Margine n anni"]],
        [1.8, 3.9], 12.5, 0.55, colori_prima={0: ROSSO_CH, 1: AMBRA_CH, 2: GRIGIO_CH})
    testo(s, 7.15, f2 + 0.15, 5.68, 6.95 - f2 - 0.15, [
        "## Nel radar",
        "- Raggio = margine + 2, fra 0,6 e 6; angolo = settore della categoria; area = P × I.",
        "- Cerchi: margine 0 (Mosca, rosso), 1, 2, 3 anni.",
        "> Origine: Michele Mosca (2015), per la migrazione alla crittografia post-quantistica. L'urgenza è una graduatoria ancorata, non una previsione."], 12)
    testo(s, 0.5, f1 + 0.2, 6.4, 6.95 - f1 - 0.2, [
        "- Cambiando le tabelle di Parametri cambia il radar, senza toccare nessun prompt.",
        "- Gli anni di adeguamento validati vivono nell'anagrafica (colonna M) e restano correggibili a mano."], 12)

    s = nuova("Livello di allerta: segnali forti, aggravanti e soglie")
    f1 = tabella(s, 0.5, 1.35, 6.05, ["", "Segnali forti", ""], [
        ["F1", "AGIRE ORA (X + Y > Z)", ""],
        ["F2", "almeno un KRI in rosso", ""],
        ["F3", "segnali netti delle notizie oltre la soglia (▲)", "soglia 2"]],
        [0.6, 4.2, 1.2], 12, 0.5)
    tabella(s, 0.5, f1 + 0.25, 6.05, ["", "Aggravanti (contano solo con margine stretto)", "Soglia"], [
        ["A1", "Probabilità × Impatto alta", "da 12"],
        ["A2", "presidio dei controlli «Assente»", "—"],
        ["A3", "quota dei KRI in vigore in ambra o rosso", "da 0,5"],
        ["—", "margine stretto", "sotto 1 anno"]],
        [0.6, 4.2, 1.2], 12, 0.5)
    f2 = tabella(s, 6.78, 1.35, 6.05, ["Livello", "Regola", "Conseguenza"], [
        ["Attivazione", "2 o più segnali forti (soglia 2)", "al comitato con un piano; se nasce nel settimanale, proposta di straordinario"],
        ["Attenzione", "1 segnale forte, oppure margine sotto 1 anno con almeno un aggravante", "l'owner conferma azioni e KRI; il rischio va nel report"],
        ["Monitoraggio", "gli altri casi", "ciclo trimestrale ordinario"]],
        [1.4, 2.4, 2.25], 12, 0.5, colori_prima={0: ROSSO_CH, 1: AMBRA_CH, 2: VERDE_CH})
    testo(s, 6.78, f2 + 0.2, 6.05, 6.95 - f2 - 0.2, [
        "> Il livello non è uno score: i segnali netti cambiano ogni settimana e con loro il livello. Il tema del "
        "supervisore non alza il livello (è evidenziato con ◆)."], 12)

    s = nuova("Registro unico: chi entra nel radar di Mosca")
    f1 = tabella(s, 0.5, 1.35, 6.05, ["Parametro", "Valore"], [
        ["Capienza del radar", "15 (fino a 20)"],
        ["Soglia di Prossimità per entrare", "2 (entro 3 anni)"],
        ["Soglia di Probabilità per entrare", "2 (non solo teorico)"],
        ["Cicli sotto la soglia prima di uscire", "2 (un ciclo «Radar, in uscita»)"],
        ["Rischi in osservazione al massimo", "20 (fisso)"]],
        [3.6, 2.45], 12.5, 0.55)
    testo(s, 6.78, 1.3, 6.05, 5.7, [
        "## La regola (la applica lo script a ogni ciclo e validazione)",
        "- **Candidato**: Prossimità ≥ 2 e Probabilità ≥ 2.",
        "- **Sempre nel radar**, anche sotto le soglie: un rischio AGIRE ORA, un tema del supervisore, un rischio che il validatore mette nel radar con la motivazione.",
        "- **Mai nel radar**: un rischio che il validatore mette in osservazione con la motivazione.",
        "- Oltre la capienza: prima le eccezioni, poi i più alti per score; gli altri restano «fuori per capienza».",
        "- Chi scende resta un ciclo **«Radar, in uscita»**; chi sale entra subito, con KRI e azioni.",
        "## Rischi in osservazione",
        "- Stessa scala e stesso score; **niente KRI né azioni**; rivalutati solo con evidenza nuova; con **orizzonte** in anni e **segnale di avvicinamento**."], 12)
    testo(s, 0.5, f1 + 0.25, 6.05, 6.95 - f1 - 0.25, [
        "- **Stato**: Radar · Radar, in uscita · Osservazione · Dismesso.",
        "- **Movimento**: Entrato · Uscito.",
        "- **Decisione sullo stato** del validatore: Automatico (vale la regola) · Radar · Osservazione."], 12)

    s = nuova("Altri parametri e limiti")
    tabella(s, 0.5, 1.35, 12.33, ["Parametro", "Valore", "Dove / chi"], [
        ["Presidio dei controlli", "Non verificato · Adeguato · Parziale · Assente", "anagrafica, colonna O · validatore"],
        ["Attenzione del supervisore", "Sì · No (◆)", "anagrafica, colonna N · validatore"],
        ["Owner", "una funzione, mai una persona; «Da assegnare» per un rischio nuovo", "anagrafica, colonna L · validatore"],
        ["Azioni per rischio · orizzonte", "da 1 a 5 · 3, 6 o 12 mesi", "A1 propone, il validatore delibera"],
        ["Anni di adeguamento", "0-30 o vuoto", "A1 stima, il validatore decide (colonna M)"],
        ["Motivazione in validazione", "almeno 10 caratteri, con l'evidenza", "obbligatoria se cambia un valore, se si respinge, con ⚑"],
        ["Note dell'analista", "al massimo 600 caratteri (sintesi di A7)", "Quadro, cella gialla"],
        ["Finestre temporali", "segnali netti e notizie: 4 settimane · grafici settimanali: 12 settimane · trend: 8 cicli", "Parametri, scritte dallo script"],
        ["Capienza dell'anagrafica", "20 posti (Rischio_01..20 per A7)", "struttura del template"],
        ["Istruzioni degli agenti", "al massimo 8.000 caratteri ciascuna", "limite di Agent Builder, controllato da componi_prompt.py"]],
        [2.8, 5.3, 4.2], 12, 0.47)

    s = nuova("Valori ammessi e codici")
    tabella(s, 0.5, 1.35, 6.3, ["Campo", "Valori"], [
        ["Esito del rischio", "Nuovo · Variato · Confermato"],
        ["Validazione", "Da validare · Accettato · Rettificato · Respinto (in riapertura: Invariato, … (riapertura))"],
        ["Stato di un'azione", "Proposta · Approvata · Respinta · In corso · Completata · Annullata"],
        ["Stato di un KRI", "Verde · Ambra · Rosso · Non rilevato · Soglie da definire"],
        ["Tipo di KRI · di fonte", "Esposizione · Evoluzione · Interna · Esterna"],
        ["Segnale di una notizia", "Peggiorativo · Migliorativo · Neutro · Non rilevante"],
        ["Livello di allerta", "Attivazione · Attenzione · Monitoraggio"],
        ["Urgenza", "AGIRE ORA · Agire entro l'anno · Margine n anni"]],
        [1.9, 4.4], 11.5, 0.6)
    tabella(s, 7.05, 1.35, 5.78, ["Oggetto", "Codice"], [
        ["Ciclo trimestrale", "AAAAQn (2026Q3)"],
        ["Ciclo straordinario", "AAAAQn-Sk (2026Q4-S1)"],
        ["Settimana", "il suo lunedì, AAAA-MM-GG"],
        ["Rischio", "NFR-nn"],
        ["KRI", "NFR-nn-Kn"],
        ["Azione", "<ciclo>-NFR-nn-An"],
        ["Notizia", "N-AAAAMMGG-nn"],
        ["Fonte", "FNT-nn · MAN"],
        ["Categoria · driver", "9.x · DRV-xxx"]],
        [2.2, 3.6], 11.5, 0.53)


def chiusura():
    s = nuova("In sintesi", "Conclusione")
    righe = [("Gli agenti giudicano e scrivono", "A1 e A8 valutano, A3 legge le notizie, A4 dà pareri, A5 controlla, A6 e A7 spiegano. Nessuno calcola né decide."),
             ("Python controlla e scrive", "un comando per ogni passo; si ferma prima di sbagliare, mai dopo; tiene backup e archivio."),
             ("Excel calcola e mostra", "score, urgenza di Mosca, livello di allerta, stato dei KRI; report al comitato e viste di lavoro."),
             ("Le persone decidono", "il validatore accetta, rettifica o respinge ogni proposta; la funzione di rischio fissa i parametri."),
             ("Ogni numero ha un perché", "nelle schede: valori validati, documenti, cambiamenti, azioni e KRI con l'effetto atteso.")]
    for i, (t, d) in enumerate(righe):
        y = 1.45 + i * 1.08
        rettangolo(s, 0.5, y, 0.18, 0.85, [NAVY, BLU, VERDE, AMBRA, ROSSO][i])
        testo(s, 0.85, y - 0.03, 11.8, 0.45, [f"**{t}**"], 18, NAVY)
        testo(s, 0.85, y + 0.4, 11.8, 0.5, [d], 14, GRIGIO)


def costruisci():
    copertina()
    indice()
    divisore("1", "Il sistema in breve", ["Componenti, ritmi di lavoro, principi."])
    sistema()
    principi()
    chi_fa_cosa()
    divisore("2", "Gli agenti e i passi Python", ["Che cosa fa ogni agente quando viene usato,",
                                                   "e quale comando Python lo precede e lo segue."])
    agenti_panoramica()
    agenti()
    comandi()
    app_nfr()
    cicli()
    schede_rischi()
    report_intro()
    report_fogli()
    parametri()
    chiusura()
    prs.save(USCITA)
    print(f"Presentazione: {USCITA} ({stato['n']} slide)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--solo-pptx", action="store_true", help="riusa le immagini gia' preparate")
    if not ap.parse_args().solo_pptx:
        prepara_immagini()
    costruisci()
