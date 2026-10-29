"""Schede dei rischi: file a parte rispetto al registro (DECISIONI D-40).

Tre passi, tutti dopo la validazione del ciclo trimestrale:
1. input_schede(): dal registro validato, un JSON per rischio con cio' che la
   scheda mostra (valori validati, urgenza, livello di allerta, variazioni,
   fonti, azioni e KRI con il tipo, storia; dal 02/10/2026, D-68, il razionale
   di A1 diviso in fenomeno e motivazioni per valore, e il perche' e l'effetto
   atteso di azioni e KRI). E' l'unico allegato dell'agente A6 (comando schede-input).
2. controlla_a6(): l'output di A6 (testi estesi) rispetta l'input? Stessi
   rischi, quattro dimensioni, le stesse fonti copiate identiche, nessun numero che
   non stia nei dati del rischio o nella rubrica. Se no, niente file.
3. scrivi_file_schede(): Schede_<ciclo>.xlsx, un indice e una scheda per
   rischio nell'ordine della classifica (comando schede). Senza l'output di A6
   le schede usano i testi di A1 dal registro.

Il passo 1 legge il workbook via Excel (COM); i passi 2 e 3 sono Python puro
(openpyxl), quindi il collaudo li prova anche senza Excel.
"""

import json
import re
from pathlib import Path

# Valori spiegati nelle schede (D-49, D-50, D-57, D-58): le quattro dimensioni dello score (la Velocity puo'
# mancare). La Confidence non c'e' piu': ogni fonte ha la sua rilevanza 1-3. Il presidio dei controlli e'
# un giudizio qualitativo della Banca, mostrato a parte.
DIMENSIONI = ["impatto", "probabilita", "prossimita", "velocity"]
NOMI_DIM = {"impatto": "Impatto", "probabilita": "Probabilita", "prossimita": "Prossimita", "velocity": "Velocity"}
COLONNA_DIM = dict(NOMI_DIM)
# Etichette anche per le variazioni dei cicli scritti prima del 30/09 sera (severity, confidence).
ETICHETTE_VARIAZIONI = dict(NOMI_DIM, severity="Impatto", confidence="Confidence")
PESI = {"impatto": "pesoImpatto", "probabilita": "pesoProbabilita", "prossimita": "pesoProssimita",
        "velocity": "pesoVelocity"}
RILEVANZA = {1: "Bassa", 2: "Media", 3: "Alta"}
NON_VALUTATA = "n.v."
STATI_CHIUSI = {"Completata", "Respinta", "Annullata"}
NOME_ANALISI = "Horizon Scan NFR"
RISERVATEZZA = "Uso interno"


def _num(v, decimali=2):
    """Numero dal workbook, arrotondato (Excel restituisce 3.8000000000000003)."""
    if v is None or v == "":
        return None
    try:
        f = round(float(v), decimali)
    except (TypeError, ValueError):
        return v
    return int(f) if f == int(f) else f


def _colonna(sessione, nome):
    v = sessione.intervallo_nome(nome).Value
    return [r[0] for r in v] if isinstance(v, tuple) else [v]


# ---------------------------------------------------------------- razionali in parti (D-68)
# A1 scrive il razionale del rischio in parti con etichette fisse, come MESSAGGIO/DECISIONI della sintesi:
# "RISCHIO: il fenomeno in termini di rischio. IMPATTO: perche' il valore, fatto e documento. PROBABILITA: ...
# PROSSIMITA: ... VELOCITY: ... ADEGUAMENTO: ...", e il razionale di azioni e KRI come "perche'. EFFETTO: ...".
# Un'etichetta conta all'inizio del testo o dopo un punto, un punto e virgola o un a capo (maiuscole o minuscole),
# oppure ovunque se e' scritta tutta in maiuscolo: "Impatto resta 4:" a meta' frase non e' un'etichetta. I razionali
# scritti prima (senza etichette) restano interi.
ETICHETTE_RAZIONALE = {"RISCHIO": "rischio", "IMPATTO": "impatto", "PROBABILITA": "probabilita",
                       "PROSSIMITA": "prossimita", "VELOCITY": "velocity", "ADEGUAMENTO": "adeguamento"}
_NOMI = r"RISCHIO|IMPATTO|PROBABILIT(?:A'|À|A)|PROSSIMIT(?:A'|À|A)|VELOCITY|ADEGUAMENTO"
_ETICHETTA = re.compile(rf"(?:(?:^|(?<=[.;\n]))\s*(?i:({_NOMI}))|(?<=\s)({_NOMI}))\s*:")
_EFFETTO = re.compile(r"(?:(?:^|(?<=[.;\n]))\s*(?i:EFFETTO(?: ATTESO)?)|(?<=\s)EFFETTO(?: ATTESO)?)\s*:")


def _pulito(testo):
    """Parte di un razionale senza spazi e punti e virgola ai bordi, con la maiuscola iniziale ("RISCHIO: le frodi"
    diventa "Le frodi"): nella scheda ogni parte sta da sola, in una cella o in una sezione."""
    t = str(testo or "").strip().strip(";").strip()
    return (t[0].upper() + t[1:]) if t else None


def parti_razionale(testo):
    """Parti del razionale di un rischio: {"rischio", "impatto", "probabilita", "prossimita", "velocity",
    "adeguamento", "altro" (testo prima della prima etichetta)}; {} se il razionale non ha etichette."""
    testo = str(testo or "")
    trovate = list(_ETICHETTA.finditer(testo))
    if not trovate:
        return {}
    parti = {}
    if _pulito(testo[:trovate[0].start()]):
        parti["altro"] = _pulito(testo[:trovate[0].start()])
    for i, m in enumerate(trovate):
        chiave = ETICHETTE_RAZIONALE[(m.group(1) or m.group(2)).upper().replace("À", "A").replace("'", "")]
        fine = trovate[i + 1].start() if i + 1 < len(trovate) else len(testo)
        valore = _pulito(testo[m.end():fine])
        if valore:
            parti[chiave] = (parti[chiave] + " " + valore) if chiave in parti else valore
    return parti


def parti_effetto(testo):
    """(perche', effetto atteso) dal razionale di un'azione o di un KRI; effetto None se manca "EFFETTO:"."""
    testo = str(testo or "")
    m = _EFFETTO.search(testo)
    if not m:
        return _pulito(testo), None
    return _pulito(testo[:m.start()]), _pulito(testo[m.end():])


# ================================================================ 1. input di A6
def input_schede(scr, sessione, ciclo, ciclo_prec, data_rif):
    """Dati di ogni rischio del ciclo, dal registro VALIDATO e dalle viste (urgenza, segnali)."""
    _, anag = scr.leggi_anagrafica()
    posizione = {r["ID"]: i for i, (_, r) in enumerate(anag)}
    reg = scr.righe("tblRegister")
    del_ciclo = {r["ID_Rischio"]: r for r in reg if r["Ciclo"] == ciclo}
    cicli = sorted({r["Ciclo"] for r in reg if r["Ciclo"] <= ciclo})[-8:]
    azioni = scr.righe("tblAzioni")
    kri = scr.righe("tblKRI")
    kri_prec = {r["KRI_ID"]: r for r in kri if r["Ciclo"] == ciclo_prec}
    viste = {n: _colonna(sessione, n) for n in ("rmID", "rmX", "rmY", "rmZ", "rmMargine", "rmUrgenza", "rmLivello",
                                                  "rmMotivoLivello", "posNetti")}
    et_codici = _colonna(sessione, "codiciL1")
    et_nomi = [r[0] for r in sessione.intervallo_nome("codiciL1").GetOffset(0, 1).Value]
    nome_et = dict(zip(et_codici, et_nomi))
    pesi = {d: _num(sessione.intervallo_nome(n).Value) for d, n in PESI.items()}
    rischi = []
    for _, a in anag:
        rid = a["ID"]
        r = del_ciclo.get(rid)
        if r is None:
            continue                           # non valutato: nessuna scheda (il controllo di copertura lo segnala)
        i = posizione[rid]
        prec = next((x for x in reg if x["Ciclo"] == ciclo_prec and x["ID_Rischio"] == rid), None)
        parti = parti_razionale(r["Razionale"])
        rischi.append({
            "id_rischio": rid, "nome": r["Nome"], "event_type": r["Event_Type"],
            "event_type_nome": nome_et.get(r["Event_Type"]), "driver": r["Driver"], "owner": a["Owner"],
            "attenzione_supervisore": a.get("Attenzione_Supervisore") == "Sì",
            "anni_adeguamento": _num(a.get("Anni_Adeguamento")), "esito": r["Esito"],
            "presidio_controlli": a.get("Presidio_Controlli") or "Non verificato",
            "descrizione": r["Descrizione"],
            # D-68: il fenomeno in termini di rischio e, a parte, il perche' di ogni valore (null se A1 non ha diviso
            # il razionale: allora razionale_rischio e' il razionale intero).
            "razionale_rischio": (" ".join(p for p in (parti.get("altro"), parti.get("rischio")) if p) or None)
                                 if parti else r["Razionale"],
            "motivazioni": {k: parti.get(k) for k in DIMENSIONI + ["adeguamento"]} if parti else None,
            "valori": {d: _num(r[COLONNA_DIM[d]]) for d in DIMENSIONI},
            "score": _num(r["Score_Composito"]),
            "score_ciclo_precedente": _num(prec["Score_Composito"]) if prec else None,
            "urgenza": {"esito": viste["rmUrgenza"][i], "x": _num(viste["rmX"][i]), "y": _num(viste["rmY"][i]),
                        "z": _num(viste["rmZ"][i]), "margine": _num(viste["rmMargine"][i])},
            # D-63: livello di allerta calcolato dal foglio Radar_Mosca (colore del radar).
            "livello_allerta": {"livello": viste["rmLivello"][i] or None, "motivo": viste["rmMotivoLivello"][i] or None},
            "segnali_netti_4_settimane": _num(viste["posNetti"][i]),
            "variazioni": json.loads(r["Variazioni"] or "[]"),
            "fonti": json.loads(r["Fonti"] or "[]"),
            "azioni": [{"id_azione": x["ID_Azione"], "descrizione": x["Descrizione"], "owner": x["Owner"],
                        "orizzonte": x["Orizzonte"], "stato": x["Stato"], "ciclo": x["Ciclo"],
                        **dict(zip(("razionale", "effetto_atteso"), parti_effetto(x.get("Razionale"))))}
                       for x in azioni if x["ID_Rischio"] == rid and x.get("ID_Azione")
                       and (x["Ciclo"] == ciclo or x["Stato"] not in STATI_CHIUSI)],
            "kri": [{"kri_id": k["KRI_ID"], "tipo": k.get("Tipo") or None, "nome": k["Nome"], "unita": k["Unita"],
                     **dict(zip(("razionale", "effetto_atteso"), parti_effetto(k.get("Razionale")))),
                     "valore": _num(k["Valore_Corrente"]),
                     "valore_ciclo_precedente": _num((kri_prec.get(k["KRI_ID"]) or {}).get("Valore_Corrente")),
                     "soglia_attenzione": _num(k["Soglia_Attenzione"]), "soglia_allerta": _num(k["Soglia_Allerta"]),
                     "stato": k["Stato"], "fonte_dati": k["Fonte_Dati"]}
                    for k in kri if k["Ciclo"] == ciclo and k["ID_Rischio"] == rid and k["In_Vigore"] == "Sì"],
            "storia": [{"ciclo": c, **{d: _num(x[COLONNA_DIM[d]]) for d in DIMENSIONI}, "score": _num(x["Score_Composito"])}
                       for c in cicli for x in reg if x["Ciclo"] == c and x["ID_Rischio"] == rid],
        })
    rischi.sort(key=lambda x: (-(x["score"] or 0), posizione[x["id_rischio"]]))
    for n, x in enumerate(rischi, start=1):
        x["posizione_in_classifica"] = n
    return {"ciclo": ciclo, "data_riferimento": data_rif, "ciclo_precedente": ciclo_prec, "pesi": pesi,
            "rischi": rischi}


# ================================================================ 2. controlli sull'output di A6
NUMERO = re.compile(r"\d+(?:[.,]\d+)?")


def numeri(testo):
    """Numeri di un testo come valori (4,2 = 4.2; 06 = 6): cosi' 10/06/2026 e 2026-06-10 coincidono."""
    return {float(n.replace(",", ".")) for n in NUMERO.findall(str(testo))}


def controlla_a6(ingresso, uscita, testo_rubrica=""):
    """Problemi dell'output di A6 rispetto al suo input (lista vuota = va bene). Ogni problema
    ferma il comando: A6 scrive testi, non cambia rischi, valori o fonti."""
    problemi = []
    if uscita.get("ciclo") != ingresso["ciclo"]:
        problemi.append(f"Ciclo {uscita.get('ciclo')} diverso da quello dell'input ({ingresso['ciclo']}).")
    attesi = {r["id_rischio"]: r for r in ingresso["rischi"]}
    visti = [s["id_rischio"] for s in uscita.get("schede", [])]
    doppi = sorted({i for i in visti if visti.count(i) > 1})
    if doppi:
        problemi.append(f"Schede ripetute: {doppi}.")
    if set(visti) - set(attesi):
        problemi.append(f"Schede di rischi non nell'input: {sorted(set(visti) - set(attesi))}.")
    if set(attesi) - set(visti):
        problemi.append(f"Rischi senza scheda: {sorted(set(attesi) - set(visti))}.")
    ammessi_sempre = {1.0, 2.0, 3.0, 4.0, 5.0} | numeri(testo_rubrica) | numeri(ingresso["ciclo"]) | \
        numeri(ingresso["data_riferimento"])
    for s in uscita.get("schede", []):
        rid = s["id_rischio"]
        r = attesi.get(rid)
        if r is None:
            continue
        dims = [d["dimensione"] for d in s["dimensioni"]]
        if sorted(dims) != sorted(DIMENSIONI):
            problemi.append(f"{rid}: servono le quattro dimensioni una volta ciascuna, trovate {dims}.")
        rif_attesi = [f["riferimento"] for f in r["fonti"]]
        rif_visti = [f["riferimento"] for f in s["fonti"]]
        if sorted(rif_visti) != sorted(rif_attesi):
            mancanti = sorted(set(rif_attesi) - set(rif_visti))
            nuove = sorted(set(rif_visti) - set(rif_attesi))
            problemi.append(f"{rid}: le fonti devono essere quelle dell'input, copiate identiche, una volta ciascuna"
                            + (f"; mancano {mancanti}" if mancanti else "") + (f"; non nell'input {nuove}" if nuove else "")
                            + ("; ripetute" if len(rif_visti) != len(set(rif_visti)) else "") + ".")
        ammessi = ammessi_sempre | numeri(json.dumps(r, ensure_ascii=False))
        testi = [("descrizione_estesa", s["descrizione_estesa"])] + \
                [(d["dimensione"], d["spiegazione"]) for d in s["dimensioni"]] + \
                [(f"fonte '{f['riferimento'][:30]}'", f["a_sostegno_di"]) for f in s["fonti"]]
        for dove, testo in testi:
            estranei = sorted(n for n in numeri(testo) if n not in ammessi)
            if estranei:
                valori = ", ".join(f"{n:g}".replace(".", ",") for n in estranei)
                problemi.append(f"{rid}, {dove}: numeri che non stanno nei dati del rischio ({valori}). A6 non puo' "
                                "introdurre dati: correggere il testo o rilanciare A6.")
    return problemi


# D-88: lunghezze dei testi di A6 (caratteri, spazi compresi): (campo, minimo, consigliato fino a). Le stesse sono nelle
# istruzioni di A6 (02_Agenti/modelli/A6.md, regole 4-6); il collaudo le confronta. Piu' lunghe di quelle di A1: la
# scheda spiega a chi non conosce il tema. Un testo piu' corto non ferma le schede: il comando lo dice e l'app propone
# di farlo allungare ad A6 nella stessa chat.
LUNGHEZZE_A6 = [("descrizione_estesa", 1000, 2000), ("spiegazione", 400, 800), ("a_sostegno_di", 100, 300)]


def testi_corti_a6(ingresso, uscita):
    """I testi dell'output di A6 sotto il minimo di LUNGHEZZE_A6: ["NFR-01, descrizione_estesa 640 caratteri (minimo
    1000)", ...]. La spiegazione di una Velocity non valutata puo' essere breve."""
    minimo = {campo: m for campo, m, _ in LUNGHEZZE_A6}
    valori = {r["id_rischio"]: r.get("valori") or {} for r in ingresso.get("rischi", [])}
    corti = []

    def misura(rid, dove, testo, campo):
        n = len(str(testo or "").strip())
        if n < minimo[campo]:
            corti.append(f"{rid}, {dove} {n} caratteri (minimo {minimo[campo]})")
    for s in uscita.get("schede", []):
        rid = s.get("id_rischio")
        misura(rid, "descrizione_estesa", s.get("descrizione_estesa"), "descrizione_estesa")
        for d in s.get("dimensioni", []):
            if not (d.get("dimensione") == "velocity" and valori.get(rid, {}).get("velocity") is None):
                misura(rid, f"spiegazione di {d.get('dimensione')}", d.get("spiegazione"), "spiegazione")
        for f in s.get("fonti", []):
            misura(rid, f"a_sostegno_di della fonte '{str(f.get('riferimento'))[:30]}'", f.get("a_sostegno_di"),
                   "a_sostegno_di")
    return corti


# ================================================================ 3. file delle schede
BLU_NOTTE, BLU_MEDIO, BLU_CHIARO, BLU_DIM = "1F3864", "2F5597", "DEEAF6", "2E75B6"
ROSSO, AMBRA, GRIGIO = "C00000", "FFC000", "A6A6A6"
TESTO, TESTO_TENUE = "262626", "595959"
ROSSO_TENUE, AMBRA_TENUE = "FBE5E5", "FFF2CC"
# Livello di allerta (D-63): gli stessi colori del radar.
COLORI_LIVELLO = {"Attivazione": ("E06666", "FFFFFF"), "Attenzione": ("F6C342", TESTO), "Monitoraggio": ("93C47D", TESTO)}
FONDI_LIVELLO = {"Attivazione": ROSSO_TENUE, "Attenzione": AMBRA_TENUE, "Monitoraggio": "E2EFDA"}
SCALA = {1: "F2F2F2", 2: "FCE4D6", 3: "F8CBAD", 4: "F4B183", 5: "C55A11"}
SCALA_RIL = {1: "F2F2F2", 2: "DEEBF7", 3: "9DC3E6"}
SEMAFORO = {"Rosso": (ROSSO, "FFFFFF"), "Ambra": (AMBRA, TESTO), "Verde": ("70AD47", "FFFFFF"),
            "Non rilevato": ("D9D9D9", TESTO), "Soglie da definire": (AMBRA_TENUE, TESTO),
            "SOGLIE INVERTITE — correggere": (ROSSO, "FFFFFF")}
LEGENDA = ("Imp = Impatto, Prob = Probabilita, Pross = Prossimita, Vel = Velocity (1-5, 5 = peggiore; n.v. = non valutata, "
           "nessuna fonte ne documenta i tempi). Rilevanza di una fonte per la previsione: Alta = fonte primaria o caso "
           "documentato, Media = fonte specializzata, Bassa = generalista. Presidio = "
           "giudizio della Banca sui controlli, fuori dallo score. X = anni per decidere e avviare la mitigazione; "
           "Y = anni per mitigare (azioni o adeguamento strutturale); Z = anni prima che il rischio possa manifestarsi; "
           "AGIRE ORA = X + Y > Z. Livello di allerta (calcolato dal registro): Attivazione = almeno due fra AGIRE ORA, "
           "KRI in rosso, notizie in peggioramento; Attenzione = uno di questi, o margine sotto l'anno con P x I alto, "
           "presidio assente o KRI in ambra; Monitoraggio = gli altri. KRI = indicatore di rischio, di esposizione (quanto "
           "la Banca e' esposta) o di evoluzione (quanto cresce il fenomeno). ◆ = tema all'attenzione del supervisore.")


def _it(v, decimali=2):
    if v is None or v == "":
        return "—"
    if isinstance(v, (int, float)):
        return f"{v:.{decimali}f}".replace(".", ",")
    return str(v)


def _data_it(s):
    s = str(s or "")
    return f"{s[8:10]}/{s[5:7]}/{s[:4]}" if re.match(r"\d{4}-\d{2}-\d{2}", s) else s


def _urgenza_breve(testo):
    t = str(testo or "")
    return "AGIRE ORA" if t.startswith("AGIRE ORA") else "Entro l'anno" if t.startswith("Agire entro") else t


def cosa_cambiato(r):
    if r["esito"] == "Nuovo":
        return "Rischio nuovo nel ciclo."
    if not r["variazioni"]:
        return "Nessun valore cambiato: confermato sulle fonti del ciclo precedente."
    return "\n".join(f"• {ETICHETTE_VARIAZIONI.get(v['dimensione'], v['dimensione'])} {_nv(v.get('da'))} → {_nv(v.get('a'))}: "
                     f"{v['evidenza']} ({_data_it(v.get('data_evidenza'))})" for v in r["variazioni"])


def _nv(v):
    return NON_VALUTATA if v is None else v


def scrivi_file_schede(ingresso, uscita_a6, percorso, validazione=""):
    """Schede_<ciclo>.xlsx: indice e una scheda per rischio. uscita_a6 = output di A6 gia'
    controllato, o None (testi di A1 dal registro)."""
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    fill = lambda c: PatternFill("solid", fgColor=c, bgColor=c)
    sottile = Border(bottom=Side(style="thin", color="D9D9D9"))
    testi = {s["id_rischio"]: s for s in (uscita_a6 or {}).get("schede", [])}
    ciclo = ingresso["ciclo"]
    origine = "testi estesi dell'agente A6" if uscita_a6 else "testi di A1 dal registro (senza A6)"
    intestazione = f"{NOME_ANALISI} — schede dei rischi · ciclo {ciclo} · dati al {_data_it(ingresso['data_riferimento'])} " \
                   f"· {validazione or 'dati validati'} · {RISERVATEZZA}"

    def cella(ws, r, c, v=None, **stile):
        x = ws.cell(row=r, column=c, value=v)
        for k, val in stile.items():
            setattr(x, k, val)
        return x

    def stampa(ws, orizzontale=False):
        ws.page_setup.orientation = "landscape" if orizzontale else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.oddHeader.center.text = RISERVATEZZA
        ws.oddHeader.right.text = NOME_ANALISI
        ws.oddFooter.left.text = "&F · &A"
        ws.oddFooter.right.text = "Pagina &P di &N"

    def chip(c, urgenza):
        u = _urgenza_breve(urgenza)
        colori = {"AGIRE ORA": (ROSSO, "FFFFFF"), "Entro l'anno": (AMBRA, TESTO)}.get(u, ("D9D9D9", TESTO))
        c.value = u
        c.fill = fill(colori[0])
        c.font = Font(size=9, bold=True, color=colori[1])
        c.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)

    def chip_livello(c, livello):
        colori = COLORI_LIVELLO.get(livello, ("D9D9D9", TESTO))
        c.value = livello or "—"
        c.fill = fill(colori[0])
        c.font = Font(size=9, bold=True, color=colori[1])
        c.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)

    def righe_testo(testo, larghezza, sz=10):
        per_riga = max(1, int(larghezza * 1.3 * 10 / sz))
        return sum(max(1, -(-len(p) // per_riga)) for p in str(testo).split("\n"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Indice"
    ws.sheet_view.showGridLines = False
    for l, w in zip("ABCDEFGH", [6, 10, 56, 9, 15, 14, 13, 30]):
        ws.column_dimensions[l].width = w
    ws.merge_cells("A1:H1")
    cella(ws, 1, 1, intestazione, font=Font(size=13, bold=True, color="FFFFFF"), fill=fill(BLU_NOTTE),
          alignment=Alignment(vertical="center", indent=1, wrap_text=True))
    ws.row_dimensions[1].height = 36
    ws.merge_cells("A2:H2")
    cella(ws, 2, 1, "Una scheda per rischio, nell'ordine della classifica (score composito), con i valori VALIDATI del "
                    f"registro. Testi: {origine}; numeri, fonti e urgenza vengono dal registro e non li cambia nessun agente. "
                    "Clic sull'ID per aprire la scheda. Prodotto dal comando  python esegui_ciclo_locale.py schede.",
          font=Font(size=9.5, color=TESTO), fill=fill("DDEBF7"), alignment=Alignment(wrap_text=True, vertical="top", indent=1))
    ws.row_dimensions[2].height = 42
    for j, t in enumerate(["#", "ID", "Rischio", "Score", "Urgenza", "Livello", "Esito", "Owner"], start=1):
        cella(ws, 4, j, t, font=Font(bold=True, color="FFFFFF"), fill=fill(BLU_NOTTE))
    for r in ingresso["rischi"]:
        riga = 4 + r["posizione_in_classifica"]
        cella(ws, riga, 1, r["posizione_in_classifica"], font=Font(bold=True, color=BLU_NOTTE))
        c = cella(ws, riga, 2, r["id_rischio"], font=Font(bold=True, color=BLU_DIM, underline="single"))
        c.hyperlink = f"#'{r['id_rischio']}'!A1"
        cella(ws, riga, 3, r["nome"] + ("  ◆" if r["attenzione_supervisore"] else ""))
        cella(ws, riga, 4, r["score"], number_format="0.00", font=Font(bold=True))
        chip(ws.cell(row=riga, column=5), r["urgenza"]["esito"])
        chip_livello(ws.cell(row=riga, column=6), (r.get("livello_allerta") or {}).get("livello"))
        cella(ws, riga, 7, r["esito"])
        cella(ws, riga, 8, r["owner"], font=Font(size=9))
        for x in range(1, 9):
            ws.cell(row=riga, column=x).border = sottile
        ws.row_dimensions[riga].height = 20
    fine = 5 + len(ingresso["rischi"])
    ws.merge_cells(start_row=fine + 1, start_column=1, end_row=fine + 1, end_column=8)
    cella(ws, fine + 1, 1, "Legenda — " + LEGENDA, font=Font(size=8, italic=True, color=TESTO_TENUE),
          alignment=Alignment(wrap_text=True, vertical="top"))
    ws.row_dimensions[fine + 1].height = righe_testo("Legenda — " + LEGENDA, 153, 8) * 11 + 6
    stampa(ws, orizzontale=True)

    for r in ingresso["rischi"]:
        t = testi.get(r["id_rischio"])
        ws = wb.create_sheet(r["id_rischio"])
        ws.sheet_view.showGridLines = False
        larghezze = [18, 12, 16, 16, 16, 16, 16, 34]
        for j, w in enumerate(larghezze, start=1):
            ws.column_dimensions[chr(64 + j)].width = w
        tot = sum(larghezze)
        ws.merge_cells("A1:H1")
        cella(ws, 1, 1, intestazione, font=Font(size=9, color="FFFFFF"), fill=fill(BLU_MEDIO),
              alignment=Alignment(vertical="center", indent=1))
        ws.merge_cells("A2:F2")
        ws.merge_cells("G2:H2")
        cella(ws, 2, 1, f"{r['posizione_in_classifica']} · {r['id_rischio']} · {r['nome']}",
              font=Font(size=14, bold=True, color="FFFFFF"), alignment=Alignment(vertical="center", indent=1, wrap_text=True))
        d = None if r["score_ciclo_precedente"] is None else round(r["score"] - r["score_ciclo_precedente"], 2)
        freccia = "nuovo" if d is None else ("▲ +" if d > 0 else "▼ " if d < 0 else "● ") + _it(d)
        liv = (r.get("livello_allerta") or {}).get("livello")
        cella(ws, 2, 7, f"score {_it(r['score'])}   {freccia}   ·   {_urgenza_breve(r['urgenza']['esito'])}"
                        + (f"   ·   {liv}" if liv else ""),
              font=Font(size=12, bold=True, color="FFFFFF"), alignment=Alignment(horizontal="right", vertical="center"))
        for j in range(1, 9):
            ws.cell(row=2, column=j).fill = fill(BLU_NOTTE)
        ws.row_dimensions[2].height = 34
        etichette = [((1, 2), f"{r['event_type']} · {r['event_type_nome'] or ''}"), ((3, 4), f"Driver: {r['driver'] or '—'}"),
                     ((5, 6), f"Owner: {r['owner'] or '—'} · Presidio: {r.get('presidio_controlli') or '—'}"),
                     ((7, 7), "◆ Attenzione del supervisore" if r["attenzione_supervisore"] else ""),
                     ((8, 8), f"Esito: {r['esito']} · validato")]
        for (da, a), testo in etichette:
            if da != a:
                ws.merge_cells(start_row=3, start_column=da, end_row=3, end_column=a)
            sup = testo.startswith("◆")
            for x in range(da, a + 1):
                ws.cell(row=3, column=x).fill = fill(AMBRA_TENUE if sup else BLU_CHIARO)
            cella(ws, 3, da, testo or None, font=Font(size=9, bold=sup, color="7F6000" if sup else TESTO_TENUE),
                  alignment=Alignment(vertical="center", wrap_text=True))
        ws.row_dimensions[3].height = 28
        riga = 5

        def sezione(testo):
            nonlocal riga
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=8)
            cella(ws, riga, 1, testo, font=Font(size=12, bold=True, color=BLU_NOTTE))
            for x in range(1, 9):
                ws.cell(row=riga, column=x).border = Border(bottom=Side(style="medium", color=BLU_NOTTE))
            riga += 1

        def testo_lungo(testo, sz=10, fondo=None):
            nonlocal riga
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=8)
            c = cella(ws, riga, 1, testo, font=Font(size=sz), alignment=Alignment(wrap_text=True, vertical="top"))
            if fondo:
                c.fill = fill(fondo)
            ws.row_dimensions[riga].height = min(409, righe_testo(testo, tot, sz) * sz * 1.35 + 6)
            riga += 1

        # D-68: prima il fenomeno (Il rischio, poi il Razionale in termini di rischio), poi i valori, ognuno con il
        # suo perche'. Con A6 la descrizione estesa e' gia' il fenomeno in termini di rischio: niente Razionale a parte.
        motivazioni = r.get("motivazioni") or {}
        sezione("Il rischio")
        testo_lungo(t["descrizione_estesa"] if t else r["descrizione"])
        if not t:
            riga += 1
            sezione("Razionale: il fenomeno in termini di rischio" if motivazioni else "Razionale (registro)")
            testo_lungo(r.get("razionale_rischio", r.get("razionale")) or "—")
        riga += 1
        sezione("Valutazione, dimensione per dimensione (valori validati)")
        for j, testo in ((1, "Dimensione"), (2, "Valore"), (3, "Perche' questo valore: informazioni e documenti")):
            cella(ws, riga, j, testo, font=Font(size=9, bold=True, color=TESTO_TENUE))
        ws.merge_cells(start_row=riga, start_column=3, end_row=riga, end_column=8)
        riga += 1
        spiegazioni = {x["dimensione"]: x["spiegazione"] for x in t["dimensioni"]} if t else motivazioni
        for dim in DIMENSIONI:
            v = r["valori"][dim]
            cella(ws, riga, 1, NOMI_DIM[dim], font=Font(size=10, bold=True, color=BLU_NOTTE), alignment=Alignment(vertical="center"))
            colore = SCALA.get(v, "FFFFFF")
            cella(ws, riga, 2, NON_VALUTATA if v is None else v, fill=fill(colore),
                  font=Font(size=14 if v is not None else 10, bold=True, color="FFFFFF" if v == 5 else TESTO),
                  alignment=Alignment(horizontal="center", vertical="center"))
            ws.merge_cells(start_row=riga, start_column=3, end_row=riga, end_column=8)
            sp = spiegazioni.get(dim) or ""
            if sp:
                cella(ws, riga, 3, sp, font=Font(size=10), alignment=Alignment(wrap_text=True, vertical="center"))
            else:
                cella(ws, riga, 3, "Motivo non diviso per dimensione nel razionale di A1: vedi il Razionale qui sopra.",
                      font=Font(size=9, italic=True, color=TESTO_TENUE), alignment=Alignment(vertical="center"))
            ws.row_dimensions[riga].height = max(26, righe_testo(sp, tot - 30) * 13.5 + 6) if sp else 22
            for x in range(1, 9):
                ws.cell(row=riga, column=x).border = sottile
            riga += 1
        pesi = ingresso["pesi"]
        presenti = [d for d in PESI if r["valori"].get(d) is not None]
        formula = " + ".join(f"{_it(pesi[d])} x {r['valori'][d]}" for d in presenti)
        somma = _it(sum(pesi[d] for d in presenti if isinstance(pesi.get(d), (int, float))))
        nota_vel = "" if "velocity" in presenti else " Velocity non valutata: i pesi si ridistribuiscono sulle altre tre."
        ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=8)
        cella(ws, riga, 1, f"Score composito {_it(r['score'])} = ({formula}) / {somma}.{nota_vel} La rilevanza delle "
                           "fonti e' nella tabella Fonti.", font=Font(size=8, italic=True, color=TESTO_TENUE))
        riga += 2
        u = r["urgenza"]
        sezione("Urgenza (disuguaglianza di Mosca) e livello di allerta")
        la = r.get("livello_allerta") or {}
        if la.get("livello"):
            testo_lungo(f"Livello di allerta: {la['livello']}" + (f" — {la['motivo']}." if la.get("motivo") else "."),
                        fondo=FONDI_LIVELLO.get(la["livello"]))
        perche_adeg = motivazioni.get("adeguamento")
        adeg = (f" (con {_it(r['anni_adeguamento'], 1)} anni di adeguamento strutturale, stimati da A1 e validati"
                + (f"; motivo: {perche_adeg[0].lower() + perche_adeg[1:].rstrip('.')}" if perche_adeg else "") + ")"
                if r["anni_adeguamento"] else "")
        if isinstance(u["x"], (int, float)) and isinstance(u["y"], (int, float)):
            testo_lungo(f"{u['esito']}. Tempo necessario X + Y = {_it(u['x'])} + {_it(u['y'])}{adeg} = "
                        f"{_it(u['x'] + u['y'])} anni; tempo disponibile Z = {_it(u['z'], 1)} anni (Prossimita "
                        f"{r['valori']['prossimita']}). Margine Z - (X + Y) = {_it(u['margine'])}.",
                        fondo=ROSSO_TENUE if _urgenza_breve(u["esito"]) == "AGIRE ORA" else None)
        else:
            testo_lungo("Urgenza non calcolabile: manca un valore del ciclo.")
        riga += 1
        sezione("Cosa e' cambiato nel ciclo")
        testo_lungo(cosa_cambiato(r))
        riga += 1
        sezione("Andamento negli ultimi trimestri")
        intest = ["Ciclo", "Imp", "Prob", "Pross", "Vel", "Score"]
        for j, testo in enumerate(intest, start=1):
            cella(ws, riga, j, testo, font=Font(size=9, bold=True, color=TESTO_TENUE), alignment=Alignment(horizontal="center"))
        prima = riga
        for s in r["storia"]:
            riga += 1
            cella(ws, riga, 1, s["ciclo"], font=Font(bold=True, color=BLU_NOTTE), alignment=Alignment(horizontal="center"))
            for j, dim in enumerate(DIMENSIONI, start=2):
                v = s.get(dim)
                cella(ws, riga, j, NON_VALUTATA if v is None else v, fill=fill(SCALA.get(v, "FFFFFF")),
                      font=Font(bold=True, color="FFFFFF" if v == 5 else TESTO), alignment=Alignment(horizontal="center"))
            cella(ws, riga, 6, s["score"], number_format="0.00", font=Font(bold=True), alignment=Alignment(horizontal="center"))
        if len(r["storia"]) >= 2:
            graf = LineChart()
            graf.add_data(Reference(ws, min_col=6, min_row=prima, max_row=riga), titles_from_data=True)
            graf.set_categories(Reference(ws, min_col=1, min_row=prima + 1, max_row=riga))
            s = graf.series[0]
            s.smooth = False
            s.marker.symbol = "circle"
            s.marker.size = 7
            s.graphicalProperties.line.solidFill = BLU_NOTTE
            s.marker.graphicalProperties = GraphicalProperties(solidFill=BLU_NOTTE)
            s.dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False,
                                    showPercent=False, dLblPos="t")
            graf.y_axis.scaling.min, graf.y_axis.scaling.max, graf.y_axis.majorUnit = 1, 5, 1
            graf.x_axis.delete = graf.y_axis.delete = False
            graf.legend = None
            graf.title = "Score composito per trimestre (scala fissa 1-5)"
            graf.title.overlay = False
            graf.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
            graf.height, graf.width = 5.5, 11
            ws.add_chart(graf, f"A{riga + 2}")
            riga += 14
        else:
            riga += 1
            cella(ws, riga, 1, f"Rischio entrato nel registro nel {r['storia'][0]['ciclo'] if r['storia'] else ciclo}: "
                               "l'andamento si vedra' dal prossimo ciclo.", font=Font(size=9, italic=True, color=TESTO_TENUE))
            riga += 1
        riga += 1
        sezione("Fonti")
        for j, testo in zip((1, 4, 5, 6, 7), ["Fonte", "Data", "Tipo", "Rilevanza", "A sostegno di" if t else ""]):
            cella(ws, riga, j, testo or None, font=Font(size=9, bold=True, color=TESTO_TENUE))
        ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=3)
        ws.merge_cells(start_row=riga, start_column=7, end_row=riga, end_column=8)
        riga += 1
        sostegno = {f["riferimento"]: f["a_sostegno_di"] for f in t["fonti"]} if t else {}
        for f in r["fonti"]:
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=3)
            ws.merge_cells(start_row=riga, start_column=7, end_row=riga, end_column=8)
            cella(ws, riga, 1, f["riferimento"], font=Font(size=10), alignment=Alignment(wrap_text=True, vertical="top"))
            cella(ws, riga, 4, _data_it(f.get("data")), font=Font(size=10), alignment=Alignment(vertical="top"))
            cella(ws, riga, 5, f.get("tipo"), font=Font(size=10, bold=True, color=BLU_DIM if f.get("tipo") == "Interna" else TESTO_TENUE),
                  alignment=Alignment(vertical="top"))
            ril = f.get("rilevanza")
            cella(ws, riga, 6, RILEVANZA.get(ril, ril), fill=fill(SCALA_RIL.get(ril, "FFFFFF")), font=Font(size=10, bold=True),
                  alignment=Alignment(horizontal="center", vertical="top"))
            so = sostegno.get(f["riferimento"])
            cella(ws, riga, 7, so, font=Font(size=10), alignment=Alignment(wrap_text=True, vertical="top"))
            for x in range(1, 9):
                ws.cell(row=riga, column=x).border = sottile
            ws.row_dimensions[riga].height = max(righe_testo(f["riferimento"], 44), righe_testo(so or "", 44)) * 13.5 + 5
            riga += 1
        riga += 1

        def perche_ed_effetto(perche, effetto, etichetta_perche):
            """Riga sotto un'azione o un KRI (D-68): a sinistra il perche', a destra l'effetto atteso (A1)."""
            nonlocal riga
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=4)
            ws.merge_cells(start_row=riga, start_column=5, end_row=riga, end_column=8)
            sx = f"{etichetta_perche}: {perche or '—'}"
            dx = f"Effetto atteso: {effetto or 'non indicato nel razionale di A1.'}"
            for col, testo in ((1, sx), (5, dx)):
                cella(ws, riga, col, testo, font=Font(size=9.5, color=TESTO), fill=fill("F7F9FC"),
                      alignment=Alignment(wrap_text=True, vertical="top", indent=1))
            for x in range(1, 9):
                ws.cell(row=riga, column=x).fill = fill("F7F9FC")
                ws.cell(row=riga, column=x).border = sottile
            ws.row_dimensions[riga].height = max(righe_testo(sx, sum(larghezze[:4]) - 2, 9.5),
                                                 righe_testo(dx, sum(larghezze[4:]) - 2, 9.5)) * 13 + 6
            riga += 1

        sezione("Azioni")
        if not r["azioni"]:
            testo_lungo("Nessuna azione del ciclo o aperta.")
        else:
            for j, testo in zip((1, 6, 8), ["Azione", "Owner", "Orizzonte · stato"]):
                cella(ws, riga, j, testo, font=Font(size=9, bold=True, color=TESTO_TENUE))
            riga += 1
        for a in r["azioni"]:
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=5)
            ws.merge_cells(start_row=riga, start_column=6, end_row=riga, end_column=7)
            cella(ws, riga, 1, f"• {a['descrizione']}", font=Font(size=10, bold=True, color=TESTO),
                  alignment=Alignment(wrap_text=True, vertical="top"))
            cella(ws, riga, 6, a["owner"], font=Font(size=9.5), alignment=Alignment(wrap_text=True, vertical="top"))
            cella(ws, riga, 8, f"{a['orizzonte']} · {a['stato']} (ciclo {a['ciclo']})", font=Font(size=9.5),
                  alignment=Alignment(wrap_text=True, vertical="top"))
            ws.row_dimensions[riga].height = max(righe_testo(a["descrizione"], sum(larghezze[:5])),
                                                 righe_testo(a["owner"] or "", sum(larghezze[5:7]), 9.5)) * 13.5 + 5
            riga += 1
            perche_ed_effetto(a.get("razionale"), a.get("effetto_atteso"), "Perche'")
        riga += 1
        sezione("KRI in vigore")
        for j, testo in zip((1, 5, 6, 7, 8), ["Indicatore", "Valore", "Soglie (attenzione / allerta)", "Ciclo prima", "Stato"]):
            cella(ws, riga, j, testo, font=Font(size=9, bold=True, color=TESTO_TENUE), alignment=Alignment(wrap_text=True))
        riga += 1
        for k in r["kri"]:
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=4)
            tipo = f"[{k['tipo']}] " if k.get("tipo") else ""
            cella(ws, riga, 1, f"{tipo}{k['nome']} ({k['unita']})", font=Font(size=10),
                  alignment=Alignment(wrap_text=True, vertical="center"))
            cella(ws, riga, 5, k["valore"] if k["valore"] is not None else "—", font=Font(size=10, bold=True),
                  alignment=Alignment(horizontal="center", vertical="center"))
            cella(ws, riga, 6, f"{_it(k['soglia_attenzione'], 1)} / {_it(k['soglia_allerta'], 1)}", font=Font(size=10),
                  alignment=Alignment(horizontal="center", vertical="center"))
            cella(ws, riga, 7, k["valore_ciclo_precedente"] if k["valore_ciclo_precedente"] is not None else "—",
                  font=Font(size=10), alignment=Alignment(horizontal="center", vertical="center"))
            f_, t_ = SEMAFORO.get(k["stato"], ("D9D9D9", TESTO))
            cella(ws, riga, 8, k["stato"], fill=fill(f_), font=Font(bold=True, color=t_), alignment=Alignment(horizontal="center", vertical="center"))
            ws.row_dimensions[riga].height = max(20, righe_testo(k["nome"], 60) * 13.5 + 5)
            riga += 1
            perche_ed_effetto(k.get("razionale"), k.get("effetto_atteso"), "Perche' lo misuriamo")
        riga += 1
        ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=8)
        cella(ws, riga, 1, "Legenda — " + LEGENDA, font=Font(size=8, italic=True, color=TESTO_TENUE),
              alignment=Alignment(wrap_text=True, vertical="top"))
        ws.row_dimensions[riga].height = righe_testo("Legenda — " + LEGENDA, tot, 8) * 11 + 6   # 36 fisso la tagliava
        stampa(ws)
    Path(percorso).parent.mkdir(parents=True, exist_ok=True)
    wb.save(percorso)
    return percorso
