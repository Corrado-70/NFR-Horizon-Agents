"""Report al comitato in HTML (D-83, tappa M4): la seconda uscita del comando report, accanto al PDF.

Un file unico per ciclo, Cicli/<ciclo>/Report_comitato_<ciclo>.html, che si apre nel browser senza installare niente e
senza risorse da internet: le viste del report (Quadro, Classifica, Radar di Mosca, Matrice, Delta, Rischi in
osservazione, Metodologia) con i grafici in SVG, le schede di tutti i rischi collegate fra loro e, con --con-lavoro, le
viste di lavoro (Registro, Storico, Indicatori, Notizie, Piano azioni, Osservazione, Validazione, Parametri,
Dizionario). In Cicli/Indice_report.html l'elenco dei report di tutti i cicli.

Le viste sono calcolate qui in Python dalle tabelle dati del registro (RiskRegister, Azioni, KRI, NewsLog,
Sintesi_Cicli, Sintesi_Settimane, Storico_Validazioni), dai parametri e dall'anagrafica, con le regole delle formule di
genera_dashboard.py; le schede riusano schede.input_schede. Prima di uscire, il CANCELLO: i numeri calcolati (X, Y, Z,
margine, urgenza, livello di allerta e motivo, segnali netti, ordine della classifica, Delta) si confrontano con
quelli che Excel ha calcolato nelle viste del registro. Se anche uno non torna l'HTML non esce (il PDF si'), e il
motivo e' nell'uscita del comando.

Due lettori dei dati con la stessa interfaccia (tab, colonne, nome, testi_foglio, righe_foglio, cella):
  - RegistroExcel: la sessione COM gia' aperta da report e valida (esegui_ciclo_locale), dopo il ricalcolo;
  - RegistroFile: il file .xlsx con openpyxl (valori salvati da Excel), per il collaudo, la demo e l'anteprima
    dell'analisi 16 (Analisi_Architettura/16_genera_anteprima_html.py usa questo modulo).

Uso (senza Excel, sul file):  python report_html.py --registro NFR_Registro.xlsx --uscita CARTELLA [--con-lavoro]
Prototipo: Analisi_Architettura/16_genera_anteprima_html.py (analisi 16, 14/10/2026).
"""

import argparse
import html
import json
import math
import re
import sys
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import schede  # parti_razionale, parti_effetto, cosa_cambiato, input_schede, LEGENDA

QUI = Path(__file__).resolve().parent
NOME_ANALISI = "Horizon Scan NFR"
RISERVATEZZA = "Uso interno"
INDICE = "Indice_report.html"                 # in Cicli/, accanto alle cartelle dei cicli
# Radar di Mosca: distanza dal centro = margine + SCOSTAMENTO, fra RAGGIO_MIN e RAGGIO_MAX (come genera_dashboard.py)
SCOSTAMENTO_RAGGIO, RAGGIO_MIN, RAGGIO_MAX = 2, 0.6, 6
NOMI_SETTORI = {"9.1": "Condotta", "9.4": "Legale", "9.5": "Fiscale", "9.6": "Tecnologico", "9.7": "Sicurezza informatica",
                "9.8": "Sicurezza fisica", "9.9": "Terze parti", "9.11": "Employer", "9.13": "Processo"}
NOMI_PROB = {5: "Quasi certa", 4: "Probabile", 3: "Possibile", 2: "Improbabile", 1: "Remota"}
NOMI_IMP = {1: "Limitato", 2: "Contenuto", 3: "Rilevante", 4: "Grave", 5: "Critico"}
DIM = [("impatto", "Impatto", "Imp"), ("probabilita", "Probabilita", "Prob"), ("prossimita", "Prossimita", "Pross"),
       ("velocity", "Velocity", "Vel")]
COL_DIM = {"impatto": "Impatto", "probabilita": "Probabilita", "prossimita": "Prossimita", "velocity": "Velocity"}


# ================================================================ formati
def esc(v):
    return html.escape("" if v is None else str(v))


def it(v, d=2):
    """Numero all'italiana (4,65); None = trattino."""
    if v is None or v == "":
        return "—"
    if isinstance(v, (int, float)):
        return f"{arrotonda(v, d):.{d}f}".replace(".", ",")       # come FIXED di Excel (trappola 111)
    return str(v)


def arrotonda(v, d=0):
    """ROUND di Excel (meta' lontano dallo zero): round() di Python porta 1,25 a 1,2, Excel a 1,3."""
    q = Decimal(1).scaleb(-d)
    return float(Decimal(str(v)).quantize(q, rounding=ROUND_HALF_UP))


def it_breve(v):
    """Numero senza zeri inutili: 3 / 0,25 / 1,5."""
    if v is None or v == "":
        return "—"
    if isinstance(v, (int, float)):
        t = f"{v:.2f}".rstrip("0").rstrip(".")
        return t.replace(".", ",")
    return str(v)


def segnato(v, d=2):
    if v is None:
        return "—"
    if abs(v) < 1e-9:
        return it(0, d)
    return ("+" if v > 0 else "−") + it(abs(v), d)


def data_it(s):
    s = str(s or "")
    return f"{s[8:10]}/{s[5:7]}/{s[:4]}" if re.match(r"\d{4}-\d{2}-\d{2}", s) else s


def a_data(s):
    return datetime.strptime(str(s)[:10], "%Y-%m-%d").date() if s else None


def piu_mesi(d, mesi):
    """EDATE di Excel."""
    m = d.month - 1 + mesi
    a, m = d.year + m // 12, m % 12 + 1
    giorni = [31, 29 if a % 4 == 0 and (a % 100 or a % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]
    return date(a, m, min(d.day, giorni))


def accorcia(t, n):
    t = str(t or "")
    return t if len(t) <= n else t[:n - 1].rstrip() + "…"


def classe(t):
    return re.sub(r"[^a-z0-9]+", "-", str(t or "").lower()).strip("-")


# ================================================================ lettura del registro (solo dati, nessuna formula)
def da_excel(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and v.startswith("#") and v.endswith(("!", "A", "?", "0")) and len(v) <= 7:
        return None                                   # #N/A, #DIV/0!...: un vuoto
    if isinstance(v, int) and not isinstance(v, bool) and v in ERRORI_COM:
        return None                                   # gli stessi errori letti via COM
    return v


ERRORI_COM = {-2146826281, -2146826246, -2146826259, -2146826288, -2146826252, -2146826265, -2146826273,
              -2146826256, -2146826238}


class Registro:
    """L'interfaccia comune dei lettori: le tabelle dati (tab, colonne), i nomi definiti (nome), i testi dei fogli
    (testi_foglio, righe_foglio) e una cella (cella, per il cancello e la regola del registro unico)."""

    def testi_foglio(self, foglio):
        """Riga 2 (come si legge) e righe che cominciano con Legenda / Limiti del foglio: testi gia' scritti dal
        generatore del registro, riusati tali e quali."""
        out = {"come": COME.get(foglio, self.cella(foglio, 2, 1) or ""), "legenda": "", "limiti": ""}
        for v in self.colonna_a(foglio, 3):
            if isinstance(v, str) and v.startswith("Legenda"):
                out["legenda"] = LEGENDE.get(foglio, pulisci(v))
            if isinstance(v, str) and v.startswith("Limiti"):
                out["limiti"] = v
        return out


class RegistroFile(Registro):
    """Il registro letto dal file con openpyxl: i valori che Excel ha salvato (collaudo, demo, anteprima)."""

    def __init__(self, percorso):
        import openpyxl
        self.wb = openpyxl.load_workbook(percorso, data_only=True)
        self.tab = {}
        self.colonne = {}
        for ws in self.wb.worksheets:
            for nome in ws.tables:
                righe = list(ws[ws.tables[nome].ref])
                intest = [c.value for c in righe[0]]
                dati = [{k: da_excel(c.value) for k, c in zip(intest, r)} for r in righe[1:]]
                self.tab[nome] = [d for d in dati if d.get(intest[0]) not in (None, "")]   # la riga segnaposto
                self.colonne[nome] = intest

    def nome(self, n):
        """Valore di un nome definito: uno scalare o la lista dei valori (per colonna)."""
        dn = self.wb.defined_names[n]
        dest = list(dn.destinations)
        if not dest:
            return costante(dn.attr_text)
        foglio, rif = dest[0]
        celle = self.wb[foglio][rif.replace("$", "")]
        if not isinstance(celle, tuple):
            return da_excel(celle.value)
        return per_colonna([[da_excel(c.value) for c in (r if isinstance(r, tuple) else (r,))] for r in celle])

    def cella(self, foglio, riga, colonna):
        return da_excel(self.wb[foglio].cell(row=riga, column=colonna).value)

    def colonna_a(self, foglio, da):
        ws = self.wb[foglio]
        return [ws.cell(row=r, column=1).value for r in range(da, ws.max_row + 1)]

    def righe_foglio(self, foglio, da, colonne):
        ws = self.wb[foglio]
        out = []
        for r in range(da, ws.max_row + 1):
            valori = [ws.cell(row=r, column=c).value for c in range(1, colonne + 1)]
            if any(v not in (None, "") for v in valori):
                out.append(valori)
        return out


class RegistroExcel(Registro):
    """Il registro letto dalla sessione COM aperta (excel_com.SessioneExcel), dopo il ricalcolo: i valori che le
    formule hanno appena calcolato, nella stessa sessione da cui esce il PDF."""

    def __init__(self, sessione):
        from excel_com import valori_tabella
        self.s = sessione
        self.tab = {}
        self.colonne = {}
        for ws in sessione.wb.Worksheets:
            for lo in ws.ListObjects:
                intest = [c.Name for c in lo.ListColumns]
                dati = [{k: da_excel(v) for k, v in r.items()} for r in valori_tabella(lo)]
                self.tab[lo.Name] = [d for d in dati if d.get(intest[0]) not in (None, "")]
                self.colonne[lo.Name] = intest

    def nome(self, n):
        nome = self.s.wb.Names(n)
        try:
            valori = nome.RefersToRange.Value
        except Exception:                            # noqa: BLE001 - un nome che e' una costante (="0,25", ="testo")
            return costante(str(nome.RefersTo).lstrip("="))
        if not isinstance(valori, tuple):
            return da_excel(valori)
        return per_colonna([[da_excel(v) for v in r] for r in valori])

    def cella(self, foglio, riga, colonna):
        return da_excel(self.s.foglio(foglio).Cells(riga, colonna).Value)

    def _usate(self, foglio):
        ur = self.s.foglio(foglio).UsedRange
        return ur.Row + ur.Rows.Count - 1

    def colonna_a(self, foglio, da):
        ultima = self._usate(foglio)
        if ultima < da:
            return []
        ws = self.s.foglio(foglio)
        valori = ws.Range(ws.Cells(da, 1), ws.Cells(ultima, 1)).Value
        return [r[0] for r in valori] if isinstance(valori, tuple) else [valori]

    def righe_foglio(self, foglio, da, colonne):
        ultima = self._usate(foglio)
        if ultima < da:
            return []
        ws = self.s.foglio(foglio)
        valori = ws.Range(ws.Cells(da, 1), ws.Cells(ultima, colonne)).Value
        righe = valori if isinstance(valori, tuple) and isinstance(valori[0], tuple) else ((valori,),)
        return [[da_excel(v) for v in r] for r in righe if any(da_excel(v) not in (None, "") for v in r)]


def costante(testo):
    t = str(testo).strip().strip('"')
    try:
        return float(t.replace(",", ".")) if re.fullmatch(r"-?\d+[.,]\d+", t) else int(t)
    except ValueError:
        return t


def per_colonna(righe):
    return [r[0] for r in righe] if all(len(r) == 1 for r in righe) else righe


# Testi "come si legge" e legende: quelli dei fogli del registro, riscritti dove citano colonne, righe o celle di Excel.
COME = {
    "Quadro": "Prima pagina per il comitato. I cinque riquadri dicono se il ciclo e' completo e validato; poi le priorita' "
              "per score, la gravita' (matrice) e l'urgenza (radar). Sotto, il commento di A1 (scritto prima della "
              "validazione: se dopo ci sono state rettifiche lo dice la riga \"Dopo il commento\" e valgono i numeri "
              "validati) e i rischi segnalati, con il perche' calcolato dai dati validati. Metodo e scale: Metodologia; "
              "ogni dato: Dizionario.",
    "Classifica": "Un rischio per barra, dal piu' alto per score composito (1-5, 5 = peggiore). Colore = urgenza di Mosca: "
                  "rosso agire ora, ambra entro l'anno, grigio c'e' margine. Sotto, la lente di Mosca spiega il colore: "
                  "tempo necessario per decidere e mitigare (X + Y) contro tempo prima che il rischio possa manifestarsi "
                  "(Z). Nessun agente ordina o stima l'urgenza: e' calcolata dai valori del ciclo.",
    "Radar_Mosca": "Quanto presto agire e quanto e' grande ogni rischio. Piu' vicino al centro = agire prima. Dentro il "
                   "cerchio rosso (X + Y > Z) si e' gia' in ritardo anche se l'evento non e' imminente. Area del cerchio = "
                   "Probabilita x Impatto; colore = livello di allerta (rosso attivazione, giallo attenzione, verde "
                   "monitoraggio); ◆ = tema all'attenzione del supervisore; settore = categoria di rischio. Z dalla "
                   "Prossimita, Y dalle azioni o dagli anni di adeguamento; i parametri sono nella pagina Parametri.",
    "Matrice": "Ogni rischio nella cella della sua Probabilita (righe) e del suo Impatto (colonne); colore = P x I (verde "
               "fino a 6, ambra 8-12, rosso da 15). Dice la gravita', non l'urgenza (quella e' nel radar). Accanto, i "
               "rischi dal piu' grave.",
    "Delta": "Una barra per rischio, dalla variazione piu' alta dello score composito: rosso in aumento, blu in "
             "diminuzione. Accanto, nello stesso ordine, il motivo con l'evidenza e la sua data; sotto, il dettaglio per "
             "dimensione (ogni dimensione cambia al massimo di un punto fra due cicli).",
    "Fenomeni": "I rischi del registro rimasti fuori dal radar: stessa scala e stesso score dei rischi del radar, niente "
                "KRI ne' azioni. Per ognuno: perche' e' fuori (quale soglia di Prossimita o Probabilita manca, o la "
                "capienza), l'orizzonte stimato e il segnale che lo farebbe salire. Dal piu' alto per score; in fondo chi "
                "e' entrato o uscito dal radar nel ciclo.",
    "Metodologia": "Scale e regole dalla rubrica del pacchetto, pesi e parametri in vigore. Nessun punteggio composito, "
                   "urgenza o stato di un KRI e' stimato da un agente: li calcola il motore, con le stesse regole per tutti.",
    "Registro": "Prima i due riquadri: se il ciclo non e' validato o e' incompleto i numeri non sono definitivi. Un rischio "
                "per riga, dal piu' alto per score. Punteggi 1-5, 5 = peggiore; n.v. = Velocity non valutata (lo score "
                "usa le altre tre). Presidio = giudizio della Banca sui controlli, fuori dallo score. Segnali netti = "
                "notizie peggiorative meno migliorative delle ultime 4 settimane (▲ ▼ oltre la soglia): e' un allarme per "
                "il prossimo trimestre, non cambia lo score. L'andamento di un rischio e' nella sua scheda.",
    "Indicatori": "In alto i KRI in vigore nel ciclo (A1 rilegge il valore a ogni ciclo), sotto la storia di ognuno. Stato: "
                  "Verde sotto la soglia di attenzione, Ambra fino all'allerta, Rosso oltre, Non rilevato senza valore "
                  "(mai zero), Soglie da definire finche' la Banca non le fissa in validazione. Tendenza sul ciclo prima: "
                  "valori crescenti = peggio (▲ peggiora).",
    "Notizie": "Per categoria quante notizie e quanto forti, nelle ultime 4 settimane. AREA NON PRESIDIATA = notizie di "
               "intensita' 2 o piu' su una categoria senza rischi nel registro: tema da portare al prossimo ciclo "
               "trimestrale; In osservazione = nessun rischio nel radar, almeno uno in osservazione (D-79). Le notizie non "
               "cambiano score. Accanto la sintesi di A3 della settimana, i flag proposti e l'eventuale proposta di ciclo "
               "straordinario (D-75).",
    "Piano_azioni": "Tutte le azioni, dalla piu' recente. La validazione e' la delibera (Accettato o Rettificato = "
                    "Approvata, Respinto = Respinta); In corso e Completata le registra chi tiene il piano, quando l'owner "
                    "lo comunica. Scadenza = data della decisione + orizzonte. Le azioni aperte passano ad A1 nel "
                    "trimestre dopo.",
    "Osservazione": "In alto i rischi fuori dal radar con le quattro dimensioni, lo score, il motivo dello stato, "
                    "l'orizzonte, il segnale di avvicinamento e le notizie che A3 ha collegato nelle ultime 4 settimane. "
                    "Sotto, dove si concentrano: rischi in osservazione per driver e categoria e, accanto, i rischi del "
                    "radar con lo stesso incrocio.",
}
LEGENDE = {"Piano_azioni": "Legenda — Orizzonte = tempo per completare l'azione; Scadenza = data della decisione + "
                           "orizzonte (per le azioni approvate o in corso)."}


def pulisci(t):
    """Toglie da una legenda i rimandi a colonne, righe e fogli del registro Excel."""
    for a, b in ((r"\s*\(colonne U-AB\)", ""), (r":\s*colonne U-AB", ""), (r":\s*foglio Radar_Mosca", ""),
                 (r"\s*\(righe \d+-\d+\)", ""), (r"\s*\(foglio [A-Za-z_]+\)", ""), (r":\s*testo nel RiskRegister", ""),
                 (r"\s*Storia completa: foglio RiskRegister[^.]*\.", ""), (r"#N/D = ", "— = ")):
        t = re.sub(a, b, t)
    return t


# ================================================================ motore delle viste (Python al posto delle formule)
def cicli_ordinati(reg):
    return sorted({r["Ciclo"] for r in reg.tab["tblRegister"]})


def trimestre_meno(ciclo, n):
    a, q = int(ciclo[:4]), int(ciclo[5])
    k = a * 4 + (q - 1) - n
    return f"{k // 4}Q{k % 4 + 1}"


class Vista:
    """Tutto quello che il report di un ciclo mostra, calcolato dai dati. Le regole sono quelle delle formule di
    genera_dashboard.py (fogli Radar_Mosca, Classifica, Delta, Registro, Notizie...), citate accanto."""

    def __init__(self, reg, ciclo):
        self.reg = reg
        self.ciclo = ciclo
        cicli = cicli_ordinati(reg)
        self.cicli = cicli[:cicli.index(ciclo) + 1]
        self.corrente = ciclo == reg.nome("cicloCorrente")
        self.prec = self.cicli[-2] if len(self.cicli) > 1 else None
        self.meno2 = trimestre_meno(ciclo, 2)
        self.trend = self.cicli[-8:]
        p = {n: reg.nome(n) for n in ("pesoImpatto", "pesoProbabilita", "pesoProssimita", "pesoVelocity", "moscaX",
                                       "moscaYDefault", "livelloPxI", "livelloQuotaKRI", "livelloMargine", "livelloForti",
                                       "sogliaDirezione", "sogliaDelta", "sogliaProssimitaRadar",
                                       "sogliaProbabilitaRadar", "cicliUscitaRadar")}
        p["anniProssimita"] = reg.nome("anniProssimita")
        p["capienzaRadar"] = reg.nome("capienzaRadar")
        p["capienzaOsservazione"] = reg.nome("capienzaOsservazione")
        self.p = p
        self.codici = reg.nome("codiciL1")
        self.nomi_cat = dict(zip(self.codici, reg.nome("nomiL1")))
        self.nomi_driver = dict(zip(reg.nome("codiciDriver"), reg.nome("nomiDriver")))
        # Snapshot: le righe dei dati fino a questo ciclo (un file per ciclo e' una fotografia).
        dentro = set(self.cicli)
        self.rr_tutte = [r for r in reg.tab["tblRegister"] if r["Ciclo"] in dentro]
        self.rr = [r for r in self.rr_tutte if r["Ciclo"] == ciclo]
        self.per_id = {r["ID_Rischio"]: r for r in self.rr}
        self.data_rif = self.rr[0]["Data_Riferimento"]
        self.azioni = [a for a in reg.tab["tblAzioni"] if a["Ciclo"] in dentro]
        self.kri_tutti = [k for k in reg.tab["tblKRI"] if k["Ciclo"] in dentro]
        self.kri = [k for k in self.kri_tutti if k["Ciclo"] == ciclo and k["In_Vigore"] == "Sì"]
        self.sintesi = next((s for s in reg.tab["tblSintesi"] if s["Ciclo"] == ciclo), {})
        self.val = [v for v in reg.tab["tblValidazione"] if v["Periodo"] == ciclo]
        # Settimana corrente: quella dei parametri per il ciclo corrente; per un ciclo passato l'ultima settimana di
        # notizie non dopo la data di riferimento.
        settimane = sorted({n["Settimana"] for n in reg.tab["tblNews"] if n["Settimana"] <= self.data_rif})
        self.settimana = reg.nome("settimanaCorrente") if self.corrente else (settimane[-1] if settimane else None)
        self.news = [n for n in reg.tab["tblNews"] if self.settimana and n["Settimana"] <= self.settimana]
        self.settimane_sintesi = [s for s in reg.tab["tblSettimane"] if self.settimana and s["Settimana"] <= self.settimana]
        if self.settimana:
            s = a_data(self.settimana)
            self.inizio4 = (s - timedelta(days=21)).isoformat()
            self.settimane12 = [(s - timedelta(days=7 * k)).isoformat() for k in range(11, -1, -1)]
        else:
            self.inizio4, self.settimane12 = None, []
        self.finestra = [n for n in self.news if self.inizio4 and self.inizio4 <= n["Settimana"] <= self.settimana]
        self._anagrafica()
        self._rischi()

    # ------------------------------------------------ anagrafica (= il radar)
    def _anagrafica(self):
        attuale = [dict(zip(["ID", "Nome", "Event_Type", "Driver", "Owner", "Anni_Adeguamento",
                             "Attenzione_Supervisore", "Presidio_Controlli"], r))
                   for r in self.reg.nome("anagrafica") if r[0]]
        per_id = {a["ID"]: a for a in attuale}
        if self.corrente:
            self.anag = attuale
        else:
            # Ciclo passato: il radar di allora dal RiskRegister; owner e anni di adeguamento dalla riga del ciclo;
            # supervisore e presidio non sono storicizzati nel registro: si usano quelli di oggi (limite dell'anteprima).
            self.anag = []
            for r in self.rr:
                if str(r.get("Stato") or "Radar").startswith("Radar"):
                    oggi = per_id.get(r["ID_Rischio"], {})
                    self.anag.append({"ID": r["ID_Rischio"], "Nome": r["Nome"], "Event_Type": r["Event_Type"],
                                      "Driver": r["Driver"], "Owner": r.get("Owner") or oggi.get("Owner"),
                                      "Anni_Adeguamento": r.get("Anni_Adeguamento"),
                                      "Attenzione_Supervisore": oggi.get("Attenzione_Supervisore"),
                                      "Presidio_Controlli": oggi.get("Presidio_Controlli")})
        self.oss_ids = [r["ID_Rischio"] for r in sorted(
            (r for r in self.rr if r.get("Stato") == "Osservazione"),
            key=lambda r: -(r.get("Ordine_Osservazione") or r.get("Score_Composito") or 0))]

    def score_di(self, rid, ciclo):
        r = next((x for x in self.rr_tutte if x["Ciclo"] == ciclo and x["ID_Rischio"] == rid), None) if ciclo else None
        if r is None:
            r = next((x for x in self.reg.tab["tblRegister"] if x["Ciclo"] == ciclo and x["ID_Rischio"] == rid), None)
        return r

    # ------------------------------------------------ un rischio del radar: Radar_Mosca, Classifica, Registro
    def _rischi(self):
        p = self.p
        conta_settore = {}
        self.rischi = []
        for pos, a in enumerate(self.anag, start=1):
            rid = a["ID"]
            r = self.per_id.get(rid)
            if r is None:
                continue
            x = {"id": rid, "pos": pos, "nome": r["Nome"], "cat": a["Event_Type"], "driver": a["Driver"],
                 "owner": a.get("Owner") or "—", "sup": a.get("Attenzione_Supervisore") == "Sì",
                 "presidio": a.get("Presidio_Controlli") or "Non verificato", "esito": r["Esito"],
                 "I": r["Impatto"], "P": r["Probabilita"], "Pr": r["Prossimita"], "V": r["Velocity"],
                 "score": r["Score_Composito"], "rilievi": r.get("Rilievi"), "riga": r,
                 "variazioni": json.loads(r.get("Variazioni") or "[]")}
            prec = self.score_di(rid, self.prec)
            m2 = self.score_di(rid, self.meno2)
            x["score_prec"] = prec["Score_Composito"] if prec else None
            x["delta"] = round(x["score"] - x["score_prec"], 2) if prec else None
            x["delta6"] = round(x["score"] - m2["Score_Composito"], 2) if m2 else None
            x["dd"] = {d: (r[COL_DIM[d]] - prec[COL_DIM[d]]) if prec and r[COL_DIM[d]] is not None
                       and prec[COL_DIM[d]] is not None else (None if prec else None) for d, _, _ in DIM}
            x["pxi"] = x["I"] * x["P"]
            # Mosca (Radar_Mosca, colonne H-L): Z dalla Prossimita, X fisso, Y = orizzonte piu' lungo delle azioni
            # non respinte del periodo in cui il rischio e' stato valutato (o la riserva), o gli anni di adeguamento.
            x["z"] = p["anniProssimita"][x["Pr"] - 1]
            x["x"] = p["moscaX"]
            periodo = r.get("Periodo_Valutazione") or self.ciclo
            mesi = [z["Mesi_Orizzonte"] for z in self.azioni if z["ID_Rischio"] == rid and z["Ciclo"] == periodo
                    and z["Stato"] != "Respinta" and (z["Mesi_Orizzonte"] or 0) > 0]
            y_azioni = max(mesi) / 12 if mesi else p["moscaYDefault"]
            x["adeg"] = a.get("Anni_Adeguamento") or 0
            x["y"] = max(y_azioni, x["adeg"])
            x["margine"] = round(x["z"] - (x["x"] + x["y"]), 6)
            m = x["margine"]
            x["urgenza"] = ("AGIRE ORA: X+Y > Z" if m < 0 else "Agire entro l'anno" if m < 1 else
                            f"Margine {it_breve(arrotonda(m, 1))} {'anno' if arrotonda(m, 1) == 1 else 'anni'}")
            x["urg"] = "agire" if m < 0 else "entro" if m < 1 else "margine"
            # KRI del ciclo
            kri = [k for k in self.kri if k["ID_Rischio"] == rid]
            x["kri_r"] = sum(k["Stato"] == "Rosso" for k in kri)
            x["kri_a"] = sum(k["Stato"] == "Ambra" for k in kri)
            x["kri_v"] = sum(k["Stato"] == "Verde" for k in kri)
            x["kri_n"] = len(kri)
            # Segnali netti delle ultime 4 settimane (Registro, colonna Q)
            x["netti"] = sum(1 if n["Segnale"] == "Peggiorativo" else -1 for n in self.finestra
                             if n["Rischio_ID"] == rid and n["Segnale"] in ("Peggiorativo", "Migliorativo"))
            # Livello di allerta (Radar_Mosca, colonne U-AB; D-63)
            f1, f2, f3 = int(m < 0), int(x["kri_r"] > 0), int(x["netti"] >= p["sogliaDirezione"])
            a1 = int(x["pxi"] >= p["livelloPxI"])
            a2 = int(x["presidio"] == "Assente")
            a3 = int(x["kri_n"] > 0 and x["kri_r"] + x["kri_a"] >= p["livelloQuotaKRI"] * x["kri_n"])
            forti, aggr = f1 + f2 + f3, a1 + a2 + a3
            x["livello"] = ("Attivazione" if forti >= p["livelloForti"] else
                            "Attenzione" if forti >= 1 or (m < p["livelloMargine"] and aggr >= 1) else "Monitoraggio")
            if x["livello"] == "Monitoraggio":
                x["motivo"] = "nessun segnale forte"
            else:
                parti = [t for f, t in ((f1, "AGIRE ORA"), (f2, "KRI in rosso"), (f3, "notizie in peggioramento")) if f]
                if forti == 0:
                    agg = [t for f, t in ((a1, "P x I alta"), (a2, "presidio dei controlli assente"),
                                          (a3, "KRI in ambra o rosso")) if f]
                    parti.append("margine sotto l'anno con " + ", ".join(agg))
                x["motivo"] = "; ".join(parti)
            # Posizione nel radar (Radar_Mosca, colonne M-S)
            x["settore"] = self.codici.index(x["cat"]) if x["cat"] in self.codici else None
            x["n_settore"] = conta_settore.get(x["cat"], 0)
            conta_settore[x["cat"]] = x["n_settore"] + 1
            self.rischi.append(x)
        for x in self.rischi:
            tot = conta_settore[x["cat"]]
            if x["settore"] is not None:
                x["angolo"] = 2 * math.pi * (x["settore"] + (x["n_settore"] + 0.5) / tot) / len(self.codici)
                x["raggio"] = max(RAGGIO_MIN, min(RAGGIO_MAX, x["margine"] + SCOSTAMENTO_RAGGIO))
        self.classifica = sorted(self.rischi, key=lambda x: (-x["score"], x["pos"]))
        for k, x in enumerate(self.classifica, start=1):
            x["rank"] = k
        self.lente = sorted(self.rischi, key=lambda x: (x["margine"], x["rank"]))
        self.per_rid = {x["id"]: x for x in self.rischi}

    # ------------------------------------------------ altre viste
    def rischi_oss(self):
        out = []
        for rid in self.oss_ids:
            r = self.per_id[rid]
            out.append({"id": rid, "riga": r, "nome": r["Nome"], "cat": r["Event_Type"], "I": r["Impatto"],
                        "P": r["Probabilita"], "Pr": r["Prossimita"], "V": r["Velocity"], "score": r["Score_Composito"],
                        "motivo": r.get("Motivo_Stato"), "orizzonte": r.get("Orizzonte_Anni"),
                        "segnale": r.get("Segnale_Avvicinamento"), "driver": r.get("Driver"),
                        "notizie": sum(n["Rischio_ID"] == rid for n in self.finestra)})
        return out

    def movimenti(self):
        return [r for r in self.rr if r.get("Movimento")]

    def notizie_per_categoria(self):
        out = []
        oss_cat = {self.per_id[i]["Event_Type"] for i in self.oss_ids}
        for cod in self.codici:
            col = "Imp_" + cod.replace(".", "_")
            vals = [n.get(col) or 0 for n in self.finestra]
            rischi = sum(a["Event_Type"] == cod for a in self.anag)
            forti = sum(v >= 2 for v in vals)
            presidio = ("Presidiata" if rischi else "In osservazione" if cod in oss_cat else
                        "AREA NON PRESIDIATA" if forti else "Nessun rischio censito")
            out.append({"cod": cod, "nome": self.nomi_cat.get(cod, ""), "max": max(vals, default=0), "forti": forti,
                        "rischi": rischi, "presidio": presidio, "i1": vals.count(1), "i2": vals.count(2),
                        "i3": vals.count(3)})
        return out

    def pressione(self):
        out = []
        for s in self.settimane12:
            ns = [n for n in self.news if n["Settimana"] == s]
            peg = sum(n["Segnale"] == "Peggiorativo" for n in ns)
            mig = sum(n["Segnale"] == "Migliorativo" for n in ns)
            out.append({"settimana": s, "notizie": len(ns), "peg": peg, "mig": mig, "netto": peg - mig})
        return out

    def piano(self):
        out = []
        for a in self.azioni:
            if not a.get("ID_Azione"):
                continue
            dec = a_data(a.get("Data_Decisione"))
            scad = piu_mesi(dec, a.get("Mesi_Orizzonte") or 0) if dec and a.get("Mesi_Orizzonte") else None
            out.append(dict(a, Scadenza=scad.isoformat() if scad else None))
        return out

    def stato_validazione(self):
        da_decidere = {(v["Oggetto"], v["ID"]) for v in self.val if v.get("Stato") == "Da decidere"}
        rettifiche = sum(v["Decisione"] == "Rettificato" and v.get("Stato") == "Applicata" for v in self.val)
        return len(da_decidere), rettifiche

    def da_validare_news(self):
        return [n for n in self.news if n.get("Validazione") == "Da validare"]

    def etichetta(self):
        tipo = "Ciclo straordinario" if "-S" in self.ciclo else "Ciclo trimestrale"
        n, _ = self.stato_validazione()
        return f"{tipo} {self.ciclo} · dati al {data_it(self.data_rif)} · {'validato' if n == 0 else 'NON VALIDATO'}"


def perche_segnalato(v, x):
    """Testo "Perche' lo segnaliamo" del Quadro (formula di foglio_quadro, colonna W)."""
    p, t = v.p, []
    if x["rank"] <= 3:
        d = x["delta"]
        t.append(f"{x['rank']}° per score composito ({it(x['score'])}"
                 + (f", {segnato(d)} sul ciclo prima" if d not in (None, 0) else "") + ")")
    if x["urg"] == "agire":
        t.append(f"AGIRE ORA: servono X + Y = {it(x['x'] + x['y'])} anni per decidere e mitigare"
                 + (f", di cui {it_breve(x['adeg'])} di adeguamento strutturale" if x["adeg"] else "")
                 + f"; ne restano Z = {it(x['z'], 1)} prima che possa manifestarsi")
    elif x["rank"] <= 3 and x["urg"] == "entro":
        t.append(f"margine di Mosca sotto l'anno ({it(x['margine'])})")
    if x["esito"] == "Variato":
        d = x["delta"]
        t.append("variato nel ciclo" + (": peggiora" if d and d > 0 else ": migliora" if d and d < 0 else ""))
    elif x["esito"] == "Nuovo":
        t.append("nuovo nel ciclo")
    if x["sup"]:
        t.append("tema all'attenzione del supervisore")
    if x["presidio"] == "Assente":
        t.append("presidio dei controlli assente")
    if x["kri_r"]:
        t.append(f"{x['kri_r']} KRI in rosso")
    if x["netti"] >= p["sogliaDirezione"]:
        t.append(f"notizie in peggioramento (+{x['netti']} segnali netti in 4 settimane)")
    elif x["netti"] <= -p["sogliaDirezione"]:
        t.append(f"notizie in miglioramento ({x['netti']} segnali netti in 4 settimane)")
    return "; ".join(t) + "." if t else ""


# ================================================================ confronto con i valori calcolati da Excel
def verifica(reg, v):
    """Il cancello (ciclo corrente): i numeri calcolati qui contro quelli calcolati dalle formule del registro.
    Restituisce (uguali, diversi)."""
    ids = reg.nome("rmID")
    excel = {n: dict(zip(ids, reg.nome(n))) for n in ("rmZ", "rmX", "rmY", "rmMargine", "rmUrgenza", "rmLivello",
                                                      "rmMotivoLivello")}
    netti = dict(zip(reg.nome("posID"), reg.nome("posNetti")))
    campi = [("rmZ", "z"), ("rmX", "x"), ("rmY", "y"), ("rmMargine", "margine"), ("rmUrgenza", "urgenza"),
             ("rmLivello", "livello"), ("rmMotivoLivello", "motivo")]
    uguali, diversi = 0, []
    for x in v.rischi:
        for n, k in campi:
            a, b = x[k], excel[n].get(x["id"])
            ok = abs(a - b) < 1e-6 if isinstance(a, (int, float)) and isinstance(b, (int, float)) else a == b
            uguali += ok
            if not ok:
                diversi.append(f"{x['id']} {n}: Python {a!r}, Excel {b!r}")
        ok = x["netti"] == netti.get(x["id"])
        uguali += ok
        if not ok:
            diversi.append(f"{x['id']} posNetti: Python {x['netti']!r}, Excel {netti.get(x['id'])!r}")
    ordine_excel = [reg.cella("Classifica", 5 + i, 2) for i in range(len(v.rischi))]
    ok = ordine_excel == [x["id"] for x in v.classifica]
    uguali += ok
    if not ok:
        diversi.append(f"classifica: Python {[x['id'] for x in v.classifica]}, Excel {ordine_excel}")
    for i in range(len(v.rischi)):
        rid, d = reg.cella("Delta", 31 + i, 1), reg.cella("Delta", 31 + i, 7)
        ok = (rid in v.per_rid and v.per_rid[rid]["delta"] is not None and isinstance(d, (int, float))
              and abs(v.per_rid[rid]["delta"] - d) < 1e-6)
        uguali += ok
        if not ok:
            diversi.append(f"{rid} Delta: Python {v.per_rid.get(rid, {}).get('delta')!r}, Excel {d!r}")
    return uguali, diversi


# ================================================================ SVG
def svg_apri(w, h, etichetta, classe_svg="grafico"):
    return (f'<svg class="{classe_svg}" viewBox="0 0 {w} {h}" role="img" aria-label="{esc(etichetta)}" '
            f'xmlns="http://www.w3.org/2000/svg">')


def barra_oriz(x0, y, lung, alt, classe_barra, raggio=4):
    """Barra orizzontale: estremo dei dati arrotondato, base squadrata."""
    if lung <= 0:
        return ""
    r = min(raggio, lung, alt / 2)
    return (f'<path class="{classe_barra}" d="M{x0:.1f},{y:.1f} h{lung - r:.1f} a{r},{r} 0 0 1 {r},{r} '
            f'v{alt - 2 * r:.1f} a{r},{r} 0 0 1 -{r},{r} h-{lung - r:.1f} z"/>')


def barra_vert(x, y_base, alt, larg, classe_barra, arrotonda=True, raggio=4):
    if alt <= 0:
        return ""
    if not arrotonda:
        return f'<rect class="{classe_barra}" x="{x:.1f}" y="{y_base - alt:.1f}" width="{larg:.1f}" height="{alt:.1f}"/>'
    r = min(raggio, alt, larg / 2)
    return (f'<path class="{classe_barra}" d="M{x:.1f},{y_base:.1f} v-{alt - r:.1f} a{r},{r} 0 0 1 {r},-{r} '
            f'h{larg - 2 * r:.1f} a{r},{r} 0 0 1 {r},{r} v{alt - r:.1f} z"/>')


def grafico_radar(v, piccolo=False):
    """Radar di Mosca (rubrica, sezione 8): distanza = margine + 2, area = P x I, settore = categoria, colore = livello."""
    s = 34 if piccolo else 44                 # pixel per unita' di raggio
    lato = 2 * (RAGGIO_MAX + 1.9) * s
    c = lato / 2
    out = [svg_apri(lato, lato + (0 if piccolo else 70), "Radar di Mosca del ciclo " + v.ciclo, "grafico radar")]
    n = len(v.codici)
    for k in range(n):                        # raggi dei settori
        a = 2 * math.pi * k / n
        out.append(f'<line class="asse" x1="{c}" y1="{c}" x2="{c + (RAGGIO_MAX + .3) * s * math.sin(a):.1f}" '
                   f'y2="{c - (RAGGIO_MAX + .3) * s * math.cos(a):.1f}"/>')
    for margine, cl, testo in ((3, "anello", "3 anni"), (2, "anello", "2 anni"), (1, "anello tratto", "1 anno"),
                               (0, "mosca", "Mosca")):
        r = (margine + SCOSTAMENTO_RAGGIO) * s
        out.append(f'<circle class="{cl}" cx="{c}" cy="{c}" r="{r}"/>')
        if not piccolo:
            out.append(f'<text class="{"et-mosca" if margine == 0 else "et-anello"}" x="{c + r * math.sin(1.4) + 4:.1f}" '
                       f'y="{c - r * math.cos(1.4) - 4:.1f}">{testo}</text>')
    for k, cod in enumerate(v.codici):        # etichette dei settori
        a = 2 * math.pi * (k + .5) / n
        rx, ry = c + (RAGGIO_MAX + 1.15) * s * math.sin(a), c - (RAGGIO_MAX + 1.15) * s * math.cos(a)
        nome = NOMI_SETTORI.get(cod, v.nomi_cat.get(cod, ""))
        out.append(f'<text class="et-settore" text-anchor="middle" x="{rx:.1f}" y="{ry - 4:.1f}">'
                   f'<tspan class="cod">{esc(cod)}</tspan><tspan x="{rx:.1f}" dy="15">{esc(nome)}</tspan></text>')
    k_area = (8.6 if piccolo else 10.5)
    etichette, occupate = [], []
    for x in sorted(v.rischi, key=lambda x: -x["pxi"]):  # i grandi sotto, i piccoli sopra
        if x.get("angolo") is None:
            continue
        px = c + x["raggio"] * s * math.sin(x["angolo"])
        py = c - x["raggio"] * s * math.cos(x["angolo"])
        r = k_area * math.sqrt(x["pxi"])
        etich = x["id"][4:] if piccolo else x["id"]
        tip = (f"{x['id']} {x['nome']}\n{x['livello']} — {x['motivo']}\n{x['urgenza']}: margine {it(x['margine'])} anni "
               f"(Z {it(x['z'], 1)} − X {it(x['x'])} − Y {it(x['y'])})\nP x I = {x['pxi']} (P {x['P']}, I {x['I']}) · score "
               f"{it(x['score'])}" + ("\n◆ attenzione del supervisore" if x["sup"] else ""))
        out.append(f'<a href="#scheda-{x["id"]}" class="punto"><g><title>{esc(tip)}</title>'
                   f'<circle class="bolla liv-{classe(x["livello"])}" cx="{px:.1f}" cy="{py:.1f}" r="{r:.1f}"/></g></a>')
        # Etichette dopo tutti i cerchi, cosi' nessun cerchio le copre; due etichette che si toccano: la seconda scende.
        ly = py + 4
        larg = (24 if piccolo else 50) + (12 if x["sup"] else 0)
        while any(abs(px - ox) < (larg + ol) / 2 and abs(ly - oy) < 13 for ox, oy, ol in occupate):
            ly += 13
        occupate.append((px, ly, larg))
        etichette.append(f'<text class="et-bolla" text-anchor="middle" x="{px:.1f}" y="{ly:.1f}">{esc(etich)}'
                         f'{" ◆" if x["sup"] else ""}</text>')
    out.extend(etichette)
    if not piccolo:                           # legenda sotto il radar
        y = lato + 22
        xs = 20
        for liv in ("Attivazione", "Attenzione", "Monitoraggio"):
            out.append(f'<circle class="bolla liv-{classe(liv)}" cx="{xs + 9}" cy="{y}" r="9"/>'
                       f'<text class="et-leg" x="{xs + 24}" y="{y + 4}">{liv}</text>')
            xs += 125
        out.append(f'<line class="mosca" x1="{xs}" y1="{y}" x2="{xs + 26}" y2="{y}"/>'
                   f'<text class="et-leg" x="{xs + 32}" y="{y + 4}">dentro il cerchio rosso: AGIRE ORA (X + Y &gt; Z)</text>')
        y2 = y + 34
        out.append(f'<text class="et-leg" x="20" y="{y2 + 4}">Area del cerchio = P x I:</text>')
        xs = 190
        for val in (4, 12, 25):
            r = k_area * math.sqrt(val) * .55
            out.append(f'<circle class="campione" cx="{xs + r}" cy="{y2}" r="{r:.1f}"/>'
                       f'<text class="et-leg" x="{xs + 2 * r + 6:.1f}" y="{y2 + 4}">{val}</text>')
            xs += 2 * r + 34
    out.append("</svg>")
    return "".join(out)


def grafico_classifica(v):
    righe = v.classifica
    sx, w_plot, alt_riga = 300, 420, 30
    h = 40 + alt_riga * len(righe) + 26
    out = [svg_apri(sx + w_plot + 70, h, "Classifica per score composito")]
    y0 = 30
    for k in range(6):
        x = sx + w_plot * k / 5
        out.append(f'<line class="griglia" x1="{x:.1f}" y1="{y0 - 6}" x2="{x:.1f}" y2="{y0 + alt_riga * len(righe)}"/>'
                   f'<text class="et-asse" text-anchor="middle" x="{x:.1f}" y="{y0 - 12}">{it(k)}</text>')
    for i, x in enumerate(righe):
        y = y0 + i * alt_riga
        freccia = "nuovo" if x["delta"] is None else "▲" if x["delta"] > 0 else "▼" if x["delta"] < 0 else "●"
        tip = (f"{x['id']} {x['nome']}\nscore {it(x['score'])} ({segnato(x['delta'])} sul ciclo prima)\n{x['urgenza']}")
        out.append(f'<a href="#scheda-{x["id"]}" class="punto"><g><title>{esc(tip)}</title>'
                   f'<rect class="hit" x="0" y="{y}" width="{sx + w_plot + 70}" height="{alt_riga}"/>'
                   f'<text class="et-cat" text-anchor="end" x="{sx - 10}" y="{y + 19}"><tspan class="id">{x["id"]}</tspan> '
                   f'<tspan class="freccia">{freccia}</tspan> {esc(accorcia(x["nome"], 30))}</text>'
                   + barra_oriz(sx, y + 6, w_plot * x["score"] / 5, 18, f"b-urg-{x['urg']}")
                   + f'<text class="et-val" x="{sx + w_plot * x["score"] / 5 + 6:.1f}" y="{y + 19}">{it(x["score"])}</text></g></a>')
    out.append(f'<line class="asse-base" x1="{sx}" y1="{y0 - 6}" x2="{sx}" y2="{y0 + alt_riga * len(righe)}"/>')
    out.append("</svg>")
    return "".join(out)


def grafico_lente(v):
    righe = v.lente
    larg_col, sx, alt = 58, 46, 300
    massimo = max(3.5, math.ceil(max(max(x["x"] + x["y"], x["z"]) for x in righe) * 2) / 2)
    w = sx + larg_col * len(righe) + 10
    h = alt + 70
    y_base = 20 + alt
    sc = alt / massimo
    out = [svg_apri(w, h, "Lente di Mosca: chi e' in ritardo, e di quanto")]
    passo = 0.5 if massimo <= 4 else 1
    t = 0
    while t <= massimo + 1e-9:
        y = y_base - t * sc
        out.append(f'<line class="griglia" x1="{sx}" y1="{y:.1f}" x2="{w - 10}" y2="{y:.1f}"/>'
                   f'<text class="et-asse" text-anchor="end" x="{sx - 6}" y="{y + 4:.1f}">{it_breve(t)}</text>')
        t += passo
    for i, x in enumerate(righe):
        cx = sx + i * larg_col + 12
        bw = larg_col - 24
        entro = min(x["y"], max(x["z"] - x["x"], 0))
        oltre = x["y"] - entro
        hx, he, ho = x["x"] * sc, entro * sc, oltre * sc
        tip = (f"{x['id']} {x['nome']}\nX {it(x['x'])} + Y {it(x['y'])} = {it(x['x'] + x['y'])} anni; Z {it(x['z'], 1)}\n"
               f"margine {segnato(x['margine'])} anni · {x['urgenza']}")
        seg = [barra_vert(cx, y_base, hx, bw, "b-x", arrotonda=False)]
        seg.append(barra_vert(cx, y_base - hx - 2, he - 2, bw, "b-y-entro", arrotonda=ho <= 0) if he > 2 else "")
        if ho > 0:
            seg.append(barra_vert(cx, y_base - hx - he - 2, ho - 2, bw, "b-y-oltre"))
        zy = y_base - x["z"] * sc
        rombo = f'<path class="rombo" d="M{cx + bw / 2:.1f},{zy - 7:.1f} l7,7 l-7,7 l-7,-7 z"/>'
        out.append(f'<a href="#scheda-{x["id"]}" class="punto"><g><title>{esc(tip)}</title>'
                   f'<rect class="hit" x="{cx - 10}" y="16" width="{bw + 20}" height="{alt + 50}"/>' + "".join(seg) + rombo
                   + f'<text class="et-asse" text-anchor="middle" x="{cx + bw / 2:.1f}" y="{y_base + 16}">{x["id"]}</text>'
                   f'<text class="et-asse tenue" text-anchor="middle" x="{cx + bw / 2:.1f}" y="{y_base + 31}">'
                   f'({segnato(x["margine"])})</text></g></a>')
    out.append(f'<line class="asse-base" x1="{sx}" y1="{y_base}" x2="{w - 10}" y2="{y_base}"/>')
    out.append(f'<text class="et-asse" transform="rotate(-90)" x="{-(20 + alt / 2)}" y="12" text-anchor="middle">anni</text>')
    out.append("</svg>")
    return "".join(out)


def grafico_delta(v):
    righe = sorted(v.rischi, key=lambda x: (-(x["delta"] if x["delta"] is not None else -9), x["pos"]))
    sx, w_plot, alt_riga, lim = 290, 440, 28, 1.5
    h = 36 + alt_riga * len(righe) + 10
    zero = sx + w_plot / 2
    sc = (w_plot / 2) / lim
    out = [svg_apri(sx + w_plot + 20, h, "Delta dello score composito sul ciclo precedente")]
    y0 = 28
    for t in (-1.5, -1, -0.5, 0, 0.5, 1, 1.5):
        x = zero + t * sc
        out.append(f'<line class="{"asse-base" if t == 0 else "griglia"}" x1="{x:.1f}" y1="{y0 - 6}" x2="{x:.1f}" '
                   f'y2="{y0 + alt_riga * len(righe)}"/><text class="et-asse" text-anchor="middle" x="{x:.1f}" '
                   f'y="{y0 - 12}">{segnato(t) if t else "0"}</text>')
    for i, x in enumerate(righe):
        y = y0 + i * alt_riga
        d = x["delta"]
        tip = f"{x['id']} {x['nome']}\nΔ score {segnato(d)}" + ("" if d is not None else " (nuovo)")
        corpo = ""
        if d:
            lung = min(abs(d), lim) * sc
            if d > 0:
                corpo = barra_oriz(zero, y + 6, lung, 16, "b-su")
                corpo += f'<text class="et-val" x="{zero + lung + 6:.1f}" y="{y + 18}">{segnato(d)}</text>'
            else:
                corpo = (f'<g transform="translate({2 * zero:.1f},0) scale(-1,1)">{barra_oriz(zero, y + 6, lung, 16, "b-giu")}</g>'
                         f'<text class="et-val" text-anchor="end" x="{zero - lung - 6:.1f}" y="{y + 18}">{segnato(d)}</text>')
        out.append(f'<a href="#scheda-{x["id"]}" class="punto"><g><title>{esc(tip)}</title>'
                   f'<rect class="hit" x="0" y="{y}" width="{sx + w_plot}" height="{alt_riga}"/>'
                   f'<text class="et-cat" text-anchor="end" x="{sx - 10}" y="{y + 18}"><tspan class="id">{x["id"]}</tspan> '
                   f'{esc(accorcia(x["nome"], 30))}</text>{corpo}</g></a>')
    out.append("</svg>")
    return "".join(out)


def sparkline(valori, etichette, larg=150, alt=40, minimo=1, massimo=5, classe_linea="linea"):
    """Andamento su scala fissa: linea 2px, punto finale, tooltip per punto."""
    punti = [(i, val) for i, val in enumerate(valori) if val is not None]
    if not punti:
        return '<span class="tenue">—</span>'
    n = max(len(valori) - 1, 1)
    X = lambda i: 4 + (larg - 8) * i / n
    Y = lambda val: alt - 4 - (alt - 8) * (val - minimo) / (massimo - minimo)
    out = [f'<svg class="spark" viewBox="0 0 {larg} {alt}" width="{larg}" height="{alt}" role="img" '
           f'aria-label="andamento">']
    out.append(f'<line class="griglia" x1="0" y1="{alt - 4}" x2="{larg}" y2="{alt - 4}"/>')
    seg = []
    for i, val in punti:
        seg.append(f"{'L' if seg else 'M'}{X(i):.1f},{Y(val):.1f}")
    if len(punti) > 1:
        out.append(f'<path class="{classe_linea}" d="{" ".join(seg)}"/>')
    for k, (i, val) in enumerate(punti):
        fine = k == len(punti) - 1
        out.append(f'<g class="punto"><title>{esc(etichette[i])}: {it(val)}</title>'
                   f'<rect class="hit" x="{X(i) - 8:.1f}" y="0" width="16" height="{alt}"/>'
                   + (f'<circle class="pallino-fine" cx="{X(i):.1f}" cy="{Y(val):.1f}" r="4"/>' if fine else
                      f'<circle class="pallino-hover" cx="{X(i):.1f}" cy="{Y(val):.1f}" r="3.5"/>') + "</g>")
    out.append("</svg>")
    return "".join(out)


def passo_tondo(x):
    """Passo d'asse tondo (1, 2, 2,5, 5 per una potenza di dieci) vicino a x."""
    if x <= 0:
        return 1
    e = 10 ** math.floor(math.log10(x))
    return next(m * e for m in (1, 2, 2.5, 5, 10) if m * e >= x)


def grafico_kri(storia, cicli, nome, unita):
    """Valore di un KRI negli ultimi cicli con le due soglie (linee di soglia tratteggiate)."""
    larg, alt, sx = 300, 120, 34
    vals = [s["valore"] for s in storia if s["valore"] is not None]
    soglie = [s[k] for s in storia for k in ("att", "all") if s[k] is not None]
    if not vals and not soglie:
        return '<p class="tenue piccolo">Nessun valore rilevato.</p>'
    passo = passo_tondo(max(vals + soglie) / 3)
    massimo = passo * math.ceil(max(vals + soglie) * 1.08 / passo) or 1
    n = max(len(cicli) - 1, 1)
    X = lambda i: sx + 8 + (larg - sx - 18) * i / n
    Y = lambda val: alt - 22 - (alt - 34) * val / massimo
    out = [svg_apri(larg, alt, f"Storia del KRI {nome}", "grafico kri")]
    for t in [k * passo for k in range(int(round(massimo / passo)) + 1)]:
        out.append(f'<line class="griglia" x1="{sx}" y1="{Y(t):.1f}" x2="{larg - 6}" y2="{Y(t):.1f}"/>'
                   f'<text class="et-asse" text-anchor="end" x="{sx - 4}" y="{Y(t) + 4:.1f}">{it_breve(round(t, 1))}</text>')
    for k, cl in (("att", "soglia-att"), ("all", "soglia-all")):
        pts = [(i, s[k]) for i, s in enumerate(storia) if s[k] is not None]
        if pts:
            out.append(f'<path class="{cl}" d="{" ".join(("M" if j == 0 else "L") + f"{X(i):.1f},{Y(val):.1f}" for j, (i, val) in enumerate(pts))}"/>')
    pts = [(i, s["valore"]) for i, s in enumerate(storia) if s["valore"] is not None]
    if len(pts) > 1:
        out.append(f'<path class="linea" d="{" ".join(("M" if j == 0 else "L") + f"{X(i):.1f},{Y(val):.1f}" for j, (i, val) in enumerate(pts))}"/>')
    for j, (i, val) in enumerate(pts):
        s = storia[i]
        tip = (f"{cicli[i]}: {it_breve(val)} {unita}\nsoglie {it_breve(s['att'])} / {it_breve(s['all'])} · {s['stato']}")
        out.append(f'<g class="punto"><title>{esc(tip)}</title><rect class="hit" x="{X(i) - 10:.1f}" y="0" width="20" '
                   f'height="{alt}"/><circle class="{"pallino-fine" if j == len(pts) - 1 else "pallino"}" cx="{X(i):.1f}" '
                   f'cy="{Y(val):.1f}" r="{4 if j == len(pts) - 1 else 3}"/></g>')
    for i, c in enumerate(cicli):
        if len(cicli) <= 5 or i in (0, len(cicli) - 1):
            out.append(f'<text class="et-asse" text-anchor="middle" x="{X(i):.1f}" y="{alt - 6}">{esc(c)}</text>')
    out.append("</svg>")
    return "".join(out)


def grafico_notizie_cat(cat):
    sx, w_plot, alt_riga = 175, 250, 28
    massimo = max(4, max(c["i1"] + c["i2"] + c["i3"] for c in cat))
    sc = w_plot / massimo
    h = 32 + alt_riga * len(cat) + 6
    out = [svg_apri(sx + w_plot + 30, h, "Notizie delle ultime 4 settimane per categoria e intensita'")]
    y0 = 26
    for t in range(massimo + 1):
        x = sx + t * sc
        out.append(f'<line class="griglia" x1="{x:.1f}" y1="{y0 - 6}" x2="{x:.1f}" y2="{y0 + alt_riga * len(cat)}"/>'
                   f'<text class="et-asse" text-anchor="middle" x="{x:.1f}" y="{y0 - 11}">{t}</text>')
    for i, c in enumerate(cat):
        y = y0 + i * alt_riga
        x = sx
        corpo = []
        for k in (1, 2, 3):
            n = c[f"i{k}"]
            if n:
                l = n * sc - (2 if k < 3 and any(c[f"i{j}"] for j in range(k + 1, 4)) else 0)
                corpo.append(f'<rect class="b-int{k}" x="{x:.1f}" y="{y + 6}" width="{max(l, 1):.1f}" height="16" rx="2"/>')
                x += n * sc
        tip = f"{c['cod']} {c['nome']}: intensita' 1 = {c['i1']}, 2 = {c['i2']}, 3 = {c['i3']} · {c['presidio']}"
        out.append(f'<g class="punto"><title>{esc(tip)}</title><rect class="hit" x="0" y="{y}" width="{sx + w_plot}" '
                   f'height="{alt_riga}"/><text class="et-cat" text-anchor="end" x="{sx - 10}" y="{y + 18}">'
                   f'<tspan class="id">{c["cod"]}</tspan> {esc(c["nome"])}</text>{"".join(corpo)}</g>')
    out.append(f'<line class="asse-base" x1="{sx}" y1="{y0 - 6}" x2="{sx}" y2="{y0 + alt_riga * len(cat)}"/>')
    out.append("</svg>")
    return "".join(out)


def grafico_pressione(press):
    if not press:
        return '<p class="tenue">Nessuna settimana di notizie prima di questo ciclo.</p>'
    larg_col, sx, alt = 52, 40, 150
    massimo = max(2, max(max(p["peg"], p["mig"]) for p in press))
    w = sx + larg_col * len(press) + 10
    zero = 16 + alt / 2
    sc = (alt / 2 - 8) / massimo
    out = [svg_apri(w, alt + 52, "Segnali per settimana, ultime 12 settimane")]
    for t in range(-massimo, massimo + 1):
        y = zero - t * sc
        out.append(f'<line class="{"asse-base" if t == 0 else "griglia"}" x1="{sx}" y1="{y:.1f}" x2="{w - 10}" y2="{y:.1f}"/>'
                   f'<text class="et-asse" text-anchor="end" x="{sx - 6}" y="{y + 4:.1f}">{t}</text>')
    for i, p in enumerate(press):
        cx = sx + i * larg_col + 14
        bw = larg_col - 28
        tip = (f"Settimana del {data_it(p['settimana'])}: {p['notizie']} notizie, {p['peg']} peggiorative, "
               f"{p['mig']} migliorative, netto {p['netto']:+d}")
        out.append(f'<g class="punto"><title>{esc(tip)}</title><rect class="hit" x="{cx - 12}" y="0" width="{bw + 24}" '
                   f'height="{alt + 40}"/>' + barra_vert(cx, zero - 1, p["peg"] * sc, bw, "b-su")
                   + (f'<g transform="translate(0,{2 * zero:.1f}) scale(1,-1)">{barra_vert(cx, zero - 1, p["mig"] * sc, bw, "b-giu")}</g>'
                      if p["mig"] else "")
                   + f'<text class="et-asse" text-anchor="middle" x="{cx + bw / 2:.1f}" y="{alt + 34}">'
                   f'{data_it(p["settimana"])[:5]}</text></g>')
    out.append("</svg>")
    return "".join(out)


# ================================================================ pezzi di HTML
def chip(testo, cl):
    return f'<span class="chip {cl}">{esc(testo)}</span>'


def chip_livello(liv):
    return chip(liv or "—", f"liv-{classe(liv)}") if liv else '<span class="tenue">—</span>'


def chip_urgenza(x):
    breve = {"agire": "AGIRE ORA", "entro": "Entro l'anno"}.get(x["urg"], x["urgenza"])
    return chip(breve, f"urg-{x['urg']}")


def chip_kri(stato):
    return chip(stato or "—", "kri-" + classe(stato or "nd"))


def cella_dim(v):
    return f'<td class="dim"><span class="dim-v dim-{v if v else "nv"}">{v if v is not None else "n.v."}</span></td>'


def link_rischio(rid):
    return f'<a class="id" href="#scheda-{esc(rid)}">{esc(rid)}</a>'


def tabella(intest, righe, cl="", allinea=None):
    allinea = allinea or {}
    t = [f'<div class="tab-scorre"><table class="{cl}"><thead><tr>']
    t += [f'<th class="{allinea.get(i, "")}">{h}</th>' for i, h in enumerate(intest)]
    t.append("</tr></thead><tbody>")
    for r in righe:
        if isinstance(r, str):
            t.append(r)
            continue
        t.append("<tr>" + "".join(c if str(c).startswith("<td") else f'<td class="{allinea.get(i, "")}">{c}</td>'
                                  for i, c in enumerate(r)) + "</tr>")
    t.append("</tbody></table></div>")
    return "".join(t)


def sezione(id_, titolo, gruppo, corpo, come="", legenda="", classe_extra=""):
    return (f'<section id="{id_}" class="tab {classe_extra}" aria-labelledby="t-{id_}">'
            f'<header class="tab-testa"><p class="gruppo">{esc(gruppo)}</p><h2 id="t-{id_}">{esc(titolo)}</h2>'
            + (f'<p class="come">{esc(come)}</p>' if come else "") + "</header>" + corpo
            + (f'<p class="legenda">{esc(legenda)}</p>' if legenda else "") + "</section>")


def come_si_legge(testo):
    """Seconda riga dei fogli: 'PROCESSO — chi produce.\nCome si legge: ...' -> la parte dopo 'Come si legge'."""
    t = str(testo or "")
    m = re.search(r"Come si legge:?\s*(.*)", t, re.S)
    return (m.group(1) if m else t).strip()


def parti_commento(testo):
    t = str(testo or "")
    m = re.match(r"\s*MESSAGGIO:\s*(.*?)\s*(?:DECISIONI:\s*(.*))?$", t, re.S)
    maiuscola = lambda x: (x[0].upper() + x[1:]) if x else x
    return (maiuscola(m.group(1)), maiuscola(m.group(2))) if m else (t, None)


# ================================================================ le sezioni
def s_quadro(v, testi):
    n_da, rett = v.stato_validazione()
    att = [x for x in v.classifica if x["livello"] == "Attivazione"]
    attz = [x for x in v.rischi if x["livello"] == "Attenzione"]
    mon = [x for x in v.rischi if x["livello"] == "Monitoraggio"]
    var = sum(x["esito"] == "Variato" for x in v.rischi)
    nuovi = sum(x["esito"] == "Nuovo" for x in v.rischi)
    mov = v.movimenti()
    entr = sum(str(r.get("Movimento")).startswith("Entrat") for r in mov)
    usc = sum(str(r.get("Movimento")).startswith("Uscit") for r in mov)
    rossi = sum(k["Stato"] == "Rosso" for k in v.kri)
    tiles = [
        ("Rischi valutati", f"{len(v.rischi)}", f"su {len(v.anag)} in anagrafica · "
         + ("completo" if len(v.rischi) == len(v.anag) else "INCOMPLETO"), ""),
        ("Livello di allerta: attivazione · attenzione", f"{len(att)} · {len(attz)}",
         f"in attivazione: {', '.join(x['id'] for x in att) or 'nessuno'} · in monitoraggio {len(mon)}",
         "tile-allerta" if att else ""),
        ("Variati · nuovi (radar)", f"{var} · {nuovi}", f"rispetto al {v.prec or '—'} · entrati nel radar {entr}, usciti {usc}"
         f" · in osservazione {len(v.oss_ids)}", ""),
        ("KRI in rosso", f"{rossi}", f"su {len(v.kri)} in vigore", "tile-allerta" if rossi else ""),
        ("Validazione", "✓" if n_da == 0 else f"{n_da}", "ciclo validato" if n_da == 0 else "righe da validare",
         "tile-ok" if n_da == 0 else "tile-attesa"),
    ]
    h = ['<div class="tiles">']
    for etich, val, sotto, cl in tiles:
        h.append(f'<div class="tile {cl}"><p class="tile-et">{esc(etich)}</p><p class="tile-val">{esc(val)}</p>'
                 f'<p class="tile-sotto">{esc(sotto)}</p></div>')
    h.append("</div>")
    straord = v.sintesi.get("Straordinario")
    if straord:
        h.append(f'<div class="riquadro avviso"><strong>Ciclo straordinario.</strong> {esc(straord)}</div>')
    # priorita' e matrice
    righe = []
    for x in v.classifica[:5]:
        righe.append([f'<b>{x["rank"]}</b>', f'{link_rischio(x["id"])} {esc(x["nome"])}{" ◆" if x["sup"] else ""}',
                      f'<span class="barretta"><span style="width:{x["score"] / 5 * 100:.0f}%" class="b-urg-{x["urg"]}">'
                      f'</span></span> <b>{it(x["score"])}</b>', segnato(x["delta"]), chip_urgenza(x)])
    h.append('<div class="griglia-2"><div class="pannello"><h3>Priorita' + "'" + ' — primi cinque per score composito</h3>'
             + tabella(["#", "Rischio", "Score", "Δ", "Urgenza"], righe, "compatta", {0: "num", 3: "num"})
             + '<p class="nota">Score 1-5, 5 = peggiore; barra su scala fissa 0-5. Δ sul ciclo precedente.</p></div>'
             '<div class="pannello"><h3>Gravita' + "'" + ': Probabilita x Impatto</h3>' + matrice_html(v, piccola=True)
             + '<p class="nota">Numeri = NFR-nn. Verde P x I fino a 6, ambra 8-12, rosso da 15.</p></div></div>')
    note = v.sintesi.get("Note_Analista")
    if note:
        h.append('<div class="riquadro"><strong>Note dell' + "'" + f'analista.</strong> {esc(note)}</div>')
    h.append('<div class="griglia-2"><div class="pannello"><h3>Urgenza: radar di Mosca</h3>'
             f'{grafico_radar(v, piccolo=True)}<p class="nota">Piu' + "'" + ' vicino al centro = agire prima; area = P x I; '
             'colore = livello di allerta; ◆ = attenzione del supervisore. Clic su un cerchio per la scheda.</p></div>'
             '<div class="pannello"><h3>I rischi del ciclo</h3><ul class="elenco-rischi">'
             + "".join(f'<li>{chip_livello(x["livello"])} {link_rischio(x["id"])} {esc(x["nome"])}{" ◆" if x["sup"] else ""}</li>'
                       for x in sorted(v.rischi, key=lambda x: x["pos"]))
             + "</ul></div></div>")
    # commento di A1
    mess, dec = parti_commento(v.sintesi.get("Sintesi_A1"))
    a5 = v.sintesi.get("Sintesi_A5") or ""
    oss = v.sintesi.get("Osservazioni_Controlli")
    righe = [("Messaggio chiave", mess), ("Decisioni richieste", dec),
             ("Controlli (A5 se usato)", a5 + (f" Osservazioni: {oss}." if oss not in (None, "") else "")),
             ("Dopo il commento", f"Il validatore ha rettificato {rett} valori." if rett else "Nessuna rettifica dopo il commento."),
             ("Segnalazioni da motivare", v.sintesi.get("Rilievi_Superati"))]
    h.append('<div class="pannello"><h3>Il commento dell' + "'" + 'analisi (agente A1, scritto prima della validazione)</h3>'
             '<dl class="commento">' + "".join(f"<dt>{esc(a)}</dt><dd>{esc(b) or '—'}</dd>" for a, b in righe) + "</dl></div>")
    # segnalati
    seg = [x for x in v.classifica if x["rank"] <= 3 or x["livello"] in ("Attivazione", "Attenzione")
           or x["esito"] in ("Variato", "Nuovo")]
    righe = []
    for x in seg:
        cambiato = "Rischio nuovo nel ciclo." if x["esito"] == "Nuovo" else schede.cosa_cambiato(
            {"esito": x["esito"], "variazioni": x["variazioni"]})
        righe.append([f'{link_rischio(x["id"])}<br><span class="piccolo">{esc(x["nome"])}</span>'
                      + ('<br><span class="sup">◆ attenzione del supervisore</span>' if x["sup"] else ""),
                      chip_livello(x["livello"]), esc(perche_segnalato(v, x)),
                      esc(cambiato).replace("\n", "<br>")])
    h.append('<div class="pannello"><h3>Rischi segnalati e perche' + "'" + '</h3><p class="nota">Fra i primi tre per score, '
             'in attivazione o in attenzione, variati o nuovi nel ciclo. Tutto calcolato dai dati validati.</p>'
             + tabella(["Rischio", "Livello", "Perche' lo segnaliamo", "Cosa e' cambiato, e su quale evidenza"], righe,
                       "segnalati") + "</div>")
    if testi["limiti"]:
        h.append(f'<p class="limiti">{esc(testi["limiti"])}</p>')
    return sezione("quadro", "Quadro d'insieme", "Report al comitato", "".join(h), come_si_legge(testi["come"]),
                   testi["legenda"])


def matrice_html(v, piccola=False):
    celle = {}
    for x in sorted(v.rischi, key=lambda x: x["pos"]):
        celle.setdefault((x["P"], x["I"]), []).append(x)
    t = ['<table class="matrice' + (" piccola" if piccola else "") + '"><tbody>']
    for p in range(5, 0, -1):
        t.append(f'<tr><th class="riga-p">{"" if piccola else NOMI_PROB[p] + " "}<b>{p}</b></th>')
        for i in range(1, 6):
            pxi = p * i
            fascia = "verde" if pxi <= 6 else "ambra" if pxi <= 12 else "rosso"
            dentro = celle.get((p, i), [])
            testo = " ".join(f'<a href="#scheda-{x["id"]}" title="{esc(x["id"] + " " + x["nome"])}">'
                             f'{x["id"][4:] if piccola else x["id"]}{"◆" if x["sup"] and not piccola else ""}</a>'
                             for x in dentro)
            t.append(f'<td class="fascia-{fascia}" title="P {p} x I {i} = {pxi}">{testo}</td>')
        t.append("</tr>")
    t.append('<tr><th></th>' + "".join(f'<th class="col-i"><b>{i}</b>{"" if piccola else "<br>" + NOMI_IMP[i]}</th>'
                                        for i in range(1, 6)) + "</tr></tbody></table>")
    if not piccola:
        t.insert(0, '<p class="asse-matrice">Righe: Probabilita · colonne: Impatto</p>')
    return "".join(t)


def s_classifica(v, testi):
    agire = [x["id"] for x in v.classifica if x["urg"] == "agire"]
    leg = ('<div class="legenda-chip">' + chip("Agire ora (X + Y > Z)", "urg-agire") + chip("Agire entro l'anno", "urg-entro")
           + chip("Con margine", "urg-margine") + "</div>")
    leg_lente = ('<div class="legenda-chip"><span class="sw sw-x"></span> X: decidere e avviare '
                 '<span class="sw sw-y-entro"></span> Y: mitigare, entro Z <span class="sw sw-y-oltre"></span> Y: mitigare, '
                 'OLTRE Z <span class="sw sw-z">◆</span> Z: tempo disponibile</div>')
    righe = [[x["rank"], link_rischio(x["id"]), esc(x["nome"]), f'<b>{it(x["score"])}</b>', it(x["score_prec"]),
              segnato(x["delta"]), esc(x["esito"]), chip_urgenza(x), segnato(x["margine"]), esc(x["owner"])]
             for x in v.classifica]
    corpo = (f'<div class="pannello"><h3>{len(agire)} da affrontare subito{": " + ", ".join(agire) if agire else ""}</h3>'
             f'{leg}{grafico_classifica(v)}</div>'
             f'<div class="pannello"><h3>Lente di Mosca: chi e' + "'" + f' in ritardo, e di quanto (anni)</h3>{leg_lente}'
             f'{grafico_lente(v)}</div><div class="pannello"><h3>Tabella</h3>'
             + tabella(["Pos.", "ID", "Nome", "Score", "Ciclo prima", "Δ", "Esito", "Urgenza (Mosca)", "Margine Z−(X+Y)",
                        "Owner"], righe, "", {0: "num", 3: "num", 4: "num", 5: "num", 8: "num"}) + "</div>")
    return sezione("classifica", "Classifica e lente di Mosca", "Report al comitato", corpo,
                   come_si_legge(testi["come"]), testi["legenda"])


def s_radar(v, testi):
    righe = []
    for x in sorted(v.rischi, key=lambda x: x["pos"]):
        righe.append([link_rischio(x["id"]), esc(x["nome"]), esc(x["cat"]), x["I"], x["P"], x["pxi"], x["Pr"],
                      it(x["z"], 1), it(x["x"]), it(x["y"]), f'<b>{segnato(x["margine"])}</b>', chip_urgenza(x),
                      "◆" if x["sup"] else "", chip_livello(x["livello"]), esc(x["motivo"])])
    corpo = (f'<div class="pannello centro">{grafico_radar(v)}</div><div class="pannello"><h3>I numeri del radar</h3>'
             + tabella(["ID", "Nome", "Cat.", "I", "P", "P x I", "Pross.", "Z anni", "X anni", "Y anni", "Margine",
                        "Urgenza", "◆", "Livello", "Motivo del livello"], righe, "compatta",
                       {3: "num", 4: "num", 5: "num", 6: "num", 7: "num", 8: "num", 9: "num", 10: "num"}) + "</div>")
    return sezione("radar", "Radar di Mosca", "Report al comitato", corpo,
                   come_si_legge(testi["come"]), testi["legenda"])


def s_matrice(v, testi):
    righe = [[link_rischio(x["id"]), esc(x["nome"]) + (" ◆" if x["sup"] else ""), f'{x["P"]} · {x["I"]}', f'<b>{x["pxi"]}</b>']
             for x in sorted(v.rischi, key=lambda x: (-x["pxi"], x["pos"]))]
    corpo = (f'<div class="griglia-2"><div class="pannello">{matrice_html(v)}</div><div class="pannello"><h3>I rischi, dal '
             f'piu' + "'" + ' grave</h3>' + tabella(["ID", "Rischio", "P · I", "P x I"], righe, "compatta", {3: "num"}) + "</div></div>")
    return sezione("matrice", "Matrice di gravita'", "Report al comitato", corpo, come_si_legge(testi["come"]),
                   testi["legenda"])


def s_delta(v, testi):
    su = sum(1 for x in v.rischi if x["delta"] and x["delta"] > 0)
    giu = sum(1 for x in v.rischi if x["delta"] and x["delta"] < 0)
    zero = sum(1 for x in v.rischi if x["delta"] == 0)
    ordinati = sorted(v.rischi, key=lambda x: (-(x["delta"] if x["delta"] is not None else -9), x["pos"]))
    motivi = []
    for x in ordinati:
        if x["esito"] == "Nuovo" or x["delta"] is None:
            m = "Rischio nuovo nel ciclo."
        elif x["variazioni"]:
            m = schede.cosa_cambiato({"esito": x["esito"], "variazioni": x["variazioni"]})
        else:
            m = "Invariato: nessuna evidenza nuova dopo il ciclo precedente; restano valide le sue fonti."
        evid = " evidenziato" if x["delta"] is not None and abs(x["delta"]) >= v.p["sogliaDelta"] else ""
        motivi.append([link_rischio(x["id"]), f'<span class="{"forte" if evid else ""}">{segnato(x["delta"])}</span>',
                       esc(m).replace("\n", "<br>")])
    dett = [[link_rischio(x["id"]), esc(x["nome"])] + [segnato(x["dd"][d], 0) if x["dd"][d] is not None else "—"
                                                       for d, _, _ in DIM] + [f'<b>{segnato(x["delta"])}</b>']
            for x in sorted(v.rischi, key=lambda x: x["pos"])]
    leg = ('<div class="legenda-chip"><span class="sw sw-su"></span> In aumento <span class="sw sw-giu"></span> '
           'In diminuzione</div>')
    corpo = (f'<div class="pannello"><h3>{su} in aumento, {giu} in diminuzione, {zero} '
             f'invariati (sul {v.prec or "—"})</h3>{leg}{grafico_delta(v)}</div><div class="pannello"><h3>Perche' + "'" +
             ' varia (o no), nello stesso ordine</h3>' + tabella(["Rischio", "Δ", "Motivo, con l'evidenza"], motivi, "compatta")
             + '</div><div class="pannello"><h3>Dettaglio per dimensione</h3>'
             + tabella(["ID", "Nome", "Δ Impatto", "Δ Probabilita", "Δ Prossimita", "Δ Velocity", "Δ Score"], dett, "compatta",
                       {2: "num", 3: "num", 4: "num", 5: "num", 6: "num"}) + "</div>")
    return sezione("delta", "Delta: che cosa e' cambiato dal ciclo precedente", "Report al comitato", corpo,
                   come_si_legge(testi["come"]), testi["legenda"])


def dims_testo(o):
    return " · ".join(str(o[k]) if o[k] is not None else "—" for k in ("I", "P", "Pr", "V"))


def nomi_driver(v, codici):
    return ", ".join(v.nomi_driver.get(c.strip(), c.strip()) for c in str(codici or "").split(",") if c.strip())


def s_fenomeni(v, testi):
    oss = v.rischi_oss()
    mov = v.movimenti()
    entr = sum(str(r.get("Movimento")).startswith("Entrat") for r in mov)
    usc = sum(str(r.get("Movimento")).startswith("Uscit") for r in mov)
    testa = (f'<p class="riquadro">Registro unico: nel radar <b>{len(v.rischi)}</b> rischi (capienza {v.p["capienzaRadar"]}), '
             f'in osservazione <b>{len(oss)}</b> (al massimo {v.p["capienzaOsservazione"]}) · nel ciclo: entrati nel radar '
             f'{entr}, usciti {usc}.</p>')
    righe = [[k, link_rischio(o["id"]), esc(o["nome"]), f'{esc(o["cat"])} {esc(v.nomi_cat.get(o["cat"], ""))}', dims_testo(o),
              f'<b>{it(o["score"])}</b>', esc(o["motivo"]), f'{it_breve(o["orizzonte"])} anni' if o["orizzonte"] else "—",
              esc(o["segnale"]), esc(nomi_driver(v, o["driver"])), o["notizie"]]
             for k, o in enumerate(oss, start=1)]
    mv = [[link_rischio(r["ID_Rischio"]), esc(r["Nome"]), esc(r["Movimento"]), esc(r.get("Motivo_Stato"))] for r in mov]
    corpo = (testa + '<div class="pannello">' + tabella(
        ["#", "ID", "Rischio", "Categoria", "Imp · Prob · Pross · Vel", "Score", "Perche' fuori dal radar", "Orizzonte",
         "Segnale che lo porterebbe nel radar", "Driver", "Notizie A3 (4 sett.)"], righe, "", {0: "num", 5: "num", 10: "num"})
             + '</div><div class="pannello"><h3>Entrati e usciti dal radar in questo ciclo</h3>'
             + (tabella(["ID", "Rischio", "Movimento", "Motivo (regola, D-79)"], mv) if mv else
                '<p class="tenue">Nessun movimento nel ciclo.</p>') + "</div>")
    return sezione("fenomeni", "Rischi in osservazione", "Report al comitato", corpo, come_si_legge(testi["come"]),
                   testi["legenda"])


def blocchi_testo(righe, colonne_intest=None):
    """Righe di un foglio di testo (Guida, Metodologia, Dizionario) in titoli e tabelle."""
    out, tab = [], []

    def chiudi():
        if tab:
            out.append('<div class="tab-scorre"><table class="testo"><tbody>' + "".join(tab) + "</tbody></table></div>")
            tab.clear()
    for r in righe:
        piene = [x for x in r if x not in (None, "")]
        if len(piene) == 1 and r[0] not in (None, ""):
            chiudi()
            t = str(r[0])
            if len(t) < 80 and (t.isupper() or " — " in t):
                out.append(f"<h3>{esc(t.capitalize() if t.isupper() else t)}</h3>")
            else:
                out.append(f'<p class="paragrafo">{esc(t)}</p>')
        else:
            celle = "".join(f"<td>{esc(it_breve(x) if isinstance(x, float) else x)}</td>" for x in r)
            if colonne_intest and list(r[:len(colonne_intest)]) == colonne_intest:
                celle = "".join(f"<th>{esc(x)}</th>" for x in r)
            tab.append(f"<tr>{celle}</tr>")
    chiudi()
    return "".join(out)


def s_metodologia(v, reg, testi):
    righe = reg.righe_foglio("Metodologia", 4, 2)
    righe = [r for r in righe if not str(r[0]).startswith("Legenda")]
    return sezione("metodologia", "Metodologia: come si leggono i numeri", "Report al comitato",
                   '<div class="pannello">' + blocchi_testo(righe) + "</div>", come_si_legge(testi["come"]))


def s_registro(v, testi):
    n_da, rett = v.stato_validazione()
    testa = ('<div class="tiles stretti">'
             f'<div class="tile {"tile-ok" if n_da == 0 else "tile-attesa"}"><p class="tile-et">Validazione del ciclo</p>'
             f'<p class="tile-val">{n_da}</p><p class="tile-sotto">righe da validare · rettifiche applicate {rett} · '
             f'{"ciclo validato" if n_da == 0 else "CICLO NON ANCORA VALIDATO"}</p></div>'
             f'<div class="tile {"tile-ok" if len(v.rischi) == len(v.anag) else "tile-allerta"}"><p class="tile-et">Controllo di '
             f'copertura</p><p class="tile-val">{len(v.rischi)} / {len(v.anag)}</p><p class="tile-sotto">valutati nel ciclo / '
             f'rischi in anagrafica</p></div></div>')
    righe = []
    for x in v.classifica:
        sim = "▲" if x["netti"] >= v.p["sogliaDirezione"] else "▼" if x["netti"] <= -v.p["sogliaDirezione"] else "●"
        righe.append([f'<td class="liv-fondo liv-{classe(x["livello"])}">{link_rischio(x["id"])}</td>', esc(x["nome"]),
                      esc(x["cat"]), cella_dim(x["I"]), cella_dim(x["P"]), cella_dim(x["Pr"]), cella_dim(x["V"]),
                      f'<b>{it(x["score"])}</b>', segnato(x["delta"]),
                      segnato(x["delta6"]) if x["delta6"] is not None else "nuovo", chip_urgenza(x), segnato(x["margine"]),
                      "◆" if x["sup"] else "", x["kri_r"], x["kri_a"], x["kri_v"], f"{sim} {x['netti']:+d}" if x["netti"]
                      else f"{sim} 0", esc(x["presidio"]), esc(x["rilievi"] or ""), esc(x["owner"])])
    corpo = testa + '<div class="pannello">' + tabella(
        ["ID", "Nome", "Cat.", "Imp", "Prob", "Pross", "Vel", "Score", "Δ ciclo prima", "Δ 6 mesi", "Urgenza",
         "Margine (anni)", "◆", "KRI R", "KRI A", "KRI V", "Segnali netti", "Presidio", "Rilievi", "Owner"], righe,
        "registro", {7: "num", 8: "num", 9: "num", 11: "num", 13: "num", 14: "num", 15: "num"}) + "</div>"
    return sezione("registro", "Registro: il ciclo corrente", "Lavoro", corpo, come_si_legge(testi["come"]), testi["legenda"])


def s_storico(v, testi):
    righe = []
    for x in sorted(v.rischi, key=lambda x: x["pos"]):
        vals = []
        for c in v.trend:
            r = v.score_di(x["id"], c)
            vals.append(r["Score_Composito"] if r else None)
        righe.append([link_rischio(x["id"]), esc(x["nome"]), sparkline(vals, v.trend)] + [it(s) if s is not None else "—"
                                                                                        for s in vals])
    corpo = '<div class="pannello">' + tabella(["ID", "Nome", "Andamento (1-5)"] + v.trend, righe, "storico",
                                               {i: "num" for i in range(3, 3 + len(v.trend))}) + "</div>"
    return sezione("storico", "Storico: score composito negli ultimi cicli", "Lavoro", corpo,
                   "Un rischio per riga, con l'andamento su scala fissa 1-5 (passare sui punti per il valore). Un punto "
                   "mancante = rischio non ancora nel registro, o in osservazione con un altro score.", testi["legenda"])


def storia_kri(v, kri_id):
    out = []
    for c in v.trend:
        k = next((r for r in v.kri_tutti if r["KRI_ID"] == kri_id and r["Ciclo"] == c), None)
        out.append({"valore": k.get("Valore_Corrente") if k else None, "att": k.get("Soglia_Attenzione") if k else None,
                    "all": k.get("Soglia_Allerta") if k else None, "stato": k.get("Stato") if k else "—"})
    return out


def s_indicatori(v, testi):
    stati = ["Verde", "Ambra", "Rosso", "Non rilevato", "Soglie da definire"]
    conta = {s: sum(k["Stato"] == s for k in v.kri) for s in stati}
    testa = '<div class="legenda-chip conteggi">' + "".join(f'{chip_kri(s)} <b>{conta[s]}</b>' for s in stati) + "</div>"
    prec = {k["KRI_ID"]: k for k in v.kri_tutti if k["Ciclo"] == v.prec}
    righe = []
    for k in v.kri:
        p = prec.get(k["KRI_ID"], {}).get("Valore_Corrente")
        val = k.get("Valore_Corrente")
        tend = "" if val is None or p is None else ("▲ peggiora" if val > p else "▼ migliora" if val < p else "● invariato")
        righe.append([f'<a href="#kri-{esc(k["KRI_ID"])}">{esc(k["KRI_ID"])}</a>', link_rischio(k["ID_Rischio"]),
                      esc(k["Nome"]), esc(k["Unita"]), it_breve(k.get("Soglia_Attenzione")), it_breve(k.get("Soglia_Allerta")),
                      f'<b>{it_breve(val)}</b>', data_it(k.get("Data_Rilevazione")), chip_kri(k["Stato"]),
                      it_breve(p) if p is not None else "—", f'<span class="tend">{tend}</span>',
                      esc(k.get("Tipo") or "")])
    piccoli = []
    for k in v.kri:
        piccoli.append(f'<figure class="multiplo" id="kri-{esc(k["KRI_ID"])}"><figcaption><b>{esc(k["KRI_ID"])}</b> '
                       f'{chip_kri(k["Stato"])}<br><span class="piccolo">{esc(k["Nome"])} ({esc(k["Unita"])})</span>'
                       f'</figcaption>{grafico_kri(storia_kri(v, k["KRI_ID"]), v.trend, k["KRI_ID"], k["Unita"])}</figure>')
    corpo = (testa + '<div class="pannello">' + tabella(
        ["KRI", "Rischio", "Nome", "Unita'", "Soglia attenzione", "Soglia allerta", "Valore", "Rilevato il", "Stato",
         "Ciclo prima", "Tendenza", "Tipo"], righe, "compatta", {4: "num", 5: "num", 6: "num", 9: "num"})
             + '</div><div class="pannello"><h3>Storia di ogni KRI, ultimi cicli</h3><p class="nota">Linea blu = valore; '
               'tratteggio giallo = soglia di attenzione, rosso = soglia di allerta. Nell' + "'" + 'Excel si sceglieva un KRI '
               'alla volta (cella N4): qui ci sono tutti.</p><div class="multipli">' + "".join(piccoli) + "</div></div>")
    return sezione("indicatori", "Indicatori: KRI del ciclo e storia di ogni KRI", "Lavoro", corpo,
                   come_si_legge(testi["come"]), testi["legenda"])


def s_notizie(v, testi):
    cat = v.notizie_per_categoria()
    righe = [[esc(c["cod"]), esc(c["nome"]), c["max"], c["forti"], c["rischi"],
              f'<span class="{"forte rosso" if c["presidio"] == "AREA NON PRESIDIATA" else ""}">{esc(c["presidio"])}</span>',
              c["i1"], c["i2"], c["i3"]] for c in cat]
    anag_ids = {a["ID"] for a in v.anag}
    scollegate = sum(1 for n in v.finestra if n.get("Rischio_ID") not in anag_ids)
    sett = next((s for s in v.settimane_sintesi if s["Settimana"] == v.settimana), None)
    press = v.pressione()
    elenco = [[data_it(n["Data"]), esc(n["ID_Notizia"]), esc(n["Titolo"]), esc(n["Fonte"]),
               link_rischio(n["Rischio_ID"]) if n.get("Rischio_ID") else "—", esc(n["Segnale"]), esc(n.get("Validazione"))]
              for n in sorted(v.finestra, key=lambda n: n["Data"], reverse=True)]
    leg = ('<div class="legenda-chip"><span class="sw sw-int1"></span> Intensita' + "'" + ' 1: nesso indiretto '
           '<span class="sw sw-int2"></span> 2: nesso diretto <span class="sw sw-int3"></span> 3: diretto e immediato</div>')
    corpo = (f'<p class="riquadro">Settimana corrente: <b>{data_it(v.settimana) or "—"}</b> · finestra di 4 settimane dal '
             f'{data_it(v.inizio4) or "—"} · notizie nella finestra: <b>{len(v.finestra)}</b> · non collegate a un rischio '
             f'del radar: {scollegate}</p>'
             f'<div class="griglia-2 larga-sx"><div class="pannello"><h3>Notizie delle ultime 4 settimane per categoria, '
             f'per intensita' + "'" + f'</h3>{leg}{grafico_notizie_cat(cat)}</div><div class="pannello"><h3>Sintesi di A3 '
             f'della settimana</h3><p>{esc(sett.get("Sintesi_A3")) if sett else "—"}</p><h4>Flag proposti da A3 (non '
             f'cambiano score)</h4><p>{esc(sett.get("Flag_Proposti")) if sett else "—"}</p><h4>Proposta di ciclo '
             f'straordinario (decide la funzione di rischio, D-75)</h4><p>{esc(sett.get("Proposta_Straordinario")) if sett and sett.get("Proposta_Straordinario") not in (None, "—") else "Nessuna."}</p></div></div>'
             '<div class="pannello">' + tabella(["Codice", "Categoria", "Intensita' max", "Notizie con intensita' >= 2",
                                                 "Rischi nel radar", "Presidio", "Int. 1", "Int. 2", "Int. 3"], righe, "compatta",
                                                {2: "num", 3: "num", 4: "num", 6: "num", 7: "num", 8: "num"}) + "</div>"
             '<div class="pannello"><h3>Pressione settimanale (12 settimane)</h3><div class="legenda-chip">'
             '<span class="sw sw-su"></span> peggiorativi (sopra lo zero) <span class="sw sw-giu"></span> migliorativi '
             f'(sotto)</div>{grafico_pressione(press)}</div>'
             '<div class="pannello"><h3>Le notizie della finestra</h3>'
             + (tabella(["Data", "ID", "Titolo", "Fonte", "Rischio", "Segnale", "Validazione"], elenco, "compatta")
                if elenco else '<p class="tenue">Nessuna notizia nelle ultime 4 settimane.</p>') + "</div>")
    return sezione("notizie", "Notizie: pressione sulle categorie di rischio", "Lavoro", corpo,
                   come_si_legge(testi["come"]), testi["legenda"])


def s_piano(v, testi):
    righe = [[esc(a["Ciclo"]), link_rischio(a["ID_Rischio"]), esc(a["ID_Azione"]), esc(a["Descrizione"]), esc(a["Owner"]),
              esc(a["Orizzonte"]), f'<span class="stato-az stato-{classe(a["Stato"])}">{esc(a["Stato"])}</span>',
              data_it(a.get("Data_Decisione")), data_it(a.get("Scadenza")), esc(a.get("Validazione"))]
             for a in reversed(v.piano())]
    corpo = '<div class="pannello">' + tabella(["Ciclo", "Rischio", "ID azione", "Azione", "Owner", "Orizzonte", "Stato",
                                                "Decisa il", "Scadenza", "Validazione"], righe, "compatta") + "</div>"
    return sezione("piano", "Piano delle azioni", "Lavoro", corpo,
                   "Tutte le azioni proposte da A1 e decise dal validatore, dalla piu' recente. La scadenza e' la data della "
                   "decisione piu' l'orizzonte.", testi["legenda"])


def s_osservazione(v, testi):
    oss = v.rischi_oss()
    righe = [[link_rischio(o["id"]), esc(o["nome"]), esc(o["cat"]), esc(o["driver"]), cella_dim(o["I"]), cella_dim(o["P"]),
              cella_dim(o["Pr"]), cella_dim(o["V"]), f'<b>{it(o["score"])}</b>', esc(o["motivo"]), it_breve(o["orizzonte"]),
              esc(o["segnale"]), esc(accorcia(o["riga"].get("Descrizione"), 160)), esc(o["riga"].get("Validazione")),
              o["notizie"]] for o in oss]
    drivers = list(v.nomi_driver.items())

    def incrocio(elementi, titolo):
        conta = {}
        for e in elementi:
            for d in str(e["driver"] or "").split(","):
                conta[(d.strip(), e["cat"])] = conta.get((d.strip(), e["cat"]), 0) + 1
        t = [f'<h3>{titolo}</h3><div class="tab-scorre"><table class="incrocio"><thead><tr><th>Driver</th>'
             + "".join(f"<th>{esc(c)}</th>" for c in v.codici) + "<th>Totale</th></tr></thead><tbody>"]
        for cod, nome in drivers:
            vals = [conta.get((cod, c), 0) for c in v.codici]
            t.append(f"<tr><th>{esc(nome)}</th>" + "".join(
                f'<td class="calore-{min(n, 3)}">{n or ""}</td>' for n in vals) + f"<td><b>{sum(vals)}</b></td></tr>")
        t.append("</tbody></table></div>")
        return "".join(t)
    corpo = ('<div class="pannello">' + tabella(
        ["ID", "Rischio", "Cat.", "Driver", "Imp.", "Prob.", "Pross.", "Vel.", "Score", "Perche' fuori dal radar (D-79)",
         "Orizz. (anni)", "Segnale che lo porterebbe nel radar", "Meccanismo (descrizione)", "Validazione",
         "Notizie A3 (4 sett.)"], righe, "compatta", {8: "num", 10: "num", 14: "num"}) + "</div>"
             '<div class="griglia-2"><div class="pannello">'
             + incrocio([{"driver": o["driver"], "cat": o["cat"]} for o in oss], "Dove si concentrano i rischi in osservazione")
             + '</div><div class="pannello">'
             + incrocio([{"driver": a["Driver"], "cat": a["Event_Type"]} for a in v.anag], "Rischi del radar con lo stesso "
                        "incrocio") + '<p class="nota">Un rischio con piu' + "'" + ' driver conta in ogni driver. Una casella '
             'piena sopra e vuota sotto e' + "'" + ' un tema che il radar non copre.</p></div></div>')
    return sezione("osservazione", "Osservazione: dove si concentrano i rischi fuori dal radar", "Lavoro", corpo,
                   come_si_legge(testi["come"]), testi["legenda"])


def s_validazione(v, testi):
    n_da, rett = v.stato_validazione()
    pend = v.da_validare_news()
    oggetti = sorted({x["Oggetto"] for x in v.val})
    decisioni = ["Accettato", "Rettificato", "Respinto", "Da decidere"]
    esito = {}
    for x in v.val:                              # esito di un oggetto: il campo piu' "forte" fra i suoi
        d = str(x["Decisione"] or "Da decidere")
        d = "Rettificato" if d.startswith("Rettificato") else "Accettato" if d.startswith("Invariato") else d
        k = (x["Oggetto"], x["ID"])
        esito[k] = max(esito.get(k, d), d, key=["Accettato", "Rettificato", "Respinto", "Da decidere"].index)
    conta = [[esc(o)] + [sum(1 for (og, _), d2 in esito.items() if og == o and d2 == d) for d in decisioni]
             for o in oggetti]
    cambi = [[esc(x["Oggetto"]), esc(x["ID"]), esc(accorcia(x.get("Nome"), 60)), esc(x["Campo"]),
              esc(x.get("Proposta_Agente")), f'<b>{esc(x.get("Valore_Validato"))}</b>', esc(x["Decisione"]),
              esc(x.get("Motivazione")), esc(x.get("Validato_Da")), data_it(x.get("Data_Validazione"))]
             for x in v.val if str(x["Decisione"]).startswith(("Rettificato", "Respinto"))]
    testa = (f'<div class="riquadro {"ok" if n_da == 0 else "avviso"}"><b>{"Ciclo validato" if n_da == 0 else "Ciclo NON validato"}'
             f'</b> · righe da validare {n_da} · rettifiche applicate {rett}'
             + (f' · <b>in attesa</b>: {len(pend)} {"notizia" if len(pend) == 1 else "notizie"} (settimana del '
                f'{data_it(pend[0]["Settimana"])})' if pend else "")
             + '</div><div class="riquadro">La validazione si fa nella griglia del registro Excel (foglio Validazione): '
               'questa pagina mostra l' + "'" + 'esito, non si compila.</div>')
    corpo = (testa + '<div class="griglia-2 larga-dx"><div class="pannello"><h3>Decisioni del ciclo, un conteggio per oggetto</h3>'
             + tabella(["Oggetto"] + decisioni, conta, "compatta", {1: "num", 2: "num", 3: "num", 4: "num"})
             + '</div><div class="pannello"><h3>Rettifiche e respingimenti, con il motivo</h3>'
             + (tabella(["Oggetto", "ID", "Nome", "Campo", "Proposta", "Validato", "Decisione", "Motivazione", "Da", "Il"],
                        cambi, "compatta") if cambi else '<p class="tenue">Nessuna rettifica: tutte le proposte accettate.</p>')
             + "</div></div>")
    return sezione("validazione", "Validazione: l'esito", "Lavoro", corpo,
                   "Una riga per oggetto e campo nello storico delle validazioni: la proposta dell'agente, il valore deciso e "
                   "il motivo. Qui l'esito del ciclo.", "")


def s_parametri(v, reg, testi):
    p = v.p
    pesi = [["Impatto", it(reg.nome("pesoImpatto"))], ["Probabilita", it(reg.nome("pesoProbabilita"))],
            ["Prossimita", it(reg.nome("pesoProssimita"))], ["Velocity (solo se valutata)", it(reg.nome("pesoVelocity"))]]
    mosca = [["X: anni per decidere e avviare la mitigazione", it(p["moscaX"])],
             ["Y di riserva: anni di mitigazione se il rischio non ha azioni", it(p["moscaYDefault"])]]
    mosca += [[f"Z per Prossimita {k} ", it_breve(z)] for k, z in enumerate(p["anniProssimita"], start=1)]
    liv = [["P x I alta da (aggravante)", p["livelloPxI"]], ["Quota dei KRI in ambra o rosso da", it_breve(p["livelloQuotaKRI"])],
           ["Margine stretto: sotto (anni)", it_breve(p["livelloMargine"])], ["Segnali forti per l'Attivazione", p["livelloForti"]],
           ["Soglia di direzione (segnali netti in 4 settimane)", p["sogliaDirezione"]],
           ["Soglia di evidenziazione del Delta", it_breve(p["sogliaDelta"])]]
    ru = [["Capienza del radar", p["capienzaRadar"]], ["Soglia di Prossimita per entrare", p["sogliaProssimitaRadar"]],
          ["Soglia di Probabilita per entrare", p["sogliaProbabilitaRadar"]],
          ["Cicli sotto la soglia prima di uscire", p["cicliUscitaRadar"]],
          ["Rischi in osservazione al massimo", p["capienzaOsservazione"]]]
    cicli = [["Ciclo", v.ciclo], ["Periodo precedente", v.prec or "—"], ["Trimestre di due trimestri prima (Δ 6 mesi)", v.meno2],
             ["Data di riferimento", data_it(v.data_rif)], ["Settimana corrente", data_it(v.settimana) or "—"],
             ["Periodi del trend", ", ".join(v.trend)]]
    anag = [[link_rischio(a["ID"]), esc(a["Nome"]), esc(a["Event_Type"]), esc(a["Driver"]), esc(a.get("Owner")),
             it_breve(a.get("Anni_Adeguamento")) if a.get("Anni_Adeguamento") else "", esc(a.get("Attenzione_Supervisore") or ""),
             esc(a.get("Presidio_Controlli"))] for a in v.anag]
    regola = reg.cella("Parametri", 71, 1) or ""
    tass = [[esc(c), esc(v.nomi_cat[c])] for c in v.codici]
    drv = [[esc(c), esc(n)] for c, n in v.nomi_driver.items()]
    t2 = lambda r: tabella(["", ""], r, "chiave-valore", {1: "num"})
    corpo = ('<div class="griglia-3">'
             f'<div class="pannello"><h3>Pesi dello score composito (D-50)</h3>{t2(pesi)}</div>'
             f'<div class="pannello"><h3>Urgenza: disuguaglianza di Mosca</h3>{t2(mosca)}</div>'
             f'<div class="pannello"><h3>Livello di allerta e soglie</h3>{t2(liv)}</div>'
             f'<div class="pannello"><h3>Registro unico (D-79)</h3>{t2(ru)}</div>'
             f'<div class="pannello"><h3>Periodi</h3>{t2(cicli)}</div>'
             f'<div class="pannello"><h3>Tassonomia interna (D-74) e driver</h3>{tabella(["Codice", "Categoria"], tass, "compatta")}'
             f'{tabella(["Codice", "Driver"], drv, "compatta")}</div></div>'
             f'<div class="pannello"><h3>La regola del registro unico</h3><p>{esc(regola)}</p></div>'
             '<div class="pannello"><h3>Anagrafica dei rischi nel radar</h3>'
             + tabella(["ID", "Nome", "Categoria", "Driver", "Owner", "Anni di adeguamento", "Attenzione del supervisore",
                        "Presidio dei controlli"], anag, "compatta") + "</div>")
    return sezione("parametri", "Parametri: pesi, soglie, urgenza e anagrafica", "Lavoro", corpo,
                   "Decisi dalla funzione di rischio e registrati in DECISIONI.md; stanno nel foglio Parametri del "
                   "registro. Qui si leggono, non si cambiano.")


def s_dizionario(v, reg, testi):
    righe = reg.righe_foglio("Dizionario", 4, 7)
    righe = [r for r in righe if not str(r[0]).startswith("Legenda")]
    intest = list(righe[0])
    return sezione("dizionario", "Dizionario dei dati", "Lavoro", '<div class="pannello">'
                   + blocchi_testo(righe, colonne_intest=intest) + "</div>",
                   "Ogni dato: chi lo scrive, quando, che cosa significa e come si calcola. La colonna Dove nomina i fogli "
                   "del registro Excel.")


# ------------------------------------------------ schede dei rischi
class LettoreSchede:
    """Il lettore che schede.input_schede si aspetta (scr + sessione), ma sui dati e sulle viste calcolate in Python:
    lo stesso codice del pacchetto produce i dati delle schede senza Excel."""

    class _Valore:
        def __init__(self, valore, destra=None):
            self.Value = valore
            self._destra = destra

        def GetOffset(self, _r, _c):
            return LettoreSchede._Valore(self._destra)

    def __init__(self, v):
        self.v = v

    def righe(self, nome):
        return {"tblRegister": self.v.rr_tutte, "tblAzioni": self.v.azioni, "tblKRI": self.v.kri_tutti}[nome]

    def leggi_anagrafica(self):
        return None, [(i, a) for i, a in enumerate(self.v.anag, start=1)]

    def intervallo_nome(self, nome):
        v = self.v
        col = lambda chiave: tuple((v.per_rid[a["ID"]][chiave] if a["ID"] in v.per_rid else None,) for a in v.anag)
        mappa = {"rmID": "id", "rmX": "x", "rmY": "y", "rmZ": "z", "rmMargine": "margine", "rmUrgenza": "urgenza",
                 "rmLivello": "livello", "rmMotivoLivello": "motivo", "posNetti": "netti"}
        if nome in mappa:
            return self._Valore(col(mappa[nome]))
        if nome == "codiciL1":
            return self._Valore(tuple((c,) for c in v.codici), tuple((v.nomi_cat[c],) for c in v.codici))
        return self._Valore(v.reg.nome(nome))


def dati_schede(v):
    lettore = LettoreSchede(v)
    return schede.input_schede(lettore, lettore, v.ciclo, v.prec, v.data_rif)


def scheda_radar(v, r, prec_id, succ_id, totale):
    x = v.per_rid[r["id_rischio"]]
    d = None if r["score_ciclo_precedente"] is None else round(r["score"] - r["score_ciclo_precedente"], 2)
    freccia = "nuovo" if d is None else ("▲ " if d > 0 else "▼ " if d < 0 else "● ") + segnato(d)
    motiv = r.get("motivazioni") or {}
    h = [f'<header class="scheda-testa liv-bordo-{classe(x["livello"])}"><p class="scheda-pos">{r["posizione_in_classifica"]} '
         f'di {totale} per score</p><h2 id="t-scheda-{r["id_rischio"]}">{esc(r["id_rischio"])} · {esc(r["nome"])}</h2>'
         f'<p class="scheda-numeri"><span class="grande">{it(r["score"])}</span> score <span class="freccia">{freccia}</span> '
         f'{chip_urgenza(x)} {chip_livello(x["livello"])}</p><p class="meta">'
         f'<span>{esc(r["event_type"])} · {esc(r["event_type_nome"] or "")}</span><span>Driver: {esc(nomi_driver(v, r["driver"]))}'
         f'</span><span>Owner: {esc(r["owner"])}</span><span>Presidio: {esc(r.get("presidio_controlli"))}</span>'
         f'<span>Esito: {esc(r["esito"])} · validato</span>'
         + ('<span class="sup">◆ Attenzione del supervisore</span>' if r["attenzione_supervisore"] else "") + "</p></header>"]
    h.append(f'<h3>Il rischio</h3><p>{esc(r["descrizione"])}</p>')
    h.append(f'<h3>{"Razionale: il fenomeno in termini di rischio" if motiv else "Razionale (registro)"}</h3>'
             f'<p>{esc(r.get("razionale_rischio") or "—")}</p>')
    righe = []
    for d_, nome, _ in DIM:
        val = r["valori"][d_]
        sp = motiv.get(d_)
        righe.append([f"<b>{nome}</b>", cella_dim(val),
                      esc(sp) if sp else '<span class="tenue">Motivo non diviso per dimensione nel razionale di A1: vedi il '
                                         'Razionale.</span>'])
    pesi = v.p
    presenti = [d_ for d_, _, _ in DIM if r["valori"][d_] is not None]
    nomi_peso = {"impatto": "pesoImpatto", "probabilita": "pesoProbabilita", "prossimita": "pesoProssimita",
                 "velocity": "pesoVelocity"}
    formula = " + ".join(f"{it(pesi[nomi_peso[d_]])} x {r['valori'][d_]}" for d_ in presenti)
    somma = it(sum(pesi[nomi_peso[d_]] for d_ in presenti))
    h.append("<h3>Valutazione, dimensione per dimensione (valori validati)</h3>"
             + tabella(["Dimensione", "Valore", "Perche' questo valore: informazioni e documenti"], righe, "dimensioni")
             + f'<p class="nota">Score composito {it(r["score"])} = ({formula}) / {somma}.'
             + ("" if "velocity" in presenti else " Velocity non valutata: i pesi si ridistribuiscono sulle altre tre.") + "</p>")
    u, la = r["urgenza"], r.get("livello_allerta") or {}
    adeg = ""
    if r["anni_adeguamento"]:
        pa = motiv.get("adeguamento")
        adeg = (f" (con {it(r['anni_adeguamento'], 1)} anni di adeguamento strutturale, stimati da A1 e validati"
                + (f"; motivo: {pa[0].lower() + pa[1:].rstrip('.')}" if pa else "") + ")")
    entro = min(x["y"], max(x["z"] - x["x"], 0))
    tot_barra = max(x["x"] + x["y"], x["z"], 0.5)
    barra = (f'<div class="barra-mosca" role="img" aria-label="X {it(x["x"])}, Y {it(x["y"])}, Z {it(x["z"], 1)} anni">'
             f'<span class="sw-x" style="width:{x["x"] / tot_barra * 100:.1f}%"></span>'
             f'<span class="sw-y-entro" style="width:{entro / tot_barra * 100:.1f}%"></span>'
             f'<span class="sw-y-oltre" style="width:{(x["y"] - entro) / tot_barra * 100:.1f}%"></span>'
             f'<i class="segno-z" style="left:{x["z"] / tot_barra * 100:.1f}%" title="Z = {it(x["z"], 1)} anni">◆</i></div>')
    h.append("<h3>Urgenza (disuguaglianza di Mosca) e livello di allerta</h3>"
             f'<p class="riquadro liv-fondo liv-{classe(la.get("livello"))}">Livello di allerta: <b>{esc(la.get("livello"))}</b>'
             f' — {esc(la.get("motivo"))}.</p>'
             f'<p>{esc(u["esito"])}. Tempo necessario X + Y = {it(u["x"])} + {it(u["y"])}{esc(adeg)} = {it(u["x"] + u["y"])} '
             f'anni; tempo disponibile Z = {it(u["z"], 1)} anni (Prossimita {r["valori"]["prossimita"]}). Margine Z − (X + Y) '
             f'= <b>{segnato(u["margine"])}</b>.</p>{barra}')
    h.append('<h3>Cosa e' + "'" + f' cambiato nel ciclo</h3><p>{esc(schede.cosa_cambiato(r)).replace(chr(10), "<br>")}</p>')
    cicli = [s["ciclo"] for s in r["storia"]]
    righe = [[f"<b>{esc(s['ciclo'])}</b>"] + [cella_dim(s.get(d_)) for d_, _, _ in DIM] + [f"<b>{it(s['score'])}</b>"]
             for s in r["storia"]]
    h.append('<h3>Andamento negli ultimi cicli</h3><div class="andamento">'
             + tabella(["Ciclo", "Imp", "Prob", "Pross", "Vel", "Score"], righe, "compatta storia", {5: "num"})
             + (sparkline([s["score"] for s in r["storia"]], cicli, 260, 70) if len(cicli) > 1 else
                f'<p class="tenue">Rischio entrato nel registro nel {esc(cicli[0] if cicli else v.ciclo)}: l' + "'" +
                'andamento si vedra' + "'" + ' dal prossimo ciclo.</p>') + "</div>")
    fonti = [[esc(f["riferimento"]), data_it(f.get("data")), esc(f.get("tipo")),
              f'<span class="ril ril-{f.get("rilevanza")}">{schede.RILEVANZA.get(f.get("rilevanza"), f.get("rilevanza"))}</span>']
             for f in r["fonti"]]
    h.append("<h3>Fonti</h3>" + tabella(["Fonte", "Data", "Tipo", "Rilevanza"], fonti, "compatta"))
    if r["azioni"]:
        az = []
        for a in r["azioni"]:
            az.append(f'<tr><td><b>{esc(a["descrizione"])}</b></td><td>{esc(a["owner"])}</td><td>{esc(a["orizzonte"])} · '
                      f'{esc(a["stato"])} (ciclo {esc(a["ciclo"])})</td></tr><tr class="perche"><td colspan="3"><span>'
                      f'<b>Perche' + "'" + f':</b> {esc(a.get("razionale") or "—")}</span><span><b>Effetto atteso:</b> '
                      f'{esc(a.get("effetto_atteso") or "non indicato nel razionale di A1.")}</span></td></tr>')
        h.append("<h3>Azioni</h3>" + tabella(["Azione", "Owner", "Orizzonte · stato"], az, "compatta azioni"))
    else:
        h.append("<h3>Azioni</h3><p class='tenue'>Nessuna azione del ciclo o aperta.</p>")
    kk = []
    for k in r["kri"]:
        tipo = f"[{k['tipo']}] " if k.get("tipo") else ""
        kk.append(f'<tr><td>{esc(tipo + k["nome"])} ({esc(k["unita"])})</td><td class="num"><b>{it_breve(k["valore"])}</b></td>'
                  f'<td class="num">{it_breve(k["soglia_attenzione"])} / {it_breve(k["soglia_allerta"])}</td>'
                  f'<td class="num">{it_breve(k["valore_ciclo_precedente"])}</td><td>{chip_kri(k["stato"])}</td></tr>'
                  f'<tr class="perche"><td colspan="5"><span><b>Perche' + "'" + f' lo misuriamo:</b> {esc(k.get("razionale") or "—")}'
                  f'</span><span><b>Effetto atteso:</b> {esc(k.get("effetto_atteso") or "non indicato nel razionale di A1.")}'
                  "</span></td></tr>")
    h.append("<h3>KRI in vigore</h3>" + (tabella(["Indicatore", "Valore", "Soglie (attenzione / allerta)", "Ciclo prima", "Stato"],
                                                kk, "compatta kri") if kk else "<p class='tenue'>Nessun KRI in vigore.</p>"))
    h.append(nav_schede(prec_id, succ_id))
    return (f'<section id="scheda-{esc(r["id_rischio"])}" class="tab scheda" aria-labelledby="t-scheda-{esc(r["id_rischio"])}">'
            + "".join(h) + "</section>")


def scheda_oss(v, o, prec_id, succ_id, totale, n):
    r = o["riga"]
    parti = schede.parti_razionale(r.get("Razionale"))
    raz = (" ".join(p for p in (parti.get("altro"), parti.get("rischio")) if p) or None) if parti else r.get("Razionale")
    h = [f'<header class="scheda-testa oss"><p class="scheda-pos">In osservazione · {n} di {totale}</p>'
         f'<h2 id="t-scheda-{o["id"]}">{esc(o["id"])} · {esc(o["nome"])}</h2><p class="scheda-numeri"><span class="grande">'
         f'{it(o["score"])}</span> score {chip("In osservazione", "stato-oss")}</p><p class="meta"><span>{esc(o["cat"])} · '
         f'{esc(v.nomi_cat.get(o["cat"], ""))}</span><span>Driver: {esc(nomi_driver(v, o["driver"]))}</span><span>Esito: '
         f'{esc(r["Esito"])} · {esc(r.get("Validazione"))}</span></p></header>']
    h.append('<p class="riquadro">Fuori dal radar (registro unico, D-79): <b>' + esc(o["motivo"]) + "</b>. Niente KRI ne' azioni "
             "finche' non entra; orizzonte " + (f"{it_breve(o['orizzonte'])} anni" if o["orizzonte"] else "—")
             + f". <b>Segnale che lo porterebbe nel radar:</b> {esc(o['segnale'])}</p>")
    h.append(f'<h3>Il rischio</h3><p>{esc(r.get("Descrizione"))}</p><h3>Razionale</h3><p>{esc(raz or "—")}</p>')
    righe = [[f"<b>{nome}</b>", cella_dim(o[k]), esc(parti.get(d_)) if parti.get(d_) else '<span class="tenue">—</span>']
             for (d_, nome, _), k in zip(DIM, ("I", "P", "Pr", "V"))]
    h.append("<h3>Valutazione</h3>" + tabella(["Dimensione", "Valore", "Perche' questo valore"], righe, "dimensioni"))
    h.append('<h3>Cosa e' + "'" + ' cambiato nel ciclo</h3><p>' + esc(schede.cosa_cambiato(
        {"esito": r["Esito"], "variazioni": json.loads(r.get("Variazioni") or "[]")})).replace("\n", "<br>") + "</p>")
    storia = [(c, v.score_di(o["id"], c)) for c in v.trend]
    storia = [(c, s) for c, s in storia if s]
    righe = [[f"<b>{esc(c)}</b>"] + [cella_dim(s[COL_DIM[d_]]) for d_, _, _ in DIM] + [f"<b>{it(s['Score_Composito'])}</b>",
                                                                                      esc(s.get("Stato"))]
             for c, s in storia]
    h.append('<h3>Andamento</h3><div class="andamento">' + tabella(["Ciclo", "Imp", "Prob", "Pross", "Vel", "Score", "Stato"],
                                                                   righe, "compatta storia", {5: "num"})
             + (sparkline([s["Score_Composito"] for _, s in storia], [c for c, _ in storia], 260, 70) if len(storia) > 1 else "")
             + "</div>")
    fonti = [[esc(f["riferimento"]), data_it(f.get("data")), esc(f.get("tipo")),
              f'<span class="ril ril-{f.get("rilevanza")}">{schede.RILEVANZA.get(f.get("rilevanza"), f.get("rilevanza"))}</span>']
             for f in json.loads(r.get("Fonti") or "[]")]
    h.append("<h3>Fonti</h3>" + tabella(["Fonte", "Data", "Tipo", "Rilevanza"], fonti, "compatta"))
    h.append(nav_schede(prec_id, succ_id))
    return (f'<section id="scheda-{esc(o["id"])}" class="tab scheda" aria-labelledby="t-scheda-{esc(o["id"])}">'
            + "".join(h) + "</section>")


def nav_schede(prec_id, succ_id):
    return ('<nav class="nav-schede">' + (f'<a href="#scheda-{prec_id}">← {prec_id}</a>' if prec_id else "<span></span>")
            + '<a href="#schede">Indice delle schede</a>'
            + (f'<a href="#scheda-{succ_id}">{succ_id} →</a>' if succ_id else "<span></span>") + "</nav>")


def s_schede(v):
    ingresso = dati_schede(v)
    radar = ingresso["rischi"]
    oss = v.rischi_oss()
    ordine = [r["id_rischio"] for r in radar] + [o["id"] for o in oss]
    righe = []
    for r in radar:
        x = v.per_rid[r["id_rischio"]]
        righe.append([r["posizione_in_classifica"], link_rischio(r["id_rischio"]),
                      esc(r["nome"]) + (" ◆" if r["attenzione_supervisore"] else ""), f'<b>{it(r["score"])}</b>',
                      chip_urgenza(x), chip_livello(x["livello"]), esc(r["esito"]), esc(r["owner"])])
    for o in oss:
        righe.append(["—", link_rischio(o["id"]), esc(o["nome"]), f'<b>{it(o["score"])}</b>', "—",
                      chip("In osservazione", "stato-oss"), esc(o["riga"]["Esito"]), "—"])
    indice = sezione("schede", "Schede dei rischi", "Schede", '<div class="pannello">' + tabella(
        ["#", "ID", "Rischio", "Score", "Urgenza", "Livello", "Esito", "Owner"], righe, "compatta", {0: "num", 3: "num"})
        + "</div>",
        f"Una scheda per rischio: prima i {len(radar)} del radar nell'ordine della classifica, poi i {len(oss)} in "
        "osservazione. Valori validati del registro; testi di A1 dal registro (senza A6). Clic sull'ID per aprire la "
        "scheda.", "Legenda — " + schede.LEGENDA)
    out = [indice]
    totale = len(radar)
    for i, r in enumerate(radar):
        out.append(scheda_radar(v, r, ordine[i - 1] if i else None, ordine[i + 1] if i + 1 < len(ordine) else None, totale))
    for j, o in enumerate(oss):
        i = len(radar) + j
        out.append(scheda_oss(v, o, ordine[i - 1], ordine[i + 1] if i + 1 < len(ordine) else None, len(oss), j + 1))
    return "".join(out), ordine


def s_dati(v, reg):
    """Le tabelle dati fino a questo ciclo: filtro e scaricamento in CSV (JavaScript), lettura anche senza."""
    sett = v.settimana or "0000"
    insiemi = [
        ("tblRegister", "RiskRegister", v.rr_tutte), ("tblAzioni", "Azioni", v.azioni), ("tblKRI", "KRI", v.kri_tutti),
        ("tblNews", "NewsLog", v.news), ("tblSintesi", "Sintesi_Cicli", [s for s in reg.tab["tblSintesi"]
                                                                        if s["Ciclo"] in set(v.cicli)]),
        ("tblSettimane", "Sintesi_Settimane", v.settimane_sintesi),
        ("tblValidazione", "Storico_Validazioni", [x for x in reg.tab["tblValidazione"]
                                                    if x["Periodo"] in set(v.cicli) or (re.match(r"\d{4}-\d{2}-\d{2}", str(x["Periodo"]))
                                                                                        and str(x["Periodo"]) <= sett)])]
    parti = ['<p class="riquadro">Le tabelle dati fino a questo ciclo: nello scenario sono il file <code>dati_registro.json</code>; '
             'qui si filtrano e si scaricano in CSV. Il filtro e lo scaricamento usano il JavaScript della pagina.</p>'
             '<p class="indice-dati">' + " · ".join(f'<a href="#dati-{n}">{t} ({len(r)})</a>' for n, t, r in insiemi) + "</p>"]
    for nome, titolo, righe in insiemi:
        col = reg.colonne[nome]
        corpo = "".join("<tr>" + "".join(f'<td><div class="cella">{esc(r.get(c))}</div></td>' for c in col) + "</tr>"
                        for r in righe)        # testo intero (filtro e CSV), tre righe a vista
        parti.append(f'<div class="pannello" id="dati-{nome}"><h3>{titolo} <span class="tenue">({len(righe)} righe)</span></h3>'
                     f'<div class="strumenti"><input type="search" placeholder="Filtra le righe…" data-filtra="tab-{nome}" '
                     f'aria-label="Filtra {titolo}"><button type="button" data-csv="tab-{nome}" data-nome="{titolo}_{v.ciclo}">'
                     f'Scarica CSV</button></div><div class="tab-scorre alta"><table class="dati" id="tab-{nome}"><thead><tr>'
                     + "".join(f"<th>{esc(c)}</th>" for c in col) + f"</tr></thead><tbody>{corpo}</tbody></table></div></div>")
    return sezione("dati", "Tabelle dati", "Dati", "".join(parti),
                   "Le sette tabelle che lo script scrive e il validatore conferma. Ogni altra pagina e' calcolata da queste.",
                   classe_extra="no-stampa")


def s_guida(v, reg, anteprima=False, sezioni=None):
    mappa = [("quadro", "Quadro"), ("classifica", "Classifica"), ("radar", "Radar_Mosca"), ("matrice", "Matrice"),
             ("delta", "Delta"), ("fenomeni", "Fenomeni"), ("metodologia", "Metodologia"), ("registro", "Registro"),
             ("storico", "Storico"), ("indicatori", "Indicatori"), ("notizie", "Notizie"), ("piano", "Piano_azioni"),
             ("osservazione", "Osservazione"), ("validazione", "Validazione"), ("parametri", "Parametri"),
             ("dizionario", "Dizionario"), ("schede", "Schede_<ciclo>.xlsx"), ("dati", "RiskRegister, Azioni, KRI, ...")]
    serve = {}
    for r in reg.righe_foglio("Guida", 8, 5):
        if r[0] and r[4]:
            for nome in str(r[0]).split(","):
                serve[nome.strip()] = r[4]
    serve.update({"Storico": "score composito negli ultimi cicli", "Indicatori": "KRI del ciclo e storia di ognuno"})
    righe = [[f'<a href="#{a}">{esc(TITOLI[a])}</a>', esc(f), esc(serve.get(f.split(",")[0].strip(), ""))] for a, f in mappa
             if sezioni is None or a in sezioni]
    indice = "index.html" if anteprima else f"../{INDICE}"
    corpo = ('<div class="pannello"><h3>Come si usa</h3><ul>'
             "<li>Un file per ciclo, che si apre nel browser: nessuna installazione, nessuna risorsa da internet. L'"
             f'<a href="{indice}">archivio dei cicli</a> elenca tutti i report.</li>'
             "<li>Il menu a sinistra porta alle pagine; i collegamenti sugli ID aprono la scheda del rischio. Il pulsante "
             "Indietro del browser torna alla pagina di prima.</li>"
             "<li>Passando sui grafici (o con il tasto Tab) compaiono i numeri; ogni grafico ha accanto la sua tabella.</li>"
             "<li>Stampa o PDF: il comando Stampa del browser mette ogni pagina e ogni scheda su un foglio a parte, senza "
             "menu e senza le tabelle dati.</li>"
             + ("" if anteprima else "<li>I numeri sono quelli del registro validato: prima di uscire il report li ha "
                "confrontati con quelli calcolati da Excel (X, Y, Z, margine, urgenza, livello di allerta, segnali netti, "
                "classifica, Delta); se uno non tornava, questo file non sarebbe stato scritto.</li>") +
             "<li>Si legge soltanto: i dati li scrive lo script dopo i controlli; il validatore decide nella griglia Excel di "
             "validazione; i parametri stanno nel loro file.</li></ul></div>"
             '<div class="pannello"><h3>Le pagine e i fogli del registro di oggi</h3>'
             + tabella(["Pagina", "Foglio dell'Excel", "A che cosa serve"], righe, "compatta") + "</div>")
    return sezione("guida", "Guida: come si usa questo report", "Guida", corpo,
                   f"{NOME_ANALISI}: registro dei rischi non finanziari emergenti della Banca, ciclo {v.ciclo}. "
                   + ("Anteprima dell'analisi 16 (report in HTML al posto dei fogli di Excel), con i dati della demo."
                      if anteprima else "Report al comitato in HTML, accanto al PDF dello stesso ciclo (D-83)."))


TITOLI = {"quadro": "Quadro d'insieme", "classifica": "Classifica", "radar": "Radar di Mosca", "matrice": "Matrice",
          "delta": "Delta", "fenomeni": "Rischi in osservazione", "metodologia": "Metodologia", "registro": "Registro",
          "storico": "Storico", "indicatori": "Indicatori (KRI)", "notizie": "Notizie", "piano": "Piano azioni",
          "osservazione": "Osservazione", "validazione": "Validazione", "parametri": "Parametri", "dizionario": "Dizionario",
          "schede": "Indice delle schede", "dati": "Tabelle dati", "guida": "Guida"}
GRUPPI = [("Report al comitato", ["quadro", "classifica", "radar", "matrice", "delta", "fenomeni", "metodologia"]),
          ("Lavoro", ["registro", "storico", "indicatori", "notizie", "piano", "osservazione", "validazione", "parametri",
                      "dizionario"]),
          ("Schede dei rischi", ["schede"]), ("Dati", ["dati"]), ("Aiuto", ["guida"])]


# ================================================================ pagina
CSS = r"""
:root{color-scheme:light;--th:#1f3864;--th-2:#2f5597;--bg:#f3f4f6;--surface:#ffffff;--surface-2:#f7f8fa;--ink:#1c2230;--ink-2:#4a5263;--muted:#6b7385;
--rule:#e1e4ea;--rule-2:#cdd2db;--brand:#1f3864;--brand-2:#2f5597;--brand-soft:#e3ecf7;--link:#2459a6;
--att:#e06666;--att-ink:#ffffff;--attz:#f6c342;--attz-ink:#1c2230;--mon:#93c47d;--mon-ink:#1c2230;
--att-soft:#fbe5e5;--attz-soft:#fff2cc;--mon-soft:#e2efda;
--agire:#c00000;--agire-ink:#ffffff;--entro:#ffc000;--entro-ink:#1c2230;--margine:#a6a6a6;--margine-ink:#1c2230;
--rosso:#c00000;--ambra:#ffc000;--verde:#70ad47;--nr:#d9d9d9;
--d1:#f2f2f2;--d2:#fce4d6;--d3:#f8cbad;--d4:#f4b183;--d5:#c55a11;--d5-ink:#ffffff;
--su:#c00000;--giu:#2e75b6;--x:#b8bcc4;--y-entro:#ed7d31;--y-oltre:#c00000;--z:#1f3864;
--i1:#fbe0cf;--i2:#f4b183;--i3:#c55a11;--linea:#1f3864;--soglia-att:#bf9000;--soglia-all:#c00000;
--fascia-verde:#e2efda;--fascia-ambra:#fff2cc;--fascia-rosso:#f8d7d7;--calore-1:#deebf7;--calore-2:#9dc3e6;--calore-3:#2e75b6}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){color-scheme:dark;--th:#2a3a57;--th-2:#33476b;--bg:#111318;--surface:#1b1e25;
--surface-2:#22262e;--ink:#eceef2;--ink-2:#b9bfcb;--muted:#9097a5;--rule:#2d323c;--rule-2:#3b414d;--brand:#9cb8e6;
--brand-2:#7fa3dc;--brand-soft:#253247;--link:#8fb4f0;--att:#d9605f;--attz:#d9aa2f;--mon:#7aa866;--att-soft:#3a2425;
--attz-soft:#3a3220;--mon-soft:#24321f;--agire:#e05252;--entro:#d9a400;--margine:#7c818b;--margine-ink:#eceef2;
--d1:#2b2f37;--d2:#4a3427;--d3:#6b4630;--d4:#93603a;--d5:#c55a11;--d1-ink:#eceef2;--su:#e05252;--giu:#5b9bd5;--x:#5f6570;
--y-entro:#e07b35;--y-oltre:#e05252;--z:#9cb8e6;--i1:#4a3427;--i2:#93603a;--i3:#d06a1c;--linea:#9cb8e6;
--soglia-att:#d9aa2f;--soglia-all:#e05252;--fascia-verde:#24321f;--fascia-ambra:#3a3220;--fascia-rosso:#3a2425;
--calore-1:#22344a;--calore-2:#2f5b8a;--calore-3:#4a86c5}}
:root[data-theme="dark"]{color-scheme:dark;--th:#2a3a57;--th-2:#33476b;--bg:#111318;--surface:#1b1e25;
--surface-2:#22262e;--ink:#eceef2;--ink-2:#b9bfcb;--muted:#9097a5;--rule:#2d323c;--rule-2:#3b414d;--brand:#9cb8e6;
--brand-2:#7fa3dc;--brand-soft:#253247;--link:#8fb4f0;--att:#d9605f;--attz:#d9aa2f;--mon:#7aa866;--att-soft:#3a2425;
--attz-soft:#3a3220;--mon-soft:#24321f;--agire:#e05252;--entro:#d9a400;--margine:#7c818b;--margine-ink:#eceef2;
--d1:#2b2f37;--d2:#4a3427;--d3:#6b4630;--d4:#93603a;--d5:#c55a11;--d1-ink:#eceef2;--su:#e05252;--giu:#5b9bd5;--x:#5f6570;
--y-entro:#e07b35;--y-oltre:#e05252;--z:#9cb8e6;--i1:#4a3427;--i2:#93603a;--i3:#d06a1c;--linea:#9cb8e6;
--soglia-att:#d9aa2f;--soglia-all:#e05252;--fascia-verde:#24321f;--fascia-ambra:#3a3220;--fascia-rosso:#3a2425;
--calore-1:#22344a;--calore-2:#2f5b8a;--calore-3:#4a86c5}
*{box-sizing:border-box}html{scroll-padding-top:16px}
body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 "Segoe UI",Aptos,Calibri,system-ui,sans-serif}
a{color:var(--link)}code{font-size:.92em}
.testata{background:#1f3864;color:#fff;padding:14px 20px;display:flex;flex-wrap:wrap;gap:6px 18px;align-items:baseline}
.testata h1{margin:0;font-size:20px;font-weight:600}.testata .etichetta{font-size:14px;opacity:.9}
.testata a{color:#cfe0ff;font-size:13px;margin-left:auto}
.impianto{display:grid;grid-template-columns:230px minmax(0,1fr);gap:0;min-height:calc(100vh - 56px)}
.menu{background:var(--surface);border-right:1px solid var(--rule);padding:12px 0;position:sticky;top:0;align-self:start;
max-height:100vh;overflow:auto}
.menu h2{font-size:11px;letter-spacing:.06em;text-transform:uppercase;color:var(--muted);margin:14px 18px 4px}
.menu a{display:block;padding:5px 18px;color:var(--ink-2);text-decoration:none;border-left:3px solid transparent;font-size:14px}
.menu a:hover{background:var(--surface-2);color:var(--ink)}
.menu .conta{color:var(--muted);font-size:12px}
main{padding:20px 24px 60px;min-width:0}
.tab{display:none}.tab:target{display:block}
body:not(:has(.tab:target)) #quadro{display:block}
.tab-testa{margin-bottom:14px}.tab-testa h2{margin:0;font-size:24px;color:var(--brand);font-weight:600}
.gruppo{margin:0;font-size:12px;text-transform:uppercase;letter-spacing:.06em;color:var(--muted)}
.come{max-width:980px;color:var(--ink-2);margin:6px 0 0}
.pannello{background:var(--surface);border:1px solid var(--rule);border-radius:8px;padding:14px 16px;margin:0 0 16px;min-width:0}
.pannello h3,.scheda h3{margin:0 0 10px;font-size:16px;color:var(--brand);font-weight:600}
.pannello h4{margin:12px 0 4px;font-size:14px}
.centro{text-align:center}
.griglia-2{display:grid;grid-template-columns:minmax(0,1fr) minmax(0,1fr);gap:16px}
.griglia-2.larga-sx{grid-template-columns:minmax(0,1.25fr) minmax(0,1fr)}.griglia-2.larga-dx{grid-template-columns:minmax(0,.7fr) minmax(0,1.3fr)}
.griglia-3{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:16px}
.tiles{display:grid;grid-template-columns:repeat(5,minmax(0,1fr));gap:12px;margin-bottom:16px}
.tiles.stretti{grid-template-columns:repeat(2,minmax(0,320px))}
.tile{background:var(--surface);border:1px solid var(--rule);border-top:4px solid var(--brand-2);border-radius:8px;padding:10px 14px}
.tile-allerta{border-top-color:var(--att)}.tile-ok{border-top-color:var(--mon)}.tile-attesa{border-top-color:var(--attz)}
.tile-et{margin:0;font-size:13px;color:var(--ink-2)}.tile-val{margin:2px 0;font-size:30px;font-weight:600;line-height:1.15}
.tile-sotto{margin:0;font-size:12.5px;color:var(--muted)}
.riquadro{background:var(--brand-soft);border-radius:8px;padding:10px 14px;margin:0 0 16px}
.riquadro.avviso{background:var(--attz-soft)}.riquadro.ok{background:var(--mon-soft)}
.nota,.piccolo{font-size:12.5px;color:var(--muted)}.nota{margin:8px 0 0}
.tenue{color:var(--muted)}.forte{font-weight:600}.rosso{color:var(--rosso)}
.legenda,.limiti{font-size:12px;color:var(--muted);max-width:1100px;margin:8px 0 0}
.tab-scorre{overflow-x:auto;max-width:100%}.tab-scorre.alta{max-height:520px;overflow:auto}
table{border-collapse:collapse;width:100%;font-size:13.5px}
th{background:var(--th);color:#fff;font-weight:600;text-align:left;padding:6px 8px;font-size:12.5px;vertical-align:bottom}
td{padding:6px 8px;border-bottom:1px solid var(--rule);vertical-align:top}
td.num,th.num{text-align:right;font-variant-numeric:tabular-nums}
table.compatta td{padding:4px 7px}
table.chiave-valore th{display:none}table.chiave-valore td:first-child{color:var(--ink-2)}
table.testo td:first-child{font-weight:600;width:200px;color:var(--brand)}table.testo th{background:var(--th-2)}
table.dati{font-size:12px}table.dati td{min-width:80px;max-width:300px}
table.dati .cella{display:-webkit-box;-webkit-line-clamp:3;-webkit-box-orient:vertical;overflow:hidden}
table.dati thead th{position:sticky;top:0}
.chip{display:inline-block;padding:1px 8px;border-radius:10px;font-size:12px;font-weight:600;white-space:nowrap;line-height:1.6}
.liv-attivazione{background:var(--att);color:var(--att-ink)}.liv-attenzione{background:var(--attz);color:var(--attz-ink)}
.liv-monitoraggio{background:var(--mon);color:var(--mon-ink)}
td.liv-fondo.liv-attivazione,p.liv-fondo.liv-attivazione{background:var(--att-soft);color:var(--ink)}
td.liv-fondo.liv-attenzione,p.liv-fondo.liv-attenzione{background:var(--attz-soft);color:var(--ink)}
td.liv-fondo.liv-monitoraggio,p.liv-fondo.liv-monitoraggio{background:var(--mon-soft);color:var(--ink)}
td.liv-fondo{border-left:4px solid var(--rule-2)}td.liv-fondo.liv-attivazione{border-left-color:var(--att)}
td.liv-fondo.liv-attenzione{border-left-color:var(--attz)}td.liv-fondo.liv-monitoraggio{border-left-color:var(--mon)}
.urg-agire{background:var(--agire);color:var(--agire-ink)}.urg-entro{background:var(--entro);color:var(--entro-ink)}
.urg-margine{background:var(--margine);color:var(--margine-ink)}
.kri-rosso{background:var(--rosso);color:#fff}.kri-ambra{background:var(--ambra);color:#1c2230}.kri-verde{background:var(--verde);color:#fff}
.kri-non-rilevato,.kri-nd,.kri-soglie-da-definire{background:var(--nr);color:#1c2230}
.stato-oss{background:var(--brand-soft);color:var(--brand)}
td.dim{text-align:center;width:44px}.dim-v{display:inline-block;min-width:26px;padding:1px 4px;border-radius:4px;
font-weight:600;font-variant-numeric:tabular-nums}
.dim-1{background:var(--d1)}.dim-2{background:var(--d2)}.dim-3{background:var(--d3)}.dim-4{background:var(--d4)}
.dim-5{background:var(--d5);color:var(--d5-ink)}.dim-nv{color:var(--muted);font-weight:400}
.ril{padding:1px 6px;border-radius:4px;font-size:12px}.ril-3{background:var(--calore-2)}.ril-2{background:var(--calore-1)}.ril-1{background:var(--d1)}
.id,.nowrap,td a[href^="#kri-"]{font-weight:600;white-space:nowrap}
.tend{white-space:nowrap}table.registro td:nth-child(2){min-width:190px}table.registro{font-size:13px}.sup{color:#7f6000;font-weight:600}
.barretta{display:inline-block;width:90px;height:10px;background:var(--surface-2);border-radius:5px;vertical-align:middle;overflow:hidden}
.barretta span{display:block;height:100%;border-radius:0 5px 5px 0}
span.b-urg-agire{background:var(--agire)}span.b-urg-entro{background:var(--entro)}span.b-urg-margine{background:var(--margine)}
.legenda-chip{display:flex;flex-wrap:wrap;gap:6px 14px;align-items:center;font-size:12.5px;color:var(--ink-2);margin:0 0 8px}
.conteggi{font-size:14px;margin-bottom:14px}
.sw{display:inline-block;width:14px;height:10px;border-radius:2px;vertical-align:middle}
.sw-x{background:var(--x)}.sw-y-entro{background:var(--y-entro)}.sw-y-oltre{background:var(--y-oltre)}
.sw-z{width:auto;height:auto;color:var(--z);background:none}.sw-su{background:var(--su)}.sw-giu{background:var(--giu)}
.sw-int1{background:var(--i1)}.sw-int2{background:var(--i2)}.sw-int3{background:var(--i3)}
.matrice{width:auto;border-collapse:separate;border-spacing:3px}.matrice td{width:96px;height:64px;text-align:center;
vertical-align:middle;border:0;border-radius:4px;font-size:12.5px;font-weight:600}.matrice.piccola td{width:52px;height:38px}
.matrice th{background:none;color:var(--ink-2);font-weight:400;text-align:right;font-size:12px;vertical-align:middle}
.matrice th.col-i{text-align:center;vertical-align:top}.matrice a{display:inline-block;margin:1px 3px;color:var(--ink);text-decoration:none}
.fascia-verde{background:var(--fascia-verde)}.fascia-ambra{background:var(--fascia-ambra)}.fascia-rosso{background:var(--fascia-rosso)}
.asse-matrice{font-size:12px;color:var(--muted);margin:0 0 4px}
.incrocio td{text-align:center;width:44px}.incrocio th{font-size:12px}
.calore-1{background:var(--calore-1)}.calore-2{background:var(--calore-2)}.calore-3{background:var(--calore-3);color:#fff}
dl.commento{display:grid;grid-template-columns:190px 1fr;gap:6px 14px;margin:0}dl.commento dt{font-weight:600;color:var(--brand)}
dl.commento dd{margin:0}
ul.elenco-rischi{list-style:none;padding:0;margin:0;columns:1}ul.elenco-rischi li{padding:3px 0;border-bottom:1px solid var(--rule)}
ul.elenco-rischi .chip{min-width:96px;text-align:center}
table.segnalati td:nth-child(3),table.segnalati td:nth-child(4){font-size:13px}
.multipli{display:grid;grid-template-columns:repeat(auto-fill,minmax(300px,1fr));gap:14px}
figure.multiplo{margin:0;border:1px solid var(--rule);border-radius:6px;padding:8px 10px;background:var(--surface)}
figure.multiplo figcaption{font-size:13px;margin-bottom:4px}
.stato-az{font-weight:600}.stato-respinta,.stato-annullata{color:var(--muted)}.stato-completata{color:var(--verde)}
.strumenti{display:flex;gap:8px;margin:0 0 8px}.strumenti input{flex:1;max-width:360px;padding:6px 10px;border:1px solid var(--rule-2);
border-radius:6px;background:var(--surface);color:var(--ink);font:inherit}
.strumenti button{padding:6px 12px;border:1px solid var(--brand-2);background:var(--surface);color:var(--brand);border-radius:6px;
font:inherit;cursor:pointer}
.indice-dati{font-size:14px}
/* schede */
.scheda{background:var(--surface);border:1px solid var(--rule);border-radius:10px;padding:0 22px 18px;max-width:1100px}
.scheda-testa{margin:0 -22px 16px;padding:14px 22px;background:#1f3864;color:#fff;border-radius:10px 10px 0 0;border-bottom:6px solid var(--rule-2)}
.scheda-testa.liv-bordo-attivazione{border-bottom-color:var(--att)}.scheda-testa.liv-bordo-attenzione{border-bottom-color:var(--attz)}
.scheda-testa.liv-bordo-monitoraggio{border-bottom-color:var(--mon)}.scheda-testa.oss{background:#2f5597}
.scheda-testa h2{margin:2px 0 6px;font-size:22px;color:#fff}.scheda-pos{margin:0;font-size:12.5px;opacity:.85}
.scheda-numeri{margin:0 0 8px;display:flex;gap:10px;align-items:center;flex-wrap:wrap}
.scheda-numeri .grande{font-size:30px;font-weight:600;line-height:1}.freccia{opacity:.9}
.meta{display:flex;flex-wrap:wrap;gap:4px 18px;margin:0;font-size:13px;opacity:.92}.scheda-testa .sup{color:#ffd966}
.scheda h3{margin-top:18px}
table.dimensioni td:first-child{width:120px}
tr.perche td{background:var(--surface-2);font-size:12.5px;display:table-cell}tr.perche span{display:inline-block;width:49%;vertical-align:top;padding-right:10px}
.andamento{display:flex;gap:18px;align-items:center;flex-wrap:wrap}.andamento .tab-scorre{flex:0 1 420px}
.barra-mosca{position:relative;display:flex;height:16px;max-width:520px;background:var(--surface-2);border-radius:4px;margin:6px 0 14px}
.barra-mosca span{display:block;height:100%}.barra-mosca span+span{border-left:2px solid var(--surface)}
.segno-z{position:absolute;top:-4px;transform:translateX(-50%);color:var(--z);font-style:normal;font-size:16px}
.nav-schede{display:flex;justify-content:space-between;margin-top:22px;padding-top:12px;border-top:1px solid var(--rule);font-size:14px}
/* grafici */
svg.grafico{width:100%;height:auto;max-width:860px;display:block;margin:0 auto;overflow:visible}
svg.radar{max-width:720px}.pannello svg.radar{max-width:520px}#radar svg.radar{max-width:760px}
svg.kri{max-width:none}svg.spark{display:block}
svg text{fill:var(--ink-2);font-size:12px}
svg .et-cat{fill:var(--ink);font-size:12.5px}svg .et-cat .id{font-weight:600}svg .et-val{fill:var(--ink);font-weight:600;font-size:12.5px}
svg .et-asse{fill:var(--muted);font-size:11.5px}svg .tenue{fill:var(--muted)}
svg .et-settore{fill:var(--brand);font-size:12.5px}svg .et-settore .cod{font-weight:700}
svg .et-bolla{fill:#1c2230;font-weight:700;font-size:11.5px;pointer-events:none;paint-order:stroke;
stroke:rgba(255,255,255,.7);stroke-width:3px;stroke-linejoin:round}
svg .et-anello{fill:var(--muted);font-size:11px}svg .et-mosca{fill:var(--agire);font-weight:700;font-size:12px}
svg .et-leg{fill:var(--ink-2);font-size:12.5px}
svg .asse{stroke:var(--rule);stroke-width:1}svg .griglia{stroke:var(--rule);stroke-width:1}svg .asse-base{stroke:var(--ink-2);stroke-width:1}
svg .anello{fill:none;stroke:var(--rule-2);stroke-width:1.5}svg .anello.tratto{stroke-dasharray:5 4}
svg .mosca{fill:none;stroke:var(--agire);stroke-width:3}
svg .bolla{fill-opacity:.82;stroke:var(--surface);stroke-width:2}
svg .bolla.liv-attivazione{fill:var(--att)}svg .bolla.liv-attenzione{fill:var(--attz)}svg .bolla.liv-monitoraggio{fill:var(--mon)}
svg .campione{fill:var(--brand-soft);stroke:var(--brand-2);stroke-width:1}
svg .b-urg-agire{fill:var(--agire)}svg .b-urg-entro{fill:var(--entro)}svg .b-urg-margine{fill:var(--margine)}
svg .b-x{fill:var(--x)}svg .b-y-entro{fill:var(--y-entro)}svg .b-y-oltre{fill:var(--y-oltre)}svg .rombo{fill:var(--z);stroke:var(--surface);stroke-width:2}
svg .b-su{fill:var(--su)}svg .b-giu{fill:var(--giu)}
svg .b-int1{fill:var(--i1)}svg .b-int2{fill:var(--i2)}svg .b-int3{fill:var(--i3)}
svg .linea{fill:none;stroke:var(--linea);stroke-width:2;stroke-linejoin:round;stroke-linecap:round}
svg .soglia-att{fill:none;stroke:var(--soglia-att);stroke-width:1.5;stroke-dasharray:5 4}
svg .soglia-all{fill:none;stroke:var(--soglia-all);stroke-width:1.5;stroke-dasharray:5 4}
svg .pallino{fill:var(--linea);stroke:var(--surface);stroke-width:2}svg .pallino-fine{fill:var(--linea);stroke:var(--surface);stroke-width:2}
svg .hit{fill:transparent}svg .pallino-hover{fill:var(--linea);opacity:0}
svg .punto:hover .pallino-hover{opacity:1}
svg .punto:hover .bolla,svg .punto:focus .bolla{fill-opacity:1;stroke:var(--ink);stroke-width:2}
svg .punto:hover path,svg .punto:focus path{filter:brightness(1.08)}
svg .punto:hover .hit{fill:var(--brand-soft);fill-opacity:.45}
svg a{cursor:pointer}svg a:focus{outline:none}svg a:focus-visible .hit{fill:var(--brand-soft);fill-opacity:.6}
#tip{position:fixed;z-index:10;pointer-events:none;background:var(--surface);color:var(--ink);border:1px solid var(--rule-2);
border-radius:6px;padding:6px 9px;font-size:12.5px;line-height:1.4;max-width:360px;box-shadow:0 4px 14px rgba(0,0,0,.18);
white-space:pre-line;display:none}
@media (max-width:1000px){.impianto{display:block}.menu{position:static;max-height:none;border-right:0;border-bottom:1px solid var(--rule);
display:flex;flex-wrap:wrap;gap:0 4px;padding:6px 10px}.menu h2{width:100%;margin:8px 6px 2px}.menu a{border-left:0;padding:3px 8px}
.tiles{grid-template-columns:repeat(2,minmax(0,1fr))}.griglia-2,.griglia-2.larga-sx,.griglia-2.larga-dx,.griglia-3{grid-template-columns:minmax(0,1fr)}
main{padding:16px}dl.commento{grid-template-columns:1fr}.scheda{padding:0 16px 16px}.scheda-testa{margin:0 -16px 14px;padding:12px 16px}
tr.perche span{width:100%}}
@media print{.menu,.testata a,.no-stampa,.nav-schede,.strumenti{display:none!important}.impianto{display:block}main{padding:0}
.tab{display:block!important;break-before:page}#quadro{break-before:auto}.pannello{break-inside:avoid;border-color:#ccc}
body{background:#fff;font-size:10pt}.scheda{border:0;max-width:none}.tab.no-stampa{display:none!important}
.tab-scorre.alta{max-height:none}svg.radar{max-width:560px}#radar svg.radar{max-width:600px}
main{zoom:.62}.griglia-2{grid-template-columns:minmax(0,1fr) minmax(0,1fr)}.griglia-2.larga-sx{grid-template-columns:minmax(0,1.25fr) minmax(0,1fr)}
.griglia-2.larga-dx{grid-template-columns:minmax(0,.7fr) minmax(0,1.3fr)}.griglia-3{grid-template-columns:repeat(3,minmax(0,1fr))}
.tiles{grid-template-columns:repeat(5,minmax(0,1fr))}dl.commento{grid-template-columns:190px 1fr}tr.perche span{width:49%}
.tab-scorre{overflow:visible}@page{size:A4;margin:10mm}}
"""

JS = r"""
(function(){
  var tip=document.getElementById('tip');
  document.querySelectorAll('svg title').forEach(function(t){
    var g=t.parentNode; g.setAttribute('data-tip',t.textContent); g.setAttribute('tabindex', g.closest('a')?'-1':'0'); t.remove();
  });
  function mostra(e,el){tip.textContent=el.getAttribute('data-tip');tip.style.display='block';
    var x=(e.clientX||el.getBoundingClientRect().right)+14,y=(e.clientY||el.getBoundingClientRect().top)+14;
    var w=tip.offsetWidth,h=tip.offsetHeight;if(x+w>innerWidth-8)x=innerWidth-w-8;if(y+h>innerHeight-8)y=y-h-28;
    tip.style.left=x+'px';tip.style.top=y+'px';}
  document.addEventListener('pointermove',function(e){var el=e.target.closest&&e.target.closest('[data-tip]');
    if(el)mostra(e,el);else tip.style.display='none';});
  document.addEventListener('focusin',function(e){var el=e.target.closest&&(e.target.closest('[data-tip]')||
    (e.target.querySelector&&e.target.querySelector('[data-tip]')));if(el)mostra({},el);});
  document.addEventListener('focusout',function(){tip.style.display='none';});
  document.querySelectorAll('[data-filtra]').forEach(function(inp){
    var righe=document.getElementById(inp.getAttribute('data-filtra')).tBodies[0].rows;
    inp.addEventListener('input',function(){var q=inp.value.toLowerCase();
      for(var i=0;i<righe.length;i++){righe[i].style.display=righe[i].textContent.toLowerCase().indexOf(q)<0?'none':'';}});
  });
  document.querySelectorAll('[data-csv]').forEach(function(b){b.addEventListener('click',function(){
    var t=document.getElementById(b.getAttribute('data-csv')),out=[];
    Array.prototype.forEach.call(t.rows,function(r){if(r.style.display==='none')return;
      out.push(Array.prototype.map.call(r.cells,function(c){return '"'+c.textContent.replace(/"/g,'""')+'"';}).join(';'));});
    var a=document.createElement('a');a.href=URL.createObjectURL(new Blob(['﻿'+out.join('\r\n')],{type:'text/csv'}));
    a.download=b.getAttribute('data-nome')+'.csv';a.click();});});
})();
"""


GRUPPI_REPORT = ["Report al comitato", "Schede dei rischi", "Aiuto"]


def gruppi(con_lavoro=False, anteprima=False):
    """Le pagine del file: il report, le schede e la guida; con_lavoro anche le viste di lavoro; l'anteprima tutto."""
    if anteprima:
        return GRUPPI
    return [(g, voci) for g, voci in GRUPPI if g in GRUPPI_REPORT or (con_lavoro and g == "Lavoro")]


def validazione_del_ciclo(v):
    """(validatore, data) dell'ultima decisione del ciclo nello storico delle validazioni."""
    righe = [x for x in v.val if x.get("Data_Validazione")]
    if not righe:
        return None, None
    ultima = max(righe, key=lambda x: str(x["Data_Validazione"]))
    return ultima.get("Validato_Da"), ultima["Data_Validazione"]


def pagina(v, reg, nome_file, con_lavoro=False, anteprima=False):
    t = {f: reg.testi_foglio(f) for f in ("Quadro", "Classifica", "Radar_Mosca", "Matrice", "Delta", "Fenomeni", "Metodologia",
                                          "Registro", "Storico", "Indicatori", "Notizie", "Piano_azioni", "Osservazione")}
    scelti = gruppi(con_lavoro, anteprima)
    voci = {voce for _, vv in scelti for voce in vv}
    schede_html, ordine = s_schede(v)
    costruttori = [("quadro", lambda: s_quadro(v, t["Quadro"])), ("classifica", lambda: s_classifica(v, t["Classifica"])),
                   ("radar", lambda: s_radar(v, t["Radar_Mosca"])), ("matrice", lambda: s_matrice(v, t["Matrice"])),
                   ("delta", lambda: s_delta(v, t["Delta"])), ("fenomeni", lambda: s_fenomeni(v, t["Fenomeni"])),
                   ("metodologia", lambda: s_metodologia(v, reg, t["Metodologia"])),
                   ("registro", lambda: s_registro(v, t["Registro"])), ("storico", lambda: s_storico(v, t["Storico"])),
                   ("indicatori", lambda: s_indicatori(v, t["Indicatori"])), ("notizie", lambda: s_notizie(v, t["Notizie"])),
                   ("piano", lambda: s_piano(v, t["Piano_azioni"])), ("osservazione", lambda: s_osservazione(v, t["Osservazione"])),
                   ("validazione", lambda: s_validazione(v, {})), ("parametri", lambda: s_parametri(v, reg, {})),
                   ("dizionario", lambda: s_dizionario(v, reg, {})), ("schede", lambda: schede_html),
                   ("dati", lambda: s_dati(v, reg)), ("guida", lambda: s_guida(v, reg, anteprima, voci))]
    sezioni = [f() for voce, f in costruttori if voce in voci]
    menu = []
    for gruppo, vv in scelti:
        menu.append(f"<h2>{esc(gruppo)}</h2>")
        for voce in vv:
            extra = f' <span class="conta">({len(ordine)})</span>' if voce == "schede" else ""
            menu.append(f'<a href="#{voce}">{esc(TITOLI[voce])}{extra}</a>')
    # voce attiva del menu: una regola per pagina (le schede accendono "Indice delle schede")
    attivi = [f'body:has(#{voce}:target) .menu a[href="#{voce}"]' for _, vv in scelti for voce in vv]
    attivi += [f'body:has(#scheda-{rid}:target) .menu a[href="#schede"]' for rid in ordine]
    attivi.append('body:not(:has(.tab:target)) .menu a[href="#quadro"]')
    css_attivo = ",".join(attivi) + "{border-left-color:var(--brand-2);background:var(--brand-soft);color:var(--brand);font-weight:600}"
    titolo = f"{NOME_ANALISI} {v.ciclo}"
    chi, quando = validazione_del_ciclo(v)
    etichetta = esc(v.etichetta())
    if quando and not anteprima and etichetta.endswith("· validato"):
        etichetta += f" il {data_it(quando)}" + (f" da {esc(chi)}" if chi else "")
    indice = "index.html" if anteprima else f"../{INDICE}"
    descrizione = (f"report del ciclo {esc(v.ciclo)}, anteprima HTML dell'analisi 16" if anteprima else
                   f"report al comitato del ciclo {esc(v.ciclo)}")
    return (f'<!doctype html><html lang="it"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,'
            f'initial-scale=1"><title>{esc(titolo)}</title><meta name="description" content="{esc(NOME_ANALISI)}: '
            f'{descrizione}"><style>{CSS}{css_attivo}</style></head><body>'
            f'<header class="testata"><h1>{esc(NOME_ANALISI)}</h1><span class="etichetta">{etichetta} · '
            f'{RISERVATEZZA}</span><a href="{indice}">Archivio dei cicli</a></header>'
            f'<div class="impianto"><nav class="menu" aria-label="Pagine del report">{"".join(menu)}</nav>'
            f'<main>{"".join(sezioni)}</main></div><div id="tip" role="tooltip"></div><script>{JS}</script></body></html>')


def riassunto(v, file):
    """La riga dell'indice di un ciclo (si salva accanto al report: l'indice si ricompone senza rileggere il registro)."""
    mess, _ = parti_commento(v.sintesi.get("Sintesi_A1"))
    n_da, _ = v.stato_validazione()
    chi, quando = validazione_del_ciclo(v)
    return {"ciclo": v.ciclo, "file": file, "tipo": "Straordinario" if "-S" in v.ciclo else "Trimestrale",
            "data_rif": v.data_rif, "validato": n_da == 0, "validatore": chi, "data_validazione": quando,
            "nel_radar": len(v.rischi), "in_osservazione": len(v.oss_ids),
            "in_attivazione": [x["id"] for x in v.classifica if x["livello"] == "Attivazione"],
            "kri_rossi": sum(k["Stato"] == "Rosso" for k in v.kri), "messaggio": accorcia(mess, 260),
            "score_max": max((x["score"] for x in v.rischi), default=None)}


def pagina_indice(righe_cicli, anteprima=False):
    """L'archivio dei cicli: una riga per report (dal piu' recente) e lo score piu' alto per ciclo."""
    cicli = sorted(righe_cicli, key=lambda r: r["ciclo"])
    righe = []
    for r in reversed(cicli):
        righe.append(f'<tr><td><a class="id" href="{esc(r["file"])}">{esc(r["ciclo"])}</a></td>'
                     f'<td>{esc(r["tipo"])}</td><td>{data_it(r["data_rif"])}</td>'
                     f'<td>{"✓ validato" if r["validato"] else "da validare"}'
                     + (f' <span class="tenue">{data_it(r["data_validazione"])}</span>' if r.get("data_validazione") else "")
                     + f'</td><td class="num">{r["nel_radar"]}</td>'
                     f'<td class="num">{r["in_osservazione"]}</td><td>{esc(", ".join(r["in_attivazione"])) or "—"}</td>'
                     f'<td class="num">{r["kri_rossi"]}</td><td class="piccolo">{esc(r["messaggio"])}</td></tr>')
    corrente = cicli[-1]
    serie = [(r["ciclo"], r["score_max"]) for r in cicli if r["score_max"] is not None]
    nota = ("Anteprima dell'analisi 16 con i dati della demo (Demo_W1_W2). Generata da "
            "Analisi_Architettura/16_genera_anteprima_html.py." if anteprima else
            "Ricomposto dal comando report (e da valida) a ogni report HTML che esce; accanto a ogni report c'e' il PDF.")
    corpo = ('<p class="come">Un report per ciclo, ognuno in un file che si apre nel browser e si conserva come evidenza. '
             'L' + "'" + 'ultimo ciclo e' + "'" + f' il <a href="{esc(corrente["file"])}">{esc(corrente["ciclo"])}</a>.</p>'
             '<div class="pannello"><div class="tab-scorre"><table><thead><tr><th>Ciclo</th><th>Tipo</th><th>Dati al</th>'
             '<th>Validazione</th><th class="num">Nel radar</th><th class="num">In osservazione</th><th>In attivazione</th>'
             '<th class="num">KRI in rosso</th><th>Messaggio di A1</th></tr></thead><tbody>' + "".join(righe)
             + '</tbody></table></div></div>'
             + ('<div class="pannello"><h3>Score composito piu' + "'" + ' alto per ciclo</h3>'
                + sparkline([s for _, s in serie], [c for c, _ in serie], 520, 90) + '</div>' if len(serie) > 1 else "")
             + f'<p class="legenda">{esc(nota)}</p>')
    return (f'<!doctype html><html lang="it"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,'
            f'initial-scale=1"><title>{NOME_ANALISI} archivio</title><style>{CSS}main{{max-width:1300px;margin:0 auto}}</style>'
            f'</head><body><header class="testata"><h1>{NOME_ANALISI}</h1><span class="etichetta">Archivio dei cicli · '
            f'{RISERVATEZZA}</span></header><main><section class="tab" id="indice" style="display:block"><header class="tab-testa">'
            f'<p class="gruppo">Archivio</p><h2>I report dei cicli</h2></header>{corpo}</section></main>'
            f'<div id="tip" role="tooltip"></div><script>{JS}</script></body></html>')


# ================================================================ il report di un ciclo, con il cancello
class Cancello(Exception):
    """I numeri dell'HTML non tornano con quelli di Excel: l'HTML non esce."""

    def __init__(self, ciclo, uguali, diversi):
        self.ciclo, self.uguali, self.diversi = ciclo, uguali, diversi
        super().__init__(f"report HTML del ciclo {ciclo} non uscito: {len(diversi)} numeri diversi da quelli di Excel "
                         f"({uguali} uguali). Primi: " + "; ".join(diversi[:5]))


def genera(reg, ciclo=None, con_lavoro=False, nome_file=None):
    """Il report HTML del ciclo (di default il corrente) dopo il cancello. Restituisce (html, riassunto, uguali).
    Solleva Cancello se un numero non torna con Excel."""
    ciclo = ciclo or reg.nome("cicloCorrente")
    v = Vista(reg, ciclo)
    uguali, diversi = verifica(reg, v)
    if diversi:
        raise Cancello(ciclo, uguali, diversi)
    nome_file = nome_file or f"Report_comitato_{ciclo}.html"
    return pagina(v, reg, nome_file, con_lavoro), riassunto(v, f"{ciclo}/{nome_file}"), uguali


def scrivi_indice(cartella_cicli):
    """Cicli/Indice_report.html dai riassunti accanto ai report (Cicli/<ciclo>/Report_comitato_<ciclo>.json)."""
    cartella_cicli = Path(cartella_cicli)
    righe = []
    for f in sorted(cartella_cicli.glob("*/Report_comitato_*.json")):
        try:
            r = json.loads(f.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        if (f.parent / Path(r.get("file", "")).name).exists():
            righe.append(r)
    if not righe:
        return None
    indice = cartella_cicli / INDICE
    indice.write_text(pagina_indice(righe), encoding="utf-8")
    return indice


def main(argv=None):
    ap = argparse.ArgumentParser(description="Report al comitato in HTML dal file del registro (senza Excel)")
    ap.add_argument("--registro", type=Path, required=True)
    ap.add_argument("--uscita", type=Path, required=True, help="cartella in cui scrivere il report e il riassunto")
    ap.add_argument("--ciclo")
    ap.add_argument("--con-lavoro", dest="con_lavoro", action="store_true")
    a = ap.parse_args(argv)
    reg = RegistroFile(a.registro)
    try:
        testo, riga, uguali = genera(reg, a.ciclo, a.con_lavoro)
    except Cancello as ex:
        print(f"FERMATO: {ex}")
        for d in ex.diversi:
            print("  DIVERSO:", d)
        return 1
    a.uscita.mkdir(parents=True, exist_ok=True)
    file = a.uscita / Path(riga["file"]).name
    file.write_text(testo, encoding="utf-8")
    (file.with_suffix(".json")).write_text(json.dumps(riga, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Report HTML del ciclo {riga['ciclo']}: {file} ({file.stat().st_size // 1024} KB); cancello: {uguali} numeri "
          "uguali a quelli di Excel.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
