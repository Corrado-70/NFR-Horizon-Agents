"""Verifica la mappa sorgente -> destinazione in entrambe le direzioni.

1. Ogni campo foglia degli schemi A1 e A3 ha una destinazione: una colonna,
   un parametro, o una voce motivata di 'non_nel_workbook'.
2. Ogni colonna ha una sorgente valida: un percorso che esiste nello schema,
   un campo derivato che lo scrittore sa calcolare, o una formula.
3. Le tabelle del template hanno esattamente le colonne della mappa, in ordine.
4. La documentazione SharePoint elenca tutte le colonne delle liste.

Esecuzione: python test_mappatura.py   (codice di uscita 0 = tutto ok)
"""

import re
import sys

from comune import KERNEL, RADICE, TEMPLATE, Esito, carica

SCHEMI = {"A1": "schema_output_A1.json", "A3": "schema_output_A3.json"}


def foglie(schema, radice=None, prefisso=""):
    """Percorsi di tutti i campi foglia (e dei nodi intermedi) di uno schema."""
    radice = radice or schema
    if "$ref" in schema:
        nodo = radice
        for parte in schema["$ref"][2:].split("/"):
            nodo = nodo[parte]
        return foglie(nodo, radice, prefisso)
    percorsi = {"foglie": set(), "nodi": set()}
    if schema.get("type") == "object" or "properties" in schema:
        for nome, sotto in schema.get("properties", {}).items():
            p = f"{prefisso}.{nome}" if prefisso else nome
            percorsi["nodi"].add(p)
            figli = foglie(sotto, radice, p)
            percorsi["foglie"] |= figli["foglie"]
            percorsi["nodi"] |= figli["nodi"]
        return percorsi
    if schema.get("type") == "array" and isinstance(schema.get("items"), dict):
        elem = schema["items"]
        if "$ref" in elem:
            nodo = radice
            for parte in elem["$ref"][2:].split("/"):
                nodo = nodo[parte]
            elem = nodo
        if elem.get("type") == "object":
            figli = foglie(elem, radice, prefisso + "[]")
            percorsi["nodi"].add(prefisso + "[]")
            percorsi["foglie"] |= figli["foglie"]
            percorsi["nodi"] |= figli["nodi"]
            return percorsi
    percorsi["foglie"].add(prefisso)
    return percorsi


def sorgenti_per_payload(mappa):
    """{payload: [(dove, sorgente, colonna_spec)]} da tabelle, anagrafica e parametri."""
    out = {"A1": [], "A3": []}
    for nome_t, t in mappa["tabelle"].items():
        for c in t["colonne"]:
            out[t["payload"]].append((f"{nome_t}.{c['nome']}", c["sorgente"], c))
    for c in mappa["anagrafica"]["colonne"]:
        out["A1"].append((f"anagrafica.{c['nome']}", c["sorgente"], c))
    out["A1"].append(("Parametri.cicloCorrente", mappa["parametri"]["cicloCorrente"], {}))
    out["A3"].append(("Parametri.settimanaCorrente", mappa["parametri"]["settimanaCorrente"], {}))
    return out


def percorso_payload(sorgente):
    return sorgente.split(":", 1)[1].split("|")[0]


def main():
    e = Esito("Mappa sorgente -> destinazione")
    mappa = carica(KERNEL / "mappatura_campi.json")
    per_payload = sorgenti_per_payload(mappa)
    non_nel_wb = {k for k in mappa["non_nel_workbook"] if not k.startswith("_")}

    try:
        import esegui_ciclo_locale as campo
        derivati_noti = set(campo.DERIVATI)
    except Exception as ex:  # lo scrittore non ancora scritto o non importabile
        derivati_noti = None
        e.verifica(False, "lo scrittore di campo e' importabile", repr(ex))

    for nome_payload, file_schema in SCHEMI.items():
        albero = foglie(carica(KERNEL / file_schema))
        coperti = set(non_nel_wb)
        for dove, sorgente, spec in per_payload[nome_payload]:
            tipo = sorgente.split(":", 1)[0]
            if tipo == "payload":
                p = percorso_payload(sorgente)
                coperti.add(p)
                e.verifica(p in albero["foglie"] or p in albero["nodi"],
                           f"{nome_payload} {dove}: la sorgente '{p}' esiste nello schema")
            elif tipo == "derivato":
                if derivati_noti is not None:
                    nome = sorgente.split(":", 1)[1]
                    e.verifica(nome in derivati_noti,
                               f"{dove}: lo scrittore sa calcolare il derivato '{nome}'")
            elif tipo == "formula":
                e.verifica(sorgente.startswith("formula:="), f"{dove}: formula ben formata")
            else:
                e.verifica(False, f"{dove}: tipo di sorgente sconosciuto '{tipo}'")
            coperti.update(spec.get("usa", []))

        # Una foglia e' coperta se e' mappata essa stessa o un suo antenato
        # (es. 'rischi[].fonti' serializzato in JSON copre 'rischi[].fonti[].data').
        # 'sintesi' in non_nel_workbook vale per la radice di entrambi i payload.
        scoperte = [f for f in sorted(albero["foglie"])
                    if not any(f == c or f.startswith(c + ".") or f.startswith(c + "[]")
                               for c in coperti)]
        e.verifica(not scoperte, f"{nome_payload}: ogni campo dello schema ha una destinazione",
                   "\n".join(scoperte))

    # Ogni colonna ha il suo significato: il generatore lo scrive sopra l'intestazione.
    senza = [f"{nome}.{c['nome']}" for nome, spec in [*mappa["tabelle"].items(), ("validazione", mappa["validazione"]),
                                                      ("anagrafica", mappa["anagrafica"])]
             for c in spec["colonne"] if len(c.get("significato", "")) < 10]
    e.verifica(not senza, "ogni colonna della mappa ha un significato (riga 3 delle tabelle del workbook)", str(senza))

    # Le tabelle del template combaciano con la mappa (se il template esiste gia').
    if TEMPLATE.exists():
        from openpyxl import load_workbook
        wb = load_workbook(TEMPLATE)
        tabelle_wb = {}
        for ws in wb.worksheets:
            for t in ws.tables.values():
                tabelle_wb[t.name] = (ws.title, [c.name for c in t.tableColumns])
        for nome_t, t in mappa["tabelle"].items():
            attese = [c["nome"] for c in t["colonne"]]
            foglio, trovate = tabelle_wb.get(nome_t, (None, None))
            e.verifica(trovate == attese, f"template: {nome_t} ha le colonne della mappa",
                       f"attese {attese}\ntrovate {trovate}")
            e.verifica(foglio == t["foglio"], f"template: {nome_t} sta nel foglio {t['foglio']}")
        v = mappa["validazione"]
        foglio, trovate = tabelle_wb.get(v["tabella"], (None, None))
        e.verifica(foglio == v["foglio"] and trovate == [c["nome"] for c in v["colonne"]],
                   f"template: {v['tabella']} ha le colonne della mappa (sezione validazione)", f"trovate {trovate}")
        for oggetto, o in v["oggetti"].items():
            colonne = [c["nome"] for c in mappa["tabelle"][o["tabella"]]["colonne"]]
            mancanti = [c for c in [o["id"], o["nome"], *o["campi"], "Validazione"] + ([o["periodo"]] if o["periodo"] else [])
                        if c not in colonne]
            e.verifica(not mancanti, f"validazione: i campi di {oggetto} esistono in {o['tabella']}", str(mancanti))
            # Griglia del foglio Validazione (D-67): lo script la legge e scrive solo con questi nomi.
            attesi = [f"VG_{oggetto}_{n}" for n in ["ID", "Nome", "Motivazione", "Periodo", "Appoggio", "righe", "titolo",
                                                     *o["campi"], *(f"{c}_P" for c in o["campi"])]
                                                    + (["Respingi"] if o["respingibile"] else [])]
            assenti = [n for n in attesi if n not in wb.defined_names]
            fuori = [n for n in attesi if n in wb.defined_names and not wb.defined_names[n].attr_text.startswith("'Validazione'!")]
            e.verifica(not assenti and not fuori, f"griglia di validazione: i nomi VG_{oggetto}_* esistono nel foglio Validazione",
                       f"assenti {assenti} fuori dal foglio {fuori}")
            if not o["respingibile"]:
                e.verifica(f"VG_{oggetto}_Respingi" not in wb.defined_names, f"griglia: {oggetto} non ha la colonna Respingi")
        e.verifica("valStato" in wb.defined_names, "griglia di validazione: la riga dei periodi da validare (valStato) esiste")
    else:
        e.verifica(False, "il template esiste (generarlo con 03_Excel/genera_dashboard.py)")

    # La documentazione delle liste cita ogni colonna.
    doc_sp = RADICE / "04_Riferimenti" / "struttura_e_liste.md"
    if e.verifica(doc_sp.exists(), "esiste 04_Riferimenti/struttura_e_liste.md"):
        testo = doc_sp.read_text(encoding="utf-8")
        for lista, spec in mappa["liste_sharepoint"].items():
            if lista.startswith("_"):
                continue
            base = [c["nome"] for c in mappa["tabelle"][spec["da_tabella"]]["colonne"]
                    if c["nome"] not in spec["colonne_escluse"]]
            attese = spec["colonne_aggiunte"] + base
            sezione = testo.split(f"### {lista}", 1)[-1].split("\n### ", 1)[0]
            mancanti = [c for c in attese if not re.search(rf"`{c}`", sezione)]
            e.verifica(f"### {lista}" in testo and not mancanti,
                       f"04_Riferimenti/struttura_e_liste.md documenta tutte le colonne di {lista}", ", ".join(mancanti))

    return 0 if e.stampa() else 1


if __name__ == "__main__":
    sys.exit(main())
