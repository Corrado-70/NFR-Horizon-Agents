"""Validatore, controlli meccanizzati di A5, dati di esempio, consistenza.

Senza Excel: gira ovunque ci sia Python. Esegue:
1. validatore su payload conformi e su difetti mirati (incluso cio' che lo
   schema NON puo' vedere: le soglie invertite passano la validazione);
2. controlli di A5 su un ciclo pulito (deve approvare) e su difetti mirati,
   uno per controllo (deve fermare il ciclo col controllo giusto);
3. i dati di esempio del template rispettano la rubrica;
4. due run sullo stesso input danno gli stessi giudizi;
5. precontrollo dell'output di A1 e allegato di A1 (D-69);
6. kit della prova di Copilot Studio: la skill precontrollo-nfr contiene il codice del pacchetto e risponde
   come il comando precontrollo (D-70);
7. schede dei rischi: schema e controlli dell'output di A6, file delle schede.
"""

import copy
import sys

from comune import DATI, EXCEL, KERNEL, Esito, carica   # per primo: aggiunge 07_Senza_Studio al percorso
import controlli_a5
from confronta import abbina, confronta
from valida_output import valida

S_A1 = carica(KERNEL / "schema_output_A1.json")
S_A3 = carica(KERNEL / "schema_output_A3.json")
S_A5 = carica(KERNEL / "schema_output_A5.json")
DM = "Da motivare"                      # D-71
RUN1 = carica(DATI / "run1_A1_2026Q3.json")
BASE_Q3 = carica(DATI / "baseline_2026Q3.json")


def rischio(p, rid):
    return next(r for r in p["rischi"] if r["id_rischio"] == rid)


def test_validatore(e):
    for nome, schema in [("run1_A1_2026Q3.json", S_A1), ("run2_A1_2026Q3.json", S_A1), ("run3_A1_2026Q4.json", S_A1),
                         ("run_A3_2026-10-12.json", S_A3), ("run_A5_2026Q3.json", S_A5), ("run_A5_2026Q4.json", S_A5)]:
        errori = valida(carica(DATI / nome), schema)
        e.verifica(not errori, f"validatore: {nome} conforme", "\n".join(errori))

    def difetto(descrizione, modifica, atteso):
        p = copy.deepcopy(RUN1)
        modifica(p)
        errori = valida(p, S_A1)
        e.verifica(any(atteso in x for x in errori), f"validatore rifiuta: {descrizione}", "\n".join(errori) or "nessun errore")

    difetto("campo score_composito aggiunto", lambda p: rischio(p, "NFR-01").update(score_composito=3.8), "campo non previsto 'score_composito'")
    difetto("punteggio interpolato 3,5", lambda p: rischio(p, "NFR-01").update(impatto=3.5), "tipo float")
    difetto("booleano al posto di un intero", lambda p: rischio(p, "NFR-01").update(probabilita=True), "tipo bool")
    difetto("confidence, tolta il 30/09 sera (D-58)", lambda p: rischio(p, "NFR-01").update(confidence=4),
            "campo non previsto 'confidence'")
    difetto("fonte senza rilevanza (D-58)", lambda p: rischio(p, "NFR-01")["fonti"][0].pop("rilevanza"),
            "manca il campo obbligatorio 'rilevanza'")
    difetto("rilevanza 4 (scala 1-3)", lambda p: rischio(p, "NFR-01")["fonti"][0].update(rilevanza=4), "> massimo 3")
    difetto("anni di adeguamento oltre 30 (D-60)", lambda p: rischio(p, "NFR-11").update(anni_adeguamento=40), "> massimo 30")
    difetto("punteggio 6", lambda p: rischio(p, "NFR-01").update(velocity=6), "> massimo 5")
    difetto("ciclo nel formato sbagliato", lambda p: p.update(ciclo="2026-Q3"), "non rispetta")
    # Un solo KRI passa lo schema dall'08/10 (D-79: un rischio in osservazione non ne ha): lo presidia C9 (sotto).
    difetto("tre KRI (due per rischio, D-64)", lambda p: rischio(p, "NFR-01")["kri_proposti"].append(
        dict(rischio(p, "NFR-01")["kri_proposti"][0], kri_id=None)), "elementi > 2")
    difetto("KRI senza tipo (D-64)", lambda p: rischio(p, "NFR-01")["kri_proposti"][0].pop("tipo"),
            "manca il campo obbligatorio 'tipo'")
    difetto("tipo di KRI fuori elenco (D-64)", lambda p: rischio(p, "NFR-01")["kri_proposti"][0].update(tipo="Controllo"),
            "fuori da")
    difetto("gap_controlli, uscito dal modello (D-49)", lambda p: rischio(p, "NFR-01").update(gap_controlli=3),
            "campo non previsto 'gap_controlli'")
    difetto("esito scritto da A1 (lo calcola lo script, D-51)", lambda p: rischio(p, "NFR-01").update(esito="Confermato"),
            "campo non previsto 'esito'")
    difetto("da/a nelle variazioni (li aggiunge lo script, D-51)",
            lambda p: rischio(p, "NFR-03")["variazioni"][0].update(da=3, a=4), "campo non previsto 'da'")
    difetto("dimensione dell'azione (D-54)",
            lambda p: rischio(p, "NFR-01")["azioni_mitigazione"][0].update(dimensione="impatto"), "campo non previsto 'dimensione'")
    difetto("kri_id malformato", lambda p: rischio(p, "NFR-01")["kri_proposti"][0].update(kri_id="K1"), "non rispetta")
    difetto("ricalcolo KRI dentro il rischio",
            lambda p: rischio(p, "NFR-01").update(ricalcolo_kri_ciclo_precedente=[]), "campo non previsto")
    difetto("categoria fuori perimetro (9.12 Frode, D-74)", lambda p: rischio(p, "NFR-01").update(event_type="9.12"), "fuori da")
    difetto("codice di Basilea in un output nuovo (D-74)", lambda p: rischio(p, "NFR-01").update(event_type="L1-02"), "fuori da")
    difetto("driver ripetuto", lambda p: rischio(p, "NFR-01").update(driver=["DRV-AI", "DRV-AI"]), "duplicati")
    difetto("manca la sintesi", lambda p: p.pop("sintesi"), "manca il campo obbligatorio 'sintesi'")

    p = copy.deepcopy(RUN1)
    rischio(p, "NFR-02")["kri_proposti"][1].pop("kri_id")
    e.verifica(not valida(p, S_A1), "validatore accetta un KRI nuovo senza kri_id (campo non obbligatorio)")
    e.verifica(rischio(RUN1, "NFR-10")["velocity"] is None and not valida(RUN1, S_A1),
               "validatore accetta la Velocity null = non valutata (D-50; NFR-10 nel run1)")
    k = rischio(RUN1, "NFR-10")["kri_proposti"][0]
    e.verifica(k["soglia_attenzione"] is None and k["soglia_allerta"] is None and not valida(RUN1, S_A1),
               "validatore accetta un KRI senza soglie: le fissa la Banca (D-53; NFR-10-K1 nel run1)")
    p = copy.deepcopy(RUN1)
    k = rischio(p, "NFR-01")["kri_proposti"][0]
    k["soglia_attenzione"], k["soglia_allerta"] = 9, 2
    e.verifica(not valida(p, S_A1), "validatore NON vede le soglie invertite (vincolo fra campi): serve C9")
    try:
        valida({}, {"type": "object", "oneOf": []})
        e.verifica(False, "un costrutto non supportato nello schema solleva errore")
    except ValueError:
        e.verifica(True, "un costrutto non supportato nello schema solleva errore")


def test_controlli_a5(e):
    esito = controlli_a5.controlla(RUN1, BASE_Q3)
    e.verifica(esito["esito"] == "Approvato", "A5 meccanico approva il ciclo pulito 2026Q3",
               "\n".join(r["descrizione"] for r in esito["rilievi"]))
    e.verifica(not valida(esito, S_A5), "l'esito dei controlli e' conforme allo schema di A5")
    esito = controlli_a5.controlla(carica(DATI / "run3_A1_2026Q4.json"), carica(DATI / "baseline_2026Q4.json"))
    e.verifica(esito["esito"] == "Approvato", "A5 meccanico approva il ciclo pulito 2026Q4",
               "\n".join(r["descrizione"] for r in esito["rilievi"]))

    def difetto(controllo, descrizione, modifica, gravita="Bloccante"):
        p = copy.deepcopy(RUN1)
        modifica(p)
        esito = controlli_a5.controlla(p, BASE_Q3)
        trovati = [r for r in esito["rilievi"] if r["controllo"] == controllo and r["gravita"] == gravita]
        atteso_esito = "Da rivedere" if gravita == "Bloccante" else esito["esito"]
        e.verifica(trovati and esito["esito"] == atteso_esito, f"A5 {controllo} ({gravita}): {descrizione}",
                   "\n".join(f"{r['controllo']} {r['gravita']} {r['descrizione']}" for r in esito["rilievi"]) or "nessun rilievo")

    def togli(p, rid):
        p["rischi"] = [r for r in p["rischi"] if r["id_rischio"] != rid]

    def varia(p, rid, **k):
        rischio(p, rid).update(**k)

    difetto("C1", "rischio in anagrafica non valutato", lambda p: togli(p, "NFR-04"))
    difetto("C1", "categoria cambiata su un rischio esistente", lambda p: varia(p, "NFR-06", event_type="9.9"))
    difetto("C1", "ciclo diverso da quello della baseline", lambda p: p.update(ciclo="2026Q4"))
    difetto("C2", "confermato con un valore cambiato", lambda p: varia(p, "NFR-04", prossimita=4))
    difetto("C2", "variazione dichiarata su un valore che non cambia",
            lambda p: rischio(p, "NFR-04")["variazioni"].append({"dimensione": "impatto", "data_evidenza": "2026-09-01",
                                                                 "evidenza": "Nuovo piano di continuita' approvato a settembre."}))
    difetto("C2", "Velocity da un valore a non valutata senza variazione", lambda p: varia(p, "NFR-01", velocity=None))
    difetto("C3", "evidenza datata prima del ciclo precedente",
            lambda p: rischio(p, "NFR-03")["variazioni"][0].update(data_evidenza="2026-06-15"), gravita=DM)
    difetto("C3", "evidenza generica",
            lambda p: rischio(p, "NFR-03")["variazioni"][0].update(evidenza="Il contesto in evoluzione e la maggiore attenzione mediatica sul tema."), gravita=DM)
    difetto("C3", "salto di due punti", lambda p: varia(p, "NFR-03", prossimita=5), gravita=DM)
    difetto("C5", "Prossimita 4 senza una fonte di rilevanza 3 (deroga da dichiarare)",
            lambda p: [f.update(rilevanza=2) for f in rischio(p, "NFR-12")["fonti"]], gravita="Osservazione")
    difetto("C5", "fonti perse con i valori invariati",
            lambda p: varia(p, "NFR-05", fonti=rischio(p, "NFR-05")["fonti"][:1]), gravita="Osservazione")
    difetto("C5", "fonte datata dopo la data di riferimento",
            lambda p: rischio(p, "NFR-04")["fonti"][0].update(data="2026-10-15"), gravita=DM)
    def inflazione(p):
        """Sette rischi su dodici salgono di un punto di Prossimita' citando solo le fonti del ciclo prima."""
        d_prec = BASE_Q3["data_riferimento_precedente"]
        scelti = [r for r in p["rischi"] if r["prossimita"] < 5 and not r["variazioni"]][:7]
        for r in scelti:
            r["prossimita"] += 1
            r["variazioni"] = [{"dimensione": "prossimita", "data_evidenza": "2026-09-01",
                                "evidenza": "Aumento dell'attenzione di settore sul tema nel trimestre, riportato dalle stesse fonti."}]
        for r in p["rischi"]:                     # nessun rischio porta fonti nuove (anche NFR-03, gia' in aumento)
            vecchie = [f for f in r["fonti"] if f["data"] <= d_prec]
            if vecchie:
                r["fonti"] = vecchie
        return len(scelti)

    difetto("C6", "piu' di meta' dei rischi in aumento senza fonti nuove (inflazione progressiva)",
            inflazione, gravita="Osservazione")
    difetto("C7", "KRI riletto due volte", lambda p: p["ricalcolo_kri_ciclo_precedente"].append(
        dict(p["ricalcolo_kri_ciclo_precedente"][0])))
    difetto("C7", "KRI riletto ma inesistente", lambda p: p["ricalcolo_kri_ciclo_precedente"].append(
        {"kri_id": "NFR-19-K1", "valore": 1, "fonte": "Sistema X", "commento": "Indicatore inesistente."}))
    difetto("C8", "owner che e' una persona (Osservazione: lo ferma la validazione, D-54)",
            lambda p: rischio(p, "NFR-01")["azioni_mitigazione"][0].update(owner="Mario Rossi"), gravita="Osservazione")
    difetto("C9", "soglie invertite (lo schema non le vede)",
            lambda p: rischio(p, "NFR-01")["kri_proposti"][0].update(soglia_attenzione=9, soglia_allerta=2), gravita=DM)
    difetto("C9", "kri_id di un altro rischio",
            lambda p: rischio(p, "NFR-01")["kri_proposti"][0].update(kri_id="NFR-02-K1"))
    difetto("C9", "kri_id valorizzato ma non in vigore",
            lambda p: rischio(p, "NFR-02")["kri_proposti"][1].update(kri_id="NFR-02-K9"))
    difetto("C9", "una soglia sola (entrambe o nessuna, D-53)",
            lambda p: rischio(p, "NFR-01")["kri_proposti"][0].update(soglia_allerta=None), gravita=DM)
    difetto("C9", "KRI in vigore sostituito da uno nuovo dello stesso tipo: il vecchio esce di vigore",
            lambda p: rischio(p, "NFR-01")["kri_proposti"][0].update(kri_id=None), gravita="Osservazione")
    difetto("C9", "un solo KRI per un rischio del radar (due per rischio, D-64; lo schema lo ammette dal D-79)",
            lambda p: rischio(p, "NFR-01")["kri_proposti"].pop(), gravita=DM)
    difetto("C9", "due KRI dello stesso tipo (uno di Esposizione e uno di Evoluzione, D-64)",
            lambda p: [k.update(tipo="Esposizione") for k in rischio(p, "NFR-01")["kri_proposti"]], gravita=DM)
    difetto("C9", "tipo cambiato su un KRI in vigore (resta quello della baseline)",
            lambda p: [k.update(tipo={"Esposizione": "Evoluzione", "Evoluzione": "Esposizione"}[k["tipo"]])
                       for k in rischio(p, "NFR-01")["kri_proposti"]], gravita="Osservazione")

    # Cose che NON devono fermare il ciclo (D-50, D-53).
    def pulito(descrizione, modifica, rid):
        p = copy.deepcopy(RUN1)
        modifica(p)
        rilievi = [r for r in controlli_a5.controlla(p, BASE_Q3)["rilievi"] if r["id_rischio"] in (rid, None)]
        e.verifica(not rilievi, f"A5 nessun rilievo: {descrizione}",
                   "\n".join(f"{r['controllo']} {r['gravita']} {r['descrizione']}" for r in rilievi))

    e.verifica(BASE_Q3["rischi"][9]["id_rischio"] == "NFR-10" and BASE_Q3["rischi"][9]["velocity"] is None,
               "baseline 2026Q3: NFR-10 con Velocity non valutata (dati d'esempio del template)")
    pulito("una Velocity prima non valutata entra con qualunque valore (qui 5), con la sua variazione",
           lambda p: varia(p, "NFR-10", velocity=5, variazioni=[{"dimensione": "velocity", "data_evidenza": "2026-09-10",
                                                                 "evidenza": "Il registro reclami mostra risposte errate che raggiungono migliaia di clienti in poche ore."}]),
           "NFR-10")
    pulito("un KRI in vigore non riletto resta Non rilevato (la rilettura riguarda solo i valori trovati)",
           lambda p: p["ricalcolo_kri_ciclo_precedente"].pop(0), "NFR-01")

    # D-71: le regole di metodo (C3, date di C5, coppia e soglie dei KRI) non fermano: vanno in validazione.
    p = copy.deepcopy(RUN1)
    rischio(p, "NFR-03")["variazioni"][0].update(evidenza="Il contesto in evoluzione e la maggiore attenzione mediatica sul tema.")
    rischio(p, "NFR-04")["fonti"][0].update(data="2026-10-15")
    rischio(p, "NFR-01")["kri_proposti"][0].update(soglia_attenzione=9, soglia_allerta=2)
    esito = controlli_a5.controlla(p, BASE_Q3)
    e.verifica(esito["esito"] == "Approvato" and sum(1 for r in esito["rilievi"] if r["gravita"] == DM) == 3,
               "D-71: C3, date di C5 e soglie di C9 sono 'Da motivare' e non fermano il ciclo", str(esito["rilievi"]))
    # A5 facoltativo (D-71): i suoi Bloccanti diventano "Da motivare" e non fermano; senza A5 vale il solo meccanico.
    agente = {"ciclo": "2026Q3", "esito": "Da rivedere", "sintesi": "Prova di arresto da parte dell'agente.",
              "rilievi": [{"controllo": "C4", "id_rischio": "NFR-01", "gravita": "Bloccante", "descrizione": "Impatto non coerente col descrittore."},
                          {"controllo": "C8", "id_rischio": "NFR-02", "gravita": "Osservazione", "descrizione": "Azione generica."}]}
    combinato = controlli_a5.combina(controlli_a5.controlla(RUN1, BASE_Q3), agente)
    e.verifica(combinato["esito"] == "Approvato" and len(combinato["da_motivare"]) == 1
               and combinato["da_motivare"][0]["origine"] == "A5" and combinato["da_motivare"][0]["controllo"] == "C4"
               and any(r["origine"] == "A5" and r["gravita"] == "Osservazione" for r in combinato["rilievi"]),
               "D-71: un Bloccante di A5 non ferma, diventa 'Da motivare'; la sua Osservazione resta tale", str(combinato))
    combinato = controlli_a5.combina(controlli_a5.controlla(RUN1, BASE_Q3))
    e.verifica(combinato["esito"] == "Approvato" and not combinato["da_motivare"],
               "D-71: senza A5 vale il solo esito dei controlli meccanizzati")
    rotto = copy.deepcopy(RUN1)
    togli(rotto, "NFR-04")
    combinato = controlli_a5.combina(controlli_a5.controlla(rotto, BASE_Q3), agente)
    e.verifica(combinato["esito"] == "Da rivedere",
               "D-71: un Bloccante dei controlli meccanizzati (integrita') ferma comunque il ciclo (D-10)")


def test_dati_esempio(e):
    """Ricostruisce ogni ciclo di esempio come se fosse un payload e lo sottopone
    ai controlli C1-C5 contro il ciclo precedente."""
    sys.path.insert(0, str(EXCEL))
    import genera_dashboard as g
    register, _, kri, _, _ = g.dati_esempio()
    import json
    per_ciclo = {}
    for r in register:
        per_ciclo.setdefault(r["Ciclo"], []).append({
            "id_rischio": r["ID_Rischio"], "impatto": r["Impatto"], "velocity": r["Velocity"],
            "prossimita": r["Prossimita"], "probabilita": r["Probabilita"],
            "variazioni": [{k: v for k, v in z.items() if k not in ("da", "a")} for z in json.loads(r["Variazioni"])],
            "fonti": json.loads(r["Fonti"]), "event_type": r["Event_Type"],
            "azioni_mitigazione": [], "kri_proposti": []})
    cicli = sorted(per_ciclo)
    for prec, corr in zip(cicli, cicli[1:]):
        base = {"ciclo_da_valutare": corr, "data_riferimento": g.DATE_CICLO[corr].isoformat(),
                "data_riferimento_precedente": g.DATE_CICLO[prec].isoformat(),
                "rischi": per_ciclo[prec], "kri_in_vigore": []}
        payload = {"ciclo": corr, "data_riferimento": g.DATE_CICLO[corr].isoformat(), "rischi": per_ciclo[corr],
                   "ricalcolo_kri_ciclo_precedente": []}
        rilievi = [r for r in controlli_a5.controlla(payload, base)["rilievi"] if r["controllo"] in ("C1", "C2", "C3", "C5")]
        # Lo script calcola esito e da/a (D-51): devono coincidere con quelli dei dati d'esempio.
        import esegui_ciclo_locale as campo
        completo = campo.completa_payload(copy.deepcopy(payload), base)
        esiti = {x["id_rischio"]: x["_esito"] for x in completo["rischi"]}
        attesi = {r["ID_Rischio"]: r["Esito"] for r in register if r["Ciclo"] == corr}
        e.verifica(esiti == attesi, f"dati di esempio {corr}: l'esito calcolato dallo script = quello scritto nel template",
                   str({k: (esiti.get(k), attesi[k]) for k in attesi if esiti.get(k) != attesi[k]}))
        e.verifica(not rilievi, f"dati di esempio {prec} -> {corr} rispettano la rubrica (C1, C2, C3, C5)",
                   "\n".join(f"{r['controllo']} {r['id_rischio']} {r['descrizione']}" for r in rilievi))
    stati = set()
    for k in kri:
        v, a, al = k["Valore_Corrente"], k["Soglia_Attenzione"], k["Soglia_Allerta"]
        e.verifica((a is None) == (al is None) and (a is None or al >= a), f"dati di esempio: {k['KRI_ID']} ha soglie coerenti")
        stati.add("Soglie da definire" if a is None else "Non rilevato" if v is None else "Rosso" if v >= al
                  else "Ambra" if v >= a else "Verde")
    e.verifica(stati == {"Verde", "Ambra", "Rosso", "Non rilevato", "Soglie da definire"},
               "dati di esempio: i KRI coprono Verde, Ambra, Rosso, Non rilevato e Soglie da definire", str(stati))


def test_rettifiche(e):
    """Una rettifica umana rispetta le stesse regole di una proposta dell'agente."""
    import esegui_ciclo_locale as campo
    cv = campo.controlla_valore
    ctx = {"precedente": {"Impatto": 3, "Velocity": None}, "rischi": {"NFR-01", "NFR-14"}}
    casi = [
        ("Rischio", "Impatto", 4, True, "rischio: un punto sopra il ciclo precedente e' ammesso"),
        ("Rischio", "Impatto", 5, False, "rischio: due punti sopra il ciclo precedente no (regola di variazione 2)"),
        ("Rischio", "Impatto", 6, False, "rischio: fuori scala 1-5 no"),
        ("Rischio", "Impatto", 3.5, False, "rischio: non intero no"),
        ("Rischio", "Velocity", "", True, "rischio: Velocity vuota = non valutata, ammessa (D-50)"),
        ("Rischio", "Velocity", 5, True, "rischio: Velocity prima non valutata entra con qualunque valore"),
        ("Rischio", "Impatto", "", False, "rischio: Impatto vuoto no (solo Velocity e anni di adeguamento sono facoltativi)"),
        ("Rischio", "Anni_Adeguamento", "", True, "rischio: anni di adeguamento vuoti ammessi (D-60)"),
        ("Rischio", "Anni_Adeguamento", "2,5", True, "rischio: anni di adeguamento decimali ammessi"),
        ("Rischio", "Anni_Adeguamento", 45, False, "rischio: anni di adeguamento oltre 30 no"),
        ("KRI", "Valore_Corrente", "", True, "KRI: valore vuoto = non rilevato, ammesso"),
        ("KRI", "Soglia_Allerta", "7,5", True, "KRI: numero con la virgola ammesso"),
        ("KRI", "Soglia_Allerta", "alta", False, "KRI: soglia non numerica no"),
        ("Azione", "Orizzonte", "2 anni", False, "azione: orizzonte fuori elenco no"),
        ("Azione", "Owner", "Mario Rossi", False, "azione: owner persona no"),
        ("Azione", "Owner", "Funzione Compliance", True, "azione: owner funzione ammesso"),
        ("Notizia", "Segnale", "Negativo", False, "notizia: segnale fuori vocabolario no"),
        ("Notizia", "Rischio_ID", "NON PRESIDIATO", True, "notizia: NON PRESIDIATO ammesso"),
        ("Notizia", "Rischio_ID", "NFR-99", False, "notizia: rischio non in anagrafica no"),
        ("Notizia", "Rischio_ID", "NFR-14", True, "notizia: rischio in osservazione ammesso (D-79)"),
        ("Rischio", "Decisione_Stato", "radar", True, "rischio: decisione sullo stato Radar (maiuscole indifferenti, D-79)"),
        ("Rischio", "Decisione_Stato", "Automatico", True, "rischio: decisione sullo stato Automatico (vale la regola)"),
        ("Rischio", "Decisione_Stato", "Forse", False, "rischio: decisione sullo stato fuori elenco no"),
    ]
    for oggetto, campo_, valore, ok, descrizione in casi:
        _, problema = cv(oggetto, campo_, valore, ctx)
        e.verifica((problema is None) == ok, f"rettifiche: {descrizione}", str(problema))

    # Variazioni dopo una rettifica: coerenti con i valori validati.
    import json
    agente = json.dumps([{"dimensione": "prossimita", "da": 4, "a": 3, "evidenza": "x" * 20, "data_evidenza": "2026-09-05"},
                         {"dimensione": "probabilita", "da": 2, "a": 3, "evidenza": "y" * 20, "data_evidenza": "2026-09-05"}])
    prec = {"Prossimita": 4, "Probabilita": 2, "Impatto": 3, "Velocity": None, "Anni_Adeguamento": None}
    v = json.loads(campo.variazioni_dopo_rettifica(agente, prec, {"Probabilita": (2, "Nessun caso nel settore.")}, "2026-10-14"))
    e.verifica([x["dimensione"] for x in v] == ["prossimita"],
               "rettifiche: riportato al valore precedente, la variazione dell'agente sparisce", str(v))
    v = json.loads(campo.variazioni_dopo_rettifica(agente, prec, {"Impatto": (4, "Perdite gia' osservate.")}, "2026-10-14"))
    nuova = [x for x in v if x["dimensione"] == "impatto"]
    e.verifica(len(v) == 3 and nuova and nuova[0]["da"] == 3 and nuova[0]["a"] == 4
               and nuova[0]["evidenza"].startswith("Rettifica del validatore") and nuova[0]["data_evidenza"] == "2026-10-14",
               "rettifiche: un valore cambiato dal validatore diventa una variazione che cita la rettifica", str(v))
    v = json.loads(campo.variazioni_dopo_rettifica("[]", None, {"Impatto": (4, "Rischio nuovo.")}, "2026-10-14"))
    e.verifica(v == [], "rettifiche: un rischio nuovo resta senza variazioni")
    v = json.loads(campo.variazioni_dopo_rettifica("[]", prec, {"Anni_Adeguamento": (2, "Migrazione pluriennale.")}, "2026-10-14"))
    e.verifica(v == [], "rettifiche: gli anni di adeguamento non sono una dimensione, nessuna variazione (D-60)", str(v))
    v = json.loads(campo.variazioni_dopo_rettifica("[]", prec, {"Velocity": (3, "Tempi documentati dai reclami.")}, "2026-10-14"))
    e.verifica(len(v) == 1 and v[0]["dimensione"] == "velocity" and v[0]["da"] is None and v[0]["a"] == 3,
               "rettifiche: una Velocity da non valutata a un valore e' una variazione 'da n.v.'", str(v))
    e.verifica(campo._uguali("7,5", 7.5) and campo._uguali("3", 3) and not campo._uguali(1, 3) and campo._uguali("Neutro", "Neutro "),
               "rettifiche: confronto fra valore scritto e proposta (numeri come testo, spazi)")
    e.verifica(campo.per_cella("=1+1") == "'=1+1" and campo.per_cella("5/6") == "'5/6" and campo.per_cella(3) == 3
               and campo.per_cella(None) is None and campo.per_cella("") == "",
               "scrittura: ogni testo va in cella come testo (niente formule o date involontarie); numeri e vuoti intatti")


def test_registro_unico(e):
    """Registro unico (D-79): la regola del radar (radar.py), i controlli dello script sui rischi in osservazione e su
    chi entra, lo schema di A1."""
    import radar
    regola = dict(radar.REGOLA_DEFAULT)

    def r(rid, imp=3, prob=3, pross=3, vel=3, **k):
        return {"id": rid, "impatto": imp, "probabilita": prob, "prossimita": pross, "velocity": vel, **k}

    # 1. Soglie: candidato con Prossimita e Probabilita almeno 2; il motivo dice quale soglia manca.
    s = radar.stati([r("A"), r("B", pross=1), r("C", prob=1, pross=1)], regola)
    e.verifica((s["A"]["stato"], s["B"]["stato"], s["C"]["stato"]) == ("Radar", "Osservazione", "Osservazione")
               and s["A"]["motivo"] == "Prossimita 3 e Probabilita 3: sopra la soglia"
               and s["B"]["motivo"] == "Prossimita 1 sotto la soglia 2"
               and s["C"]["motivo"] == "Prossimita 1 sotto la soglia 2; Probabilita 1 sotto la soglia 2",
               "radar: soglie di Prossimita e Probabilita, con il motivo", str(s))
    # 2. Eccezioni: AGIRE ORA (margine < 0), tema del supervisore, decisione Radar; decisione Osservazione esclude.
    s = radar.stati([r("A", prob=1, pross=4, y_azioni=1.0), r("B", prob=1, supervisore=True),
                     r("C", prob=1, decisione="Radar"), r("D", decisione="Osservazione")], regola)
    e.verifica(s["A"]["stato"] == "Radar" and "AGIRE ORA" in s["A"]["motivo"] and s["A"]["margine"] < 0
               and s["B"]["stato"] == "Radar" and "tema del supervisore" in s["B"]["motivo"]
               and s["C"]["stato"] == "Radar" and "decisione del validatore" in s["C"]["motivo"]
               and s["D"]["stato"] == "Osservazione" and s["D"]["motivo"] == "in osservazione per decisione del validatore",
               "radar: eccezioni (AGIRE ORA, supervisore, decisione Radar) e decisione Osservazione", str(s))
    # 3. Capienza: le eccezioni prima, poi i candidati per score; gli altri fuori per capienza.
    piccola = dict(regola, capienza=2)
    s = radar.stati([r("A", imp=5), r("B", imp=4), r("C", imp=2), r("D", imp=1, supervisore=True)], piccola)
    e.verifica([i for i in "ABCD" if s[i]["stato"] == "Radar"] == ["A", "D"]
               and s["B"]["motivo"].startswith("fuori per capienza: score ")
               and s["C"]["motivo"].startswith("fuori per capienza: score "),
               "radar: capienza 2 con un'eccezione: entrano l'eccezione e il primo per score; gli altri fuori per capienza",
               str(s))
    # 4. Uscita in due cicli: il primo ciclo sotto la regola e' "Radar, in uscita", il secondo esce; con 1 esce subito.
    s1 = radar.stati([r("A", prob=1, stato_prima="Radar")], regola)
    s2 = radar.stati([r("A", prob=1, stato_prima="Radar, in uscita")], regola)
    s3 = radar.stati([r("A", prob=1, stato_prima="Radar")], dict(regola, cicli_uscita=1))
    e.verifica(s1["A"]["stato"] == "Radar, in uscita" and "(primo ciclo: esce se lo e' anche al ciclo dopo)" in s1["A"]["motivo"]
               and s1["A"]["movimento"] == "" and s2["A"]["stato"] == "Osservazione" and s2["A"]["movimento"] == "Uscito"
               and s3["A"]["stato"] == "Osservazione" and s3["A"]["movimento"] == "Uscito",
               "radar: uscita dopo due cicli sotto la regola (Radar, in uscita, poi Uscito); con cicli_uscita 1 subito",
               f"{s1} {s2} {s3}")
    s = radar.stati([r("A", stato_prima="Osservazione"), r("B", stato_prima="Radar, in uscita")], regola)
    e.verifica(s["A"]["movimento"] == "Entrato" and s["B"]["stato"] == "Radar" and s["B"]["movimento"] == "",
               "radar: un rischio che sale entra subito (Entrato); uno in uscita che risale resta nel radar", str(s))
    # 5. Oltre il massimo delle viste nessun "in uscita": esce chi ha lo score piu' basso.
    s = radar.stati([r(f"R{i:02d}", imp=1 + i % 5, supervisore=True) for i in range(22)], regola)
    fuori = [i for i, x in s.items() if x["stato"] == "Osservazione"]
    e.verifica(sum(1 for x in s.values() if x["stato"] == "Radar") == 20 and len(fuori) == 2
               and all("oltre i 20 posti del radar" in s[i]["motivo"] for i in fuori),
               "radar: 22 eccezioni, 20 posti nelle viste: i due con lo score piu' basso restano fuori", str(fuori))
    e.verifica(radar.y_azioni(["6 mesi", "12 mesi", "3 mesi"]) == 1.0 and radar.y_azioni([]) is None
               and round(radar.score({"impatto": 3, "probabilita": 2, "prossimita": 1, "velocity": None}), 6) == 2.1875,
               "radar: Y dalle azioni (orizzonte piu' lungo, in anni) e score D-50 senza Velocity")

    # 6. Controlli dello script sui dati di prova (baseline del 2026Q4: NFR-13, 14, 15, 16 in osservazione).
    base = carica(DATI / "baseline_2026Q4.json")
    run3 = carica(DATI / "run3_A1_2026Q4.json")
    previsti = controlli_a5.stati_previsti(run3, base)
    e.verifica(sorted(x["id_rischio"] for x in base["rischi_in_osservazione"]) == ["NFR-13", "NFR-14", "NFR-15", "NFR-16"]
               and previsti["NFR-16"]["stato"] == "Radar" and previsti["NFR-16"]["movimento"] == "Entrato"
               and previsti["NFR-14"]["stato"] == "Osservazione" and previsti["NFR-15"]["stato"] == "Osservazione"
               and sum(1 for x in previsti.values() if x["stato"] == "Radar") == 13,
               "controlli: lo stato previsto con i valori di run3 (NFR-16 entra, NFR-15 non toccato resta fuori)", str(previsti))

    def rilievi(modifica, controllo, gravita, rid):
        p = copy.deepcopy(run3)
        modifica(p)
        return [x for x in controlli_a5.controlla(p, base)["rilievi"]
                if x["controllo"] == controllo and x["gravita"] == gravita and x["id_rischio"] == rid]

    O, M, B = controlli_a5.OSSERVAZIONE, controlli_a5.DA_MOTIVARE, controlli_a5.BLOCCANTE
    kri16 = copy.deepcopy(rischio(run3, "NFR-16")["kri_proposti"])
    casi = [
        ("rischio in osservazione con KRI: osservazione (lo script non li scrive)",
         lambda p: rischio(p, "NFR-14").update(kri_proposti=kri16), "C9", O, "NFR-14"),
        ("rischio in osservazione riscritto senza una fonte nuova: osservazione",
         lambda p: rischio(p, "NFR-13")["fonti"][0].update(data="2026-02-10"), "C5", O, "NFR-13"),
        ("rischio in osservazione senza orizzonte: osservazione",
         lambda p: rischio(p, "NFR-14").update(orizzonte_anni=None), "C1", O, "NFR-14"),
        ("rischio che entra nel radar senza KRI: da motivare (KRI al ciclo dopo)",
         lambda p: rischio(p, "NFR-16").update(kri_proposti=[]), "C9", M, "NFR-16"),
        ("rischio del radar senza azioni: bloccante",
         lambda p: rischio(p, "NFR-05").update(azioni_mitigazione=[]), "C9", B, "NFR-05"),
    ]
    for descrizione, modifica, controllo, gravita, rid in casi:
        e.verifica(rilievi(modifica, controllo, gravita, rid), f"controlli D-79: {descrizione}")
    e.verifica(not [x for x in controlli_a5.controlla(run3, base)["rilievi"] if x["id_rischio"] in ("NFR-13", "NFR-14")],
               "controlli D-79: i rischi in osservazione di run3 (fonte nuova, orizzonte, segnale, niente KRI) sono puliti")
    # Oltre 20 rischi in osservazione: bloccante (A1 lascia fuori i meno rilevanti).
    p = copy.deepcopy(run3)
    modello = rischio(p, "NFR-14")
    for n in range(18):
        p["rischi"].append(dict(copy.deepcopy(modello), id_rischio=f"NFR-{30 + n}", variazioni=[]))
    troppi = [x for x in controlli_a5.controlla(p, base)["rilievi"] if x["controllo"] == "C1" and x["gravita"] == B
              and "oltre il massimo di 20" in x["descrizione"]]
    p["rischi"].pop()
    e.verifica(troppi and not [x for x in controlli_a5.controlla(p, base)["rilievi"] if "oltre il massimo" in x["descrizione"]],
               "controlli D-79: 21 rischi in osservazione fermano, 20 no", str(troppi))
    # Schema: orizzonte 1-30 anni, al massimo due KRI.
    for descrizione, modifica, atteso in [
            ("orizzonte oltre 30 anni", lambda p: rischio(p, "NFR-14").update(orizzonte_anni=40), "massimo 30"),
            ("tre KRI", lambda p: rischio(p, "NFR-16")["kri_proposti"].append(copy.deepcopy(kri16[0])), "elementi > 2")]:
        p = copy.deepcopy(run3)
        modifica(p)
        errori = valida(p, S_A1)
        e.verifica(any(atteso in x for x in errori), f"validatore rifiuta: {descrizione} (D-79)", "\n".join(errori) or "nessun errore")


def test_tassonomia(e):
    """D-74: riclassificazione dei codici di Basilea della baseline e rischi fuori perimetro (FP)."""
    base = copy.deepcopy(BASE_Q3)
    for x in base["rischi"]:
        if x["id_rischio"] == "NFR-06":
            x["event_type"] = "L1-01"                 # registro di prima del 08/10/2026
    c1 = lambda p: [x for x in controlli_a5.controlla(p, base)["rilievi"] if x["controllo"] == "C1"]
    r = c1(RUN1)
    e.verifica(any(x["id_rischio"] == "NFR-06" and x["gravita"] == "Osservazione" and "Riclassificato da L1-01" in x["descrizione"]
                   for x in r) and not any(x["gravita"] == "Bloccante" for x in r),
               "tassonomia: un rischio con il codice di Basilea riclassificato in una categoria interna e' un'Osservazione, "
               "non un fermo (D-74)", str(r))
    p = copy.deepcopy(RUN1)
    rischio(p, "NFR-09")["event_type"] = "FP"
    nuovo = copy.deepcopy(rischio(p, "NFR-01"))
    nuovo.update(id_rischio="NFR-20", event_type="FP", variazioni=[])
    for k in nuovo["kri_proposti"]:
        k["kri_id"] = None
    p["rischi"].append(nuovo)
    r = c1(p)
    e.verifica(any(x["id_rischio"] == "NFR-09" and x["gravita"] == DM and "Fuori perimetro" in x["descrizione"] for x in r)
               and any(x["id_rischio"] == "NFR-20" and x["gravita"] == "Bloccante" for x in r),
               "tassonomia: un rischio esistente fuori perimetro (FP) e' da motivare (poi dismetti); un rischio nuovo FP "
               "ferma il ciclo", str(r))
    e.verifica(not valida(dict(copy.deepcopy(carica(DATI / "run3_A1_2026Q4.json")), fenomeni_non_censiti=[
        {"fenomeno": "Truffe agli sportelli automatici", "event_type": None, "fonti": [], "motivo": "fuori_perimetro",
         "spiegazione": "E' una frode: fuori dal perimetro della tassonomia interna."}]), S_A1),
               "tassonomia: un fenomeno fuori perimetro si registra fra i non censiti (motivo fuori_perimetro)")


def test_straordinario(e):
    """Ciclo straordinario AAAAQn-Sk (D-75): controlli sui soli rischi indicati, allegato di A1, proposta di A3."""
    import input_a1
    import precontrollo
    run3, run4 = carica(DATI / "run3_A1_2026Q4.json"), carica(DATI / "run4_A1_2026Q4-S1.json")
    base = copy.deepcopy(carica(DATI / "baseline_2026Q4.json"))
    valori_q4 = {x["id_rischio"]: x for x in run3["rischi"]}
    for x in base["rischi"]:
        x.update({k: valori_q4[x["id_rischio"]][k] for k in ("impatto", "probabilita", "prossimita", "velocity", "fonti")})
    base.update(ciclo_precedente="2026Q4", data_riferimento_precedente="2026-12-31", ciclo_da_valutare="2026Q4-S1",
                data_riferimento="2027-01-20",
                straordinario={"rischi_da_rivalutare": ["NFR-01"], "motivo": "Primo caso di frode con voce clonata presso la Banca.",
                               "decisore": "Comitato Rischi", "data_decisione": "2027-01-20"})
    e.verifica(not valida(run4, S_A1), "straordinario: lo schema di A1 accetta il periodo 2026Q4-S1 (D-75)")
    c1 = lambda p: [x for x in controlli_a5.controlla(p, base)["rilievi"] if x["controllo"] == "C1"]
    bloccanti = lambda p: [x for x in c1(p) if x["gravita"] == controlli_a5.BLOCCANTE]
    e.verifica(not bloccanti(run4), "straordinario: solo NFR-01, il rischio indicato, nessun fermo di completezza", str(c1(run4)))
    p = copy.deepcopy(run4)
    p["rischi"] = []
    p["rischi"].append(copy.deepcopy(valori_q4["NFR-02"]))
    e.verifica(any("assente" in x["descrizione"] and x["id_rischio"] == "NFR-01" for x in c1(p))
               and any("non indicato" in x["descrizione"] and x["id_rischio"] == "NFR-02" for x in c1(p)),
               "straordinario: manca il rischio indicato e c'e' un rischio esistente non indicato: due C1 Bloccanti", str(c1(p)))
    testo, _ = input_a1.componi(base, KERNEL, "2027-01-20")
    e.verifica("CICLO STRAORDINARIO 2026Q4-S1" in testo.split(input_a1.INIZIO_BASELINE)[0]
               and "Valuta SOLO questi rischi: NFR-01" in testo
               and "Ciclo straordinario" in testo and input_a1.nome_file(base) == "Input_A8_2026Q4-S1.txt"
               and testo.startswith("INPUT DI A8 PER IL CICLO 2026Q4-S1")
               and precontrollo.ciclo_di(run4, base) == "2026Q4-S1",
               "straordinario: l'allegato (di A8, D-84) dice in testa che cosa rivalutare e con quale motivo; nome del file e ciclo "
               "del precontrollo con il periodo straordinario")
    a3 = carica(DATI / "run_A3_2026-10-12.json")
    a3["proposta_straordinario"] = {"rischi": ["NFR-01"], "motivo": "Kit di clonazione vocale in vendita: intensita' 3 su NFR-01, Impatto 4."}
    a3_rotto = copy.deepcopy(a3)
    a3_rotto["proposta_straordinario"] = {"rischi": [], "motivo": "breve"}
    e.verifica(not valida(a3, S_A3) and valida(a3_rotto, S_A3) and not valida(dict(a3, proposta_straordinario=None), S_A3),
               "straordinario: proposta di A3 nello schema (rischi e motivo, oppure null)")
    test_a8(e, base, run3, run4, c1)


def test_a8(e, base, run3, run4, c1):
    """A8 Evento critico (D-77): analisi dell'evento e straordinario di solo censimento."""
    import input_a1
    import spiegazioni
    analisi = run4["analisi_evento"]
    rotta = copy.deepcopy(run4)
    rotta["analisi_evento"]["rischi_toccati"] = []
    del rotta["analisi_evento"]["limiti"]
    e.verifica(not valida(run4, S_A1) and valida(rotta, S_A1),
               "A8: analisi_evento nello schema (almeno un rischio toccato, limiti obbligatori)", str(valida(rotta, S_A1)))
    oss = lambda p: [(x["id_rischio"], x["descrizione"]) for x in c1(p) if x["gravita"] == controlli_a5.OSSERVAZIONE]
    e.verifica(oss(run4) == [("NFR-07", oss(run4)[0][1])] and "non indicato per lo straordinario" in oss(run4)[0][1],
               "A8: un rischio toccato dall'evento ma non indicato (NFR-07) arriva al validatore come osservazione, senza "
               "fermare", str(oss(run4)))
    senza = {k: v for k, v in run4.items() if k != "analisi_evento"}
    e.verifica([x[1][:40] for x in oss(senza)] == ["Ciclo straordinario senza analisi_evento"],
               "A8: uno straordinario fatto senza analisi (con A1) e' un'osservazione, non un fermo", str(oss(senza)))
    p = copy.deepcopy(run4)
    p["analisi_evento"]["rischi_toccati"] = [{"id_rischio": "NFR-99", "effetto": "Rischio che non esiste da nessuna parte."}]
    p["analisi_evento"]["documenti"].append({"riferimento": "Documento del futuro", "data": "2027-02-01", "uso": "prova"})
    e.verifica({x[0] for x in oss(p)} == {"NFR-99", "NFR-01", None},
               "A8: rischio toccato inesistente, rischio scritto assente dai toccati, documento dopo la data di riferimento: "
               "tre osservazioni", str(oss(p)))
    trim = copy.deepcopy(run3)
    trim["analisi_evento"] = analisi
    base_q4 = {k: v for k, v in base.items() if k != "straordinario"}
    e.verifica(any("analisi_evento" in x["descrizione"] for x in controlli_a5.controlla(trim, dict(
                   base_q4, ciclo_da_valutare=run3["ciclo"], data_riferimento=run3["data_riferimento"]))["rilievi"]),
               "A8: analisi_evento in un ciclo trimestrale: osservazione (resta solo nell'archivio)")
    # Straordinario di solo censimento: nessun rischio da rivalutare, almeno un rischio nuovo.
    censimento = dict(base, straordinario=dict(base["straordinario"], rischi_da_rivalutare=[]))
    nuovo = copy.deepcopy(next(x for x in run3["rischi"] if x["id_rischio"] == "NFR-16"))
    nuovo.update(id_rischio="NFR-17", variazioni=[])
    solo_nuovo = dict(copy.deepcopy(run4), rischi=[nuovo])
    solo_nuovo["analisi_evento"]["rischi_toccati"] = [{"id_rischio": "NFR-17", "effetto": "Rischio nuovo creato dall'evento."}]
    b = lambda p: [x for x in controlli_a5.controlla(p, censimento)["rilievi"] if x["controllo"] == "C1"
                   and x["gravita"] == controlli_a5.BLOCCANTE]
    e.verifica(not b(solo_nuovo),
               "A8: straordinario di solo censimento con un rischio nuovo: nessun fermo", str(b(solo_nuovo)))
    con_vecchio = dict(copy.deepcopy(solo_nuovo), rischi=[copy.deepcopy(run4["rischi"][0])])
    e.verifica(any("solo censimento" in x["descrizione"] for x in b(con_vecchio))
               and any("non indicato" in x["descrizione"] for x in b(con_vecchio)),
               "A8: censimento senza rischi nuovi (solo un rischio esistente): fermo", str(b(con_vecchio)))
    # KRI riletti nello straordinario: run4 rilegge i due KRI di NFR-01; senza rilettura, C7 da motivare.
    c7 = lambda p: [x for x in controlli_a5.controlla(p, base)["rilievi"] if x["controllo"] == "C7"]
    k_base = {k["kri_id"] for k in base["kri_in_vigore"] if k["id_rischio"] == "NFR-01"}
    non_riletto = dict(copy.deepcopy(run4), ricalcolo_kri_ciclo_precedente=[])
    e.verifica(not c7(run4) and {z["kri_id"] for z in run4["ricalcolo_kri_ciclo_precedente"]} == k_base
               and [(x["id_rischio"], x["gravita"]) for x in c7(non_riletto)] == [("NFR-01", controlli_a5.DA_MOTIVARE)]
               and all(k in c7(non_riletto)[0]["descrizione"] for k in k_base),
               "A8: nello straordinario i KRI in vigore dei rischi rivalutati si rileggono; se no, C7 da motivare (D-77)",
               str(c7(non_riletto)))
    testo, _ = input_a1.componi(censimento, KERNEL, "2027-01-20")
    e.verifica("censisci SOLO i rischi nuovi" in testo and "analisi_evento" in testo,
               "A8: l'allegato di uno straordinario di solo censimento lo dice in testa e chiede l'analisi dell'evento")
    righe = spiegazioni.testo_analisi_evento(analisi)
    e.verifica(righe[0].startswith("ANALISI DELL'EVENTO DI A8") and any(r.startswith("- NFR-07:") for r in righe)
               and analisi["limiti"] in righe[-1] and "non presente" in spiegazioni.testo_analisi_evento(None)[0],
               "A8: knowledge di A7 (Ciclo.txt) con l'analisi dell'evento, o la frase che manca")


def test_consistenza(e):
    diff, uguali, totali = confronta(RUN1, carica(DATI / "run2_A1_2026Q3.json"))
    e.verifica(not diff, f"consistenza fra due run sullo stesso input: {uguali}/{totali} punteggi identici", "\n".join(diff))

    # Primo ciclo da registro vuoto: lo stesso rischio con id diversi nei due run.
    primo = copy.deepcopy(RUN1)
    for r in primo["rischi"]:
        r.update(esito="Nuovo", variazioni=[])
    secondo = copy.deepcopy(primo)
    secondo["rischi"].reverse()
    for n, r in enumerate(secondo["rischi"], start=1):
        r["id_rischio"] = f"NFR-{n:02d}"
    diff, _, _ = confronta(primo, secondo)
    e.verifica(diff, "primo ciclo: per id, due run con numerazione diversa sembrano diversi (serve --primo-ciclo)")
    rinumerato, coppie = abbina(primo, secondo)
    diff, uguali, totali = confronta(primo, rinumerato)
    e.verifica(not diff and len(coppie) == len(primo["rischi"]),
               f"primo ciclo: abbinati per event type e nome, {uguali}/{totali} punteggi identici", "\n".join(diff))
    tolto = copy.deepcopy(secondo)
    tolto["rischi"].pop()
    diff, _, _ = confronta(primo, abbina(primo, tolto)[0])
    e.verifica(any("presente in un solo run" in d for d in diff),
               "primo ciclo: un rischio individuato in un solo run e' segnalato", "\n".join(diff))


def test_versione_mappa(e):
    """Lo script rifiuta una mappa delle colonne di un'altra versione (es. Kernel non aggiornato)."""
    import json
    import tempfile
    from pathlib import Path
    import esegui_ciclo_locale as campo
    with tempfile.TemporaryDirectory() as tmp:
        mappa = carica(KERNEL / "mappatura_campi.json")
        e.verifica(campo.carica_mappa({"kernel": KERNEL}) == mappa, "versione: la mappa del pacchetto e' accettata")
        for descrizione, versione in (("senza versione (precedente al 25/09)", None), ("di un'altra versione", "2026-01-01")):
            vecchia = dict(mappa)
            vecchia.pop("versione_template")
            if versione:
                vecchia["versione_template"] = versione
            (Path(tmp) / "mappatura_campi.json").write_text(json.dumps(vecchia), encoding="utf-8")
            try:
                campo.carica_mappa({"kernel": Path(tmp)})
                fermato = False
            except campo.Fermo as ex:
                fermato = "mappatura_campi.json del pacchetto" in str(ex)
            e.verifica(fermato, f"versione: una mappa {descrizione} ferma lo script")


def test_precontrollo(e):
    """D-69: azioni nuove (C8), precontrollo dell'output di A1 prima di A5, messaggio di correzione, allegato di A1."""
    import contextlib
    import io
    import json
    import re
    import tempfile
    from pathlib import Path
    import esegui_ciclo_locale as campo
    import input_a1
    from comune import RADICE
    run3, base_q4 = carica(DATI / "run3_A1_2026Q4.json"), carica(DATI / "baseline_2026Q4.json")
    righe = lambda rilievi: "\n".join(f"{r['controllo']} {r['gravita']} {r['id_rischio']} {r['descrizione']}" for r in rilievi)

    # Azioni nuove: nei cicli di prova A1 propone il passo successivo, non l'azione gia' aperta.
    for nome, p, b in (("2026Q3", RUN1, BASE_Q3), ("2026Q4", run3, base_q4)):
        doppie = [r for r in controlli_a5.controlla(p, b)["rilievi"] if "uguale all'azione aperta" in r["descrizione"]]
        e.verifica(not doppie, f"C8: nel ciclo di prova {nome} nessuna azione ripete un'azione aperta (D-69)", righe(doppie))
    p = copy.deepcopy(RUN1)
    aperta = next(z for z in BASE_Q3["azioni_aperte"] if z["id_rischio"] == "NFR-04")
    rischio(p, "NFR-04")["azioni_mitigazione"][0]["descrizione"] = "  " + aperta["descrizione"].upper().rstrip(".") + "  "
    esito = controlli_a5.controlla(p, BASE_Q3)
    doppie = [r for r in esito["rilievi"] if r["controllo"] == "C8" and aperta["id_azione"] in r["descrizione"]]
    e.verifica(len(doppie) == 1 and doppie[0]["gravita"] == "Osservazione" and esito["esito"] == "Approvato",
               "C8: un'azione uguale a un'azione aperta dello stesso rischio (maiuscole, spazi, punto a parte) e' un'Osservazione",
               righe(esito["rilievi"]))

    # Precontrollo: schema, controlli meccanizzati, parti del razionale.
    pulito = campo.precontrolla(run3, base_q4, S_A1)
    e.verifica(not pulito["schema"] and not any(r["gravita"] == "Bloccante" for r in pulito["rilievi"]) and not pulito["formato"],
               "precontrollo: il ciclo di prova 2026Q4 (razionale in parti, EFFETTO) non ha niente da correggere",
               str(pulito)[:800])
    vecchio = campo.precontrolla(RUN1, BASE_Q3, S_A1)
    e.verifica(not vecchio["schema"] and not any(r["gravita"] == "Bloccante" for r in vecchio["rilievi"])
               and {x["id_rischio"] for x in vecchio["formato"]} == {r["id_rischio"] for r in RUN1["rischi"]}
               and any("RISCHIO:, IMPATTO:" in x["descrizione"] for x in vecchio["formato"]),
               "precontrollo: un razionale senza le parti (run1, scritto prima del 02/10) e' un punto di forma, non un fermo",
               str(vecchio["formato"][:3]))
    difettoso = copy.deepcopy(RUN1)
    rischio(difettoso, "NFR-01")["kri_proposti"][0].update(kri_id="NFR-02-K1")       # integrita' (Bloccante, D-71)
    rischio(difettoso, "NFR-03")["variazioni"][0]["evidenza"] = "Il contesto in evoluzione e la maggiore attenzione mediatica sul tema."
    testo = campo.messaggio_correzione("2026Q3", campo.precontrolla(difettoso, BASE_Q3, S_A1))
    e.verifica("Da correggere (fermerebbero il ciclo):" in testo and "NFR-01 · C9: " in testo and "NFR-03 · C3: " in testo
               and "l'intero JSON" in testo and "Forma dei testi" in testo,
               "precontrollo: il messaggio per A1 elenca rischio, controllo e cosa correggere, poi i punti di forma", testo[:900])
    rotto = copy.deepcopy(run3)
    rischio(rotto, "NFR-01").update(score_composito=3.8)
    esito = campo.precontrolla(rotto, base_q4, S_A1)
    testo = campo.messaggio_correzione("2026Q4", esito)
    e.verifica(esito["schema"] and not esito["rilievi"] and "score_composito" in testo and "Schema" in testo,
               "precontrollo: un campo in piu' e' un errore di schema, detto ad A1 con il nome del campo", testo[:500])

    # Il comando, senza Excel e senza registro.
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)

        def lancia(payload, baseline):
            if baseline is None:
                (tmp / "Baseline_corrente.json").unlink(missing_ok=True)
            else:
                (tmp / "Baseline_corrente.json").write_text(json.dumps(baseline, ensure_ascii=False), encoding="utf-8")
            (tmp / "A1.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                codice = campo.main(["--workbook", str(tmp / "non_esiste.xlsx"), "--kernel", str(KERNEL),
                                     "--cartella-uscita", str(tmp), "precontrollo", "--payload", str(tmp / "A1.json")])
            return codice, out.getvalue()
        codice, out = lancia(RUN1, BASE_Q3)
        a5 = tmp / "Input_A5_2026Q3.txt"
        e.verifica(codice == 0 and "0 punti da correggere" in out and "si puo' lanciare trimestrale" in out
                   and a5.exists() and input_a1.leggi_a5(a5.read_text(encoding="utf-8")) == (RUN1, BASE_Q3)
                   and str(a5) in out,
                   "precontrollo: con soli punti di forma o da verificare il codice e' 0, e scrive l'allegato di A5 "
                   "(Input_A5_<ciclo>.txt: output di A1 e baseline, D-84)", out[:400])
        codice, out = lancia(difettoso, BASE_Q3)
        file = tmp / "Cicli" / "2026Q3" / "Correzione_A1_2026Q3.txt"
        e.verifica(codice == 1 and file.exists() and "NFR-01 · C9: " in file.read_text(encoding="utf-8")
                   and "2 punti da correggere" in out and not (tmp / "non_esiste.xlsx").exists() and not a5.exists(),
                   "precontrollo: comando senza registro; codice 1 e messaggio in Cicli/<ciclo>/Correzione_A1_<ciclo>.txt; "
                   "l'allegato di A5 di un output da correggere non resta", out[:600])
        codice, out = lancia(run3, base_q4)
        e.verifica(codice == 0 and "nessun rilievo" in out and not (tmp / "Cicli" / "2026Q4" / "Correzione_A1_2026Q4.txt").exists()
                   and f"Impronta dell'output di A1: {campo.impronta(run3)}" in out,
                   "precontrollo: output pulito, nessun messaggio; l'impronta identifica l'output controllato (D-70)", out[:400])
        (tmp / "Cicli" / "2026Q4").mkdir(parents=True, exist_ok=True)
        (tmp / "Cicli" / "2026Q4" / "Correzione_A1_2026Q4.txt").write_text("messaggio di prima", encoding="utf-8")
        codice, out = lancia(run3, base_q4)
        e.verifica(codice == 0 and not (tmp / "Cicli" / "2026Q4" / "Correzione_A1_2026Q4.txt").exists(),
                   "precontrollo: a output corretto il messaggio di un output precedente sparisce")
        codice, out = lancia(run3, None)
        e.verifica(codice == 1 and "Lanciare prima: python esegui_ciclo_locale.py baseline" in out,
                   "precontrollo: senza Baseline_corrente.json si ferma e dice di lanciare baseline", out[:300])

    # Allegato di A1: baseline, controlli (lo stesso testo del prompt di A5), guida.
    controlli = input_a1.testo_controlli(KERNEL)
    e.verifica(re.findall(r"^- (C\d) ", controlli, re.M) == [f"C{i}" for i in range(1, 10)] and ">" not in controlli[:1],
               "allegato di A1: i nove controlli C1-C9 dal kernel, senza le note", controlli[:200])
    a5 = (RADICE / "02_Agenti" / "A5_Challenger.md").read_text(encoding="utf-8")
    e.verifica(controlli in a5, "allegato di A1: i controlli sono lo stesso testo delle istruzioni di A5 (un solo file, D-69)")
    testo, mancanti = input_a1.componi(BASE_Q3, KERNEL, "2026-10-02")
    letta = json.loads(testo.split("=== 1. BASELINE ===\n", 1)[1].split("\n\n=== 2. CONTROLLI ===", 1)[0])
    e.verifica(not mancanti and letta == BASE_Q3 and controlli in testo and "=== 3. GUIDA DI SCRITTURA ===" in testo
               and "perche' non il valore vicino" in testo and "inventati" in testo
               and not any(r.startswith(">") for r in testo.splitlines()),
               "allegato di A1: baseline identica, controlli e guida con l'esempio, senza le note del kernel", testo[:300])
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        (tmp / "Input_A1_2026Q2.txt").write_text("ciclo prima", encoding="utf-8")
        (tmp / "Input_A1_appunti.txt").write_text("file di qualcuno", encoding="utf-8")
        percorso, _ = input_a1.scrivi(tmp, BASE_Q3, KERNEL, "2026-10-02")
        e.verifica(percorso.name == "Input_A1_2026Q3.txt" and not (tmp / "Input_A1_2026Q2.txt").exists()
                   and (tmp / "Input_A1_appunti.txt").exists(),
                   "allegato di A1: Input_A1_<ciclo>.txt sostituisce quello del ciclo prima e non tocca altri file")
        (tmp / "kernel_vecchio").mkdir()
        testo, mancanti = input_a1.componi(BASE_Q3, tmp / "kernel_vecchio", "2026-10-02")
        e.verifica(mancanti == ["controlli_A5.md", "guida_scrittura_A1.md"] and "=== 2." not in testo
                   and '"ciclo_da_valutare": "2026Q3"' in testo,
                   "allegato di A1: con un kernel senza i file nuovi resta la sola baseline (lo script avvisa)")
    e.verifica(input_a1.leggi_baseline(input_a1.componi(base_q4, KERNEL, "2026-10-05")[0]) == base_q4,
               "allegato di A1: leggi_baseline restituisce la baseline scritta nella parte 1 (skill di Copilot Studio, D-70)")
    try:
        input_a1.leggi_baseline("=== 2. CONTROLLI ===\n- C1 ...")
        rifiutato = False
    except ValueError as ex:
        rifiutato = "1. BASELINE" in str(ex)
    e.verifica(rifiutato, "allegato di A1: un testo senza la parte 1 non e' preso per una baseline")


def test_kit_copilot_studio(e):
    """D-70: la skill precontrollo-nfr e lo script unico del kit di Copilot Studio contengono il codice del pacchetto
    e danno lo stesso esito, la stessa impronta e lo stesso messaggio del comando precontrollo sul PC."""
    import json
    import re
    import subprocess
    import tempfile
    import zipfile
    from pathlib import Path
    import input_a1
    import precontrollo
    from comune import RADICE
    sys.path.insert(0, str(RADICE / "09_Copilot_Studio"))
    import prepara_kit as kit

    e.verifica(kit.SCRIPT.exists() and kit.contenuto_script(kit.SCRIPT) == kit.sorgenti(),
               "kit: lo script della skill contiene esattamente i moduli e lo schema del pacchetto (se no: python "
               "09_Copilot_Studio/prepara_kit.py)")
    e.verifica(kit.ALLEGATO.read_bytes() == kit.SCRIPT.read_bytes(),
               "kit: precontrollo_nfr.txt (opzione 1, code interpreter) e' lo stesso script della skill (opzione 2)")
    skill_md = (kit.SKILL / "SKILL.md").read_bytes()
    for file_zip, prefisso in ((kit.ZIP_CARTELLA, kit.NOME_SKILL + "/"), (kit.ZIP_PIATTO, "")):
        with zipfile.ZipFile(file_zip) as z:
            e.verifica(sorted(z.namelist()) == [prefisso + "SKILL.md", prefisso + "scripts/precontrollo_nfr.py"]
                       and z.read(prefisso + "SKILL.md") == skill_md
                       and z.read(prefisso + "scripts/precontrollo_nfr.py") == kit.SCRIPT.read_bytes(),
                       f"kit: {file_zip.name} contiene SKILL.md e lo script attuali", str(z.namelist()))
    testa = re.match(r"---\n(.*?)\n---\n", skill_md.decode("utf-8"), re.S)
    campi = dict(re.findall(r"^(\w+): (.+)$", testa.group(1), re.M)) if testa else {}
    e.verifica(campi.get("name") == kit.NOME_SKILL == kit.SKILL.name and re.fullmatch(r"[a-z0-9]+(-[a-z0-9]+)*", campi["name"])
               and 0 < len(campi.get("description", "")) <= 1024 and "scripts/precontrollo_nfr.py" in skill_md.decode("utf-8"),
               "kit: SKILL.md con name uguale alla cartella (minuscole e trattini) e description di al massimo 1024 caratteri",
               str(campi)[:300])

    run3, base_q4 = carica(DATI / "run3_A1_2026Q4.json"), carica(DATI / "baseline_2026Q4.json")
    difettoso = copy.deepcopy(RUN1)
    rischio(difettoso, "NFR-01")["kri_proposti"][0].update(kri_id="NFR-02-K1")       # integrita' (Bloccante, D-71)
    rischio(difettoso, "NFR-03")["variazioni"][0]["evidenza"] = "Il contesto in evoluzione e la maggiore attenzione mediatica sul tema."
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)

        def lancia(*argomenti):
            r = subprocess.run([sys.executable, str(kit.ALLEGATO), *map(str, argomenti)], capture_output=True,
                               encoding="utf-8", env={**__import__("os").environ, "PYTHONIOENCODING": "utf-8"})
            return r.returncode, r.stdout

        def scrivi(nome, testo):
            (tmp / nome).write_text(testo, encoding="utf-8")
            return tmp / nome
        in_q3 = scrivi("Input_A1_2026Q3.txt", input_a1.componi(BASE_Q3, KERNEL, "2026-10-05")[0])
        in_q4 = scrivi("Input_A1_2026Q4.txt", input_a1.componi(base_q4, KERNEL, "2026-10-05")[0])

        codice, out = lancia(scrivi("A1_2026Q4.json", json.dumps(run3, ensure_ascii=False)), in_q4)
        e.verifica(codice == 0 and out.startswith("Precontrollo dell'output di A1 (2026Q4): nessun rilievo.\n")
                   and f"Impronta dell'output di A1: {precontrollo.impronta(run3)}" in out
                   and not (tmp / "Correzione_A1_2026Q4.txt").exists(),
                   "kit: output pulito, codice 0, nessun messaggio; impronta uguale a quella del comando sul PC", out[:300])
        esito = precontrollo.precontrolla(difettoso, BASE_Q3, S_A1)
        codice, out = lancia(scrivi("A1_2026Q3.json", json.dumps(difettoso, ensure_ascii=False)), in_q3)
        corr = tmp / "Correzione_A1_2026Q3.txt"
        e.verifica(codice == 1 and out.startswith(precontrollo.riga_esito("2026Q3", esito) + "\n")
                   and corr.exists() and corr.read_text(encoding="utf-8") == precontrollo.messaggio_correzione("2026Q3", esito),
                   "kit: punti da correggere, codice 1; stessa prima riga e stesso messaggio di correzione del comando", out[:400])
        codice, out = lancia(scrivi("risposta.txt", "Ecco l'output:\n```json\n" + json.dumps(run3, indent=2) + "\n```\n"), in_q4)
        e.verifica(codice == 0 and "nessun rilievo" in out and precontrollo.impronta(run3) in out,
                   "kit: JSON con recinti e testo attorno (risposta in chat) letto come il file pulito", out[:200])
        codice, out = lancia(tmp / "A1_2026Q4.json", scrivi("Baseline_corrente.json", json.dumps(base_q4)))
        e.verifica(codice == 0 and "nessun rilievo" in out, "kit: la baseline puo' essere anche Baseline_corrente.json", out[:200])
        codice, out = lancia(scrivi("vuoto.json", "nessun json qui"), in_q4)
        codice2, out2 = lancia(tmp / "A1_2026Q4.json", scrivi("altro.txt", "=== 2. CONTROLLI ===\n- C1"))
        e.verifica(codice == 2 and codice2 == 2 and "File illeggibile" in out and "1. BASELINE" in out2,
                   "kit: output di A1 o allegato illeggibili, codice 2 e il motivo", out[:200] + out2[:200])
        codice, out = lancia("--verifica")
        e.verifica(codice == 0 and len(re.findall(r"^  \S+: [0-9a-f]{12}$", out, re.M)) == len(kit.MODULI) + 1,
                   "kit: --verifica elenca l'impronta di ogni modulo e dello schema", out[:300])


def test_copilot_studio_pieno(e):
    """D-82: istruzioni dei sotto-agenti e testi dei prompt generati dagli A*.md attuali, argomenti YAML che passano i
    controlli, busta del motore (K9) che restituisce il codice identico e la riga attesa."""
    import contextlib
    import io
    import json
    from comune import RADICE
    sys.path.insert(0, str(RADICE / "09_Copilot_Studio"))
    import genera_CS as cs
    import prepara_busta as busta

    uscita = io.StringIO()
    with contextlib.redirect_stdout(uscita):
        codice = cs.main(["genera_CS.py", "--verifica"])
    e.verifica(codice == 0, "Copilot Studio pieno: file generati aggiornati, sotto i limiti, argomenti YAML in ordine "
                            "(se no: python 09_Copilot_Studio/genera_CS.py)", uscita.getvalue()[-600:])
    for nome in cs.SOTTO_AGENTI:
        testo = cs.sotto_agente(nome)
        attese = {cs.VAR_IN: 1, cs.VAR_OUT: 1, cs.VAR_PERIODO: 1 if nome == "A3_CS.md" else 0}
        e.verifica(all(testo.count(v) == n for v, n in attese.items()) and len(testo) <= cs.LIMITE - cs.MARGINE,
                   f"Copilot Studio pieno: {nome} con un segnaposto per variabile, sotto il limite con il margine",
                   f"{len(testo)} caratteri")
    ingressi = {n: cs.ingressi(cs.prompt(n)) for n in ("A5.txt", "A6.txt", "Correggi.txt", "Motore.txt")}
    e.verifica(ingressi == {"A5.txt": ["input_a5"], "A6.txt": ["input_a6"],   # D-84: un ingresso per agente
                            "Correggi.txt": ["correzione", "json_a1", "allegato"], "Motore.txt": ["comando", "busta"]},
               "Copilot Studio pieno: ogni prompt ha in fondo i suoi ingressi", str(ingressi))
    noti = cs.segnaposto_noti()
    buono = (cs.ARGOMENTI / "Correggi_ciclo.yaml").read_text(encoding="utf-8")
    difetti = {"segnaposto": buono.replace("<<ID_FLUSSO_NFR_Motore>>", "<<ID_FLUSSO_Altro>>"),
               "globale": buono.replace("Topic.periodo", "Global.allegato", 1),
               "id": buono.replace("id: sendActivity_motore", "id: cancelAllDialogs_fine"),
               "trigger": buono.replace("kind: OnActivity", "kind: OnRecognizedIntent")}
    e.verifica(not cs.controlla_argomento(buono, noti) and all(cs.controlla_argomento(t, noti) for t in difetti.values()),
               "Copilot Studio pieno: il controllo degli argomenti trova segnaposto ignoti, globali fuori elenco, id "
               "ripetuti e trigger diversi da OnActivity",
               str({k: cs.controlla_argomento(t, noti) for k, t in difetti.items()}))

    codice_busta, ingressi_busta = busta.leggi_busta(busta.PROVA)
    e.verifica(codice_busta == busta.codice_motore() and set(ingressi_busta) == {"A1", "Input_A1"},
               "busta K9: il foglio _Motore contiene esattamente il codice del motore attuale (se no: python "
               "09_Copilot_Studio/prepara_busta.py)")
    _, codice_uscita, riga = busta.esegui_come_il_prompt(busta.PROVA)
    attesa = busta.ATTESO.read_text(encoding="utf-8").strip().splitlines()[-1]
    e.verifica(codice_uscita == 0 and json.loads(riga) == json.loads(attesa)
               and json.loads(riga)["impronta_codice"] == busta.impronta(codice_busta),
               "busta K9: il motore riunito dai blocchi stampa la riga di K9_atteso.txt, con l'impronta del suo codice",
               riga[:300])
    codice_vuota, ingressi_vuota = busta.leggi_busta(busta.VUOTA)
    e.verifica(codice_vuota == codice_busta and ingressi_vuota == {},
               "busta K9: la busta vuota ha lo stesso codice e nessun ingresso")
    e.verifica(busta.IMPRONTA.read_bytes() == busta.impronta(codice_busta).encode("ascii"),
               "busta K9: Impronta.txt contiene solo l'impronta del codice della busta, senza a capo")


MOTIVAZIONI_NFR01 = {"impatto": "Grave: importi medi triplicati nel Report trimestrale frodi.",
                     "probabilita": "Probabile: casi ripetuti in banche UE secondo l'Analisi interna.",
                     "prossimita": "Entro dodici mesi: casi documentati nel Report trimestrale frodi.",
                     "velocity": "Da uno a tre mesi: le campagne crescono in poche settimane (Analisi interna).",
                     "adeguamento": None}


def ingresso_schede():
    """Input delle schede di due rischi, come lo scrive il comando schede-input: NFR-01 con il razionale di A1
    diviso in parti (D-68), NFR-12 con un razionale scritto prima, intero."""
    rischio = lambda rid, nome, score, fonti: {
        "id_rischio": rid, "nome": nome, "event_type": "9.7", "event_type_nome": "Sicurezza informatica", "driver": "DRV-AI",
        "owner": "Funzione Antifrode", "attenzione_supervisore": rid == "NFR-01", "anni_adeguamento": None,
        "presidio_controlli": "Parziale" if rid == "NFR-01" else "Non verificato",
        "esito": "Variato" if rid == "NFR-01" else "Confermato",
        "descrizione": "Frodi a danno dei clienti tramite voci sintetiche che inducono bonifici istantanei irrevocabili.",
        "razionale_rischio": ("Le voci sintetiche convincono il cliente a pagare: il bonifico istantaneo e' irrevocabile e "
                              "la perdita ricade sulla Banca." if rid == "NFR-01" else
                              "Impatto 4 per importi medi triplicati; KRI K1 a 4,2 tentativi ogni 10.000 bonifici."),
        "motivazioni": dict(MOTIVAZIONI_NFR01) if rid == "NFR-01" else None,
        "valori": {"impatto": 4, "probabilita": 4, "prossimita": 4, "velocity": 4 if rid == "NFR-01" else None},
        "score": score, "score_ciclo_precedente": 3.5,
        "urgenza": {"esito": "AGIRE ORA: X+Y > Z", "x": 0.25, "y": 0.5, "z": 0.5, "margine": -0.25},
        "segnali_netti_4_settimane": 3,
        "variazioni": [{"dimensione": "impatto", "da": 3, "a": 4, "evidenza": "Importi medi triplicati.",
                        "data_evidenza": "2026-06-10"}] if rid == "NFR-01" else [],
        "fonti": [{"riferimento": f, "data": "2026-06-10", "tipo": "Interna", "rilevanza": 3} for f in fonti],
        "azioni": [{"id_azione": f"2026Q2-{rid}-A1", "descrizione": "Verifica fuori banda.", "owner": "Funzione Antifrode",
                    "orizzonte": "6 mesi", "stato": "Approvata", "ciclo": "2026Q2",
                    "razionale": "Interrompe la frode mentre il cliente sta per pagare.",
                    "effetto_atteso": "Meno bonifici fraudolenti eseguiti." if rid == "NFR-01" else None}],
        "kri": [{"kri_id": f"{rid}-K1", "tipo": "Evoluzione", "nome": "Tentativi di frode", "unita": "n/10k", "valore": 4.2,
                 "valore_ciclo_precedente": 3.4, "soglia_attenzione": 3, "soglia_allerta": 6, "stato": "Ambra",
                 "fonte_dati": "Sistema antifrode", "razionale": "Primo segnale di nuove campagne contro i clienti.",
                 "effetto_atteso": "Smette di crescere con la verifica fuori banda." if rid == "NFR-01" else None}],
        "storia": [{"ciclo": c, "impatto": 4, "probabilita": 4, "velocity": 4 if rid == "NFR-01" else None, "prossimita": 4,
                    "score": score} for c in ("2026Q1", "2026Q2")],
        "posizione_in_classifica": 1 if rid == "NFR-01" else 2}
    return {"ciclo": "2026Q2", "data_riferimento": "2026-06-30", "ciclo_precedente": "2026Q1",
            "pesi": {"impatto": 0.35, "probabilita": 0.25, "prossimita": 0.2, "velocity": 0.2},
            "rischi": [rischio("NFR-01", "Frodi con voce sintetica", 3.8, ["Report trimestrale frodi", "Analisi interna"]),
                       rischio("NFR-12", "Errori nei pagamenti istantanei", 3.5, ["Report operativo di maggio"])]}


def test_schede(e):
    """Agente A6 e file delle schede, senza Excel: schema, controlli, file."""
    import tempfile
    from pathlib import Path
    import esegui_ciclo_locale as campo
    import schede
    from openpyxl import load_workbook
    ingresso = ingresso_schede()
    buono = {"ciclo": "2026Q2", "schede": [
        {"id_rischio": r["id_rischio"],
         "descrizione_estesa": ("Il truffatore imita con una voce sintetica una persona nota al cliente e lo induce a "
                                "disporre un bonifico istantaneo, che e' irrevocabile. " * 4)[:900],
         "dimensioni": [{"dimensione": d, "spiegazione": f"Valore {v}: importi medi triplicati, KRI K1 a 4,2 contro una "
                                                        "soglia di attenzione di 3 (report del 10/06/2026)." if v is not None else
                                                        "Velocity non valutata: nessuna fonte ne documenta i tempi."}
                        for d, v in r["valori"].items()],
         "fonti": [{"riferimento": f["riferimento"], "a_sostegno_di": "Sostiene la impatto: importi per caso."} for f in r["fonti"]]}
        for r in ingresso["rischi"]]}
    schema = carica(KERNEL / "schema_output_A6.json")
    e.verifica(not valida(buono, schema), "schede: un output di A6 corretto e' conforme allo schema", "\n".join(valida(buono, schema)))
    rubrica = (KERNEL / "rubrica_scoring.md").read_text(encoding="utf-8")
    e.verifica(schede.controlla_a6(ingresso, buono, rubrica) == [], "schede: output corretto, nessun problema",
               "\n".join(schede.controlla_a6(ingresso, buono, rubrica)))

    def difetto(descrizione, modifica, atteso):
        p = copy.deepcopy(buono)
        modifica(p)
        problemi = schede.controlla_a6(ingresso, p, rubrica)
        e.verifica(any(atteso in x for x in problemi), f"schede: fermato {descrizione}", "\n".join(problemi) or "nessun problema")

    difetto("un ciclo diverso", lambda p: p.update(ciclo="2026Q3"), "Ciclo 2026Q3 diverso")
    difetto("un rischio senza scheda", lambda p: p["schede"].pop(), "Rischi senza scheda")
    difetto("un rischio in piu'", lambda p: p["schede"].append(dict(p["schede"][0], id_rischio="NFR-99")), "non nell'input")
    difetto("una scheda ripetuta", lambda p: p["schede"].append(copy.deepcopy(p["schede"][0])), "Schede ripetute")
    difetto("una dimensione mancante e una doppia",
            lambda p: p["schede"][0]["dimensioni"].__setitem__(3, dict(p["schede"][0]["dimensioni"][0])), "quattro dimensioni")
    difetto("una fonte riscritta", lambda p: p["schede"][0]["fonti"][0].update(riferimento="Report frodi"), "mancano")
    difetto("una fonte omessa", lambda p: p["schede"][0]["fonti"].pop(), "mancano")
    difetto("un numero inventato", lambda p: p["schede"][1]["dimensioni"][2].update(
        spiegazione="Il controllo copre solo il 37% delle transazioni, secondo una stima dell'agente."), "37")
    difetto("una data inventata", lambda p: p["schede"][0].update(
        descrizione_estesa=p["schede"][0]["descrizione_estesa"] + " Primo caso il 14/03/2025."), "14")
    e.verifica(schede.numeri("10/06/2026") == schede.numeri("2026-06-10") and schede.numeri("4,2") == {4.2},
               "schede: una data scritta in italiano o ISO e un decimale con la virgola sono lo stesso numero")
    ordinato = campo.ordina_chiavi([{"data_evidenza": "2026-06-10", "evidenza": "x", "a": 4, "da": 3, "dimensione": "impatto"}])
    e.verifica(list(ordinato[0]) == ["dimensione", "da", "a", "evidenza", "data_evidenza"],
               "scrittura: le chiavi delle variazioni hanno un ordine fisso (il foglio le rende leggibili con SUBSTITUTE)")
    # D-68: razionali di A1 in parti. Un'etichetta a meta' frase ("Impatto resta 4:") non divide; il formato vecchio resta intero.
    nuovo = ("RISCHIO: voci sintetiche e bonifici irrevocabili. IMPATTO: grave (Report frodi). Probabilità: probabile; "
             "PROSSIMITA': casi interni. VELOCITY: non valutata ADEGUAMENTO: non serve.")
    e.verifica(schede.parti_razionale(nuovo) == {"rischio": "Voci sintetiche e bonifici irrevocabili.",
                                                 "impatto": "Grave (Report frodi).", "probabilita": "Probabile",
                                                 "prossimita": "Casi interni.", "velocity": "Non valutata",
                                                 "adeguamento": "Non serve."},
               "schede: razionale in parti letto per etichetta (anche con accento, apostrofo, minuscole dopo un punto), "
               "ogni parte con la maiuscola",
               str(schede.parti_razionale(nuovo)))
    e.verifica(schede.parti_razionale("Prossimita' sale a 5: casi interni. Impatto resta 4: rimborso dal 2027.") == {},
               "schede: un razionale scritto prima (senza etichette) resta intero")
    e.verifica(schede.parti_effetto("Chiude la lacuna. EFFETTO: meno frodi.") == ("Chiude la lacuna.", "Meno frodi.")
               and schede.parti_effetto("Misura i tentativi.") == ("Misura i tentativi.", None),
               "schede: razionale di azioni e KRI diviso in perche' ed effetto atteso (EFFETTO:)")
    with tempfile.TemporaryDirectory() as tmp:
        for descrizione, uscita in (("con i testi di A6", buono), ("senza A6", None)):
            f = schede.scrivi_file_schede(ingresso, uscita, Path(tmp) / "Schede_2026Q2.xlsx")
            wb = load_workbook(f)
            testo = " ".join(str(c.value) for c in wb["NFR-01"]["A"] if c.value)
            atteso = buono["schede"][0]["descrizione_estesa"] if uscita else ingresso["rischi"][0]["razionale_rischio"]
            e.verifica(wb.sheetnames == ["Indice", "NFR-01", "NFR-12"] and atteso in testo
                       and "Andamento negli ultimi trimestri" in testo,
                       f"schede: file {descrizione}: indice, una scheda per rischio in ordine di classifica, andamento",
                       f"{wb.sheetnames} {testo[:200]}")
            # Tabella dei valori: il perche' accanto a ogni valore (colonna C, righe di Impatto ... Velocity).
            ws = wb["NFR-01"]
            righe_dim = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=3).value for i in range(1, ws.max_row + 1)
                         if ws.cell(row=i, column=1).value in ("Impatto", "Probabilita", "Prossimita", "Velocity")}
            attesi = ({d["dimensione"]: d["spiegazione"] for d in buono["schede"][0]["dimensioni"]} if uscita
                      else MOTIVAZIONI_NFR01)
            e.verifica(all(righe_dim.get(n) == attesi[d] for d, n in schede.NOMI_DIM.items()),
                       f"schede: file {descrizione}: accanto a ogni valore il suo perche' "
                       + ("(spiegazione di A6)" if uscita else "(motivazione di A1)"), str(righe_dim))
            razionale = "Razionale: il fenomeno in termini di rischio" in testo
            e.verifica(razionale != bool(uscita),
                       f"schede: file {descrizione}: " + ("niente Razionale a parte (la descrizione di A6 e' gia' il "
                       "fenomeno)" if uscita else "Razionale con il fenomeno, senza i motivi dei valori"))
            e.verifica("Perche': Interrompe la frode mentre il cliente sta per pagare." in testo
                       and "Perche' lo misuriamo: Primo segnale di nuove campagne contro i clienti." in testo
                       and " ".join(str(c.value) for c in ws["E"] if c.value).count("Effetto atteso:") == 2
                       and "Effetto atteso: Meno bonifici fraudolenti eseguiti." in [c.value for c in ws["E"]],
                       f"schede: file {descrizione}: azioni e KRI con il perche' e l'effetto atteso")
            if not uscita:
                vecchio = " ".join(str(c.value) for c in wb["NFR-12"]["A"] if c.value) + \
                          " ".join(str(c.value) for c in wb["NFR-12"]["C"] if c.value)
                e.verifica("Razionale (registro)" in vecchio and ingresso["rischi"][1]["razionale_rischio"] in vecchio
                           and "Motivo non diviso per dimensione" in vecchio
                           and "Effetto atteso: non indicato nel razionale di A1." in
                           [c.value for c in wb["NFR-12"]["E"]],
                           "schede: razionale scritto prima, intero nel Razionale (registro); tabella e azioni lo dicono",
                           vecchio[:300])


def test_spiegazioni(e):
    """D-72: fenomeni non censiti nello schema di A1 e knowledge di A7 (testo per rischio e per ciclo)."""
    import tempfile
    from pathlib import Path
    import esegui_ciclo_locale as campo
    import spiegazioni
    run3 = carica(DATI / "run3_A1_2026Q4.json")
    e.verifica(not valida(run3, S_A1) and len(run3["fenomeni_non_censiti"]) == 1 and not valida(RUN1, S_A1)
               and "fenomeni_non_censiti" not in RUN1,
               "schema A1: fenomeni_non_censiti facoltativo (run3 con una voce, run1 senza) (D-72)")
    rotto = copy.deepcopy(run3)
    rotto["fenomeni_non_censiti"][0]["motivo"] = "non mi interessa"
    e.verifica(any("motivo" in x for x in valida(rotto, S_A1)), "schema A1: un motivo fuori elenco e' un errore di schema")
    rischio = {"id_rischio": "NFR-01", "nome": "Frodi", "event_type": "9.7", "event_type_nome": "Sicurezza informatica",
               "driver": "DRV-AI", "owner": "Funzione Antifrode", "attenzione_supervisore": True, "anni_adeguamento": None,
               "esito": "Variato", "presidio_controlli": "Parziale", "descrizione": "Descrizione di prova.",
               "razionale_rischio": "Il fenomeno.", "motivazioni": {"impatto": "4: grave (Documento B).", "probabilita": None,
                                                                   "prossimita": None, "velocity": None, "adeguamento": None},
               "valori": {"impatto": 4, "probabilita": 4, "prossimita": 5, "velocity": None}, "score": 4.25,
               "score_ciclo_precedente": 3.95, "urgenza": {"esito": "AGIRE ORA: X+Y > Z", "x": 0.25, "y": 1.0, "z": 0.0,
                                                           "margine": -1.25},
               "livello_allerta": {"livello": "Attivazione", "motivo": "AGIRE ORA; KRI in rosso"},
               "segnali_netti_4_settimane": 2, "variazioni": [{"dimensione": "prossimita", "da": 4, "a": 5,
                                                              "evidenza": "Primi casi.", "data_evidenza": "2026-11-02"}],
               "fonti": [{"riferimento": "Documento B", "data": "2026-11-02", "tipo": "Esterna", "rilevanza": 3}],
               "azioni": [], "kri": [], "storia": [], "posizione_in_classifica": 1}
    dati = {"ciclo": "2026Q4", "ciclo_precedente": "2026Q3", "data_riferimento": "2026-12-31", "validato": True,
            "rischi": [rischio], "pesi": {"impatto": 0.35, "probabilita": 0.25, "prossimita": 0.2, "velocity": 0.2},
            "decisioni": {"NFR-01": [{"Oggetto": "Rischio", "ID": "NFR-01", "Campo": "Prossimita", "Decisione": "Rettificato",
                                      "Proposta_Agente": 4, "Valore_Validato": 5, "Validato_Da": "Comitato Rischi",
                                      "Motivazione": "Casi presso la Banca."}]},
            "rilievi": {"NFR-01": "Osservazione · C8 · A5: Azione generica."},
            "sintesi": {"Sintesi_A1": "MESSAGGIO: prova.", "Sintesi_A5": None, "Rilievi_Ciclo": None, "Rilievi_Superati": None},
            "parametri": [("X di Mosca, anni per decidere e avviare", 0.25)], "z_prossimita": [3.5, 2.5, 1.5, 0.5, 0],
            "scartati": run3["fenomeni_non_censiti"], "dismessi": [], "nota_quadro": None, "note_passate": [],
            "osservazione": [{"ID_Rischio": "NFR-14", "Nome": "Interruzione di cavi sottomarini", "Event_Type": "9.6",
                              "Driver": "DRV-GEO, DRV-TPR", "Impatto": 4, "Probabilita": 3, "Prossimita": 1, "Velocity": 4,
                              "Score_Composito": 3.15, "Stato": "Osservazione", "Motivo_Stato": "Prossimita 1 sotto la soglia 2",
                              "Decisione_Stato": "Automatico", "Orizzonte_Anni": 4, "Descrizione": "Taglio di cavi.",
                              "Segnale_Avvicinamento": "Un cavo che serve l'Italia danneggiato.",
                              "Razionale": "Cavo danneggiato nel Mediterraneo orientale.", "Periodo_Valutazione": "2026Q4",
                              "Fonti": '[{"riferimento": "Comunicato del gestore del cavo", "data": "2026-11-20"}]',
                              "Ciclo": "2026Q4", "Validazione": "Rettificato"}],
            "a3_osservazione": {"NFR-14": {"notizie_a3": [{"id_notizia": "N-20261201-01", "titolo": "Cavo danneggiato",
                                                          "data": "2026-12-01", "segnale": "Peggiorativo"}], "flag_a3": []}},
            "movimenti": [{"ID_Rischio": "NFR-16", "Nome": "Vertenze", "Movimento": "Entrato", "Stato": "Radar",
                           "Motivo_Stato": "Prossimita 2 e Probabilita 3: sopra la soglia"}]}
    dati["decisioni"]["NFR-14"] = [{"Campo": "Probabilita", "Decisione": "Rettificato", "Proposta_Agente": 4,
                                    "Valore_Validato": 3, "Motivazione": "Rotte italiane non ancora toccate."}]
    testo = spiegazioni.testo_rischio(rischio, dati, "2026-10-07")
    e.verifica("RISCHIO NFR-01" in testo and "VALIDATI" in testo and "Velocity non valutata" in testo
               and "Prossimita: Rettificato (proposta 4, deciso 5)" in testo and "X 0,25 + Y 1,00 contro Z 0,00" in testo
               and "Documento B (2026-11-02, Esterna, rilevanza 3)" in testo,
               "knowledge di A7: file del rischio con valori, urgenza, fonti e decisione del validatore", testo[:1500])
    ciclo = spiegazioni.testo_ciclo(dict(dati, validato=False), "2026-10-07")
    e.verifica("PROPOSTI" in ciclo and all(f["fenomeno"] in ciclo for f in run3["fenomeni_non_censiti"])
               and "RISCHI IN OSSERVAZIONE" in ciclo and "- NFR-14 Interruzione di cavi sottomarini" in ciclo
               and "Perche' in osservazione (CALCOLATO dalla regola): Prossimita 1 sotto la soglia 2" in ciclo
               and "Orizzonte stimato da A1: 4 anni" in ciclo and "notizia collegata da A3 dopo il ciclo: N-20261201-01" in ciclo
               and "decisione del validatore su Probabilita: Rettificato, proposta di A1 4" in ciclo
               and "- NFR-16 Vertenze: Entrato (Radar)" in ciclo and "5 -> 0" in ciclo,
               "knowledge di A7: Ciclo.txt con stato del ciclo, fenomeni non censiti e motivo, rischi in osservazione con il "
               "perche', le notizie di A3 e le decisioni del validatore, entrati e usciti dal radar (D-79), Z per Prossimita",
               ciclo[-1500:])
    with tempfile.TemporaryDirectory() as tmp:
        cartella = Path(tmp) / "Spiegazioni"
        cartella.mkdir()
        (cartella / "NFR-09.txt").write_text("rischio dismesso", encoding="utf-8")
        (cartella / "Appunti.txt").write_text("di qualcuno", encoding="utf-8")
        (cartella / "Rischio_05.txt").write_text("posto oltre la capienza", encoding="utf-8")
        file = spiegazioni.scrivi(cartella, dict(dati, posti=3), "2026-10-07")
        e.verifica(file == ["Ciclo.txt", "Rischio_01.txt", "Rischio_02.txt", "Rischio_03.txt"]
                   and "RISCHIO NFR-01" in (cartella / "Rischio_01.txt").read_text(encoding="utf-8")
                   and "POSTO LIBERO" in (cartella / "Rischio_03.txt").read_text(encoding="utf-8")
                   and not (cartella / "NFR-09.txt").exists() and not (cartella / "Rischio_05.txt").exists()
                   and (cartella / "Appunti.txt").exists(),
                   "knowledge di A7: Ciclo.txt e un file per posto dell'anagrafica con nomi fissi (posti liberi dichiarati), "
                   "tolti i file vecchi e i posti oltre la capienza, nessun altro file toccato")
        conf = {"kernel": KERNEL, "cartella_uscita": Path(tmp)}
        e.verifica(campo.cartella_spiegazioni(conf) == Path(tmp) / "Spiegazioni"
                   and campo.cartella_spiegazioni({"kernel": Path(tmp) / "Kernel", "cartella_uscita": Path(tmp)})
                   == Path(tmp).resolve() / "Kernel" / "Spiegazioni",
                   "knowledge di A7: nella library Kernel; con la 01_Kernel del pacchetto (collaudo) nella cartella di uscita")


def test_diagnosi(e):
    """Analisi 17 (17.10): diagnosi delle valutazioni (valori all'ingresso, verso delle rettifiche, rischi lontani) e
    istruzione di A1 sugli anni di adeguamento dei rischi in osservazione."""
    import diagnosi
    import input_a1
    rr = lambda ciclo, rid, esito, i, p, pr, v=None, stato="Radar", anni=None, orizz=None: {
        "Ciclo": ciclo, "ID_Rischio": rid, "Nome": rid, "Esito": esito, "Impatto": i, "Probabilita": p, "Prossimita": pr,
        "Velocity": v, "Stato": stato, "Anni_Adeguamento": anni, "Orizzonte_Anni": orizz}
    registro = [rr("2026Q3", "NFR-01", "Nuovo", 4, 3, 3), rr("2026Q3", "NFR-02", "Nuovo", 5, 1, 1, stato="Osservazione",
                                                                  orizz=8),
                rr("2026Q4", "NFR-01", "Variato", 4, 4, 3), rr("2026Q4", "NFR-02", "Confermato", 5, 1, 1,
                                                               stato="Osservazione", orizz=8),
                rr("2026Q4", "NFR-03", "Nuovo", 3, 2, 2), rr("2026Q4", "NFR-04", "Nuovo", 5, 1, 1, stato="Osservazione",
                                                            anni=5.0, orizz=6)]
    sv = lambda periodo, rid, campo, prop, dec, val: {"Periodo": periodo, "Oggetto": "Rischio", "ID": rid, "Campo": campo,
                                                      "Proposta_Agente": prop, "Decisione": dec, "Valore_Validato": val}
    storico = [sv("2026Q3", "NFR-01", "Probabilita", 3, "Rettificato", 4), sv("2026Q4", "NFR-03", "Probabilita", 2,
                                                                                 "Rettificato", 3),
               sv("2026Q4", "NFR-01", "Impatto", 4, "Rettificato (riapertura)", 3),
               sv("2026Q4", "NFR-01", "Probabilita", 4, "Accettato", 4),
               sv("2026Q4", "NFR-01", "Velocity", None, "Rettificato", 3),
               {"Periodo": "2026Q4", "Oggetto": "KRI", "ID": "NFR-01-K1", "Campo": "Valore_Corrente",
                "Proposta_Agente": 1, "Decisione": "Rettificato", "Valore_Validato": 9}]
    d = diagnosi.diagnosi(registro, storico, "2026Q4")
    e.verifica(d["rischi_nuovi"] == 4 and d["ingresso"]["probabilita"][3] == 1 and d["ingresso"]["probabilita"][1] == 2
               and d["ingresso"]["velocity"][diagnosi.NV] == 4,
               "diagnosi: valori all'ingresso contati sui soli rischi nuovi, Velocity vuota = n.v.", d["ingresso"])
    r = d["rettifiche"]
    e.verifica(r["probabilita"]["nuovi"]["su"] == 2 and r["impatto"]["riapertura"]["giu"] == 1
               and not r["velocity"]["nuovi"] and not r["velocity"]["altri"] and d["accettati"]["probabilita"] == 1,
               "diagnosi: verso delle rettifiche per rischi nuovi e riapertura; da n.v. nessun verso; KRI esclusi", r)
    lont = {x["id"]: x for x in d["lontani"]}
    e.verifica(set(lont) == {"NFR-02", "NFR-04"} and lont["NFR-02"]["margine"] == 2.25 and lont["NFR-04"]["margine"] == -1.75,
               "diagnosi: rischi lontani con il margine di Mosca (Y di riserva senza anni, AGIRE ORA con 5 anni)", lont)
    e.verifica(any("NFR-02" in x and "anni di adeguamento" in x for x in d["letture"])
               and not any("NFR-04" in x for x in d["letture"]) and "nessuno con 4 o 5" in " ".join(d["letture"]),
               "diagnosi: segnala il rischio lontano con Impatto alto senza anni di adeguamento, non quello che li ha",
               d["letture"])
    e.verifica("3. Rischi lontani" in diagnosi.testo(d) and "NFR-04" in diagnosi.testo(d), "diagnosi: testo del rapporto")
    guida = (KERNEL / "guida_scrittura_A1.md").read_text(encoding="utf-8")
    allegato, _ = input_a1.componi(BASE_Q3, KERNEL, "2026-10-19")
    e.verifica("Rischi in osservazione con Impatto 4 o 5: stima anche anni_adeguamento" in guida
               and "Rischi in osservazione con Impatto 4 o 5" in allegato,
               "guida di scrittura di A1: anni di adeguamento dei rischi in osservazione con Impatto 4-5, nell'allegato (17.10)")


def test_esposizione(e):
    """D-81 (analisi 17, 17.10): l'esposizione della Banca non nota alle fonti va in validazione con i due valori."""
    import diagnosi
    import input_a1
    esp = lambda rilievi: [x for x in rilievi if x["controllo"] == "C5" and ("D-81" in x["descrizione"]
                                                                          or "ESPOSIZIONE NON NOTA" in x["descrizione"])]
    e.verifica(not esp(controlli_a5.controlla(RUN1, BASE_Q3)["rilievi"]), "esposizione: nessuna segnalazione senza la frase (run1)")
    run = copy.deepcopy(RUN1)
    x = run["rischi"][0]
    p, prec = x["probabilita"], next(b for b in BASE_Q3["rischi"] if b["id_rischio"] == x["id_rischio"])
    n = min(5, p + 2)
    x["razionale"] += f" ESPOSIZIONE NON NOTA: con l'esposizione documentata sarebbe {n}."
    trovati = esp(controlli_a5.controlla(run, BASE_Q3)["rilievi"])
    e.verifica(len(trovati) == 1 and trovati[0]["gravita"] == controlli_a5.DA_MOTIVARE and trovati[0]["id_rischio"] == x["id_rischio"]
               and f"Probabilita {p} proposta, {n} con l'esposizione documentata" in trovati[0]["descrizione"]
               and f"al massimo {min(5, prec['probabilita'] + 1)}" in trovati[0]["descrizione"],
               "esposizione: la frase porta i due valori in validazione (Da motivare), con il limite di un punto per un "
               "rischio gia' nel registro", trovati)
    conta_b = lambda rr: sum(1 for r in rr if r["gravita"] == controlli_a5.BLOCCANTE)
    e.verifica(conta_b(controlli_a5.controlla(run, BASE_Q3)["rilievi"]) == conta_b(controlli_a5.controlla(RUN1, BASE_Q3)["rilievi"]),
               "esposizione: la segnalazione non ferma il ciclo (nessun Bloccante in piu')")
    x["razionale"] = x["razionale"].replace(f"sarebbe {n}.", f"sarebbe {p}.")
    trovati = esp(controlli_a5.controlla(run, BASE_Q3)["rilievi"])
    e.verifica(len(trovati) == 1 and trovati[0]["gravita"] == controlli_a5.OSSERVAZIONE,
               "esposizione: un valore non maggiore della Probabilita proposta e' solo un'Osservazione sulla forma", trovati)
    righe = [{"Ciclo": "2026Q4", "ID_Rischio": "NFR-01", "Esito": "Nuovo", "Probabilita": 4, "Prossimita": 3, "Stato": "Radar",
              "Razionale": "PROBABILITA: 3. ESPOSIZIONE NON NOTA: con l'esposizione documentata sarebbe 4."},
             {"Ciclo": "2026Q4", "ID_Rischio": "NFR-02", "Esito": "Confermato", "Probabilita": 3, "Prossimita": 3,
              "Stato": "Radar", "Razionale": "PROBABILITA: 3."}]
    d = diagnosi.diagnosi(righe, [], "2026Q4")
    e.verifica(d["esposizioni"] == [{"id": "NFR-01", "probabilita": 4, "se_esposta": 4, "nuovo": True}]
               and any("Esposizione non nota in 1 rischi su 2" in t and "in 1." in t for t in d["letture"])
               and "4. Esposizione non nota" in diagnosi.testo(d),
               "diagnosi: conta le note di esposizione e quante volte il validatore ha dato il valore se esposta", d["letture"])
    allegato, _ = input_a1.componi(BASE_Q3, KERNEL, "2026-10-19")
    e.verifica("ESPOSIZIONE NON NOTA: con l'esposizione documentata sarebbe N." in allegato,
               "guida di scrittura di A1 nell'allegato: la frase dell'esposizione non nota (D-81)")


def main():
    e = Esito("Validatore, controlli A5, dati di esempio, consistenza")
    test_validatore(e)
    test_controlli_a5(e)
    test_dati_esempio(e)
    test_rettifiche(e)
    test_registro_unico(e)
    test_tassonomia(e)
    test_straordinario(e)
    test_consistenza(e)
    test_versione_mappa(e)
    test_precontrollo(e)
    test_kit_copilot_studio(e)
    test_copilot_studio_pieno(e)
    test_schede(e)
    test_spiegazioni(e)
    test_diagnosi(e)
    test_esposizione(e)
    return 0 if e.stampa() else 1


if __name__ == "__main__":
    sys.exit(main())
