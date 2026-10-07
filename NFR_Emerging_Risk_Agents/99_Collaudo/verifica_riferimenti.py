"""Riferimenti incrociati fra i file del pacchetto, verificati dal codice.

Il punto debole storico di questo pacchetto e' la coerenza fra file: un numero
di passo citato che non corrisponde, una cella spostata, un codice di
tassonomia cambiato in un posto solo. Qui ogni riferimento che puo' divergere
e' confrontato con la sua fonte:

1. prompt composti aggiornati rispetto al kernel e sotto il limite di caratteri
2. codici di tassonomia e driver identici in kernel, schemi, generatore, mappa
3. controlli C1-C9 identici in schema A5, kernel (testo unico, D-69), prompt
   A5, documentazione, governance e codice (C4 assente dal codice, per disegno)
4. numeri di passo e di sezione dei _DOCUMENTATO = quelli dei modelli
5. campi derivati documentati = quelli implementati
6. nomi definiti e formule citati nel manuale = quelli del template
7. file elencati in 00_LEGGIMI = file presenti, in entrambe le direzioni
8. comandi e opzioni citati nella guida e nei README = quelli degli script
9. versione della struttura del workbook uguale in generatore, script, mappa, template
"""

import re
import sys

from comune import CAMPO, EXCEL, KERNEL, RADICE, TEMPLATE, Esito, carica

AGENTI = RADICE / "02_Agenti"


def testo(percorso):
    return percorso.read_text(encoding="utf-8")


def verifica_prompt(e):
    sys.path.insert(0, str(AGENTI))
    import componi_prompt
    import contextlib
    import io
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        codice = componi_prompt.main(["", "--verifica"])
    e.verifica(codice == 0, "1. prompt composti aggiornati e sotto gli 8.000 caratteri", buf.getvalue())


def verifica_tassonomia(e):
    t = testo(KERNEL / "tassonomia_nfr.md")
    l1 = sorted(set(re.findall(r"^\| (\d+\.\d+) \|", t, re.M)))          # D-74: categorie della tassonomia interna
    drv = sorted(set(re.findall(r"^\| (DRV-[A-Z]+) \|", t, re.M)))
    s1 = carica(KERNEL / "schema_output_A1.json")["$defs"]["rischio"]["properties"]
    s3 = carica(KERNEL / "schema_output_A3.json")["properties"]["notizie"]["items"]["properties"]["impatti"]["items"]["properties"]
    sys.path.insert(0, str(EXCEL))
    import genera_dashboard as g
    mappa = carica(KERNEL / "mappatura_campi.json")
    imp = sorted(c["nome"].replace("Imp_", "").replace("_", ".") for c in mappa["tabelle"]["tblNews"]["colonne"] if c["nome"].startswith("Imp_"))
    import controlli_a5
    fonti = {
        "schema A1 event_type (senza FP)": sorted(c for c in s1["event_type"]["enum"] if c != "FP"),
        "controlli_a5 CATEGORIE": sorted(controlli_a5.CATEGORIE),
        "schema A3 impatti": sorted(s3["event_type"]["enum"]),
        "genera_dashboard TASSONOMIA": sorted(c for c, _ in g.TASSONOMIA),
        "mappatura Imp_": imp,
    }
    for nome, codici in fonti.items():
        e.verifica(codici == l1 and len(l1) == 9, f"2. categorie: {nome} = tassonomia_nfr.md (nove, D-74)", f"{codici} contro {l1}")
    for nome, codici in {"schema A1 driver": sorted(s1["driver"]["items"]["enum"]),
                         "genera_dashboard DRIVER": sorted(c for c, _ in g.DRIVER)}.items():
        e.verifica(codici == drv, f"2. driver: {nome} = tassonomia_nfr.md", f"{codici} contro {drv}")


def verifica_controlli(e):
    attesi = [f"C{i}" for i in range(1, 10)]
    fonti = {
        "schema A5": carica(KERNEL / "schema_output_A5.json")["properties"]["rilievi"]["items"]["properties"]["controllo"]["enum"],
        "kernel (controlli_A5.md, testo unico, D-69)": re.findall(r"^- (C\d) ", testo(KERNEL / "controlli_A5.md"), re.M),
        "prompt A5 composto": re.findall(r"^- (C\d) ", testo(AGENTI / "A5_Challenger.md"), re.M),
        "A3_A4_A5_DOCUMENTATO": re.findall(r"^### (C\d) ", testo(AGENTI / "COMMENTATI" / "A3_A4_A5_DOCUMENTATO.md"), re.M),
        "governance": re.findall(r"^\| (C\d) ", testo(RADICE / "06_Governance" / "controlli_e_model_risk.md"), re.M),
    }
    for nome, trovati in fonti.items():
        e.verifica(trovati == attesi, f"3. controlli C1-C9 in {nome}", str(trovati))
    modello = testo(AGENTI / "modelli" / "A5.md")
    e.verifica("{{CONTROLLI}}" in modello and not re.findall(r"^- C\d ", modello, re.M),
               "3. il modello di A5 prende i controlli dal kernel ({{CONTROLLI}}), senza una copia")
    codice = sorted(set(re.findall(r'_rilievo\("(C\d)"', testo(CAMPO / "controlli_a5.py"))))
    e.verifica(codice == [c for c in attesi if c != "C4"], "3. controlli_a5.py implementa tutti i controlli tranne C4", str(codice))


def passi_modello(nome, sezione):
    corpo = testo(AGENTI / "modelli" / nome).split(f"# {sezione}", 1)[1].split("\n# ", 1)[0]
    return [int(n) for n in re.findall(r"^(\d+)\. ", corpo, re.M)]


def verifica_passi(e):
    doc_a1 = testo(AGENTI / "COMMENTATI" / "A1_Horizon_Scanner_DOCUMENTATO.md")
    doc_345 = testo(AGENTI / "COMMENTATI" / "A3_A4_A5_DOCUMENTATO.md")
    a3_doc = doc_345.split("# A4 ", 1)[0]
    a4_doc = doc_345.split("# A4 ", 1)[1].split("# A5 ", 1)[0]
    casi = [
        ("A1", passi_modello("A1.md", "Procedura"), [int(n) for n in re.findall(r"^### Passo (\d+) ", doc_a1, re.M)]),
        ("A3", passi_modello("A3.md", "Procedura"), [int(n) for n in re.findall(r"^### Passo (\d+) ", a3_doc, re.M)]),
        ("A4", passi_modello("A4.md", "Struttura della risposta (sempre queste sezioni, in quest'ordine)"),
         [int(n) for n in re.findall(r"^### Sezione (\d+) ", a4_doc, re.M)]),
    ]
    for agente, nel_modello, nel_doc in casi:
        e.verifica(nel_modello == list(range(1, len(nel_modello) + 1)) and nel_doc == nel_modello,
                   f"4. {agente}: passi documentati = passi del modello ({len(nel_modello)})", f"modello {nel_modello}, doc {nel_doc}")
    e.verifica("Passo 10" in doc_a1 and max(passi_modello("A1.md", "Procedura")) == 10,
               "4. A1: l'autocontrollo e' il Passo 10 in modello e documentazione")


def verifica_derivati(e):
    sys.path.insert(0, str(CAMPO))
    import esegui_ciclo_locale as campo
    doc = testo(RADICE / "05_Workflow" / "workflow.md").split("## Campi derivati", 1)[1].split("\n## ", 1)[0]
    documentati = sorted(re.findall(r"^\| `(\w+)` \|", doc, re.M))
    e.verifica(documentati == sorted(campo.DERIVATI), "5. campi derivati documentati = implementati",
               f"doc {documentati}\ncodice {sorted(campo.DERIVATI)}")


def verifica_manuale(e):
    from openpyxl import load_workbook
    wb = load_workbook(TEMPLATE)
    nomi = set(wb.defined_names)
    formule = {str(c.value).replace("_xlfn.", "") for ws in wb.worksheets for riga in ws.iter_rows()
               for c in riga if isinstance(c.value, str) and c.value.startswith("=")}
    manuale = testo(EXCEL / "costruzione_manuale_template.md")
    citati = re.findall(r"^(=.+)$", "\n".join(re.findall(r"```excel\n(.*?)```", manuale, re.S)), re.M)
    mancanti = [f for f in citati if f not in formule]
    e.verifica(citati and not mancanti, f"6. le {len(citati)} formule del manuale esistono identiche nel template", "\n".join(mancanti))
    nomi_citati = set(re.findall(r"`((?:RR|AZ|KRI|NW|ANA|SI|SE|VAL|VG)_\w+|peso\w+|soglia\w+|ciclo\w+|cicli\w+|settiman\w+|inizio\w+|codiciL1|anagrafica"
                                 r"|classifica\w+|lente[XYZ]\w*|pos[A-Z]\w*|rm[A-Z]\w*|delta[A-Z]\w*|quadro\w+|presidio\w+|copertura\w+|stato[A-Z]\w*"
                                 r"|decisioniDaPrendere|etichettaCiclo|capienza\w+|elencoKRI|versioneTemplate|mosca\w+|anniProssimita|valStato"
                                 r"|livello\w+)`", manuale))
    e.verifica(nomi_citati and nomi_citati <= nomi, f"6. i {len(nomi_citati)} nomi definiti citati nel manuale esistono", str(sorted(nomi_citati - nomi)))
    ana, val = (wb.defined_names[n].attr_text for n in ("coperturaAnagrafica", "coperturaValutati"))
    celle = [wb[f.split("!")[0]][f.split("!")[1].replace("$", "")].value for f in (ana, val)]
    e.verifica(str(celle[0]).startswith('=COUNTIF(ANA_ID,"NFR-*")') and str(celle[1]).startswith("=SUMPRODUCT"),
               f"6. il controllo di copertura sta davvero nei nomi coperturaAnagrafica ({ana}) e coperturaValutati ({val})")


def verifica_leggimi(e):
    """L'indice di 00_LEGGIMI ha una riga per file: percorso completo, poi la descrizione."""
    albero = testo(RADICE / "00_LEGGIMI.md").split("```text", 1)[1].split("```", 1)[0]
    elencati = {re.split(r"\s{2,}", riga.strip())[0] for riga in albero.splitlines() if riga.strip()}   # nomi anche con spazi
    presenti = {p.relative_to(RADICE).as_posix() for p in RADICE.rglob("*")
                if p.is_file() and "__pycache__" not in p.parts and p.name != "configurazione.json"}
    e.verifica(elencati == presenti, "7. 00_LEGGIMI elenca esattamente i file presenti",
               f"elencati ma assenti: {sorted(elencati - presenti)}\n"
               f"presenti non elencati: {sorted(presenti - elencati)}")


def verifica_guida(e):
    """Ogni comando e ogni opzione citati nella guida e nei README esistono davvero
    nello script: una guida che cita un'opzione inesistente lascia l'utente bloccato."""
    sys.path.insert(0, str(CAMPO))
    import argparse
    import esegui_ciclo_locale as campo
    import raccogli_notizie
    import confronta
    comandi, opzioni = set(), set()
    parser_ciclo = None
    orig = argparse.ArgumentParser.parse_args

    def cattura(self, *a, **k):
        nonlocal parser_ciclo
        parser_ciclo = self
        raise SystemExit(0)
    argparse.ArgumentParser.parse_args = cattura
    try:
        for modulo in (campo, raccogli_notizie, confronta):
            try:
                modulo.main(["confronta.py"] if modulo is confronta else [])
            except SystemExit:
                pass
            for azione in parser_ciclo._actions:
                opzioni.update(o for o in azione.option_strings)
                if isinstance(azione, argparse._SubParsersAction):
                    for nome, sub in azione.choices.items():
                        comandi.add(nome)
                        opzioni.update(o for a in sub._actions for o in a.option_strings)
    finally:
        argparse.ArgumentParser.parse_args = orig
    for doc in (RADICE / "00_GUIDA_PASSO_PASSO.md", CAMPO / "README_variante_locale.md", RADICE / "05_Workflow" / "workflow.md"):
        t = testo(doc)
        # Il nome del comando, anche dopo le opzioni comuni (--workbook "..." --cartella-uscita ...); anche aggiorna-registro.
        citati = set(re.findall(r'esegui_ciclo_locale\.py (?:--[\w-]+ (?:"[^"]*"|\S+) )*([\w-]+)', t))
        e.verifica(citati and citati <= comandi, f"8. comandi citati in {doc.name} esistono", str(sorted(citati - comandi)))
        opz = set(re.findall(r"(?<![\w-])(--[a-z][a-z-]+)", t)) - {"--version", "--upgrade", "--force-reinstall"}   # opzioni di python/pip, non degli script
        e.verifica(opz <= opzioni, f"8. opzioni citate in {doc.name} esistono", str(sorted(opz - opzioni)))


def verifica_versione(e):
    """La versione della struttura del workbook e' scritta in quattro posti: devono coincidere."""
    from openpyxl import load_workbook
    sys.path.insert(0, str(EXCEL))
    sys.path.insert(0, str(CAMPO))
    import genera_dashboard as g
    import esegui_ciclo_locale as campo
    wb = load_workbook(TEMPLATE)
    versioni = {"genera_dashboard.py": g.VERSIONE_TEMPLATE, "esegui_ciclo_locale.py": campo.VERSIONE_TEMPLATE,
                "mappatura_campi.json": carica(KERNEL / "mappatura_campi.json").get("versione_template"),
                "template (versioneTemplate)": str(wb.defined_names["versioneTemplate"].attr_text).strip('"')}
    e.verifica(len(set(versioni.values())) == 1, "9. versione della struttura uguale in generatore, script, mappa e template",
               str(versioni))


def main():
    e = Esito("Riferimenti incrociati")
    for verifica in (verifica_prompt, verifica_tassonomia, verifica_controlli, verifica_passi,
                     verifica_derivati, verifica_manuale, verifica_leggimi, verifica_guida, verifica_versione):
        try:
            verifica(e)
        except Exception as ex:
            e.verifica(False, f"{verifica.__name__} eseguita", f"{type(ex).__name__}: {ex}")
    return 0 if e.stampa() else 1


if __name__ == "__main__":
    sys.exit(main())
