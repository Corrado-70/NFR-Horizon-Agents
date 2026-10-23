"""Precontrollo NFR dell'output di A1 per Copilot Studio (skill precontrollo-nfr o code interpreter).

GENERATO da 09_Copilot_Studio/prepara_kit.py: non modificarlo a mano (il collaudo lo confronta con il pacchetto).
Contiene, copiati senza modifiche, i moduli valida_output.py, radar.py, controlli_a5.py, schede.py, input_a1.py, precontrollo.py di 07_Senza_Studio e lo schema
01_Kernel/schema_output_A1.json: e' lo stesso codice del comando "python esegui_ciclo_locale.py precontrollo"
che la Banca lancia sul PC (D-69, D-70).

Uso:
    python precontrollo_nfr.py <output di A1, A1_<ciclo>.json> <Input_A1_<ciclo>.txt oppure Baseline_corrente.json>
    python precontrollo_nfr.py --verifica      versione e impronta di ogni modulo contenuto
Stampa l'esito (prima riga) e l'impronta dell'output di A1 (seconda riga, la stessa del comando sul PC); se trova
qualcosa scrive accanto all'output Correzione_A1_<ciclo>.txt, il messaggio di correzione per A1. Codice di uscita:
0 nessun punto da correggere, 1 punti da correggere, 2 file illeggibili. Solo libreria standard: niente rete,
niente Excel, non scrive nel registro.
"""

import hashlib
import json
import sys
import types
from pathlib import Path

VERSIONE = "13c0ee35bae8"

SCHEMA_A1 = r'''{
  "$schema": "https://json-schema.org/draft/2020-12/schema",
  "$id": "schema_output_A1",
  "title": "Output di A1 Horizon Scanner — contratto A1 -> A5 -> scrittura",
  "type": "object",
  "additionalProperties": false,
  "required": ["ciclo", "data_riferimento", "rischi", "ricalcolo_kri_ciclo_precedente", "sintesi"],
  "properties": {
    "ciclo": {"type": "string", "pattern": "^[0-9]{4}Q[1-4](-S[1-9])?$"},
    "data_riferimento": {"$ref": "#/$defs/data"},
    "rischi": {
      "type": "array",
      "minItems": 1,
      "items": {"$ref": "#/$defs/rischio"}
    },
    "ricalcolo_kri_ciclo_precedente": {
      "type": "array",
      "items": {"$ref": "#/$defs/ricalcolo_kri"}
    },
    "sintesi": {"type": "string", "minLength": 100, "maxLength": 3000},
    "fenomeni_non_censiti": {
      "type": "array",
      "maxItems": 15,
      "items": {"$ref": "#/$defs/fenomeno_non_censito"}
    },
    "analisi_evento": {"$ref": "#/$defs/analisi_evento"}
  },
  "$defs": {
    "data": {"type": "string", "pattern": "^[0-9]{4}-[0-9]{2}-[0-9]{2}$"},
    "analisi_evento": {
      "type": "object",
      "additionalProperties": false,
      "required": ["evento", "catena_causale", "rischi_toccati", "rischi_non_toccati", "documenti", "limiti"],
      "properties": {
        "evento": {"type": "string", "minLength": 50, "maxLength": 800},
        "catena_causale": {"type": "string", "minLength": 50, "maxLength": 1500},
        "rischi_toccati": {
          "type": "array",
          "minItems": 1,
          "maxItems": 10,
          "items": {
            "type": "object",
            "additionalProperties": false,
            "required": ["id_rischio", "effetto"],
            "properties": {
              "id_rischio": {"type": "string", "pattern": "^NFR-[0-9]{2}$"},
              "effetto": {"type": "string", "minLength": 10, "maxLength": 400}
            }
          }
        },
        "rischi_non_toccati": {"type": "string", "maxLength": 800},
        "documenti": {
          "type": "array",
          "minItems": 1,
          "maxItems": 15,
          "items": {
            "type": "object",
            "additionalProperties": false,
            "required": ["riferimento", "data", "uso"],
            "properties": {
              "riferimento": {"type": "string", "minLength": 3, "maxLength": 300},
              "data": {"$ref": "#/$defs/data"},
              "uso": {"type": "string", "minLength": 5, "maxLength": 300}
            }
          }
        },
        "limiti": {"type": "string", "minLength": 10, "maxLength": 600}
      }
    },
    "fenomeno_non_censito": {
      "type": "object",
      "additionalProperties": false,
      "required": ["fenomeno", "event_type", "fonti", "motivo", "spiegazione"],
      "properties": {
        "fenomeno": {"type": "string", "minLength": 5, "maxLength": 200},
        "event_type": {"type": ["string", "null"], "enum": ["9.1", "9.4", "9.5", "9.6", "9.7", "9.8", "9.9", "9.11", "9.13", null]},
        "fonti": {
          "type": "array",
          "maxItems": 3,
          "items": {
            "type": "object",
            "additionalProperties": false,
            "required": ["riferimento", "data"],
            "properties": {
              "riferimento": {"type": "string", "minLength": 3, "maxLength": 300},
              "data": {"$ref": "#/$defs/data"}
            }
          }
        },
        "motivo": {"enum": ["nessun_meccanismo", "fonte_non_specifica", "gia_coperto", "troppo_lontano", "fuori_perimetro", "altro"]},
        "spiegazione": {"type": "string", "minLength": 10, "maxLength": 300}
      }
    },
    "punteggio": {"type": "integer", "minimum": 1, "maximum": 5},
    "soglia": {"type": ["number", "null"]},
    "rischio": {
      "type": "object",
      "additionalProperties": false,
      "required": ["id_rischio", "nome", "event_type", "driver", "descrizione",
                   "impatto", "velocity", "prossimita", "probabilita", "anni_adeguamento",
                   "variazioni", "fonti", "razionale",
                   "azioni_mitigazione", "kri_proposti"],
      "properties": {
        "id_rischio": {"type": "string", "pattern": "^NFR-[0-9]{2}$"},
        "nome": {"type": "string", "minLength": 5, "maxLength": 120},
        "event_type": {"type": "string", "enum": ["9.1", "9.4", "9.5", "9.6", "9.7", "9.8", "9.9", "9.11", "9.13", "FP"]},
        "driver": {
          "type": "array", "minItems": 1, "uniqueItems": true,
          "items": {"type": "string", "enum": ["DRV-CYB", "DRV-TPR", "DRV-AI", "DRV-CLI", "DRV-CON", "DRV-GEO", "DRV-TEC", "DRV-HUM"]}
        },
        "descrizione": {"type": "string", "minLength": 20, "maxLength": 1000},
        "impatto": {"$ref": "#/$defs/punteggio"},
        "velocity": {"type": ["integer", "null"], "minimum": 1, "maximum": 5},
        "prossimita": {"$ref": "#/$defs/punteggio"},
        "probabilita": {"$ref": "#/$defs/punteggio"},
        "anni_adeguamento": {"type": ["number", "null"], "minimum": 0, "maximum": 30},
        "orizzonte_anni": {"type": ["integer", "null"], "minimum": 1, "maximum": 30},
        "segnale_avvicinamento": {"type": ["string", "null"], "minLength": 10, "maxLength": 300},
        "variazioni": {
          "type": "array",
          "items": {
            "type": "object",
            "additionalProperties": false,
            "required": ["dimensione", "evidenza", "data_evidenza"],
            "properties": {
              "dimensione": {"type": "string", "enum": ["impatto", "velocity", "prossimita", "probabilita"]},
              "evidenza": {"type": "string", "minLength": 20, "maxLength": 600},
              "data_evidenza": {"$ref": "#/$defs/data"}
            }
          }
        },
        "fonti": {
          "type": "array",
          "minItems": 1,
          "items": {
            "type": "object",
            "additionalProperties": false,
            "required": ["riferimento", "data", "tipo", "rilevanza"],
            "properties": {
              "riferimento": {"type": "string", "minLength": 5, "maxLength": 300},
              "data": {"$ref": "#/$defs/data"},
              "tipo": {"type": "string", "enum": ["Interna", "Esterna"]},
              "rilevanza": {"type": "integer", "minimum": 1, "maximum": 3}
            }
          }
        },
        "razionale": {"type": "string", "minLength": 50, "maxLength": 2500},
        "azioni_mitigazione": {
          "type": "array",
          "maxItems": 5,
          "items": {
            "type": "object",
            "additionalProperties": false,
            "required": ["descrizione", "owner", "orizzonte", "razionale"],
            "properties": {
              "descrizione": {"type": "string", "minLength": 20, "maxLength": 500},
              "owner": {"type": "string", "minLength": 3, "maxLength": 100},
              "orizzonte": {"type": "string", "enum": ["3 mesi", "6 mesi", "12 mesi"]},
              "razionale": {"type": "string", "minLength": 20, "maxLength": 700}
            }
          }
        },
        "kri_proposti": {
          "type": "array",
          "maxItems": 2,
          "items": {
            "type": "object",
            "additionalProperties": false,
            "required": ["tipo", "nome", "unita", "fonte_dati", "soglia_attenzione", "soglia_allerta", "razionale"],
            "properties": {
              "kri_id": {"type": ["string", "null"], "pattern": "^NFR-[0-9]{2}-K[0-9]+$"},
              "tipo": {"type": "string", "enum": ["Esposizione", "Evoluzione"]},
              "nome": {"type": "string", "minLength": 5, "maxLength": 150},
              "unita": {"type": "string", "minLength": 1, "maxLength": 30},
              "fonte_dati": {"type": "string", "minLength": 5, "maxLength": 200},
              "soglia_attenzione": {"$ref": "#/$defs/soglia"},
              "soglia_allerta": {"$ref": "#/$defs/soglia"},
              "razionale": {"type": "string", "minLength": 20, "maxLength": 700}
            }
          }
        }
      }
    },
    "ricalcolo_kri": {
      "type": "object",
      "additionalProperties": false,
      "required": ["kri_id", "valore", "fonte", "commento"],
      "properties": {
        "kri_id": {"type": "string", "pattern": "^NFR-[0-9]{2}-K[0-9]+$"},
        "valore": {"type": ["number", "null"]},
        "fonte": {"type": "string", "minLength": 5, "maxLength": 200},
        "commento": {"type": "string", "minLength": 10, "maxLength": 500}
      }
    }
  }
}
'''

MODULI = {
    "valida_output": r'''"""Validatore JSON Schema minimo e severo, solo libreria standard.

Unica implementazione del pacchetto: la usano lo script di campo
(esegui_ciclo_locale.py) e la suite di collaudo (99_Collaudo), che la importa
da qui invece di tenerne una copia. Le regole NON sono scritte qui: stanno
negli schemi di 01_Kernel/, che questo modulo interpreta. Se cambia una regola
si cambia solo lo schema.

Costrutti supportati: type (anche lista, con "null"), required, properties,
additionalProperties (solo false), items, minItems, maxItems, uniqueItems,
enum, pattern, minLength, maxLength, minimum, maximum, $ref locale (#/$defs/..).
Un costrutto non supportato presente nello schema e' un ERRORE, non viene
ignorato: un validatore che salta in silenzio una regola la rende finta.

Uso da riga di comando:
    python valida_output.py <schema.json> <payload.json>
Codice di uscita 0 se conforme, 1 se non conforme, 2 se lo schema e' illeggibile.
"""

import json
import re
import sys

# Parole chiave che il validatore sa applicare, piu' quelle solo descrittive.
SUPPORTATE = {
    "type", "required", "properties", "additionalProperties", "items",
    "minItems", "maxItems", "uniqueItems", "enum", "pattern",
    "minLength", "maxLength", "minimum", "maximum", "$ref",
}
DESCRITTIVE = {"$schema", "$id", "title", "description", "$defs"}


def _tipo_ok(valore, tipo):
    # bool e' sottoclasse di int in Python: va escluso esplicitamente,
    # altrimenti true passerebbe come punteggio 1.
    if tipo == "null":
        return valore is None
    if tipo == "boolean":
        return isinstance(valore, bool)
    if tipo == "integer":
        return isinstance(valore, int) and not isinstance(valore, bool)
    if tipo == "number":
        return isinstance(valore, (int, float)) and not isinstance(valore, bool)
    if tipo == "string":
        return isinstance(valore, str)
    if tipo == "array":
        return isinstance(valore, list)
    if tipo == "object":
        return isinstance(valore, dict)
    raise ValueError(f"tipo non supportato nello schema: {tipo}")


def _risolvi(radice, ref):
    if not ref.startswith("#/"):
        raise ValueError(f"$ref non locale non supportato: {ref}")
    nodo = radice
    for parte in ref[2:].split("/"):
        nodo = nodo[parte]
    return nodo


def _valida(valore, schema, radice, percorso, errori):
    sconosciute = set(schema) - SUPPORTATE - DESCRITTIVE
    if sconosciute:
        raise ValueError(f"{percorso}: costrutti non supportati {sorted(sconosciute)}")

    if "$ref" in schema:
        _valida(valore, _risolvi(radice, schema["$ref"]), radice, percorso, errori)
        return

    if "type" in schema:
        tipi = schema["type"] if isinstance(schema["type"], list) else [schema["type"]]
        if not any(_tipo_ok(valore, t) for t in tipi):
            errori.append(f"{percorso}: tipo {type(valore).__name__}, atteso {tipi}")
            return  # i controlli successivi non hanno senso su un tipo sbagliato

    if "enum" in schema and valore not in schema["enum"]:
        errori.append(f"{percorso}: valore {valore!r} fuori da {schema['enum']}")

    if isinstance(valore, str):
        if "minLength" in schema and len(valore) < schema["minLength"]:
            errori.append(f"{percorso}: lunghezza {len(valore)} < {schema['minLength']}")
        if "maxLength" in schema and len(valore) > schema["maxLength"]:
            errori.append(f"{percorso}: lunghezza {len(valore)} > {schema['maxLength']}")
        if "pattern" in schema and not re.search(schema["pattern"], valore):
            errori.append(f"{percorso}: {valore!r} non rispetta {schema['pattern']}")

    if isinstance(valore, (int, float)) and not isinstance(valore, bool):
        if "minimum" in schema and valore < schema["minimum"]:
            errori.append(f"{percorso}: {valore} < minimo {schema['minimum']}")
        if "maximum" in schema and valore > schema["maximum"]:
            errori.append(f"{percorso}: {valore} > massimo {schema['maximum']}")

    if isinstance(valore, list):
        if "minItems" in schema and len(valore) < schema["minItems"]:
            errori.append(f"{percorso}: {len(valore)} elementi < {schema['minItems']}")
        if "maxItems" in schema and len(valore) > schema["maxItems"]:
            errori.append(f"{percorso}: {len(valore)} elementi > {schema['maxItems']}")
        if schema.get("uniqueItems"):
            visti = [json.dumps(v, sort_keys=True) for v in valore]
            if len(visti) != len(set(visti)):
                errori.append(f"{percorso}: elementi duplicati")
        if "items" in schema:
            for i, elemento in enumerate(valore):
                _valida(elemento, schema["items"], radice, f"{percorso}[{i}]", errori)

    if isinstance(valore, dict):
        for campo in schema.get("required", []):
            if campo not in valore:
                errori.append(f"{percorso}: manca il campo obbligatorio '{campo}'")
        proprieta = schema.get("properties", {})
        if "additionalProperties" in schema:
            if schema["additionalProperties"] is not False:
                raise ValueError(f"{percorso}: additionalProperties supportato solo come false")
            for campo in valore:
                if campo not in proprieta:
                    errori.append(f"{percorso}: campo non previsto '{campo}'")
        for campo, sotto in proprieta.items():
            if campo in valore:
                _valida(valore[campo], sotto, radice, f"{percorso}.{campo}", errori)


def valida(payload, schema):
    """Restituisce la lista degli errori (vuota = conforme)."""
    errori = []
    _valida(payload, schema, schema, "$", errori)
    return errori


def carica_json(percorso):
    # utf-8-sig: accetta anche "UTF-8 con BOM", una delle due codifiche UTF-8
    # che il Blocco note propone; senza, il file sarebbe rifiutato come illeggibile.
    with open(percorso, encoding="utf-8-sig") as f:
        return json.load(f)


def main(argv):
    if len(argv) != 3:
        print(__doc__)
        return 2
    try:
        schema = carica_json(argv[1])
    except (OSError, json.JSONDecodeError) as e:
        print(f"Schema illeggibile: {e}")
        return 2
    try:
        payload = carica_json(argv[2])
    except (OSError, json.JSONDecodeError) as e:
        print(f"NON CONFORME: payload illeggibile ({e})")
        return 1
    errori = valida(payload, schema)
    if errori:
        print(f"NON CONFORME: {len(errori)} errori")
        for e in errori:
            print("  -", e)
        return 1
    print("CONFORME")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
''',
    "radar": r'''"""Regola del registro unico (D-79, 08/10/2026): chi sta nel radar di Mosca e chi in osservazione.

Un solo registro di rischi NFR-nn, tutti valutati con la rubrica. I primi, secondo questa regola, sono **nel radar**
(gestione completa: KRI, azioni, urgenza, report, schede; sono l'anagrafica di Parametri); gli altri sono **in
osservazione** (stessa scala e stesso score, niente KRI ne' azioni). Rubrica, sezione 9.

La regola, in quattro passi:
1. candidato al radar: Prossimita >= soglia e Probabilita >= soglia;
2. sempre nel radar (eccezioni): AGIRE ORA (margine di Mosca < 0), tema del supervisore, decisione "Radar" del
   validatore; mai nel radar: decisione "Osservazione" del validatore;
3. se i candidati superano la capienza: prima le eccezioni, poi i piu' alti per score composito; gli altri restano
   fuori "per capienza";
4. un rischio del radar che non soddisfa piu' la regola diventa "Radar, in uscita" e passa in osservazione solo se lo
   e' anche al ciclo dopo (cicli_uscita = 2; con 1 esce subito). Un rischio che sale entra subito.
Oltre il massimo delle viste (max_radar, 20) nessun "in uscita": esce chi ha lo score piu' basso.

Un solo punto del codice per la regola: la usano esegui_ciclo_locale.py (trimestrale sui valori proposti, valida e la
riapertura sui valori validati, straordinario, aggiorna-registro), controlli_a5.py e il precontrollo (anche nella skill
di Copilot Studio). Solo libreria standard.
"""

RADAR, IN_USCITA, OSSERVAZIONE = "Radar", "Radar, in uscita", "Osservazione"
STATI = (RADAR, IN_USCITA, OSSERVAZIONE)
AUTOMATICO = "Automatico"
DECISIONI = (AUTOMATICO, RADAR, OSSERVAZIONE)
ENTRATO, USCITO = "Entrato", "Uscito"

# Valori del template (Parametri); la baseline porta quelli del registro (chiave regola_radar).
REGOLA_DEFAULT = {
    "capienza": 15, "soglia_prossimita": 2, "soglia_probabilita": 2, "cicli_uscita": 2,
    "max_radar": 20, "max_osservazione": 20,
    "pesi": {"impatto": 0.35, "probabilita": 0.25, "prossimita": 0.20, "velocity": 0.20},
    "anni_prossimita": [3.5, 2.5, 1.5, 0.5, 0.0],          # Z per Prossimita 1..5 (rubrica, sezione 8)
    "mosca_x": 0.25, "mosca_y_riserva": 1.0,
}
DIMENSIONI = ("impatto", "probabilita", "prossimita", "velocity")


def nel_radar(stato):
    return stato in (RADAR, IN_USCITA)


def _num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def score(r, regola=REGOLA_DEFAULT):
    """Media pesata delle dimensioni valutate (D-50): una Velocity non valutata non conta."""
    pesi = regola["pesi"]
    usate = {d: _num(r.get(d)) for d in DIMENSIONI if _num(r.get(d)) is not None}
    if not usate:
        return None
    return sum(pesi[d] * v for d, v in usate.items()) / sum(pesi[d] for d in usate)


def margine(r, regola=REGOLA_DEFAULT):
    """Margine di Mosca Z - (X + Y), come il foglio Radar_Mosca: Z dalla Prossimita; Y = orizzonte piu' lungo delle
    azioni non respinte del periodo di valutazione (y_azioni, in anni) o la riserva senza azioni, e gli anni di
    adeguamento se maggiori."""
    p = _num(r.get("prossimita"))
    if p is None:
        return None
    z = regola["anni_prossimita"][int(p) - 1]
    y_azioni = _num(r.get("y_azioni"))
    y = max(y_azioni if y_azioni is not None else regola["mosca_y_riserva"], _num(r.get("anni_adeguamento")) or 0)
    return z - (regola["mosca_x"] + y)


def _it(v):
    return "—" if v is None else f"{v:.2f}".replace(".", ",")


def stati(rischi, regola=REGOLA_DEFAULT):
    """Stato di ogni rischio del ciclo. rischi: dict con id, impatto, probabilita, prossimita, velocity, supervisore
    (bool), anni_adeguamento, y_azioni (anni o None), stato_prima (None per un rischio nuovo), decisione (Automatico,
    Radar, Osservazione). Restituisce {id: {stato, motivo, movimento, score, margine}}."""
    sp, sb = regola["soglia_prossimita"], regola["soglia_probabilita"]
    dati = {}
    for r in rischi:
        d = dict(r)
        d["score"], d["margine"] = score(r, regola), margine(r, regola)
        d["decisione"] = r.get("decisione") or AUTOMATICO
        dati[r["id"]] = d

    def candidato(d):
        return (_num(d.get("prossimita")) or 0) >= sp and (_num(d.get("probabilita")) or 0) >= sb

    def eccezione(d):
        if d["decisione"] == RADAR:
            return "decisione del validatore"
        if d["margine"] is not None and d["margine"] < 0:
            return "AGIRE ORA (radar di Mosca)"
        if d.get("supervisore"):
            return "tema del supervisore"
        return None

    def ordine(d):
        return (-(d["score"] or 0), -(_num(d.get("impatto")) or 0), d["id"])

    esclusi = {i for i, d in dati.items() if d["decisione"] == OSSERVAZIONE}
    eccezioni = sorted((d for i, d in dati.items() if i not in esclusi and eccezione(d)), key=ordine)
    candidati = sorted((d for i, d in dati.items() if i not in esclusi and not eccezione(d) and candidato(d)), key=ordine)
    posti = max(regola["capienza"] - len(eccezioni), 0)
    dentro, fuori_capienza = candidati[:posti], {d["id"] for d in candidati[posti:]}
    ultimo = min((d["score"] for d in eccezioni + dentro if d["score"] is not None), default=None)

    esito = {}
    for d in eccezioni + dentro:
        sopra = f"Prossimita {d['prossimita']} e Probabilita {d['probabilita']}: sopra la soglia"
        e = eccezione(d)
        if e and candidato(d):
            motivo = f"{sopra}; posto garantito: {e}"
        elif e:
            motivo = f"sotto la soglia, nel radar per eccezione: {e}"
        else:
            motivo = sopra
        esito[d["id"]] = (RADAR, motivo)
    for i, d in dati.items():
        if i in esito:
            continue
        if i in esclusi:
            motivo = "in osservazione per decisione del validatore"
        elif i in fuori_capienza:
            motivo = f"fuori per capienza: score {_it(d['score'])}, ultimo nel radar {_it(ultimo)}"
        else:
            sotto = [f"{nome} {d.get(chiave)} sotto la soglia {s}" for nome, chiave, s in
                     (("Prossimita", "prossimita", sp), ("Probabilita", "probabilita", sb))
                     if (_num(d.get(chiave)) or 0) < s]
            motivo = "; ".join(sotto)
        if d.get("stato_prima") == RADAR and regola["cicli_uscita"] >= 2 and i not in esclusi:
            esito[i] = (IN_USCITA, f"{motivo} (primo ciclo: esce se lo e' anche al ciclo dopo)")
        else:
            esito[i] = (OSSERVAZIONE, motivo)
    # Oltre il massimo delle viste nessun "in uscita": esce chi ha lo score piu' basso.
    radar = sorted((dati[i] for i, (s, _) in esito.items() if nel_radar(s)), key=ordine)
    for d in radar[regola["max_radar"]:]:
        esito[d["id"]] = (OSSERVAZIONE, esito[d["id"]][1] + f"; oltre i {regola['max_radar']} posti del radar")

    out = {}
    for i, (stato, motivo) in esito.items():
        prima = dati[i].get("stato_prima")
        movimento = (ENTRATO if nel_radar(stato) and prima == OSSERVAZIONE else
                     USCITO if not nel_radar(stato) and nel_radar(prima) else "")
        out[i] = {"stato": stato, "motivo": motivo, "movimento": movimento, "score": dati[i]["score"],
                  "margine": dati[i]["margine"]}
    return out


def y_azioni(azioni):
    """Y delle azioni per il margine di Mosca: orizzonte piu' lungo, in anni, delle azioni non respinte (None senza
    azioni). azioni: elenco di orizzonti '3 mesi' / '6 mesi' / '12 mesi' (o numeri di mesi)."""
    mesi = []
    for a in azioni:
        try:
            mesi.append(int(str(a).split()[0]))
        except (ValueError, IndexError):
            continue
    return max(mesi) / 12 if mesi else None
''',
    "controlli_a5": r'''"""Controlli meccanizzabili di A5 Challenger — controparte deterministica.

Stessi nove controlli (C1-C9) del prompt 02_Agenti/A5_Challenger.md, nella
parte esprimibile come regola. Copertura DICHIARATA, non sottintesa:
  - C1, C2, C3, C7, C9: per intero (C3 = date, frasi generiche, max 1 punto).
  - C5: date delle fonti per intero (Da motivare, D-71); fonte di rilevanza 3 per le
    valutazioni alte e fonti perse come Osservazioni (D-58); dal 19/10/2026 (D-81)
    l'esposizione della Banca non nota alle fonti: la frase "ESPOSIZIONE NON NOTA: con
    l'esposizione documentata sarebbe N" nel razionale porta i due valori in validazione
    (Da motivare), dove decide il validatore.
  - C6, C8: euristica a stringhe, solo "Osservazione" (C8 dal 30/09/2026,
    D-54: il controllo dell'owner "persona" e' passato alla validazione; dal
    02/10/2026, D-69, anche un'azione uguale a un'azione aperta).
  - C4 (plausibilita' rispetto ai descrittori): NON coperto. Richiede il
    giudizio di A5. Questo modulo non sostituisce A5: lo affianca.

Dal 30/09/2026 (D-49..D-56) A1 non scrive esito ne' i "da/a" delle
variazioni (li calcola lo script: esegui_ciclo_locale.completa_payload), non
valuta Gap_Controlli, puo' lasciare la Velocity null e propone KRI con soglie
facoltative. Dal 30/09 sera (D-57, D-58, D-60) Severity si chiama Impatto, la
Confidence del rischio non c'e' piu' (ogni fonte ha una rilevanza 1-3) e A1
propone gli anni di adeguamento; i KRI sono due per rischio, uno di
Esposizione e uno di Evoluzione (D-64). I controlli lavorano sul payload COSI' COME L'HA
SCRITTO A1, confrontato con la baseline.

L'esito ha lo stesso formato dell'output di A5 (01_Kernel/schema_output_A5.json),
cosi' i due si leggono e si combinano allo stesso modo.

Dal 06/10/2026 (D-71) tre gravita':
  - "Bloccante": solo l'integrita' del registro (C1, C2, C7 e gli identificativi dei KRI di C9): ferma la
    scrittura, perche' senza queste regole il registro diventerebbe incoerente;
  - "Da motivare": regole di metodo (C3, date delle fonti di C5, coppia e soglie dei KRI di C9, KRI non riletti
    in un ciclo straordinario di C7, D-77) e i Bloccanti di A5. Non fermano: vanno nella colonna "Segnalazioni dei controlli" del foglio Validazione e chiedono una
    Motivazione al validatore; le soglie e il salto di piu' di un punto li corregge la validazione stessa;
  - "Osservazione": dubbi da considerare, senza obblighi.
A5 e' facoltativo: senza il suo esito valgono i soli controlli di questo modulo.

Dall'08/10/2026 (D-79, registro unico) i rischi sono nel radar o in osservazione: lo stato si prevede qui con la regola
della baseline (radar.stati, la stessa dello script). Completezza (C1) solo per il radar; KRI e azioni (C8, C9) solo per
chi sara' nel radar: mancano a un rischio che entra dall'osservazione = da motivare, a un altro rischio del radar =
Bloccante; scritti per un rischio in osservazione = Osservazione (lo script non li registra). Piu' di 20 rischi in
osservazione = C1 Bloccante. Le regole le legge A1 nella guida di scrittura (allegato).
"""

import re

import radar

# Dimensioni valutate da A1 (rubrica, sezione 1). La Velocity puo' essere null.
DIMENSIONI = ["impatto", "probabilita", "prossimita", "velocity"]
VALORI = DIMENSIONI
TIPI_KRI = ["Esposizione", "Evoluzione"]          # D-64: esattamente un KRI per tipo
# D-74: categorie della tassonomia interna della Banca in perimetro (01_Kernel/tassonomia_nfr.md; il collaudo verifica che
# coincidano con il kernel). FP = rischio esistente fuori perimetro, da togliere con dismetti.
CATEGORIE = ["9.1", "9.4", "9.5", "9.6", "9.7", "9.8", "9.9", "9.11", "9.13"]
FUORI_PERIMETRO = "FP"
BLOCCANTE, DA_MOTIVARE, OSSERVAZIONE = "Bloccante", "Da motivare", "Osservazione"     # D-71

# D-81 (analisi 17, 17.10): A1 non conosce l'esposizione interna della Banca (D-49) e la Probabilita 4-5 vale solo con la
# Banca esposta (D-65). Invece di scegliere in silenzio il valore piu' basso, A1 lo scrive nel razionale con il valore che
# darebbe se l'esposizione fosse documentata; il validatore, che l'esposizione la conosce, sceglie fra i due.
ETICHETTA_ESPOSIZIONE = "ESPOSIZIONE NON NOTA"
_ESPOSIZIONE = re.compile(r"ESPOSIZIONE NON NOTA\s*:[^.;\n]*?sarebbe\D{0,25}?([1-5])\b", re.IGNORECASE)


def esposizione_non_nota(razionale):
    """(dichiarata, N): se il razionale dice che l'esposizione della Banca non e' nota e con quale Probabilita N sarebbe
    valutato il rischio con l'esposizione documentata (N None se la frase non ha un valore leggibile)."""
    testo = str(razionale or "")
    if ETICHETTA_ESPOSIZIONE.lower() not in testo.lower():
        return False, None
    m = _ESPOSIZIONE.search(testo)
    return True, (int(m.group(1)) if m else None)

# Frasi che non sono evidenza (rubrica, regola di variazione 3). Confronto
# in minuscolo, per sottostringa. Estendere qui se in esercizio ne emergono altre.
FRASI_GENERICHE = [
    "contesto in evoluzione", "maggiore attenzione mediatica", "trend di mercato",
    "scenario in evoluzione", "crescente attenzione", "clima di incertezza",
    "contesto generale", "evoluzione del contesto", "attenzione dei media",
]

# Parole che indicano una funzione aziendale (C8), confrontate come parole
# intere (altrimenti "cio" riconoscerebbe "Lucio"). Euristica: un owner senza
# nessuna di queste parole produce un'Osservazione. La stessa regola (sembra_persona)
# e' un fermo nella validazione umana (esegui_ciclo_locale.py, comando valida).
PAROLE_FUNZIONE = [
    "funzione", "direzione", "ufficio", "servizio", "area", "unita", "unità",
    "comitato", "team", "struttura", "dipartimento", "divisione", "presidio",
    "risk", "compliance", "audit", "operations", "operational", "ict", "it",
    "legale", "hr", "risorse umane", "sicurezza", "security", "cyber",
    "information", "resilience", "governance", "management", "comunicazione",
    "istituzionale", "antifrode", "continuity", "pmo", "cro", "cio",
    "ciso", "coo", "cfo", "dpo", "customer", "experience", "clienti",
    "procurement", "acquisti", "immobili", "infrastrutture", "tecnologia",
    "organizzazione", "operativa", "operativo", "controlli", "business",
]
NOME_COGNOME = re.compile(r"^[A-ZÀ-Ý][a-zà-ÿ']+(\s+[A-ZÀ-Ý][a-zà-ÿ']+){1,2}$")


def ha_funzione(owner):
    return any(re.search(r"\b" + re.escape(p) + r"\b", owner.lower()) for p in PAROLE_FUNZIONE)


def sembra_persona(owner):
    """Owner con la forma 'Nome Cognome' e nessuna parola di funzione.
    La usano C8 (Osservazione) e la validazione umana (fermo): stessa regola."""
    owner = owner.strip()
    return bool(NOME_COGNOME.match(owner)) and not ha_funzione(owner)


def _rilievo(controllo, id_rischio, gravita, descrizione):
    return {"controllo": controllo, "id_rischio": id_rischio, "gravita": gravita, "descrizione": descrizione}


def _diversi(a, b):
    return a != b                       # None contro None = uguali; None contro 3 = diversi


def _stesso_testo(testo):
    """Descrizione di un'azione a meno di maiuscole, spazi e punteggiatura finale (C8, D-69)."""
    return re.sub(r"\s+", " ", testo.lower()).strip(" .;")


def _controlla_analisi_evento(payload, base, ids, d_rif):
    """Ciclo straordinario (D-77): l'analisi dell'evento di A8. Tutte Osservazioni: l'analisi spiega, non scrive valori
    nel registro; un ciclo straordinario fatto con A1 (senza analisi) resta possibile."""
    O = OSSERVAZIONE
    analisi = payload.get("analisi_evento")
    if not analisi:
        return [_rilievo("C1", None, O, "Ciclo straordinario senza analisi_evento: manca l'analisi approfondita dell'evento "
                                        "(agente A8); il validatore la trova solo nei razionali e nella sintesi.")]
    r = []
    toccati = [t["id_rischio"] for t in analisi.get("rischi_toccati") or []]
    for rid in sorted(set(toccati) - set(ids) - set(base)):
        r.append(_rilievo("C1", rid, O, "analisi_evento: rischio toccato che non e' ne' nel registro ne' fra i rischi nuovi."))
    for rid in sorted({t for t in toccati if t in base} - set(ids)):
        r.append(_rilievo("C1", rid, O, "analisi_evento: rischio toccato dall'evento ma non indicato per lo straordinario: la "
                                        "funzione di rischio valuti se rivalutarlo (un altro straordinario o il trimestre)."))
    for rid in sorted(set(ids) - set(toccati)):
        r.append(_rilievo("C1", rid, O, "analisi_evento: rischio scritto nel ciclo ma assente da rischi_toccati: dire che "
                                        "cosa gli fa l'evento."))
    for d in analisi.get("documenti") or []:
        if d_rif and str(d.get("data")) > d_rif:
            r.append(_rilievo("C1", None, O, f"analisi_evento: documento datato dopo la data di riferimento ({d.get('data')}): "
                                             f"{d.get('riferimento')}."))
    return r


def stati_previsti(payload, baseline):
    """D-79: stato che ogni rischio avra' con i valori proposti da A1, con la regola della baseline (la stessa che applica
    lo script alla scrittura). Entrano i rischi scritti e quelli che lo script riporta: in osservazione non toccati e, in
    uno straordinario, quelli del radar non rivalutati. None per una baseline precedente al registro unico."""
    regola = baseline.get("regola_radar")
    if not regola:
        return None
    nel = {x["id_rischio"]: x for x in baseline.get("rischi", [])}
    oss = {x["id_rischio"]: x for x in baseline.get("rischi_in_osservazione", [])}
    scritti = {x["id_rischio"] for x in payload.get("rischi", [])}

    def ingresso(rid, valori, prima, y):
        return {"id": rid, **{d: valori.get(d) for d in radar.DIMENSIONI},
                "supervisore": bool((nel.get(rid) or {}).get("attenzione_supervisore")),
                "anni_adeguamento": (nel[rid].get("anni_adeguamento") if rid in nel else valori.get("anni_adeguamento")),
                "y_azioni": y, "stato_prima": (prima or {}).get("stato") or (radar.RADAR if rid in nel else None),
                "decisione": (prima or {}).get("decisione_stato")}
    ingressi = [ingresso(x["id_rischio"], x, nel.get(x["id_rischio"]) or oss.get(x["id_rischio"]),
                         radar.y_azioni([a["orizzonte"] for a in x.get("azioni_mitigazione") or []]))
                for x in payload.get("rischi", [])]
    riportati = {**oss, **(nel if baseline.get("straordinario") else {})}
    for rid, p in riportati.items():
        if rid not in scritti:
            y = radar.y_azioni([a.get("orizzonte") for a in baseline.get("azioni_aperte", []) if a.get("id_rischio") == rid])
            ingressi.append(ingresso(rid, p, p, y if rid in nel else None))
    return radar.stati(ingressi, regola)


def controlla(payload, baseline):
    """Applica i controlli meccanizzabili. Restituisce un dict conforme allo schema A5."""
    r = []
    B, M, O = BLOCCANTE, DA_MOTIVARE, OSSERVAZIONE
    base = {x["id_rischio"]: x for x in baseline.get("rischi", [])}          # il radar (D-79)
    osservati = {x["id_rischio"]: x for x in baseline.get("rischi_in_osservazione", [])}
    tutti = {**base, **osservati}
    previsti = stati_previsti(payload, baseline)
    nel_radar = lambda rid: previsti is None or radar.nel_radar(previsti[rid]["stato"])
    kri_base = {k["kri_id"]: k for k in baseline.get("kri_in_vigore", [])}
    d_prec = baseline.get("data_riferimento_precedente") or "0000-00-00"
    d_rif = payload.get("data_riferimento", "")
    rischi = payload.get("rischi", [])

    # ---------------- C1 Completezza
    if payload.get("ciclo") != baseline.get("ciclo_da_valutare"):
        r.append(_rilievo("C1", None, B, f"ciclo {payload.get('ciclo')} diverso dal ciclo da valutare della baseline ({baseline.get('ciclo_da_valutare')})."))
    if d_rif != baseline.get("data_riferimento"):
        r.append(_rilievo("C1", None, B, f"data_riferimento {d_rif} diversa da quella della baseline ({baseline.get('data_riferimento')})."))
    ids = [x["id_rischio"] for x in rischi]
    straord = baseline.get("straordinario")
    if straord:                                      # D-75: solo i rischi indicati, piu' quelli nuovi
        da_rivalutare = set(straord.get("rischi_da_rivalutare") or [])
        for rid in sorted(da_rivalutare - set(ids)):
            r.append(_rilievo("C1", rid, B, "Rischio da rivalutare nel ciclo straordinario assente dall'output."))
        for rid in sorted({i for i in ids if i in tutti} - da_rivalutare):
            r.append(_rilievo("C1", rid, B, "Rischio esistente non indicato per il ciclo straordinario: non va riscritto "
                                            "(lo script lo riporta com'era)."))
        if not da_rivalutare and not any(i not in tutti for i in ids):
            r.append(_rilievo("C1", None, B, "Ciclo straordinario di solo censimento (nessun rischio da rivalutare): serve "
                                             "almeno un rischio nuovo creato dall'evento."))
        r += _controlla_analisi_evento(payload, tutti, ids, d_rif)
    else:
        for rid in sorted(set(base) - set(ids)):
            r.append(_rilievo("C1", rid, B, "Rischio in anagrafica non valutato nel ciclo."))
        if payload.get("analisi_evento"):
            r.append(_rilievo("C1", None, O, "analisi_evento e' per i cicli straordinari (A8): in un ciclo trimestrale resta "
                                             "nell'output archiviato ma non va nel registro."))
    for rid in sorted({i for i in ids if ids.count(i) > 1}):
        r.append(_rilievo("C1", rid, B, "id_rischio ripetuto nel payload."))
    for x in rischi:
        prec = tutti.get(x["id_rischio"])
        et, et_prec = x["event_type"], (prec or {}).get("event_type")
        if et == FUORI_PERIMETRO:
            if prec is None:
                r.append(_rilievo("C1", x["id_rischio"], B, "Un rischio nuovo non e' mai fuori perimetro (FP): un tema fuori "
                                                           "perimetro va fra i fenomeni non censiti (motivo fuori_perimetro)."))
            elif et_prec != FUORI_PERIMETRO:
                r.append(_rilievo("C1", x["id_rischio"], M, f"Fuori perimetro (D-74): A1 non trova una categoria in perimetro per "
                                                           f"il meccanismo (era {et_prec}). Resta nel registro, fuori dal radar, "
                                                           "finche' la funzione di rischio non lo toglie con dismetti; motivare la "
                                                           "decisione."))
        elif prec is not None and et != et_prec:
            if et_prec not in CATEGORIE:             # codice di Basilea o FP di prima: riclassificazione (D-74)
                r.append(_rilievo("C1", x["id_rischio"], O, f"Riclassificato da {et_prec} alla categoria {et} (D-74): verificare "
                                                           "che corrisponda al meccanismo di perdita; se no, correggere la "
                                                           "categoria nell'anagrafica (Parametri, colonna J)."))
            else:
                r.append(_rilievo("C1", x["id_rischio"], B, f"Categoria {et} diversa da quella in anagrafica ({et_prec}): un "
                                                           "rischio esistente non cambia meccanismo; se e' un altro rischio, "
                                                           "serve un id nuovo."))

    for x in rischi:
        rid = x["id_rischio"]
        prec = tutti.get(rid)
        var = x.get("variazioni", [])

        # ---------------- C2 Variazioni: una voce per ogni valore cambiato, e solo per quelli
        if prec is not None:
            cambiati = [v for v in VALORI if _diversi(x.get(v), prec.get(v))]
            for v in cambiati:
                voci = [z for z in var if z["dimensione"] == v]
                if len(voci) != 1:
                    r.append(_rilievo("C2", rid, B, f"{v} passa da {prec.get(v)} a {x.get(v)}: serve una (e una sola) "
                                                    "voce in variazioni con l'evidenza."))
            for z in var:
                if z["dimensione"] not in cambiati:
                    r.append(_rilievo("C2", rid, B, f"Variazione dichiarata su {z['dimensione']} che non cambia rispetto alla baseline."))
            # C3b sul valore effettivo: al massimo un punto (una Velocity prima null entra con qualunque valore).
            for v in VALORI:
                a, b = prec.get(v), x.get(v)
                if a is not None and b is not None and abs(b - a) > 1:
                    r.append(_rilievo("C3", rid, M, f"{v} passa da {a} a {b}: piu' di un punto fra cicli consecutivi "
                                                    "(in validazione va riportato a un punto: valida non lo accetta)."))
        elif var:
            r.append(_rilievo("C2", rid, B, "Un rischio nuovo non ha variazioni: non esiste un valore precedente."))

        # ---------------- C3 Evidenza delle variazioni
        for z in var:
            if not (d_prec < z["data_evidenza"] <= d_rif):
                r.append(_rilievo("C3", rid, M, f"Evidenza su {z['dimensione']} datata {z['data_evidenza']}: deve essere dopo {d_prec} e non dopo {d_rif}."))
            testo = z["evidenza"].lower()
            generiche = [f for f in FRASI_GENERICHE if f in testo]
            if generiche:
                r.append(_rilievo("C3", rid, M, f"Evidenza su {z['dimensione']} generica ({', '.join(generiche)}): serve un fatto specifico."))

        # ---------------- C5 Fonti (D-58: rilevanza per fonte al posto della Confidence)
        fonti = x.get("fonti", [])
        for f in fonti:
            if f["data"] > d_rif:
                r.append(_rilievo("C5", rid, M, f"Fonte '{f['riferimento'][:60]}' datata {f['data']}, dopo la data di riferimento {d_rif}."))
        if (x.get("probabilita") == 5 or (x.get("prossimita") or 0) >= 4) and not any(f.get("rilevanza") == 3 for f in fonti):
            r.append(_rilievo("C5", rid, O, f"Probabilita {x.get('probabilita')} e Prossimita {x.get('prossimita')} senza una fonte "
                                            "di rilevanza 3: ammesso solo con la deroga per un fenomeno ad altissimo impatto, "
                                            "dichiarata nel razionale."))
        if prec is not None and len(fonti) < len(prec.get("fonti", [])) and not any(_diversi(x.get(v), prec.get(v)) for v in VALORI):
            r.append(_rilievo("C5", rid, O, f"Fonti scese da {len(prec['fonti'])} a {len(fonti)} con i valori invariati: "
                                            "verificare che restino sostenuti (regola di variazione 5)."))
        dichiarata, n_esp = esposizione_non_nota(x.get("razionale"))
        p_prop = x.get("probabilita")
        if dichiarata and isinstance(p_prop, int) and n_esp is not None and n_esp > p_prop:
            limite = ("" if prec is None else
                      f" (rischio gia' nel registro: in validazione al massimo {min(5, int(prec.get('probabilita') or 0) + 1)}, "
                      "cioe' un punto dal ciclo prima; oltre, dopo la validazione, con riapri)")
            r.append(_rilievo("C5", rid, M, f"Esposizione della Banca non nota alle fonti (D-81): Probabilita {p_prop} proposta, "
                                            f"{n_esp} con l'esposizione documentata. Decide il validatore: se la Banca e' esposta "
                                            f"(ha il prodotto, il processo o il fornitore colpito) scrivere {n_esp} nella cella "
                                            f"gialla{limite}; altrimenti motivare {p_prop}."))
        elif dichiarata:
            r.append(_rilievo("C5", rid, O, "ESPOSIZIONE NON NOTA senza un valore maggiore della Probabilita proposta: la frase "
                                            "va scritta \"con l'esposizione documentata sarebbe N\", con N maggiore di "
                                            f"{p_prop} (guida di scrittura, D-81)."))

        # ---------------- D-79 Registro unico: KRI e azioni solo per chi sara' nel radar
        if not nel_radar(rid):
            if x.get("kri_proposti") or x.get("azioni_mitigazione"):
                r.append(_rilievo("C9", rid, O, f"Rischio in osservazione ({previsti[rid]['motivo']}): KRI e azioni non si "
                                                "scrivono, lo script non li registra."))
            if rid in osservati and not any(d_prec < f["data"] for f in x.get("fonti", [])):
                r.append(_rilievo("C5", rid, O, f"Rischio in osservazione riscritto senza una fonte dopo {d_prec}: senza "
                                                "evidenza nuova non si riscrive (resta com'era)."))
            if not x.get("orizzonte_anni") or not x.get("segnale_avvicinamento"):
                r.append(_rilievo("C1", rid, O, "Rischio in osservazione senza orizzonte_anni o segnale_avvicinamento: "
                                                "servono a chi legge per capire quanto e' lontano e che cosa lo farebbe salire."))
            continue
        entra = previsti is not None and previsti[rid]["movimento"] == radar.ENTRATO
        if not x.get("kri_proposti") or not x.get("azioni_mitigazione"):
            manca = " e ".join(m for m, c in (("i due KRI", not x.get("kri_proposti")),
                                               ("le azioni", not x.get("azioni_mitigazione"))) if c)
            r.append(_rilievo("C9", rid, M if entra else B,
                              f"Rischio che entra nel radar senza {manca}: scriverli (o restano da definire al ciclo dopo)."
                              if entra else f"Rischio del radar senza {manca}: ogni rischio del radar ha due KRI e da 1 a 5 "
                                            "azioni."))

        # ---------------- C8 Azioni (euristica, solo Osservazioni)
        aperte = {_stesso_testo(z["descrizione"]): z["id_azione"] for z in baseline.get("azioni_aperte", [])
                  if z.get("id_rischio") == rid}
        for a in x.get("azioni_mitigazione", []):
            owner = a["owner"].strip()
            if sembra_persona(owner):
                r.append(_rilievo("C8", rid, O, f"Owner '{owner}' sembra una persona: deve essere una funzione "
                                                "(la validazione non accetta il nome di una persona)."))
            elif not ha_funzione(owner):
                r.append(_rilievo("C8", rid, O, f"Owner '{owner}': verificare che sia una funzione aziendale."))
            gemella = aperte.get(_stesso_testo(a["descrizione"]))
            if gemella:
                r.append(_rilievo("C8", rid, O, f"Azione '{a['descrizione'][:60]}' uguale all'azione aperta {gemella}: "
                                                "non riproporla; se resta valida, richiamala nel razionale e proponi il "
                                                "passo successivo."))

        # ---------------- C9 KRI
        kri = x.get("kri_proposti", [])
        if kri and sorted(k.get("tipo") or "" for k in kri) != TIPI_KRI:
            r.append(_rilievo("C9", rid, M, f"{len(kri)} KRI proposti ({', '.join(k.get('tipo') or 'senza tipo' for k in kri)}): "
                                            "servono due KRI, uno di Esposizione e uno di Evoluzione."))
        for k in kri:
            nome = k["nome"][:50]
            att, all_ = k.get("soglia_attenzione"), k.get("soglia_allerta")
            if (att is None) != (all_ is None):
                r.append(_rilievo("C9", rid, M, f"KRI '{nome}': soglie entrambe o nessuna (in validazione: scriverle "
                                                "entrambe o svuotarle; valida non accetta una sola)."))
            elif att is not None and all_ < att:
                r.append(_rilievo("C9", rid, M, f"KRI '{nome}': soglia_allerta {all_} < soglia_attenzione {att} (soglie "
                                                "invertite: in validazione vanno corrette, valida non le accetta)."))
            kid = k.get("kri_id")
            if kid:
                if not kid.startswith(rid + "-K"):
                    r.append(_rilievo("C9", rid, B, f"KRI {kid} appartiene a un altro rischio."))
                elif kid not in kri_base:
                    r.append(_rilievo("C9", rid, B, f"KRI {kid} non e' in vigore nella baseline: un KRI nuovo ha kri_id null."))
                elif kri_base[kid].get("tipo") and k.get("tipo") != kri_base[kid]["tipo"]:
                    r.append(_rilievo("C9", rid, O, f"KRI {kid}: tipo {k.get('tipo')}, nella baseline {kri_base[kid]['tipo']} "
                                                    "(un KRI in vigore conserva il tipo)."))
        if prec is not None:
            riproposti = {k.get("kri_id") for k in kri}
            usciti = sorted(k for k, v in kri_base.items() if v["id_rischio"] == rid and k not in riproposti)
            if usciti:
                r.append(_rilievo("C9", rid, O, f"KRI in vigore non riproposti, escono di vigore: {', '.join(usciti)}."))

    tutti_kid = [k["kri_id"] for x in rischi for k in x.get("kri_proposti", []) if k.get("kri_id")]
    for kid in sorted({k for k in tutti_kid if tutti_kid.count(k) > 1}):
        r.append(_rilievo("C9", None, B, f"kri_id {kid} usato piu' di una volta."))

    # ---------------- C6 Inflazione progressiva (euristica, livello ciclo)
    esistenti = [x for x in rischi if x["id_rischio"] in base]
    def aumenta(x):
        p = base[x["id_rischio"]]
        return any(x.get(v) is not None and p.get(v) is not None and x[v] > p[v] for v in DIMENSIONI)
    in_aumento = [x for x in esistenti if aumenta(x)]
    if esistenti and len(in_aumento) > len(esistenti) / 2:
        nuove = any(f["data"] > d_prec for x in in_aumento for f in x["fonti"])
        if not nuove:
            r.append(_rilievo("C6", None, O, f"{len(in_aumento)} rischi su {len(esistenti)} in aumento senza fonti nuove: possibile inflazione progressiva."))

    # ---------------- C7 Rilettura dei KRI del ciclo precedente: solo KRI in vigore, una volta ciascuno
    riletti = [z["kri_id"] for z in payload.get("ricalcolo_kri_ciclo_precedente", [])]
    for kid in sorted(set(riletti) - set(kri_base)):
        r.append(_rilievo("C7", None, B, f"KRI {kid} riletto ma non in vigore nella baseline."))
    for kid in sorted({k for k in riletti if riletti.count(k) > 1}):
        r.append(_rilievo("C7", None, B, f"KRI {kid} riletto piu' di una volta."))
    if baseline.get("straordinario"):
        # D-77: nello straordinario A8 rilegge ogni KRI in vigore dei rischi rivalutati che ripropone; un valore che
        # manca lascia il KRI Non rilevato e toglie al livello di allerta l'aggravante proprio quando il rischio sale.
        # Da motivare, non Bloccante: il dato puo' non essere nei documenti e il validatore lo scrive nella griglia.
        for x in rischi:
            if x["id_rischio"] not in base:
                continue
            mancanti = sorted({k.get("kri_id") for k in x.get("kri_proposti", [])}
                              & {kid for kid, k in kri_base.items() if k.get("id_rischio") == x["id_rischio"]}
                              - set(riletti))
            if mancanti:
                r.append(_rilievo("C7", x["id_rischio"], M, f"Ciclo straordinario: KRI in vigore non riletti "
                                                            f"({', '.join(mancanti)}): restano Non rilevati. Rileggerli, o "
                                                            "scrivere il valore nella griglia di validazione."))

    # ---------------- D-79 Al massimo 20 rischi in osservazione (Parametri, capienzaOsservazione)
    if previsti is not None:
        quanti = sum(1 for s in previsti.values() if s["stato"] == radar.OSSERVAZIONE)
        massimo = baseline["regola_radar"].get("max_osservazione", radar.REGOLA_DEFAULT["max_osservazione"])
        if quanti > massimo:
            r.append(_rilievo("C1", None, B, f"{quanti} rischi in osservazione dopo il ciclo, oltre il massimo di {massimo}: "
                                             "lasciare fuori i meno rilevanti (fra i fenomeni non censiti, troppo_lontano)."))

    bloccanti = sum(1 for x in r if x["gravita"] == B)
    da_motivare = sum(1 for x in r if x["gravita"] == M)
    return {
        "ciclo": payload.get("ciclo") or "0000Q1",
        "esito": "Da rivedere" if bloccanti else "Approvato",
        "rilievi": r,
        "sintesi": (f"Controlli meccanizzati: {bloccanti} rilievi bloccanti, {da_motivare} da motivare in validazione, "
                    f"{len(r) - bloccanti - da_motivare} osservazioni. Bloccanti solo per l'integrita' (C1, C2, C7, "
                    "identificativi dei KRI); C3, date di C5, coppia e soglie dei KRI da motivare; C5 (rilevanza), C6 e "
                    "C8 come Osservazioni; C4 non coperto: resta al giudizio di A5, se usato."),
    }


def combina(meccanico, agente=None):
    """Esito del ciclo (D-71): ferma solo un Bloccante dei controlli meccanizzati, cioe' un difetto di integrita'.
    A5 e' facoltativo; i suoi rilievi Bloccanti diventano "Da motivare": vanno nella colonna Segnalazioni del foglio
    Validazione e chiedono una Motivazione al validatore. I rilievi restano tutti, con l'origine."""
    rilievi = [dict(x, origine="script") for x in meccanico["rilievi"]] + \
              [dict(x, origine="A5", gravita=DA_MOTIVARE if x["gravita"] == BLOCCANTE else x["gravita"])
               for x in (agente["rilievi"] if agente else [])]
    fermo = meccanico["esito"] != "Approvato"
    return {"esito": "Da rivedere" if fermo else "Approvato", "rilievi": rilievi,
            "da_motivare": [x for x in rilievi if x["gravita"] == DA_MOTIVARE]}
''',
    "schede": r'''"""Schede dei rischi: file a parte rispetto al registro (DECISIONI D-40).

Tre passi, tutti dopo la validazione del ciclo trimestrale:
1. input_schede(): dal registro validato, un JSON per rischio con cio' che la
   scheda mostra (valori validati, urgenza, livello di allerta, variazioni,
   fonti, azioni e KRI con il tipo, storia; dal 02/10/2026, D-68, il razionale
   di A1 diviso in fenomeno e motivazioni per valore, e il perche' e l'effetto
   atteso di azioni e KRI). E' l'unico allegato dell'agente A6 (comando schede-input).
2. controlla_a6(): l'output di A6 (testi estesi) rispetta l'input? Stessi
   rischi, quattro dimensioni, le stesse fonti copiate identiche, nessun numero che
   non stia nei dati del rischio o nella rubrica. Se no, niente file.
3. scrivi_file_schede(): Schede_<ciclo>.xlsx, un indice e una scheda per
   rischio nell'ordine della classifica (comando schede). Senza l'output di A6
   le schede usano i testi di A1 dal registro.

Il passo 1 legge il workbook via Excel (COM); i passi 2 e 3 sono Python puro
(openpyxl), quindi il collaudo li prova anche senza Excel.
"""

import json
import re
from pathlib import Path

# Valori spiegati nelle schede (D-49, D-50, D-57, D-58): le quattro dimensioni dello score (la Velocity puo'
# mancare). La Confidence non c'e' piu': ogni fonte ha la sua rilevanza 1-3. Il presidio dei controlli e'
# un giudizio qualitativo della Banca, mostrato a parte.
DIMENSIONI = ["impatto", "probabilita", "prossimita", "velocity"]
NOMI_DIM = {"impatto": "Impatto", "probabilita": "Probabilita", "prossimita": "Prossimita", "velocity": "Velocity"}
COLONNA_DIM = dict(NOMI_DIM)
# Etichette anche per le variazioni dei cicli scritti prima del 30/09 sera (severity, confidence).
ETICHETTE_VARIAZIONI = dict(NOMI_DIM, severity="Impatto", confidence="Confidence")
PESI = {"impatto": "pesoImpatto", "probabilita": "pesoProbabilita", "prossimita": "pesoProssimita",
        "velocity": "pesoVelocity"}
RILEVANZA = {1: "Bassa", 2: "Media", 3: "Alta"}
NON_VALUTATA = "n.v."
STATI_CHIUSI = {"Completata", "Respinta", "Annullata"}
NOME_ANALISI = "Horizon Scan NFR"
RISERVATEZZA = "Uso interno"


def _num(v, decimali=2):
    """Numero dal workbook, arrotondato (Excel restituisce 3.8000000000000003)."""
    if v is None or v == "":
        return None
    try:
        f = round(float(v), decimali)
    except (TypeError, ValueError):
        return v
    return int(f) if f == int(f) else f


def _colonna(sessione, nome):
    v = sessione.intervallo_nome(nome).Value
    return [r[0] for r in v] if isinstance(v, tuple) else [v]


# ---------------------------------------------------------------- razionali in parti (D-68)
# A1 scrive il razionale del rischio in parti con etichette fisse, come MESSAGGIO/DECISIONI della sintesi:
# "RISCHIO: il fenomeno in termini di rischio. IMPATTO: perche' il valore, fatto e documento. PROBABILITA: ...
# PROSSIMITA: ... VELOCITY: ... ADEGUAMENTO: ...", e il razionale di azioni e KRI come "perche'. EFFETTO: ...".
# Un'etichetta conta all'inizio del testo o dopo un punto, un punto e virgola o un a capo (maiuscole o minuscole),
# oppure ovunque se e' scritta tutta in maiuscolo: "Impatto resta 4:" a meta' frase non e' un'etichetta. I razionali
# scritti prima (senza etichette) restano interi.
ETICHETTE_RAZIONALE = {"RISCHIO": "rischio", "IMPATTO": "impatto", "PROBABILITA": "probabilita",
                       "PROSSIMITA": "prossimita", "VELOCITY": "velocity", "ADEGUAMENTO": "adeguamento"}
_NOMI = r"RISCHIO|IMPATTO|PROBABILIT(?:A'|À|A)|PROSSIMIT(?:A'|À|A)|VELOCITY|ADEGUAMENTO"
_ETICHETTA = re.compile(rf"(?:(?:^|(?<=[.;\n]))\s*(?i:({_NOMI}))|(?<=\s)({_NOMI}))\s*:")
_EFFETTO = re.compile(r"(?:(?:^|(?<=[.;\n]))\s*(?i:EFFETTO(?: ATTESO)?)|(?<=\s)EFFETTO(?: ATTESO)?)\s*:")


def _pulito(testo):
    """Parte di un razionale senza spazi e punti e virgola ai bordi, con la maiuscola iniziale ("RISCHIO: le frodi"
    diventa "Le frodi"): nella scheda ogni parte sta da sola, in una cella o in una sezione."""
    t = str(testo or "").strip().strip(";").strip()
    return (t[0].upper() + t[1:]) if t else None


def parti_razionale(testo):
    """Parti del razionale di un rischio: {"rischio", "impatto", "probabilita", "prossimita", "velocity",
    "adeguamento", "altro" (testo prima della prima etichetta)}; {} se il razionale non ha etichette."""
    testo = str(testo or "")
    trovate = list(_ETICHETTA.finditer(testo))
    if not trovate:
        return {}
    parti = {}
    if _pulito(testo[:trovate[0].start()]):
        parti["altro"] = _pulito(testo[:trovate[0].start()])
    for i, m in enumerate(trovate):
        chiave = ETICHETTE_RAZIONALE[(m.group(1) or m.group(2)).upper().replace("À", "A").replace("'", "")]
        fine = trovate[i + 1].start() if i + 1 < len(trovate) else len(testo)
        valore = _pulito(testo[m.end():fine])
        if valore:
            parti[chiave] = (parti[chiave] + " " + valore) if chiave in parti else valore
    return parti


def parti_effetto(testo):
    """(perche', effetto atteso) dal razionale di un'azione o di un KRI; effetto None se manca "EFFETTO:"."""
    testo = str(testo or "")
    m = _EFFETTO.search(testo)
    if not m:
        return _pulito(testo), None
    return _pulito(testo[:m.start()]), _pulito(testo[m.end():])


# ================================================================ 1. input di A6
def input_schede(scr, sessione, ciclo, ciclo_prec, data_rif):
    """Dati di ogni rischio del ciclo, dal registro VALIDATO e dalle viste (urgenza, segnali)."""
    _, anag = scr.leggi_anagrafica()
    posizione = {r["ID"]: i for i, (_, r) in enumerate(anag)}
    reg = scr.righe("tblRegister")
    del_ciclo = {r["ID_Rischio"]: r for r in reg if r["Ciclo"] == ciclo}
    cicli = sorted({r["Ciclo"] for r in reg if r["Ciclo"] <= ciclo})[-8:]
    azioni = scr.righe("tblAzioni")
    kri = scr.righe("tblKRI")
    kri_prec = {r["KRI_ID"]: r for r in kri if r["Ciclo"] == ciclo_prec}
    viste = {n: _colonna(sessione, n) for n in ("rmID", "rmX", "rmY", "rmZ", "rmMargine", "rmUrgenza", "rmLivello",
                                                  "rmMotivoLivello", "posNetti")}
    et_codici = _colonna(sessione, "codiciL1")
    et_nomi = [r[0] for r in sessione.intervallo_nome("codiciL1").GetOffset(0, 1).Value]
    nome_et = dict(zip(et_codici, et_nomi))
    pesi = {d: _num(sessione.intervallo_nome(n).Value) for d, n in PESI.items()}
    rischi = []
    for _, a in anag:
        rid = a["ID"]
        r = del_ciclo.get(rid)
        if r is None:
            continue                           # non valutato: nessuna scheda (il controllo di copertura lo segnala)
        i = posizione[rid]
        prec = next((x for x in reg if x["Ciclo"] == ciclo_prec and x["ID_Rischio"] == rid), None)
        parti = parti_razionale(r["Razionale"])
        rischi.append({
            "id_rischio": rid, "nome": r["Nome"], "event_type": r["Event_Type"],
            "event_type_nome": nome_et.get(r["Event_Type"]), "driver": r["Driver"], "owner": a["Owner"],
            "attenzione_supervisore": a.get("Attenzione_Supervisore") == "Sì",
            "anni_adeguamento": _num(a.get("Anni_Adeguamento")), "esito": r["Esito"],
            "presidio_controlli": a.get("Presidio_Controlli") or "Non verificato",
            "descrizione": r["Descrizione"],
            # D-68: il fenomeno in termini di rischio e, a parte, il perche' di ogni valore (null se A1 non ha diviso
            # il razionale: allora razionale_rischio e' il razionale intero).
            "razionale_rischio": (" ".join(p for p in (parti.get("altro"), parti.get("rischio")) if p) or None)
                                 if parti else r["Razionale"],
            "motivazioni": {k: parti.get(k) for k in DIMENSIONI + ["adeguamento"]} if parti else None,
            "valori": {d: _num(r[COLONNA_DIM[d]]) for d in DIMENSIONI},
            "score": _num(r["Score_Composito"]),
            "score_ciclo_precedente": _num(prec["Score_Composito"]) if prec else None,
            "urgenza": {"esito": viste["rmUrgenza"][i], "x": _num(viste["rmX"][i]), "y": _num(viste["rmY"][i]),
                        "z": _num(viste["rmZ"][i]), "margine": _num(viste["rmMargine"][i])},
            # D-63: livello di allerta calcolato dal foglio Radar_Mosca (colore del radar).
            "livello_allerta": {"livello": viste["rmLivello"][i] or None, "motivo": viste["rmMotivoLivello"][i] or None},
            "segnali_netti_4_settimane": _num(viste["posNetti"][i]),
            "variazioni": json.loads(r["Variazioni"] or "[]"),
            "fonti": json.loads(r["Fonti"] or "[]"),
            "azioni": [{"id_azione": x["ID_Azione"], "descrizione": x["Descrizione"], "owner": x["Owner"],
                        "orizzonte": x["Orizzonte"], "stato": x["Stato"], "ciclo": x["Ciclo"],
                        **dict(zip(("razionale", "effetto_atteso"), parti_effetto(x.get("Razionale"))))}
                       for x in azioni if x["ID_Rischio"] == rid and x.get("ID_Azione")
                       and (x["Ciclo"] == ciclo or x["Stato"] not in STATI_CHIUSI)],
            "kri": [{"kri_id": k["KRI_ID"], "tipo": k.get("Tipo") or None, "nome": k["Nome"], "unita": k["Unita"],
                     **dict(zip(("razionale", "effetto_atteso"), parti_effetto(k.get("Razionale")))),
                     "valore": _num(k["Valore_Corrente"]),
                     "valore_ciclo_precedente": _num((kri_prec.get(k["KRI_ID"]) or {}).get("Valore_Corrente")),
                     "soglia_attenzione": _num(k["Soglia_Attenzione"]), "soglia_allerta": _num(k["Soglia_Allerta"]),
                     "stato": k["Stato"], "fonte_dati": k["Fonte_Dati"]}
                    for k in kri if k["Ciclo"] == ciclo and k["ID_Rischio"] == rid and k["In_Vigore"] == "Sì"],
            "storia": [{"ciclo": c, **{d: _num(x[COLONNA_DIM[d]]) for d in DIMENSIONI}, "score": _num(x["Score_Composito"])}
                       for c in cicli for x in reg if x["Ciclo"] == c and x["ID_Rischio"] == rid],
        })
    rischi.sort(key=lambda x: (-(x["score"] or 0), posizione[x["id_rischio"]]))
    for n, x in enumerate(rischi, start=1):
        x["posizione_in_classifica"] = n
    return {"ciclo": ciclo, "data_riferimento": data_rif, "ciclo_precedente": ciclo_prec, "pesi": pesi,
            "rischi": rischi}


# ================================================================ 2. controlli sull'output di A6
NUMERO = re.compile(r"\d+(?:[.,]\d+)?")


def numeri(testo):
    """Numeri di un testo come valori (4,2 = 4.2; 06 = 6): cosi' 10/06/2026 e 2026-06-10 coincidono."""
    return {float(n.replace(",", ".")) for n in NUMERO.findall(str(testo))}


def controlla_a6(ingresso, uscita, testo_rubrica=""):
    """Problemi dell'output di A6 rispetto al suo input (lista vuota = va bene). Ogni problema
    ferma il comando: A6 scrive testi, non cambia rischi, valori o fonti."""
    problemi = []
    if uscita.get("ciclo") != ingresso["ciclo"]:
        problemi.append(f"Ciclo {uscita.get('ciclo')} diverso da quello dell'input ({ingresso['ciclo']}).")
    attesi = {r["id_rischio"]: r for r in ingresso["rischi"]}
    visti = [s["id_rischio"] for s in uscita.get("schede", [])]
    doppi = sorted({i for i in visti if visti.count(i) > 1})
    if doppi:
        problemi.append(f"Schede ripetute: {doppi}.")
    if set(visti) - set(attesi):
        problemi.append(f"Schede di rischi non nell'input: {sorted(set(visti) - set(attesi))}.")
    if set(attesi) - set(visti):
        problemi.append(f"Rischi senza scheda: {sorted(set(attesi) - set(visti))}.")
    ammessi_sempre = {1.0, 2.0, 3.0, 4.0, 5.0} | numeri(testo_rubrica) | numeri(ingresso["ciclo"]) | \
        numeri(ingresso["data_riferimento"])
    for s in uscita.get("schede", []):
        rid = s["id_rischio"]
        r = attesi.get(rid)
        if r is None:
            continue
        dims = [d["dimensione"] for d in s["dimensioni"]]
        if sorted(dims) != sorted(DIMENSIONI):
            problemi.append(f"{rid}: servono le quattro dimensioni una volta ciascuna, trovate {dims}.")
        rif_attesi = [f["riferimento"] for f in r["fonti"]]
        rif_visti = [f["riferimento"] for f in s["fonti"]]
        if sorted(rif_visti) != sorted(rif_attesi):
            mancanti = sorted(set(rif_attesi) - set(rif_visti))
            nuove = sorted(set(rif_visti) - set(rif_attesi))
            problemi.append(f"{rid}: le fonti devono essere quelle dell'input, copiate identiche, una volta ciascuna"
                            + (f"; mancano {mancanti}" if mancanti else "") + (f"; non nell'input {nuove}" if nuove else "")
                            + ("; ripetute" if len(rif_visti) != len(set(rif_visti)) else "") + ".")
        ammessi = ammessi_sempre | numeri(json.dumps(r, ensure_ascii=False))
        testi = [("descrizione_estesa", s["descrizione_estesa"])] + \
                [(d["dimensione"], d["spiegazione"]) for d in s["dimensioni"]] + \
                [(f"fonte '{f['riferimento'][:30]}'", f["a_sostegno_di"]) for f in s["fonti"]]
        for dove, testo in testi:
            estranei = sorted(n for n in numeri(testo) if n not in ammessi)
            if estranei:
                valori = ", ".join(f"{n:g}".replace(".", ",") for n in estranei)
                problemi.append(f"{rid}, {dove}: numeri che non stanno nei dati del rischio ({valori}). A6 non puo' "
                                "introdurre dati: correggere il testo o rilanciare A6.")
    return problemi


# ================================================================ 3. file delle schede
BLU_NOTTE, BLU_MEDIO, BLU_CHIARO, BLU_DIM = "1F3864", "2F5597", "DEEAF6", "2E75B6"
ROSSO, AMBRA, GRIGIO = "C00000", "FFC000", "A6A6A6"
TESTO, TESTO_TENUE = "262626", "595959"
ROSSO_TENUE, AMBRA_TENUE = "FBE5E5", "FFF2CC"
# Livello di allerta (D-63): gli stessi colori del radar.
COLORI_LIVELLO = {"Attivazione": ("E06666", "FFFFFF"), "Attenzione": ("F6C342", TESTO), "Monitoraggio": ("93C47D", TESTO)}
FONDI_LIVELLO = {"Attivazione": ROSSO_TENUE, "Attenzione": AMBRA_TENUE, "Monitoraggio": "E2EFDA"}
SCALA = {1: "F2F2F2", 2: "FCE4D6", 3: "F8CBAD", 4: "F4B183", 5: "C55A11"}
SCALA_RIL = {1: "F2F2F2", 2: "DEEBF7", 3: "9DC3E6"}
SEMAFORO = {"Rosso": (ROSSO, "FFFFFF"), "Ambra": (AMBRA, TESTO), "Verde": ("70AD47", "FFFFFF"),
            "Non rilevato": ("D9D9D9", TESTO), "Soglie da definire": (AMBRA_TENUE, TESTO),
            "SOGLIE INVERTITE — correggere": (ROSSO, "FFFFFF")}
LEGENDA = ("Imp = Impatto, Prob = Probabilita, Pross = Prossimita, Vel = Velocity (1-5, 5 = peggiore; n.v. = non valutata, "
           "nessuna fonte ne documenta i tempi). Rilevanza di una fonte per la previsione: Alta = fonte primaria o caso "
           "documentato, Media = fonte specializzata, Bassa = generalista. Presidio = "
           "giudizio della Banca sui controlli, fuori dallo score. X = anni per decidere e avviare la mitigazione; "
           "Y = anni per mitigare (azioni o adeguamento strutturale); Z = anni prima che il rischio possa manifestarsi; "
           "AGIRE ORA = X + Y > Z. Livello di allerta (calcolato dal registro): Attivazione = almeno due fra AGIRE ORA, "
           "KRI in rosso, notizie in peggioramento; Attenzione = uno di questi, o margine sotto l'anno con P x I alto, "
           "presidio assente o KRI in ambra; Monitoraggio = gli altri. KRI = indicatore di rischio, di esposizione (quanto "
           "la Banca e' esposta) o di evoluzione (quanto cresce il fenomeno). ◆ = tema all'attenzione del supervisore.")


def _it(v, decimali=2):
    if v is None or v == "":
        return "—"
    if isinstance(v, (int, float)):
        return f"{v:.{decimali}f}".replace(".", ",")
    return str(v)


def _data_it(s):
    s = str(s or "")
    return f"{s[8:10]}/{s[5:7]}/{s[:4]}" if re.match(r"\d{4}-\d{2}-\d{2}", s) else s


def _urgenza_breve(testo):
    t = str(testo or "")
    return "AGIRE ORA" if t.startswith("AGIRE ORA") else "Entro l'anno" if t.startswith("Agire entro") else t


def cosa_cambiato(r):
    if r["esito"] == "Nuovo":
        return "Rischio nuovo nel ciclo."
    if not r["variazioni"]:
        return "Nessun valore cambiato: confermato sulle fonti del ciclo precedente."
    return "\n".join(f"• {ETICHETTE_VARIAZIONI.get(v['dimensione'], v['dimensione'])} {_nv(v.get('da'))} → {_nv(v.get('a'))}: "
                     f"{v['evidenza']} ({_data_it(v.get('data_evidenza'))})" for v in r["variazioni"])


def _nv(v):
    return NON_VALUTATA if v is None else v


def scrivi_file_schede(ingresso, uscita_a6, percorso, validazione=""):
    """Schede_<ciclo>.xlsx: indice e una scheda per rischio. uscita_a6 = output di A6 gia'
    controllato, o None (testi di A1 dal registro)."""
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    fill = lambda c: PatternFill("solid", fgColor=c, bgColor=c)
    sottile = Border(bottom=Side(style="thin", color="D9D9D9"))
    testi = {s["id_rischio"]: s for s in (uscita_a6 or {}).get("schede", [])}
    ciclo = ingresso["ciclo"]
    origine = "testi estesi dell'agente A6" if uscita_a6 else "testi di A1 dal registro (senza A6)"
    intestazione = f"{NOME_ANALISI} — schede dei rischi · ciclo {ciclo} · dati al {_data_it(ingresso['data_riferimento'])} " \
                   f"· {validazione or 'dati validati'} · {RISERVATEZZA}"

    def cella(ws, r, c, v=None, **stile):
        x = ws.cell(row=r, column=c, value=v)
        for k, val in stile.items():
            setattr(x, k, val)
        return x

    def stampa(ws, orizzontale=False):
        ws.page_setup.orientation = "landscape" if orizzontale else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.oddHeader.center.text = RISERVATEZZA
        ws.oddHeader.right.text = NOME_ANALISI
        ws.oddFooter.left.text = "&F · &A"
        ws.oddFooter.right.text = "Pagina &P di &N"

    def chip(c, urgenza):
        u = _urgenza_breve(urgenza)
        colori = {"AGIRE ORA": (ROSSO, "FFFFFF"), "Entro l'anno": (AMBRA, TESTO)}.get(u, ("D9D9D9", TESTO))
        c.value = u
        c.fill = fill(colori[0])
        c.font = Font(size=9, bold=True, color=colori[1])
        c.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)

    def chip_livello(c, livello):
        colori = COLORI_LIVELLO.get(livello, ("D9D9D9", TESTO))
        c.value = livello or "—"
        c.fill = fill(colori[0])
        c.font = Font(size=9, bold=True, color=colori[1])
        c.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)

    def righe_testo(testo, larghezza, sz=10):
        per_riga = max(1, int(larghezza * 1.3 * 10 / sz))
        return sum(max(1, -(-len(p) // per_riga)) for p in str(testo).split("\n"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Indice"
    ws.sheet_view.showGridLines = False
    for l, w in zip("ABCDEFGH", [6, 10, 56, 9, 15, 14, 13, 30]):
        ws.column_dimensions[l].width = w
    ws.merge_cells("A1:H1")
    cella(ws, 1, 1, intestazione, font=Font(size=13, bold=True, color="FFFFFF"), fill=fill(BLU_NOTTE),
          alignment=Alignment(vertical="center", indent=1, wrap_text=True))
    ws.row_dimensions[1].height = 36
    ws.merge_cells("A2:H2")
    cella(ws, 2, 1, "Una scheda per rischio, nell'ordine della classifica (score composito), con i valori VALIDATI del "
                    f"registro. Testi: {origine}; numeri, fonti e urgenza vengono dal registro e non li cambia nessun agente. "
                    "Clic sull'ID per aprire la scheda. Prodotto dal comando  python esegui_ciclo_locale.py schede.",
          font=Font(size=9.5, color=TESTO), fill=fill("DDEBF7"), alignment=Alignment(wrap_text=True, vertical="top", indent=1))
    ws.row_dimensions[2].height = 42
    for j, t in enumerate(["#", "ID", "Rischio", "Score", "Urgenza", "Livello", "Esito", "Owner"], start=1):
        cella(ws, 4, j, t, font=Font(bold=True, color="FFFFFF"), fill=fill(BLU_NOTTE))
    for r in ingresso["rischi"]:
        riga = 4 + r["posizione_in_classifica"]
        cella(ws, riga, 1, r["posizione_in_classifica"], font=Font(bold=True, color=BLU_NOTTE))
        c = cella(ws, riga, 2, r["id_rischio"], font=Font(bold=True, color=BLU_DIM, underline="single"))
        c.hyperlink = f"#'{r['id_rischio']}'!A1"
        cella(ws, riga, 3, r["nome"] + ("  ◆" if r["attenzione_supervisore"] else ""))
        cella(ws, riga, 4, r["score"], number_format="0.00", font=Font(bold=True))
        chip(ws.cell(row=riga, column=5), r["urgenza"]["esito"])
        chip_livello(ws.cell(row=riga, column=6), (r.get("livello_allerta") or {}).get("livello"))
        cella(ws, riga, 7, r["esito"])
        cella(ws, riga, 8, r["owner"], font=Font(size=9))
        for x in range(1, 9):
            ws.cell(row=riga, column=x).border = sottile
        ws.row_dimensions[riga].height = 20
    fine = 5 + len(ingresso["rischi"])
    ws.merge_cells(start_row=fine + 1, start_column=1, end_row=fine + 1, end_column=8)
    cella(ws, fine + 1, 1, "Legenda — " + LEGENDA, font=Font(size=8, italic=True, color=TESTO_TENUE),
          alignment=Alignment(wrap_text=True, vertical="top"))
    ws.row_dimensions[fine + 1].height = righe_testo("Legenda — " + LEGENDA, 153, 8) * 11 + 6
    stampa(ws, orizzontale=True)

    for r in ingresso["rischi"]:
        t = testi.get(r["id_rischio"])
        ws = wb.create_sheet(r["id_rischio"])
        ws.sheet_view.showGridLines = False
        larghezze = [18, 12, 16, 16, 16, 16, 16, 34]
        for j, w in enumerate(larghezze, start=1):
            ws.column_dimensions[chr(64 + j)].width = w
        tot = sum(larghezze)
        ws.merge_cells("A1:H1")
        cella(ws, 1, 1, intestazione, font=Font(size=9, color="FFFFFF"), fill=fill(BLU_MEDIO),
              alignment=Alignment(vertical="center", indent=1))
        ws.merge_cells("A2:F2")
        ws.merge_cells("G2:H2")
        cella(ws, 2, 1, f"{r['posizione_in_classifica']} · {r['id_rischio']} · {r['nome']}",
              font=Font(size=14, bold=True, color="FFFFFF"), alignment=Alignment(vertical="center", indent=1, wrap_text=True))
        d = None if r["score_ciclo_precedente"] is None else round(r["score"] - r["score_ciclo_precedente"], 2)
        freccia = "nuovo" if d is None else ("▲ +" if d > 0 else "▼ " if d < 0 else "● ") + _it(d)
        liv = (r.get("livello_allerta") or {}).get("livello")
        cella(ws, 2, 7, f"score {_it(r['score'])}   {freccia}   ·   {_urgenza_breve(r['urgenza']['esito'])}"
                        + (f"   ·   {liv}" if liv else ""),
              font=Font(size=12, bold=True, color="FFFFFF"), alignment=Alignment(horizontal="right", vertical="center"))
        for j in range(1, 9):
            ws.cell(row=2, column=j).fill = fill(BLU_NOTTE)
        ws.row_dimensions[2].height = 34
        etichette = [((1, 2), f"{r['event_type']} · {r['event_type_nome'] or ''}"), ((3, 4), f"Driver: {r['driver'] or '—'}"),
                     ((5, 6), f"Owner: {r['owner'] or '—'} · Presidio: {r.get('presidio_controlli') or '—'}"),
                     ((7, 7), "◆ Attenzione del supervisore" if r["attenzione_supervisore"] else ""),
                     ((8, 8), f"Esito: {r['esito']} · validato")]
        for (da, a), testo in etichette:
            if da != a:
                ws.merge_cells(start_row=3, start_column=da, end_row=3, end_column=a)
            sup = testo.startswith("◆")
            for x in range(da, a + 1):
                ws.cell(row=3, column=x).fill = fill(AMBRA_TENUE if sup else BLU_CHIARO)
            cella(ws, 3, da, testo or None, font=Font(size=9, bold=sup, color="7F6000" if sup else TESTO_TENUE),
                  alignment=Alignment(vertical="center", wrap_text=True))
        ws.row_dimensions[3].height = 28
        riga = 5

        def sezione(testo):
            nonlocal riga
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=8)
            cella(ws, riga, 1, testo, font=Font(size=12, bold=True, color=BLU_NOTTE))
            for x in range(1, 9):
                ws.cell(row=riga, column=x).border = Border(bottom=Side(style="medium", color=BLU_NOTTE))
            riga += 1

        def testo_lungo(testo, sz=10, fondo=None):
            nonlocal riga
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=8)
            c = cella(ws, riga, 1, testo, font=Font(size=sz), alignment=Alignment(wrap_text=True, vertical="top"))
            if fondo:
                c.fill = fill(fondo)
            ws.row_dimensions[riga].height = min(409, righe_testo(testo, tot, sz) * sz * 1.35 + 6)
            riga += 1

        # D-68: prima il fenomeno (Il rischio, poi il Razionale in termini di rischio), poi i valori, ognuno con il
        # suo perche'. Con A6 la descrizione estesa e' gia' il fenomeno in termini di rischio: niente Razionale a parte.
        motivazioni = r.get("motivazioni") or {}
        sezione("Il rischio")
        testo_lungo(t["descrizione_estesa"] if t else r["descrizione"])
        if not t:
            riga += 1
            sezione("Razionale: il fenomeno in termini di rischio" if motivazioni else "Razionale (registro)")
            testo_lungo(r.get("razionale_rischio", r.get("razionale")) or "—")
        riga += 1
        sezione("Valutazione, dimensione per dimensione (valori validati)")
        for j, testo in ((1, "Dimensione"), (2, "Valore"), (3, "Perche' questo valore: informazioni e documenti")):
            cella(ws, riga, j, testo, font=Font(size=9, bold=True, color=TESTO_TENUE))
        ws.merge_cells(start_row=riga, start_column=3, end_row=riga, end_column=8)
        riga += 1
        spiegazioni = {x["dimensione"]: x["spiegazione"] for x in t["dimensioni"]} if t else motivazioni
        for dim in DIMENSIONI:
            v = r["valori"][dim]
            cella(ws, riga, 1, NOMI_DIM[dim], font=Font(size=10, bold=True, color=BLU_NOTTE), alignment=Alignment(vertical="center"))
            colore = SCALA.get(v, "FFFFFF")
            cella(ws, riga, 2, NON_VALUTATA if v is None else v, fill=fill(colore),
                  font=Font(size=14 if v is not None else 10, bold=True, color="FFFFFF" if v == 5 else TESTO),
                  alignment=Alignment(horizontal="center", vertical="center"))
            ws.merge_cells(start_row=riga, start_column=3, end_row=riga, end_column=8)
            sp = spiegazioni.get(dim) or ""
            if sp:
                cella(ws, riga, 3, sp, font=Font(size=10), alignment=Alignment(wrap_text=True, vertical="center"))
            else:
                cella(ws, riga, 3, "Motivo non diviso per dimensione nel razionale di A1: vedi il Razionale qui sopra.",
                      font=Font(size=9, italic=True, color=TESTO_TENUE), alignment=Alignment(vertical="center"))
            ws.row_dimensions[riga].height = max(26, righe_testo(sp, tot - 30) * 13.5 + 6) if sp else 22
            for x in range(1, 9):
                ws.cell(row=riga, column=x).border = sottile
            riga += 1
        pesi = ingresso["pesi"]
        presenti = [d for d in PESI if r["valori"].get(d) is not None]
        formula = " + ".join(f"{_it(pesi[d])} x {r['valori'][d]}" for d in presenti)
        somma = _it(sum(pesi[d] for d in presenti if isinstance(pesi.get(d), (int, float))))
        nota_vel = "" if "velocity" in presenti else " Velocity non valutata: i pesi si ridistribuiscono sulle altre tre."
        ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=8)
        cella(ws, riga, 1, f"Score composito {_it(r['score'])} = ({formula}) / {somma}.{nota_vel} La rilevanza delle "
                           "fonti e' nella tabella Fonti.", font=Font(size=8, italic=True, color=TESTO_TENUE))
        riga += 2
        u = r["urgenza"]
        sezione("Urgenza (disuguaglianza di Mosca) e livello di allerta")
        la = r.get("livello_allerta") or {}
        if la.get("livello"):
            testo_lungo(f"Livello di allerta: {la['livello']}" + (f" — {la['motivo']}." if la.get("motivo") else "."),
                        fondo=FONDI_LIVELLO.get(la["livello"]))
        perche_adeg = motivazioni.get("adeguamento")
        adeg = (f" (con {_it(r['anni_adeguamento'], 1)} anni di adeguamento strutturale, stimati da A1 e validati"
                + (f"; motivo: {perche_adeg[0].lower() + perche_adeg[1:].rstrip('.')}" if perche_adeg else "") + ")"
                if r["anni_adeguamento"] else "")
        if isinstance(u["x"], (int, float)) and isinstance(u["y"], (int, float)):
            testo_lungo(f"{u['esito']}. Tempo necessario X + Y = {_it(u['x'])} + {_it(u['y'])}{adeg} = "
                        f"{_it(u['x'] + u['y'])} anni; tempo disponibile Z = {_it(u['z'], 1)} anni (Prossimita "
                        f"{r['valori']['prossimita']}). Margine Z - (X + Y) = {_it(u['margine'])}.",
                        fondo=ROSSO_TENUE if _urgenza_breve(u["esito"]) == "AGIRE ORA" else None)
        else:
            testo_lungo("Urgenza non calcolabile: manca un valore del ciclo.")
        riga += 1
        sezione("Cosa e' cambiato nel ciclo")
        testo_lungo(cosa_cambiato(r))
        riga += 1
        sezione("Andamento negli ultimi trimestri")
        intest = ["Ciclo", "Imp", "Prob", "Pross", "Vel", "Score"]
        for j, testo in enumerate(intest, start=1):
            cella(ws, riga, j, testo, font=Font(size=9, bold=True, color=TESTO_TENUE), alignment=Alignment(horizontal="center"))
        prima = riga
        for s in r["storia"]:
            riga += 1
            cella(ws, riga, 1, s["ciclo"], font=Font(bold=True, color=BLU_NOTTE), alignment=Alignment(horizontal="center"))
            for j, dim in enumerate(DIMENSIONI, start=2):
                v = s.get(dim)
                cella(ws, riga, j, NON_VALUTATA if v is None else v, fill=fill(SCALA.get(v, "FFFFFF")),
                      font=Font(bold=True, color="FFFFFF" if v == 5 else TESTO), alignment=Alignment(horizontal="center"))
            cella(ws, riga, 6, s["score"], number_format="0.00", font=Font(bold=True), alignment=Alignment(horizontal="center"))
        if len(r["storia"]) >= 2:
            graf = LineChart()
            graf.add_data(Reference(ws, min_col=6, min_row=prima, max_row=riga), titles_from_data=True)
            graf.set_categories(Reference(ws, min_col=1, min_row=prima + 1, max_row=riga))
            s = graf.series[0]
            s.smooth = False
            s.marker.symbol = "circle"
            s.marker.size = 7
            s.graphicalProperties.line.solidFill = BLU_NOTTE
            s.marker.graphicalProperties = GraphicalProperties(solidFill=BLU_NOTTE)
            s.dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False,
                                    showPercent=False, dLblPos="t")
            graf.y_axis.scaling.min, graf.y_axis.scaling.max, graf.y_axis.majorUnit = 1, 5, 1
            graf.x_axis.delete = graf.y_axis.delete = False
            graf.legend = None
            graf.title = "Score composito per trimestre (scala fissa 1-5)"
            graf.title.overlay = False
            graf.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
            graf.height, graf.width = 5.5, 11
            ws.add_chart(graf, f"A{riga + 2}")
            riga += 14
        else:
            riga += 1
            cella(ws, riga, 1, f"Rischio entrato nel registro nel {r['storia'][0]['ciclo'] if r['storia'] else ciclo}: "
                               "l'andamento si vedra' dal prossimo ciclo.", font=Font(size=9, italic=True, color=TESTO_TENUE))
            riga += 1
        riga += 1
        sezione("Fonti")
        for j, testo in zip((1, 4, 5, 6, 7), ["Fonte", "Data", "Tipo", "Rilevanza", "A sostegno di" if t else ""]):
            cella(ws, riga, j, testo or None, font=Font(size=9, bold=True, color=TESTO_TENUE))
        ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=3)
        ws.merge_cells(start_row=riga, start_column=7, end_row=riga, end_column=8)
        riga += 1
        sostegno = {f["riferimento"]: f["a_sostegno_di"] for f in t["fonti"]} if t else {}
        for f in r["fonti"]:
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=3)
            ws.merge_cells(start_row=riga, start_column=7, end_row=riga, end_column=8)
            cella(ws, riga, 1, f["riferimento"], font=Font(size=10), alignment=Alignment(wrap_text=True, vertical="top"))
            cella(ws, riga, 4, _data_it(f.get("data")), font=Font(size=10), alignment=Alignment(vertical="top"))
            cella(ws, riga, 5, f.get("tipo"), font=Font(size=10, bold=True, color=BLU_DIM if f.get("tipo") == "Interna" else TESTO_TENUE),
                  alignment=Alignment(vertical="top"))
            ril = f.get("rilevanza")
            cella(ws, riga, 6, RILEVANZA.get(ril, ril), fill=fill(SCALA_RIL.get(ril, "FFFFFF")), font=Font(size=10, bold=True),
                  alignment=Alignment(horizontal="center", vertical="top"))
            so = sostegno.get(f["riferimento"])
            cella(ws, riga, 7, so, font=Font(size=10), alignment=Alignment(wrap_text=True, vertical="top"))
            for x in range(1, 9):
                ws.cell(row=riga, column=x).border = sottile
            ws.row_dimensions[riga].height = max(righe_testo(f["riferimento"], 44), righe_testo(so or "", 44)) * 13.5 + 5
            riga += 1
        riga += 1

        def perche_ed_effetto(perche, effetto, etichetta_perche):
            """Riga sotto un'azione o un KRI (D-68): a sinistra il perche', a destra l'effetto atteso (A1)."""
            nonlocal riga
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=4)
            ws.merge_cells(start_row=riga, start_column=5, end_row=riga, end_column=8)
            sx = f"{etichetta_perche}: {perche or '—'}"
            dx = f"Effetto atteso: {effetto or 'non indicato nel razionale di A1.'}"
            for col, testo in ((1, sx), (5, dx)):
                cella(ws, riga, col, testo, font=Font(size=9.5, color=TESTO), fill=fill("F7F9FC"),
                      alignment=Alignment(wrap_text=True, vertical="top", indent=1))
            for x in range(1, 9):
                ws.cell(row=riga, column=x).fill = fill("F7F9FC")
                ws.cell(row=riga, column=x).border = sottile
            ws.row_dimensions[riga].height = max(righe_testo(sx, sum(larghezze[:4]) - 2, 9.5),
                                                 righe_testo(dx, sum(larghezze[4:]) - 2, 9.5)) * 13 + 6
            riga += 1

        sezione("Azioni")
        if not r["azioni"]:
            testo_lungo("Nessuna azione del ciclo o aperta.")
        else:
            for j, testo in zip((1, 6, 8), ["Azione", "Owner", "Orizzonte · stato"]):
                cella(ws, riga, j, testo, font=Font(size=9, bold=True, color=TESTO_TENUE))
            riga += 1
        for a in r["azioni"]:
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=5)
            ws.merge_cells(start_row=riga, start_column=6, end_row=riga, end_column=7)
            cella(ws, riga, 1, f"• {a['descrizione']}", font=Font(size=10, bold=True, color=TESTO),
                  alignment=Alignment(wrap_text=True, vertical="top"))
            cella(ws, riga, 6, a["owner"], font=Font(size=9.5), alignment=Alignment(wrap_text=True, vertical="top"))
            cella(ws, riga, 8, f"{a['orizzonte']} · {a['stato']} (ciclo {a['ciclo']})", font=Font(size=9.5),
                  alignment=Alignment(wrap_text=True, vertical="top"))
            ws.row_dimensions[riga].height = max(righe_testo(a["descrizione"], sum(larghezze[:5])),
                                                 righe_testo(a["owner"] or "", sum(larghezze[5:7]), 9.5)) * 13.5 + 5
            riga += 1
            perche_ed_effetto(a.get("razionale"), a.get("effetto_atteso"), "Perche'")
        riga += 1
        sezione("KRI in vigore")
        for j, testo in zip((1, 5, 6, 7, 8), ["Indicatore", "Valore", "Soglie (attenzione / allerta)", "Ciclo prima", "Stato"]):
            cella(ws, riga, j, testo, font=Font(size=9, bold=True, color=TESTO_TENUE), alignment=Alignment(wrap_text=True))
        riga += 1
        for k in r["kri"]:
            ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=4)
            tipo = f"[{k['tipo']}] " if k.get("tipo") else ""
            cella(ws, riga, 1, f"{tipo}{k['nome']} ({k['unita']})", font=Font(size=10),
                  alignment=Alignment(wrap_text=True, vertical="center"))
            cella(ws, riga, 5, k["valore"] if k["valore"] is not None else "—", font=Font(size=10, bold=True),
                  alignment=Alignment(horizontal="center", vertical="center"))
            cella(ws, riga, 6, f"{_it(k['soglia_attenzione'], 1)} / {_it(k['soglia_allerta'], 1)}", font=Font(size=10),
                  alignment=Alignment(horizontal="center", vertical="center"))
            cella(ws, riga, 7, k["valore_ciclo_precedente"] if k["valore_ciclo_precedente"] is not None else "—",
                  font=Font(size=10), alignment=Alignment(horizontal="center", vertical="center"))
            f_, t_ = SEMAFORO.get(k["stato"], ("D9D9D9", TESTO))
            cella(ws, riga, 8, k["stato"], fill=fill(f_), font=Font(bold=True, color=t_), alignment=Alignment(horizontal="center", vertical="center"))
            ws.row_dimensions[riga].height = max(20, righe_testo(k["nome"], 60) * 13.5 + 5)
            riga += 1
            perche_ed_effetto(k.get("razionale"), k.get("effetto_atteso"), "Perche' lo misuriamo")
        riga += 1
        ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=8)
        cella(ws, riga, 1, "Legenda — " + LEGENDA, font=Font(size=8, italic=True, color=TESTO_TENUE),
              alignment=Alignment(wrap_text=True, vertical="top"))
        ws.row_dimensions[riga].height = righe_testo("Legenda — " + LEGENDA, tot, 8) * 11 + 6   # 36 fisso la tagliava
        stampa(ws)
    Path(percorso).parent.mkdir(parents=True, exist_ok=True)
    wb.save(percorso)
    return percorso
''',
    "input_a1": r'''"""Allegato di A1: baseline, controlli di A5 e guida di scrittura in un file solo.

D-69 (02/10/2026). Le istruzioni di A1 sono al limite degli 8.000 caratteri:
i controlli che A5 applichera' al suo output (01_Kernel/controlli_A5.md) e la
guida di scrittura (01_Kernel/guida_scrittura_A1.md) viaggiano nell'allegato
che A1 riceve comunque a ogni ciclo, accanto alla baseline. Un file solo, un
solo gesto per chi lancia A1; testo semplice, che gli allegati e la knowledge
di Agent Builder leggono (i .json la knowledge non li legge).

Un solo lettore dei controlli per tutti: questo modulo (allegato di A1),
02_Agenti/componi_prompt.py (segnaposto {{CONTROLLI}} nel prompt di A5) e il
collaudo. Dal 05/10/2026 (D-70) e' anche dentro lo script della skill di
Copilot Studio, che legge la baseline dall'allegato (leggi_baseline). Solo
libreria standard.
"""

import json
import re
from pathlib import Path

FILE_CONTROLLI = "controlli_A5.md"
FILE_GUIDA = "guida_scrittura_A1.md"
PREFISSO = "Input_A1"
INIZIO_BASELINE = "=== 1. BASELINE ==="


def testo_kernel(percorso):
    """Il file senza le note (righe ">") e senza il titolo ("# ")."""
    righe = [r.rstrip() for r in Path(percorso).read_text(encoding="utf-8").splitlines()
             if not r.startswith(">") and not r.startswith("# ")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(righe)).strip()


def testo_controlli(kernel):
    return testo_kernel(Path(kernel) / FILE_CONTROLLI)


def nome_file(baseline):
    ciclo = baseline.get("ciclo_da_valutare")
    return f"{PREFISSO}_{ciclo}.txt" if ciclo else f"{PREFISSO}.txt"


def componi(baseline, kernel, oggi):
    """(testo dell'allegato, file del kernel mancanti). Un file mancante non
    ferma niente, perche' la baseline basta ad A1 per lavorare, ma va detto:
    senza controlli e guida A1 risponde come prima del 02/10/2026."""
    kernel = Path(kernel)
    mancanti = [n for n in (FILE_CONTROLLI, FILE_GUIDA) if not (kernel / n).exists()]
    ciclo = baseline.get("ciclo_da_valutare") or "da valutare"
    parti = [f"INPUT DI A1 PER IL CICLO {ciclo}: scritto da esegui_ciclo_locale.py il {oggi}, da non modificare a mano.",
             "Tre parti: 1. BASELINE (JSON); 2. CONTROLLI che lo script (e A5 Challenger, se usato) applicano al tuo "
             "output; 3. GUIDA DI SCRITTURA.",
             "", INIZIO_BASELINE, json.dumps(baseline, ensure_ascii=False, indent=2)]
    s = baseline.get("straordinario")
    if s:                                            # D-75: in testa, perche' cambia che cosa A1 deve scrivere
        cosa = (f"Valuta SOLO questi rischi: {', '.join(s['rischi_da_rivalutare'])} (e un rischio nuovo solo se l'evento "
                "lo crea)" if s["rischi_da_rivalutare"] else
                "Nessun rischio da rivalutare: censisci SOLO i rischi nuovi che l'evento crea (almeno uno)")
        parti[2:2] = ["", f"CICLO STRAORDINARIO {ciclo}, deciso da {s['decisore']} il {s['data_decisione']}: {s['motivo']}",
                      f"{cosa}; gli altri restano com'erano. Scrivi anche analisi_evento. Regole nella guida di scrittura, "
                      "sezione \"Ciclo straordinario\"."]
    if FILE_CONTROLLI not in mancanti:
        parti += ["", "=== 2. CONTROLLI ===",
                  "Prima di rispondere verifica ogni rischio con questi controlli e correggi. Li applicano lo script e, se "
                  "usato, A5 Challenger: un difetto di completezza o di coerenza (C1, C2, C7, identificativi dei KRI) "
                  "ferma il ciclo e torna a te; ogni altro rilievo arriva al validatore, che dovra' motivarlo uno per uno.",
                  testo_controlli(kernel)]
    if FILE_GUIDA not in mancanti:
        parti += ["", "=== 3. GUIDA DI SCRITTURA ===", testo_kernel(kernel / FILE_GUIDA)]
    return "\n".join(parti) + "\n", mancanti


def leggi_baseline(testo):
    """La baseline dal testo di un allegato Input_A1 (il JSON della parte 1): la skill di Copilot Studio riceve solo
    l'allegato (D-70). ValueError se il testo non e' un allegato di A1."""
    inizio = testo.find(INIZIO_BASELINE)
    graffa = testo.find("{", inizio) if inizio >= 0 else -1
    if graffa < 0:
        raise ValueError(f"non e' un allegato {PREFISSO}: manca la parte \"{INIZIO_BASELINE}\" con il JSON della baseline")
    return json.JSONDecoder().raw_decode(testo, graffa)[0]


# AU7 (23/10/2026): A3 e A4 ricevono la baseline come allegato della chat, come A1, e non piu' dal file
# Baseline_corrente.txt nella knowledge (indicizzazione lenta, file singolo saltato: "Baseline mancante" del 13/10).
ALLEGATI_BASELINE = {"A3": "Input_A3.txt", "A4": "Input_A4.txt"}


def componi_baseline(baseline, agente, oggi):
    """Allegato di A3 o A4: intestazione con ciclo e data di riferimento, poi la baseline (JSON)."""
    prec = baseline.get("ciclo_precedente") or "nessuno"
    data = baseline.get("data_riferimento_precedente") or "n.d."
    return (f"INPUT DI {agente}: baseline del registro NFR, ciclo validato {prec} (data di riferimento {data}), scritta da "
            f"esegui_ciclo_locale.py il {oggi}, da non modificare a mano. Rischi del radar in \"rischi\", in osservazione "
            "in \"rischi_in_osservazione\".\n\n"
            f"{INIZIO_BASELINE}\n{json.dumps(baseline, ensure_ascii=False, indent=2)}\n")


def scrivi_allegati_baseline(cartella, baseline, oggi):
    """Scrive Input_A3.txt e Input_A4.txt accanto a Baseline_corrente; restituisce i percorsi."""
    out = []
    for agente, nome in ALLEGATI_BASELINE.items():
        percorso = Path(cartella) / nome
        percorso.write_text(componi_baseline(baseline, agente, oggi), encoding="utf-8")
        out.append(percorso)
    return out


def scrivi(cartella, baseline, kernel, oggi):
    """Scrive l'allegato accanto a Baseline_corrente e toglie quelli dei cicli
    passati (file generati, rigenerabili: uno solo in cartella, quello giusto).
    Restituisce (percorso, file del kernel mancanti)."""
    cartella = Path(cartella)
    testo, mancanti = componi(baseline, kernel, oggi)
    percorso = cartella / nome_file(baseline)
    percorso.write_text(testo, encoding="utf-8")
    for vecchio in cartella.glob(f"{PREFISSO}*.txt"):
        if vecchio != percorso and re.fullmatch(rf"{PREFISSO}(_\d{{4}}Q[1-4](-S[1-9])?)?\.txt", vecchio.name):
            try:
                vecchio.unlink()
            except OSError:
                pass                    # aperto da qualcuno: resta, il ciclo nel nome dice che e' vecchio
    return percorso, mancanti
''',
    "precontrollo": r'''"""Precontrollo dell'output di A1 prima di A5 (D-69): cio' che lo script verifica senza Excel.

Schema, controlli meccanizzati (gli stessi di trimestrale, controlli_a5.py) e parti del razionale che servono alle
schede (D-68), con il messaggio di correzione da dare ad A1. Solo libreria standard e moduli del pacchetto che non
usano Excel: lo usano il comando precontrollo di esegui_ciclo_locale.py e, dal 05/10/2026 (D-70), lo script della
skill per Copilot Studio, che 09_Copilot_Studio/prepara_kit.py compone copiando senza modifiche questo modulo e quelli
che importa. Un solo codice: quello che il precontrollo trova in Copilot Studio, il comando lo trova sul PC.
"""

import hashlib
import json
import re

import controlli_a5
import schede
from valida_output import valida


def precontrolla(payload, baseline, schema):
    """Cio' che lo script verifica dell'output di A1 prima di A5: schema, controlli
    meccanizzati (gli stessi di trimestrale) e parti del razionale che servono alle
    schede (D-68). Con lo schema violato i controlli non girano: lavorano su campi
    che lo schema garantisce."""
    errori = valida(payload, schema)
    if errori:
        return {"schema": errori, "rilievi": [], "formato": []}
    formato = []                                     # una riga per rischio: il messaggio resta corto anche con A1 vecchio
    for x in payload["rischi"]:
        parti = schede.parti_razionale(x["razionale"])
        attese = [e for e, k in schede.ETICHETTE_RAZIONALE.items()
                  if k != "adeguamento" or x.get("anni_adeguamento") is not None]
        mancano = [e for e in attese if not parti.get(schede.ETICHETTE_RAZIONALE[e])]
        azioni = [str(n) for n, a in enumerate(x["azioni_mitigazione"], start=1) if schede.parti_effetto(a["razionale"])[1] is None]
        kri = [k["tipo"] for k in x["kri_proposti"] if schede.parti_effetto(k["razionale"])[1] is None]
        senza = ([("azione " if len(azioni) == 1 else "azioni ") + ", ".join(azioni)] if azioni else []) + \
                (["KRI di " + " e di ".join(kri)] if kri else [])
        problemi = ((["razionale senza le parti " + ", ".join(e + ":" for e in mancano)] if mancano else [])
                    + (["EFFETTO: mancante nel razionale di " + "; ".join(senza)] if senza else []))
        if problemi:
            formato.append({"id_rischio": x["id_rischio"], "descrizione": "; ".join(problemi) + " (parte 3 dell'allegato)."})
    return {"schema": errori, "rilievi": controlli_a5.controlla(payload, baseline)["rilievi"], "formato": formato}


def conta(esito):
    """(punti da correggere, punti da verificare): i primi fermerebbero il ciclo."""
    da_correggere = len(esito["schema"]) + sum(1 for x in esito["rilievi"] if x["gravita"] == "Bloccante")
    da_verificare = sum(1 for x in esito["rilievi"] if x["gravita"] != "Bloccante") + len(esito["formato"])
    return da_correggere, da_verificare


def ciclo_di(payload, baseline):
    """Il ciclo dell'output di A1; se manca o non e' un ciclo, quello da valutare della baseline."""
    ciclo = str(payload.get("ciclo") or "") if isinstance(payload, dict) else ""
    return ciclo if re.fullmatch(r"\d{4}Q[1-4](-S[1-9])?", ciclo) else (baseline.get("ciclo_da_valutare") or "senza_ciclo")


def agente(ciclo):
    """Chi ha scritto l'output: A1 per un trimestre, A8 Evento critico per uno straordinario AAAAQn-Sk (D-77)."""
    return "A8" if "-S" in str(ciclo) else "A1"


def riga_esito(ciclo, esito):
    """Prima riga dell'esito, uguale nel comando sul PC e nella skill di Copilot Studio."""
    da_correggere, da_verificare = conta(esito)
    if not da_correggere and not da_verificare:
        return f"Precontrollo dell'output di {agente(ciclo)} ({ciclo}): nessun rilievo."
    return (f"Precontrollo dell'output di {agente(ciclo)} ({ciclo}): {da_correggere} punti da correggere, {da_verificare} "
            "da verificare.")


def impronta(payload):
    """12 caratteri che identificano il contenuto dell'output di A1, a meno di spazi, a capo e ordine dei campi:
    la stessa impronta sul PC e in Copilot Studio vuol dire che i due precontrolli hanno visto lo stesso output."""
    testo = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(testo.encode("utf-8")).hexdigest()[:12]


def messaggio_correzione(ciclo, esito):
    """Testo da incollare nella chat di A1: i punti del precontrollo come richiesta
    di correzione, con le parole dei controlli (le stesse dell'allegato di A1)."""
    righe = [f"Il precontrollo dello script ha trovato questi punti nel tuo output del ciclo {ciclo}. Correggili e "
             "restituisci di nuovo l'intero JSON, solo JSON, senza cambiare cio' che non e' citato. I controlli sono "
             "quelli della parte 2 dell'allegato Input_A1, la forma dei testi e' nella parte 3."]

    def blocco(titolo, voci):
        if voci:
            righe.extend(["", titolo] + [f"{n}. {v}" for n, v in enumerate(voci, start=1)])
    voce = lambda x: f"{x['id_rischio'] or 'ciclo'} · {x['controllo']}: {x['descrizione']}"
    blocco("Schema (lo script non puo' leggere il JSON finche' non e' corretto):", esito["schema"])
    blocco("Da correggere (fermerebbero il ciclo):", [voce(x) for x in esito["rilievi"] if x["gravita"] == "Bloccante"])
    blocco("Da verificare (correggi se fondati; se confermi, spiega il perche' nel razionale):",
           [voce(x) for x in esito["rilievi"] if x["gravita"] != "Bloccante"])
    blocco("Forma dei testi (serve alle schede dei rischi):", [f"{x['id_rischio']} · {x['descrizione']}" for x in esito["formato"]])
    return "\n".join(righe) + "\n"
''',
}

def _carica():
    """I moduli del pacchetto, eseguiti in memoria nell'ordine di MODULI, come se fossero file accanto allo script."""
    for nome, sorgente in MODULI.items():
        modulo = types.ModuleType(nome)
        modulo.__file__ = f"precontrollo_nfr/{nome}.py"
        sys.modules[nome] = modulo
        exec(compile(sorgente, modulo.__file__, "exec"), modulo.__dict__)
    return sys.modules["precontrollo"], sys.modules["input_a1"]


def _leggi_json(percorso):
    """JSON da un file; se c'e' testo attorno (recinti ```json di una risposta in chat), dalla prima { all'ultima }."""
    testo = Path(percorso).read_text(encoding="utf-8-sig")
    try:
        return json.loads(testo)
    except json.JSONDecodeError:
        inizio, fine = testo.find("{"), testo.rfind("}")
        if inizio < 0 or fine < inizio:
            raise
        return json.loads(testo[inizio:fine + 1])


def main(argv):
    if argv[1:] == ["--verifica"]:
        print(f"precontrollo_nfr {VERSIONE}")
        for nome, sorgente in {**MODULI, "schema_output_A1.json": SCHEMA_A1}.items():
            print(f"  {nome}: {hashlib.sha256(sorgente.encode('utf-8')).hexdigest()[:12]}")
        return 0
    if len(argv) != 3:
        print(__doc__)
        return 2
    precontrollo, input_a1 = _carica()
    try:
        payload = _leggi_json(argv[1])
    except (OSError, ValueError) as ex:              # JSONDecodeError e' un ValueError
        print(f"File illeggibile: l'output di A1 {argv[1]} non e' un JSON ({ex}).")
        return 2
    try:
        testo = Path(argv[2]).read_text(encoding="utf-8-sig")
        baseline = input_a1.leggi_baseline(testo) if input_a1.INIZIO_BASELINE in testo else _leggi_json(argv[2])
    except (OSError, ValueError) as ex:
        print(f"File illeggibile: {argv[2]} non e' un allegato Input_A1 (manca la parte \"{input_a1.INIZIO_BASELINE}\") "
              f"ne' una baseline JSON ({ex}).")
        return 2
    if not isinstance(baseline, dict):
        print("File illeggibile: la baseline deve essere un oggetto JSON (Input_A1_<ciclo>.txt o Baseline_corrente.json).")
        return 2
    esito = precontrollo.precontrolla(payload, baseline, json.loads(SCHEMA_A1))
    ciclo = precontrollo.ciclo_di(payload, baseline)
    da_correggere, da_verificare = precontrollo.conta(esito)
    print(precontrollo.riga_esito(ciclo, esito))
    print(f"Impronta dell'output di A1: {precontrollo.impronta(payload)} · script precontrollo_nfr {VERSIONE}")
    if da_correggere or da_verificare:
        messaggio = precontrollo.messaggio_correzione(ciclo, esito)
        file = Path(argv[1]).with_name(f"Correzione_A1_{ciclo}.txt")
        file.write_text(messaggio, encoding="utf-8")
        print(f"Messaggio di correzione: {file}\n\n{messaggio}")
    return 1 if da_correggere else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
