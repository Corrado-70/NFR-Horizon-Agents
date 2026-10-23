"""Kit della prova K9 (Copilot Studio pieno, D-82): il motore eseguito da un prompt con code interpreter.

Un prompt con code interpreter riceve un solo file e il codice lo scrive il modello (Microsoft Learn, 30/09/2026).
La "busta" e' quel file: una cartella di lavoro Excel con il codice del motore a blocchi nel foglio _Motore e gli
ingressi nella tabella tblIngresso del foglio Ingresso. Il prompt (Da_caricare/prompt/Motore.txt) chiede solo di
riunire i blocchi in motore.py ed eseguirlo; il programma stampa una riga JSON con l'impronta del proprio codice,
che il flusso confronta con quella attesa: la prova che e' girato il nostro codice e non uno scritto dal modello.

Per chi mantiene il pacchetto (serve openpyxl, come genera_dashboard.py): rilanciarlo dopo prepara_kit.py. Scrive in
Da_caricare/prove/:
  K9_Busta_vuota.xlsx   codice e tabella degli ingressi vuota: la riempie il flusso NFR_Motore (passo T4.2)
  K9_Busta_prova.xlsx   gli stessi, con un output di A1 e il suo allegato dai dati del collaudo (passo T4.1)
  K9_atteso.txt         l'impronta attesa del codice e la riga che il motore deve stampare sulla busta di prova

Il codice della busta e' lo script del precontrollo (prepara_kit.py) con, al posto del suo avvio, il lettore della
busta: legge l'.xlsx con la sola libreria standard (zipfile e xml), senza dipendere da cio' che il sandbox installa.
Comando della prova: precontrollo. I comandi del motore vero (trimestrale, valida, ...) si aggiungono dopo l'esito di
K8/K9.

Uso:  python 09_Copilot_Studio/prepara_busta.py
"""

import hashlib
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

QUI = Path(__file__).resolve().parent
PACCHETTO = QUI.parent
PROVE = QUI / "Da_caricare" / "prove"
VUOTA, PROVA, ATTESO = PROVE / "K9_Busta_vuota.xlsx", PROVE / "K9_Busta_prova.xlsx", PROVE / "K9_atteso.txt"
SCRIPT = QUI / "precontrollo-nfr" / "scripts" / "precontrollo_nfr.py"
DATI = PACCHETTO / "99_Collaudo" / "dati"
BLOCCO = 30000                     # caratteri per cella (Excel ne tiene al massimo 32.767)
AVVIO = 'if __name__ == "__main__":\n    sys.exit(main(sys.argv))\n'

LETTORE = r'''

# ================================================================ busta (Copilot Studio pieno, prova K9)
# Al posto dell'avvio dello script: python motore.py <busta.xlsx> <comando>. Legge la busta con la sola libreria
# standard e stampa UNA riga JSON. impronta_codice e' calcolata sul file in esecuzione, impronta_busta sui blocchi
# del foglio _Motore: devono essere uguali fra loro e a quella attesa (K9_atteso.txt).
import contextlib as _contextlib
import io as _io
import re as _re
import tempfile as _tempfile
import zipfile as _zipfile
import xml.etree.ElementTree as _ET

_M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _testo(nodo):
    """Testo di una stringa di Excel: le t dirette e quelle delle r, senza le letture fonetiche (rPh)."""
    parti = []
    for figlio in (nodo if nodo is not None else []):
        if figlio.tag == _M + "t":
            parti.append(figlio.text or "")
        elif figlio.tag == _M + "r":
            parti.extend(t.text or "" for t in figlio.iter(_M + "t"))
    return _re.sub(r"_x([0-9A-Fa-f]{4})_", lambda m: chr(int(m.group(1), 16)), "".join(parti))


def _righe(percorso, foglio):
    """{numero di riga: {colonna: testo}} di un foglio dell'.xlsx."""
    with _zipfile.ZipFile(percorso) as z:
        nomi = z.namelist()
        condivise = ([_testo(si) for si in _ET.fromstring(z.read("xl/sharedStrings.xml"))]
                     if "xl/sharedStrings.xml" in nomi else [])
        rel = {r.get("Id"): r.get("Target") for r in _ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
        fogli = {s.get("name"): rel[s.get(_R + "id")] for s in _ET.fromstring(z.read("xl/workbook.xml")).iter(_M + "sheet")}
        if foglio not in fogli:
            raise ValueError(f"la busta non ha il foglio {foglio}")
        bersaglio = fogli[foglio].lstrip("/")
        bersaglio = bersaglio if bersaglio.startswith("xl/") else "xl/" + bersaglio
        righe = {}
        for c in _ET.fromstring(z.read(bersaglio)).iter(_M + "c"):
            col, riga = _re.match(r"([A-Z]+)(\d+)", c.get("r")).groups()
            v, tipo = c.find(_M + "v"), c.get("t")
            if tipo == "s":
                valore = condivise[int(v.text)]
            elif tipo == "inlineStr":
                valore = _testo(c.find(_M + "is"))
            else:
                valore = v.text if v is not None else None
            if valore not in (None, ""):
                righe.setdefault(int(riga), {})[col] = valore
        return righe


def leggi_busta(percorso):
    """(codice del motore, {nome: testo} degli ingressi riuniti per parte)."""
    motore = _righe(percorso, "_Motore")
    codice = "".join(motore[n]["A"] for n in sorted(motore) if n >= 2 and "A" in motore[n])
    ingressi = {}
    for n, r in sorted(_righe(percorso, "Ingresso").items()):
        if n >= 2 and r.get("A"):
            ingressi.setdefault(r["A"], []).append((int(float(r.get("B", 0))), r.get("C", "")))
    return codice, {nome: "".join(t for _, t in sorted(parti)) for nome, parti in ingressi.items()}


def _impronta(testo):
    return hashlib.sha256(testo.encode("utf-8")).hexdigest()[:12]


def main_busta(argv):
    esito = {"motore": VERSIONE, "comando": argv[2] if len(argv) > 2 else ""}
    try:
        esito["impronta_codice"] = _impronta(Path(__file__).read_text(encoding="utf-8"))
        codice, ingressi = leggi_busta(argv[1])
        esito["impronta_busta"] = _impronta(codice)
        if esito["comando"] != "precontrollo":
            raise ValueError(f"comando sconosciuto: {esito['comando']!r} (in questa versione solo precontrollo)")
        mancano = [n for n in ("A1", "Input_A1") if n not in ingressi]
        if mancano:
            raise ValueError(f"mancano nella tabella degli ingressi: {', '.join(mancano)}")
        with _tempfile.TemporaryDirectory() as cartella:
            a1, allegato = Path(cartella) / "A1.json", Path(cartella) / "Input_A1.txt"
            a1.write_text(ingressi["A1"], encoding="utf-8")
            allegato.write_text(ingressi["Input_A1"], encoding="utf-8")
            stampa = _io.StringIO()
            with _contextlib.redirect_stdout(stampa):
                esito["codice_uscita"] = main(["motore", str(a1), str(allegato)])
            righe = stampa.getvalue().splitlines()
            esito["esito"] = righe[0] if righe else ""
            esito["impronta_output"] = next((r.split(": ", 1)[1].split(" ")[0] for r in righe
                                             if r.startswith("Impronta dell'output")), "")
            correzioni = list(Path(cartella).glob("Correzione_A1_*.txt"))
            esito["correzione"] = correzioni[0].read_text(encoding="utf-8") if correzioni else ""
    except Exception as ex:                          # una riga JSON anche in errore: la legge il flusso
        esito["errore"] = f"{type(ex).__name__}: {ex}"
    print(json.dumps(esito, ensure_ascii=False))
    return 0 if "errore" not in esito else 2


if __name__ == "__main__":
    sys.exit(main_busta(sys.argv))
'''


def codice_motore():
    """Il codice della busta: lo script del precontrollo con il lettore della busta al posto dell'avvio."""
    script = SCRIPT.read_text(encoding="utf-8").replace("\r\n", "\n")
    if not script.endswith(AVVIO):
        raise SystemExit(f"{SCRIPT.name} non finisce con l'avvio atteso: rilanciare prepara_kit.py e controllare.")
    return script[:-len(AVVIO)].rstrip("\n") + "\n" + LETTORE


def impronta(testo):
    return hashlib.sha256(testo.encode("utf-8")).hexdigest()[:12]


def a_blocchi(testo):
    return [testo[i:i + BLOCCO] for i in range(0, len(testo), BLOCCO)] or [""]


def ingressi_prova():
    """Un output di A1 e il suo allegato, dai dati del collaudo (ciclo 2026Q4, senza rilievi)."""
    sys.path.insert(0, str(PACCHETTO / "07_Senza_Studio"))
    import input_a1
    baseline = json.loads((DATI / "baseline_2026Q4.json").read_text(encoding="utf-8"))
    allegato, _ = input_a1.componi(baseline, PACCHETTO / "01_Kernel", "2026-10-23")
    return {"A1": (DATI / "run3_A1_2026Q4.json").read_text(encoding="utf-8"), "Input_A1": allegato}


def scrivi_busta(percorso, codice, ingressi):
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb = Workbook()
    leggimi = wb.active
    leggimi.title = "Leggimi"
    for riga in ("Busta del motore NFR (Copilot Studio pieno, D-82). Non modificare a mano.",
                 "_Motore: il codice del motore a blocchi, colonna A dalla riga 2; impronta attesa in K9_atteso.txt.",
                 "Ingresso: tabella tblIngresso, una riga per blocco (Nome, Parte, Testo); la riempie il flusso NFR_Motore."):
        leggimi.append([riga])
    motore = wb.create_sheet("_Motore")
    motore.append(["codice"])
    for blocco in a_blocchi(codice):
        motore.append([blocco])
    motore.sheet_state = "hidden"
    ingresso = wb.create_sheet("Ingresso")
    ingresso.append(["Nome", "Parte", "Testo"])
    righe = [(nome, n, blocco) for nome, testo in ingressi.items() for n, blocco in enumerate(a_blocchi(testo), start=1)]
    for riga in righe or [("", None, "")]:          # una tabella di Excel vuole almeno una riga dati
        ingresso.append(list(riga))
    tabella = Table(displayName="tblIngresso", ref=f"A1:C{max(len(righe), 1) + 1}")
    tabella.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ingresso.add_table(tabella)
    wb.save(percorso)


def leggi_busta(busta):
    """(codice, ingressi) di una busta, con lo stesso lettore che sta nel codice del motore."""
    spazio = {"__name__": "lettore", "hashlib": hashlib, "json": json, "sys": sys, "Path": Path}
    exec(compile(LETTORE, "lettore", "exec"), spazio)
    return spazio["leggi_busta"](busta)


def esegui_come_il_prompt(busta, comando="precontrollo"):
    """Cio' che il prompt deve fare: riunire i blocchi di _Motore in motore.py ed eseguirlo sulla busta."""
    codice, _ = leggi_busta(busta)
    with tempfile.TemporaryDirectory() as cartella:
        motore = Path(cartella) / "motore.py"
        motore.write_text(codice, encoding="utf-8", newline="\n")
        r = subprocess.run([sys.executable, str(motore), str(busta), comando], capture_output=True, encoding="utf-8",
                           env={**os.environ, "PYTHONIOENCODING": "utf-8"})
    return codice, r.returncode, r.stdout.strip()


def main():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise SystemExit("Serve openpyxl (pip install openpyxl), come per genera_dashboard.py.")
    PROVE.mkdir(parents=True, exist_ok=True)
    codice = codice_motore()
    scrivi_busta(VUOTA, codice, {})
    scrivi_busta(PROVA, codice, ingressi_prova())
    letto, uscita_codice, riga = esegui_come_il_prompt(PROVA)
    if letto != codice:
        raise SystemExit("La busta di prova non restituisce il codice identico: non usarla.")
    esito = json.loads(riga)
    if esito.get("impronta_codice") != impronta(codice) or esito.get("impronta_busta") != impronta(codice) \
            or "errore" in esito:
        raise SystemExit(f"Il motore sulla busta di prova non da' l'esito atteso: {riga[:400]}")
    ATTESO.write_text(
        "PROVA K9: COSA DEVE STAMPARE IL MOTORE (Copilot Studio pieno, D-82)\n\n"
        f"Impronta attesa del codice (impronta_codice e impronta_busta): {impronta(codice)}\n"
        f"Blocchi del foglio _Motore: {len(a_blocchi(codice))}, {len(codice):,} caratteri\n\n"
        "Riga attesa sulla busta di prova K9_Busta_prova.xlsx, comando precontrollo (identica, a parte l'ordine dei "
        "campi):\n" + riga + "\n", encoding="utf-8", newline="\n")
    print(f"Busta del motore: {len(codice):,} caratteri in {len(a_blocchi(codice))} blocchi, impronta {impronta(codice)}")
    for f in (VUOTA, PROVA, ATTESO):
        print(f"  {f.relative_to(PACCHETTO)}")
    print(f"  esito sulla busta di prova (codice {uscita_codice}): {esito.get('esito')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
